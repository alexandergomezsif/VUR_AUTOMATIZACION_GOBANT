"""
Automatización VUR - Consulta de gravámenes por valorización.

El aplicativo usa Chrome (Selenium) solo para que el usuario inicie sesión en
VUR. Las consultas se hacen a los mismos servicios JSON que usa la página de
VUR, dentro de esa sesión:
    webresources/consultaIndices/datosBasicos/tierras/<círculo>-<número>
    webresources/consultaIndices/estadoJuridico/tierras/<círculo>-<número>
Así no se depende de leer tablas en pantalla.

Flujo por matrícula:
  1. Datos básicos: propietarios (todos), dirección, referencia catastral,
     municipio, Nupre, fecha de apertura del folio y matrículas matriz.
  2. Clasificación del municipio con la hoja "Proy x Mpio" del Excel de
     proyectos (incorporada al final de este archivo, sección DATOS DE PROYECTOS):
       000 / NO APLICA  -> no se consulta estado jurídico.
       001 / CONDONADA  -> se consulta estado jurídico.
       Sin clasificar   -> se consulta estado jurídico por precaución.
  3. Estado jurídico: se analizan las anotaciones. Es de valorización la que
     tiene alguna PALABRAS_VALORIZACION en su especificación. Una anotación
     posterior con "Se cancela anotación No: N" la marca como cancelada.
  4. Si no hay anotación de valorización y el folio se abrió DESPUÉS de la
     fecha límite del municipio (proyecto más antiguo), se revisan sus
     matrículas matriz como consultas nuevas, nivel por nivel y sin tope de
     niveles, hasta encontrar la valorización (y a quién se impuso) o llegar a
     folios sin matriz o anteriores a la fecha límite.

Todo el aplicativo está en ESTE archivo: código, datos de proyectos por
municipio, logo, icono, tipografía Poppins y pruebas automáticas. El logo, el
icono y las fuentes se recrean solos en %LOCALAPPDATA%\\VUR_Valorizacion\\recursos.

Uso:
  pythonw vur_valorizacion.py              abre el aplicativo
  python  vur_valorizacion.py --probar     ejecuta las pruebas automáticas
  python  vur_valorizacion.py --recursos   recrea logo, icono y fuentes
Instalación: Instalar_VUR.bat (descarga el ZIP del repositorio y ejecútelo).
Datos de proyectos: se regeneran con desarrollo/generar_datos_valorizacion.py.

Requisitos:
  pip install --user selenium pandas openpyxl xlrd
  (Selenium >= 4.6 gestiona ChromeDriver automáticamente.)
"""

import html
import json
import math
import os
import queue
import re
import threading
import tkinter as tk
import unicodedata
from collections import deque
from datetime import date, datetime
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

import pandas as pd
from selenium import webdriver
from selenium.common.exceptions import (
    InvalidSessionIdException,
    JavascriptException,
    NoAlertPresentException,
    NoSuchWindowException,
    TimeoutException,
    WebDriverException,
)

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

URL_INICIO = "https://www.vur.gov.co/portal/pages/vur/inicio.jsf"
URL_API = "https://www.vur.gov.co/portal/PantallasVUR/webresources/consultaIndices/"
TIPO_CONSULTA = "tierras"  # el mismo que usan los módulos "...Tierras" del menú de VUR

# Datos propios de cada usuario de Windows (no van en la carpeta de trabajo):
# perfil de Chrome del programa (sesión de VUR) y recursos (logo, icono, fuentes).
CARPETA_USUARIO = Path(os.environ.get("LOCALAPPDATA") or Path.home()) / "VUR_Valorizacion"
_PERFIL_ANTERIOR = Path(__file__).resolve().parent / "selenium_profile"  # ubicación de versiones previas
PERFIL_CHROME = (_PERFIL_ANTERIOR if _PERFIL_ANTERIOR.is_dir() and not (CARPETA_USUARIO / "perfil_chrome").is_dir()
                 else CARPETA_USUARIO / "perfil_chrome")

MAX_INTENTOS = 3
GUARDADO_PARCIAL_CADA = 10
TIMEOUT_API = 120  # segundos por consulta a VUR

# Palabras que identifican una anotación de valorización en su ESPECIFICACION.
# "GRAVAMEN" no se usa sola: también aparece en hipotecas ("NATURALEZA JURIDICA
# DEL ACTO: GRAVAMEN"), lo que producía falsos positivos.
PALABRAS_VALORIZACION = ["VALORIZA", "INVIAS"]

# Revisión de matrículas matriz.
# No hay tope de niveles: se sube hasta que el folio no tenga matriz o sea
# anterior a la fecha límite. MAX_FOLIOS_SEGURIDAD solo protege contra un
# recorrido descontrolado (cada folio son 2 consultas del cupo diario de VUR).
MAX_FOLIOS_SEGURIDAD = 300
# Si la revisión no se puede completar (tope de seguridad o una matriz sin
# respuesta), el estado queda "NO GRAVADA - REVISION INCOMPLETA".
ESTADO_INCOMPLETO = "NO GRAVADA - REVISION INCOMPLETA"
# Fecha límite del municipio: la más antigua de sus proyectos, tomando para cada
# proyecto la Fecha de Inicio (hoja Vencimiento) o, si no la tiene, la fecha de
# la resolución (hoja Proyectos Valorización) o el 1 de enero de su año.
# Si el folio se abrió después de esa fecha, se revisa su matriz.

# Orden de columnas del Excel de resultados.
COLUMNAS_SALIDA = [
    "Matricula_Original", "Matricula_Limpia",
    "Propietario", "Total_Propietarios", "Tipo_Identificacion", "Numero_Identificacion",
    "Propietarios_Actuales",
    "Direccion_Inmueble", "Matricula_VUR", "Referencia_Catastral", "Departamento", "Municipio",
    "Nupre", "Tipo_Predio", "Estado_Folio", "Fecha_Apertura_Folio", "Matriculas_Matriz",
    "Clasificacion_Municipio", "Proyectos_Municipio", "Fecha_Limite_Revision",
    "Estado", "Confiabilidad", "Motivos_Revision", "Segunda_Pasada",
    "Matricula_Con_Anotacion", "Nivel_Matriz", "Detalle_Valorizacion", "Propietario_Gravado",
    "Matricula_Cancelacion", "Observaciones", "Afectaciones_Valorizacion",
    "Cadena_Revisada", "Folios_Revisados", "Notas_Proceso", "Intentos", "Fecha_Consulta", "Version_Aplicativo",
    "Fuente", "Funcionario", "Equipo",
]

# Orden del Excel de resultados (pedido de los profesionales, 27/10/2026): primero lo
# CONCLUYENTE (si tiene valorización, en qué folio se halló y si requiere revisión),
# luego lo que ellos usan; el resto queda AGRUPADO y oculto (botón [+] encima de las
# columnas). No se elimina: lo usan el histórico, "Comparar" y la trazabilidad.
COLUMNAS_VISIBLES = [
    "Matricula_Original", "Resultado", "Confiabilidad", "Motivos_Revision", "Estado",
    "Matricula_Con_Anotacion", "Nivel_Matriz", "Matricula_Cancelacion", "Propietario_Gravado", "Municipio",
    "Propietario", "Tipo_Identificacion", "Numero_Identificacion", "Propietarios_Actuales",
    "Observaciones", "Afectaciones_Valorizacion", "Cadena_Revisada", "Folios_Revisados", "Intentos",
    "Fuente", "Fecha_Consulta",
]
COLUMNAS_EXCEL = COLUMNAS_VISIBLES + [c for c in COLUMNAS_SALIDA if c not in COLUMNAS_VISIBLES]

# Versión de las reglas de análisis. Se escribe en cada fila del Excel para saber
# con qué reglas se obtuvo cada resultado.
VERSION_APLICATIVO = "2026.10.27-4"
# Versión de las REGLAS de análisis. Se cambia solo cuando cambia cómo se clasifica una
# matrícula; así los resultados guardados en el histórico con reglas anteriores no se
# reutilizan y se vuelven a consultar. Los cambios de pantalla no la modifican.
VERSION_REGLAS = "2026.10.26-4"

# Confiabilidad de cada resultado:
#   ALTA    -> todo lo que sustenta el estado viene explícito en VUR.
#   REVISAR -> algún dato se dedujo, faltó información o hay una alerta; los
#              motivos quedan en Motivos_Revision. Debe verlo un profesional.
# Segunda pasada: al terminar el lote, las matrículas cuyo resultado pudo deberse a
# una falla momentánea de VUR se consultan de nuevo tras una pausa. No se repiten
# las marcadas REVISAR por el contenido (propietario deducido, afectación...),
# porque consultar otra vez daría lo mismo.
ESPERA_SEGUNDA_PASADA = 60  # segundos de pausa antes de la segunda pasada


def necesita_segunda_pasada(reg):
    """True si el resultado pudo deberse a una respuesta fallida o vacía de VUR."""
    if reg.get("Estado") in ("SIN DATOS EN VUR", "ERROR"):
        return "Formato de matrícula inválido" not in (reg.get("Notas_Proceso") or "")
    if reg.get("Estado") == ESTADO_INCOMPLETO:
        notas = reg.get("Notas_Proceso") or ""
        return "no devolvió" in notas or "no se pudo revisar" in notas
    return "no devolvió el estado jurídico" in (reg.get("Notas_Proceso") or "")


CONFIABILIDAD_ALTA = "ALTA"
CONFIABILIDAD_REVISAR = "REVISAR"

# ---------------------------------------------------------------------------
# Datos de valorización por municipio (hoja "Proy x Mpio")
# ---------------------------------------------------------------------------

# Los datos (FUENTE_DATOS y MUNICIPIOS_VALORIZACION) están al final del archivo.

# Nombre como lo escribe VUR -> nombre en la hoja "Proy x Mpio".
ALIAS_MUNICIPIOS = {
    "SAN ANDRES DE CUERQUIA": "SAN ANDRES DE C.",
    "SAN JOSE DE LA MONTAÑA": "SAN JOSE DE LA M.",
    "SAN PEDRO DE LOS MILAGROS": "SAN PEDRO DE LOS M.",
    "SANTA FE DE ANTIOQUIA": "SANTA FE DE ANT",
    "SANTA ROSA DE OSOS": "SANTA ROSA",
    "DONMATIAS": "DON MATIAS",
    "EL CARMEN DE VIBORAL": "CARMEN DE VIBORAL",
    "EL SANTUARIO": "SANTUARIO",
    "EL RETIRO": "RETIRO",
    "EL PEÑOL": "PEÑOL",
    "CAROLINA DEL PRINCIPE": "CAROLINA",
    "SAN VICENTE FERRER": "SAN VICENTE",
}

ETIQUETAS_ESTADO_MUNICIPIO = {
    "000": "000 - NO GRAVADO",
    "001": "001 - BUSCAR SI TIENE GRAVAMEN",
    "NO_APLICA": "NO APLICA",
    "CONDONADA": "CONDONADA",
}

# Estados en los que NO se consulta el estado jurídico. MEDELLÍN figura como
# "NO APLICA" en la hoja; ARBOLETES (CONDONADA) sí se consulta para dejar
# constancia de la anotación y de su cancelación.
ESTADOS_SIN_CONSULTA = {"000", "NO_APLICA"}

# ---------------------------------------------------------------------------
# Utilidades puras
# ---------------------------------------------------------------------------


def normalizar(texto):
    """Mayúsculas, sin tildes ni diéresis, espacios colapsados."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFD", str(texto))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()


def clave_municipio(texto):
    """Clave de comparación: normalizado, sin código numérico inicial ni punto final."""
    # VUR envía la Ñ dañada ("PE\ufffdOL"); es la única letra no ASCII de los municipios
    # que llega así, por lo que se reemplaza por N (igual que la Ñ normalizada).
    t = re.sub(r"^[\d\s\-]+", "", normalizar(str(texto or "").replace("\ufffd", "N")))
    return t.rstrip(".").strip()


def _construir_clasificacion():
    clasificacion = {clave_municipio(n): info for n, info in MUNICIPIOS_VALORIZACION.items()}
    for alias, nombre in ALIAS_MUNICIPIOS.items():
        destino = clave_municipio(nombre)
        if destino not in clasificacion:
            raise ValueError(f"Alias apunta a municipio inexistente: {nombre}")
        clasificacion[clave_municipio(alias)] = clasificacion[destino]
    return clasificacion


_CLASIFICACION_MUNICIPIOS = None


def clasificar_municipio(municipio):
    """Devuelve (estado, etiqueta, info) o (None, '', None) si no está en la hoja."""
    global _CLASIFICACION_MUNICIPIOS
    if _CLASIFICACION_MUNICIPIOS is None:  # se arma la primera vez (los datos están al final del archivo)
        _CLASIFICACION_MUNICIPIOS = _construir_clasificacion()
    info = _CLASIFICACION_MUNICIPIOS.get(clave_municipio(municipio))
    if not info:
        return None, "", None
    return info["estado"], ETIQUETAS_ESTADO_MUNICIPIO[info["estado"]], info


def describir_proyectos(info):
    """'008 URRAO - BETULIA (res. 1187, 1998, vence 2007-05-01)' separados por ' | '."""
    if not info:
        return ""
    partes = []
    for p in info.get("proyectos", []):
        detalle = [x for x in (
            f"res. {p['resolucion']}" if p.get("resolucion") else "",
            str(p["anio"]) if p.get("anio") else "",
            f"vence {p['fecha_vencimiento']}" if p.get("fecha_vencimiento") else "",
        ) if x]
        texto = f"{p.get('codigo') or ''} {p.get('obra', '')}".strip()
        partes.append(texto + (f" ({', '.join(detalle)})" if detalle else ""))
    return " | ".join(partes)


def parsear_fecha(texto):
    """Acepta dd/mm/aaaa, dd-mm-aaaa y aaaa-mm-dd. Devuelve date o None."""
    t = (texto or "").strip()
    for patron, orden in ((r"(\d{4})-(\d{1,2})-(\d{1,2})", "amd"), (r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", "dma")):
        m = re.search(patron, t)
        if m:
            a, b, c = (int(x) for x in m.groups())
            try:
                return date(c, b, a) if orden == "dma" else date(a, b, c)
            except ValueError:
                return None
    return None


def fecha_limite_municipio(info):
    """Fecha más antigua entre los proyectos del municipio (ver configuración)."""
    fechas = []
    for p in (info or {}).get("proyectos", []):
        f = parsear_fecha(p.get("fecha_inicio")) or parsear_fecha(p.get("fecha_resolucion"))
        if not f and p.get("anio"):
            f = date(int(p["anio"]), 1, 1)
        if f:
            fechas.append(f)
    return min(fechas) if fechas else None


def limpiar_matricula(valor):
    t = re.sub(r"^MI\s*", "", normalizar(valor))
    return t.replace(" ", "")


def separar_matricula(matricula):
    """'01N-0005202' -> ('01N', '5202'). El servicio de VUR no acepta ceros a la izquierda."""
    m = re.fullmatch(r"([0-9A-Z]{3})-(\d+)", matricula or "")
    return (m.group(1), str(int(m.group(2)))) if m else None


def html_a_texto(valor):
    t = re.sub(r"<br\s*/?>", "\n", valor or "", flags=re.IGNORECASE)
    t = re.sub(r"<[^>]+>", " ", t)
    t = html.unescape(t).replace("\r", "\n")
    return "\n".join(re.sub(r"[ \t]+", " ", l).strip() for l in t.split("\n") if l.strip())


RE_ANOTACION = re.compile(r"ANOTACION:\s*Nro\s*(\d+)\s*Fecha:\s*([\d/-]+)", re.IGNORECASE)
RE_CANCELA = re.compile(r"Se cancela anotaci\S*n No:\s*([\d,\s]+)", re.IGNORECASE)


def parsear_anotacion(valor_html):
    """Convierte una entrada de textoAnotaciones de VUR en un dict."""
    texto = html_a_texto(valor_html)
    m = RE_ANOTACION.search(texto)
    espec = re.search(r"ESPECIFICACION:\s*(.+)", texto)
    naturaleza = re.search(r"NATURALEZA JURIDICA DEL ACTO:\s*(.+)", texto)
    cancela = []
    for grupo in RE_CANCELA.findall(texto):
        cancela.extend(int(n) for n in re.findall(r"\d+", grupo))
    lineas = texto.split("\n")
    especificacion = espec.group(1).strip() if espec else ""
    if not especificacion and naturaleza:
        # Algunas anotaciones antiguas no traen ESPECIFICACION y solo dicen la naturaleza
        # del acto; se usa esa para clasificarlas y no dejar pasar una valorización.
        especificacion = naturaleza.group(1).strip()
    return {
        "nro": int(m.group(1)) if m else None,
        "fecha": m.group(2) if m else "",
        "especificacion": especificacion,
        "sin_especificacion": not espec,
        "naturaleza": naturaleza.group(1).strip() if naturaleza else "",
        "cancela": cancela,
        "de": [l[3:].strip() for l in lineas if l.startswith("DE:")],
        "a": [l[2:].strip() for l in lineas if l.startswith("A:")],
        "texto": texto,
    }


# Códigos de especificación de VUR observados en gravámenes de valorización reales:
#   430 "GRAVAMEN POR CONTRIBUCION DE VALORIZACION" (tabla de códigos anterior)
#   0212 "VALORIZACION ..." (tabla actual)
# Una anotación que menciona valorización con otro código no se descarta: se
# cuenta como gravamen pero el resultado queda marcado para revisión.
CODIGOS_GRAVAMEN_VALORIZACION = {"430", "212"}


def _codigo(especificacion):
    m = re.match(r"\s*(\d{3,4})\b", especificacion or "")
    return m.group(1).lstrip("0") if m else ""


def es_cancelacion(anotacion):
    """Anotación que cancela otra: código 08xx, '(CANCELACION)' o 'Se cancela anotación No:'."""
    espec = normalizar(anotacion["especificacion"])
    return bool(anotacion["cancela"]) or "(CANCELACION)" in espec or bool(re.match(r"0?8\d\d\b", espec)) \
        or bool(re.match(r"(\d{3,4}\s+)?CANCELACION", espec))  # p. ej. "0723 CANCELACION VALORIZACION"


def _menciona_valorizacion(anotacion):
    espec = normalizar(anotacion["especificacion"])
    return any(normalizar(p) in espec for p in PALABRAS_VALORIZACION)


def _menciona_afectacion(anotacion):
    texto = normalizar(anotacion["especificacion"] + " " + anotacion["naturaleza"])
    return "AFECTACION" in texto or "INENAJENAB" in texto or "MEDIDA CAUTELAR" in texto


def tipo_valorizacion(anotacion):
    """None si no tiene relación con valorización; si no: 'cancelacion', 'afectacion',
    'gravamen' (código conocido) o 'dudosa' (menciona valorización con código no reconocido)."""
    if not _menciona_valorizacion(anotacion):
        return None
    if es_cancelacion(anotacion):
        return "cancelacion"
    if _menciona_afectacion(anotacion):
        return "afectacion"
    if _codigo(anotacion["especificacion"]) in CODIGOS_GRAVAMEN_VALORIZACION:
        return "gravamen"
    return "dudosa"


def es_valorizacion(anotacion):
    """Gravamen por valorización (no sus cancelaciones ni las afectaciones de inenajenabilidad)."""
    return tipo_valorizacion(anotacion) in ("gravamen", "dudosa")


def _referencias(texto):
    """Números de obra/resolución citados ('OBRA 808', 'RESOLUCION 939') para vincular cancelaciones."""
    return set(re.findall(r"(?:OBRA|RESOLUCION|RES\.?|OFICIO)\s*(?:N[O°.]*\s*)?(\d{2,})", normalizar(texto)))


def analizar_folio(anotaciones):
    """Clasifica las anotaciones de un folio.

    Devuelve dict con:
      valorizacion: gravámenes por valorización (con 'cancelada_por' y 'vinculo':
                    'numero' si la cancelación cita su número, 'texto' si se dedujo).
      afectaciones: afectaciones de inenajenabilidad por valorización (no deciden el
                    estado; se informan al profesional), con 'cancelada_por'.
      cancelaciones: cancelaciones de los gravámenes, o de una valorización
                    inscrita en otro folio (la cita no existe en este folio).
      estado: 'GRAVADA', 'GRAVAMEN CANCELADO' o 'NO GRAVADA'.

    Algunas cancelaciones de VUR no traen "Se cancela anotación No:" (p. ej.
    01N-5052127, anotación 6 "CANCELACION ... DE VALORIZACION OBRA 808"). En
    ese caso se vinculan a las anotaciones vigentes anteriores del mismo tipo
    que citen la misma obra o resolución; si no comparten referencia, a todas
    las vigentes anteriores de ese tipo. El vínculo queda marcado 'texto'.
    """
    valorizacion = [dict(a, cancelada_por=None, vinculo="") for a in anotaciones if es_valorizacion(a)]
    afectaciones = [dict(a, cancelada_por=None, vinculo="") for a in anotaciones
                    if tipo_valorizacion(a) == "afectacion"]
    por_nro = {a["nro"]: a for a in valorizacion + afectaciones}
    nros_folio = {a["nro"] for a in anotaciones}
    cancelaciones = []
    for a in anotaciones:
        if not es_cancelacion(a):
            continue
        objetivos = [n for n in a["cancela"] if n in por_nro]
        vinculo = "numero"
        if not objetivos and a["cancela"]:
            if _menciona_valorizacion(a) and not any(n in nros_folio for n in a["cancela"]):
                cancelaciones.append(a)  # cita una anotación de otro folio (gravamen en la matriz)
            continue
        if not objetivos and _menciona_valorizacion(a):
            grupo = afectaciones if _menciona_afectacion(a) else valorizacion
            previas = [x for x in grupo if x["cancelada_por"] is None and (x["nro"] or 0) < (a["nro"] or 10**9)]
            refs = _referencias(a["texto"])
            comunes = [x for x in previas if refs & _referencias(x["texto"])]
            objetivos = [x["nro"] for x in (comunes or previas)]
            vinculo = "texto"
            if not objetivos:
                if not _menciona_afectacion(a):
                    cancelaciones.append(a)  # cancela una valorización inscrita en otro folio
                continue
        if not objetivos:
            continue
        if any(por_nro[n] in valorizacion for n in objetivos):
            cancelaciones.append(a)
        for n in objetivos:
            por_nro[n]["cancelada_por"] = a["nro"]
            por_nro[n]["vinculo"] = vinculo
            por_nro[n]["cancelacion"] = a  # se informa junto a la afectación cancelada
    if not valorizacion:
        estado = "GRAVAMEN CANCELADO" if cancelaciones else "NO GRAVADA"
    elif any(a["cancelada_por"] is None for a in valorizacion):
        estado = "GRAVADA"
    else:
        estado = "GRAVAMEN CANCELADO"
    return {"valorizacion": valorizacion, "afectaciones": afectaciones,
            "cancelaciones": cancelaciones, "estado": estado}


def analizar_anotaciones(anotaciones):
    """Compatibilidad: (valorizacion, cancelaciones, estado)."""
    r = analizar_folio(anotaciones)
    return r["valorizacion"], r["cancelaciones"], r["estado"]


def describir_no_reconocidas(etiqueta, anotaciones):
    """Explica qué anotaciones no se pudieron leer y si podrían ocultar una valorización.

    Devuelve (motivos, notas). Las entradas vacías que a veces trae VUR se ignoran.
    """
    motivos, notas = [], []
    for i, a in enumerate(anotaciones, 1):
        if not a["texto"].strip():
            continue
        if a["nro"] is None:
            falla = f"entrada {i} sin número de anotación"
        elif not a["especificacion"]:
            falla = f"anotación {a['nro']} sin especificación ni naturaleza del acto"
        else:
            if a.get("sin_especificacion"):
                notas.append(f"{etiqueta}: la anotación {a['nro']} no trae ESPECIFICACION; "
                             "se clasificó por la naturaleza del acto")
            continue
        menciona = any(normalizar(p) in normalizar(a["texto"]) for p in PALABRAS_VALORIZACION)
        aviso = " y MENCIONA VALORIZACIÓN: revisarla en VUR" if menciona else ""
        motivos.append(f"{etiqueta}: {falla}{aviso}")
        notas.append(f"Texto de la {falla} de {etiqueta}: «{a['texto'][:300]}»")
    return motivos, notas


def es_adquisicion(anotacion):
    return "MODO DE ADQUISICION" in normalizar(anotacion["naturaleza"] + " " + anotacion["especificacion"])


def propietarios_en_fecha(anotaciones, nro_limite=None):
    """Personas 'A:' de la última anotación de adquisición (compraventa, adjudicación...)
    anterior a la anotación nro_limite (o la última del folio si no se indica).
    Devuelve (personas, nro_anotacion) o ([], None)."""
    candidatas = [a for a in anotaciones if es_adquisicion(a) and a["a"]
                  and (nro_limite is None or (a["nro"] or 0) < nro_limite)]
    if not candidatas:
        return [], None
    ultima = max(candidatas, key=lambda a: a["nro"] or 0)
    return ultima["a"], ultima["nro"]


def propietarios_hasta_fecha(anotaciones, fecha):
    """Personas 'A:' de la última adquisición inscrita en o antes de 'fecha'. ([], None) si no hay."""
    candidatas = []
    for a in anotaciones:
        f = parsear_fecha(a["fecha"])
        if es_adquisicion(a) and a["a"] and f and f <= fecha:
            candidatas.append((f, a["nro"] or 0, a))
    if not candidatas:
        return [], None
    ultima = max(candidatas, key=lambda x: x[:2])[2]
    return ultima["a"], ultima["nro"]


MAX_MATRICES_PROPIETARIO = 3  # en un englobe se miran como máximo las 3 primeras matrices


def propietario_desde_matrices(api, folio, db_folio, fecha_gravamen):
    """Deduce a quién se le impuso un gravamen que no nombra al propietario.

    Caso: el gravamen está en un folio sin compra anterior (p. ej. 01N-5052127 de
    Bello). Ese folio nació de su matriz, así que el dueño al abrirse es el dueño
    de la matriz en esa fecha. Se busca la última adquisición de la matriz hasta
    la apertura del folio (o hasta la fecha del gravamen si no hay apertura).

    Devuelve una lista de textos "PERSONA (dueño de la matriz X según anotación N)".
    Es siempre una deducción: el resultado debe quedar para revisión.
    """
    corte = parsear_fecha((db_folio or {}).get("fechaAperturaFolio")) or parsear_fecha(fecha_gravamen)
    if not corte:
        return []
    hallados = []
    for m in ((db_folio or {}).get("matriculasMatriz") or [])[:MAX_MATRICES_PROPIETARIO]:
        partes = separar_matricula(m)
        if not partes:
            continue
        ej = api.estado_juridico(*partes)
        anotaciones = [parsear_anotacion(t) for t in (ej or {}).get("textoAnotaciones") or []]
        personas, nro = propietarios_hasta_fecha(anotaciones, corte)
        hallados.extend(f"{limpiar_persona(p)} (dueño de la matriz {m} según su anotación {nro}, "
                        f"deducido)" for p in personas)
    return hallados


def limpiar_persona(texto):
    """'MONTOYA RAMIREZ JOSE   CC 123  X' -> 'MONTOYA RAMIREZ JOSE (CC 123) X'."""
    t = re.sub(r"\s+", " ", texto or "").strip()
    m = re.search(r"\b(CC|NIT|TI|CE|PA)\s*([\d.\-]+)", t)
    if m:
        nombre, resto = t[:m.start()].strip(), t[m.end():].strip()
        t = f"{nombre} ({m.group(1)} {m.group(2)})" + (f" {resto}" if resto else "")
    # En VUR, "X" al final indica titular del derecho real de dominio e "I", dominio incompleto.
    t = re.sub(r"\s+X$", " [titular]", t)
    return re.sub(r"\s+I$", " [dominio incompleto]", t)


def _tipo_documento(p):
    tipo = p.get("tipoDocumento")
    return (tipo.get("nombre") if isinstance(tipo, dict) else tipo) or ""


def formatear_propietarios(propietarios):
    """Lista JSON de VUR -> 'NOMBRE (CC 123) 50% [titular]; ...'."""
    partes = []
    for p in propietarios or []:
        nombre = re.sub(r"\s+", " ", p.get("nombre") or "").strip()
        ident = " ".join(x for x in (_tipo_documento(p), (p.get("numeroDocumento") or "").strip()) if x)
        texto = nombre + (f" ({ident})" if ident else "")
        pct = p.get("porcentajeParticipacion")
        if pct not in (None, ""):
            texto += f" {pct}" if str(pct).endswith("%") else f" {pct}%"
        if (p.get("titular") or "").strip().upper() == "X":
            texto += " [titular]"
        partes.append(texto)
    return "; ".join(partes)


def extraer_campos(db):
    """Columnas de datos básicos a partir del JSON de VUR."""
    props = db.get("propietarios") or []
    primero = props[0] if props else {}
    return {
        "Propietario": re.sub(r"\s+", " ", primero.get("nombre") or "").strip(),
        "Total_Propietarios": len(props),
        "Tipo_Identificacion": _tipo_documento(primero),
        "Numero_Identificacion": (primero.get("numeroDocumento") or "").strip(),
        "Propietarios_Actuales": formatear_propietarios(props),
        "Direccion_Inmueble": re.sub(r"\s+", " ", db.get("direccion") or "").strip(),
        "Matricula_VUR": f"{db.get('circulo', '')}-{db.get('numeroMatricula', '')}",
        "Referencia_Catastral": db.get("referenciaCatastral") or "",
        "Departamento": (db.get("departamento") or {}).get("nombre", ""),
        "Municipio": ((db.get("municipio") or {}).get("nombre") or "").replace("\ufffd", "Ñ"),
        "Nupre": db.get("nupre") or "",
        "Tipo_Predio": db.get("tipoPredio") or "",
        "Estado_Folio": db.get("estadoFolio") or "",
        "Fecha_Apertura_Folio": db.get("fechaAperturaFolio") or "",
        "Matriculas_Matriz": ", ".join(db.get("matriculasMatriz") or []),
    }


LIMITE_CELDA_EXCEL = 32000  # Excel admite 32.767 caracteres por celda.


def recortar(texto):
    return texto if len(texto) <= LIMITE_CELDA_EXCEL else texto[:LIMITE_CELDA_EXCEL] + " …[recortado]"


def leer_excel(ruta):
    df = pd.read_excel(ruta, dtype=str)
    columna = next((c for c in df.columns if "MATRICULA" in normalizar(c)), None)
    if columna is None:
        raise ValueError("No se encontró una columna cuyo encabezado contenga 'Matricula'.")
    datos = df[[columna]].rename(columns={columna: "Matricula_Original"})
    datos["Matricula_Original"] = datos["Matricula_Original"].fillna("").astype(str).str.strip()
    datos = datos[(datos["Matricula_Original"] != "") & (datos["Matricula_Original"].str.lower() != "nan")]
    datos["Matricula_Limpia"] = datos["Matricula_Original"].map(limpiar_matricula)
    return datos.reset_index(drop=True)


def mensaje_corto(error):
    texto = getattr(error, "msg", None) or str(error)
    return texto.strip().split("\n")[0][:200]


# ---------------------------------------------------------------------------
# Revisión de una matrícula y su cadena de matrices (independiente del navegador)
# ---------------------------------------------------------------------------


def revisar_cadena(api, circulo, numero, db_inicial, fecha_limite, log=lambda m: None):
    """Determina si la matrícula tiene o tuvo valorización, subiendo por sus matrices.

    `api` debe ofrecer datos_basicos(circ, num) y estado_juridico(circ, num).
    Recorre las matrices por niveles (todas las del nivel 1, luego las del
    nivel 2...) sin tope de niveles. Una rama se deja de subir solo cuando:
      - el folio no tiene matriz, o
      - el folio se abrió en o antes de la fecha límite del municipio (ningún
        proyecto pudo inscribirse en un folio anterior a él), o
      - ya se encontró la anotación de valorización.
    MAX_FOLIOS_SEGURIDAD solo evita un recorrido descontrolado; si se alcanza,
    el resultado se marca como revisión incompleta.

    Si un folio solo tiene la CANCELACION de una valorización inscrita antes
    (en la matriz), se sigue subiendo para encontrar el gravamen original y
    a quién se le impuso.

    Devuelve dict con estado, matricula (folio con el gravamen), valorizacion,
    cancelaciones, matricula_cancelacion, cadena, niveles y notas.
    """
    cola = deque([(circulo, numero, db_inicial, 0)])
    vistos = {f"{circulo}-{numero}"}
    cadena, notas = [], []
    folios = 0
    cancelacion_previa = None   # (folio, anotaciones) de una cancelación sin su gravamen
    anotaciones_por_folio = {}
    datos_por_folio = {}   # datos básicos de cada folio revisado (para buscar dueños en su matriz)
    afectaciones = []      # (folio, anotación) de afectaciones de inenajenabilidad por valorización
    no_reconocidas = []    # folios con anotaciones que no se pudieron leer
    aviso_sin_fecha = False
    incompleta = False
    nivel_max = 0

    def resultado(estado, matricula="", valorizacion=(), cancelaciones=(), matricula_canc="", nivel=None):
        return {"estado": estado, "matricula": matricula, "nivel": nivel, "valorizacion": list(valorizacion),
                "cancelaciones": list(cancelaciones), "matricula_cancelacion": matricula_canc,
                "cadena": cadena, "niveles": nivel_max, "notas": list(dict.fromkeys(notas)),
                "anotaciones_folio": anotaciones_por_folio.get(matricula, []),
                "afectaciones": afectaciones, "no_reconocidas": no_reconocidas,
                "anotaciones_inicial": anotaciones_por_folio.get(f"{circulo}-{numero}", []),
                "datos_folio": datos_por_folio.get(matricula, {})}

    while cola:
        circ, num, db, nivel = cola.popleft()
        nivel_max = max(nivel_max, nivel)
        etiqueta = f"{circ}-{num}"
        apertura = parsear_fecha(db.get("fechaAperturaFolio"))
        ej = api.estado_juridico(circ, num)
        if ej is None:
            ej = api.estado_juridico(circ, num)  # un reintento ante respuesta vacía
        if ej is None:
            notas.append(f"{etiqueta}: VUR no devolvió el estado jurídico (dos intentos)")
            incompleta = True
        anotaciones = [parsear_anotacion(t) for t in (ej or {}).get("textoAnotaciones") or []]
        anotaciones_por_folio[etiqueta] = anotaciones
        datos_por_folio[etiqueta] = db
        analisis = analizar_folio(anotaciones)
        valorizacion, cancelaciones, estado = analisis["valorizacion"], analisis["cancelaciones"], analisis["estado"]
        afectaciones.extend((etiqueta, a) for a in analisis["afectaciones"])
        motivos_formato, notas_formato = describir_no_reconocidas(etiqueta, anotaciones)
        no_reconocidas.extend(motivos_formato)
        notas.extend(notas_formato)
        folios += 1
        cadena.append(f"{etiqueta} (nivel {nivel}, apertura {db.get('fechaAperturaFolio') or '?'}, "
                      f"{len(anotaciones)} anotaciones: {estado})")
        log(f"  nivel {nivel} {etiqueta}: {estado} ({len(anotaciones)} anotaciones)")

        if valorizacion:
            if cancelacion_previa and estado == "GRAVADA":
                # Gravamen vigente en la matriz pero cancelado en el folio hijo.
                canc_folio, canc = cancelacion_previa
                return resultado("GRAVAMEN CANCELADO", etiqueta, valorizacion, canc, canc_folio, nivel)
            return resultado(estado, etiqueta, valorizacion, cancelaciones, etiqueta if cancelaciones else "", nivel)
        if cancelaciones and not cancelacion_previa:
            cancelacion_previa = (etiqueta, cancelaciones)
            notas.append(f"{etiqueta}: tiene cancelación de valorización sin el gravamen original; "
                         "se busca el gravamen en sus matrices")

        matrices = db.get("matriculasMatriz") or []
        if not matrices:
            continue
        if fecha_limite and apertura and apertura <= fecha_limite and not cancelacion_previa:
            notas.append(f"{etiqueta}: folio abierto el {apertura:%d/%m/%Y}, en o antes de la fecha límite "
                         f"{fecha_limite:%d/%m/%Y}; no se suben sus matrices")
            continue
        if not apertura and not aviso_sin_fecha:
            notas.append(f"{etiqueta}: sin fecha de apertura; se revisan sus matrices por precaución")
        if not fecha_limite and not aviso_sin_fecha:
            notas.append("Municipio sin fecha de proyecto: se suben todas las matrices")
            aviso_sin_fecha = True
        for m in matrices:
            if m in vistos:
                continue
            if folios + len(cola) >= MAX_FOLIOS_SEGURIDAD:
                notas.append(f"Tope de seguridad de {MAX_FOLIOS_SEGURIDAD} folios alcanzado; "
                             f"sin revisar: {', '.join(x for x in matrices if x not in vistos)}")
                incompleta = True
                cola.clear()
                break
            partes = separar_matricula(m)
            if not partes:
                notas.append(f"Matriz con formato no reconocido: {m}")
                continue
            vistos.add(m)
            db_m = api.datos_basicos(*partes) or api.datos_basicos(*partes)  # un reintento
            if not db_m:
                notas.append(f"Matriz {m}: VUR no devolvió datos (dos intentos); no se pudo revisar")
                incompleta = True
                continue
            cola.append((partes[0], partes[1], db_m, nivel + 1))

    if cancelacion_previa:
        canc_folio, canc = cancelacion_previa
        notas.append("No se encontró el gravamen original en las matrices revisadas")
        return resultado("GRAVAMEN CANCELADO", "", (), canc, canc_folio)
    return resultado(ESTADO_INCOMPLETO if incompleta else "NO GRAVADA")


# ---------------------------------------------------------------------------
# Excepciones de control
# ---------------------------------------------------------------------------


class SesionExpirada(Exception):
    pass


class NavegadorCerrado(Exception):
    pass


class ProcesoDetenido(Exception):
    pass


class LimiteDiario(Exception):
    pass


class ErrorVUR(Exception):
    pass


# ---------------------------------------------------------------------------
# Cliente VUR (sesión de Chrome + servicios JSON)
# ---------------------------------------------------------------------------

JS_FETCH = r"""
const url = arguments[0];
const listo = arguments[arguments.length - 1];
fetch(url, {credentials: 'include'})
  .then(async r => listo({status: r.status, url: r.url,
                          tipo: r.headers.get('content-type') || '', texto: await r.text()}))
  .catch(e => listo({error: String(e)}));
"""


class VURClient:
    def __init__(self, log, notificar=lambda estado: None):
        self.log = log
        self.notificar = notificar  # informa el estado de la sesión al semáforo de la ventana
        self.driver = None
        self.consultas = 0          # peticiones reales a VUR (cuentan para el cupo diario)
        self.consultas_cache = 0    # respuestas reutilizadas sin consultar a VUR
        self._cache = {}

    # --- Navegador ---------------------------------------------------------

    def iniciar_navegador(self):
        opciones = webdriver.ChromeOptions()
        opciones.add_argument(f"--user-data-dir={PERFIL_CHROME}")
        opciones.add_argument("--start-maximized")
        opciones.set_capability("unhandledPromptBehavior", "accept")
        # Mantiene activo el administrador de contraseñas de Chrome en este perfil, para que
        # rellene el usuario y la clave si el usuario decide guardarlos en ese perfil.
        opciones.add_experimental_option("prefs", {
            "credentials_enable_service": True,
            "profile.password_manager_enabled": True,
        })
        self.driver = webdriver.Chrome(options=opciones)
        self.driver.set_script_timeout(TIMEOUT_API)
        self.notificar("esperando")

    def nueva_ejecucion(self):
        """Al reutilizar la sesión en otro proceso: memoria y contadores desde cero
        (los datos de VUR pueden haber cambiado desde el proceso anterior)."""
        self._cache.clear()
        self.consultas = self.consultas_cache = 0

    def cerrar_navegador(self):
        if self.driver:
            try:
                self.driver.quit()
            except WebDriverException:
                pass
            self.driver = None
        self.notificar("cerrada")

    def vivo(self):
        if not self.driver:
            return False
        try:
            _ = self.driver.window_handles
            return True
        except (InvalidSessionIdException, NoSuchWindowException, WebDriverException):
            return False

    def abrir_inicio(self):
        self.driver.switch_to.default_content()
        self.driver.get(URL_INICIO)

    def cerrar_alerta(self):
        try:
            self.driver.switch_to.alert.accept()
        except (NoAlertPresentException, WebDriverException):
            pass

    def recuperar(self):
        if not self.vivo():
            raise NavegadorCerrado("El navegador se cerró o perdió la sesión de Selenium.")
        self.cerrar_alerta()

    # --- Servicios JSON ------------------------------------------------------

    def _asegurar_origen(self):
        """fetch() necesita estar en una página de www.vur.gov.co para enviar la sesión."""
        url = (self.driver.current_url or "").lower()
        if "siteminder" in url or "login" in url:
            self.notificar("vencida")
            raise SesionExpirada("VUR muestra la página de inicio de sesión.")
        if "vur.gov.co" not in url:
            self.abrir_inicio()

    def api_get(self, ruta):
        """Consulta un servicio; si la sesión parece vencida, recarga el inicio y reintenta una vez.

        Se comprobó en VUR que volver a cargar inicio.jsf renueva la sesión de
        los servicios sin pedir usuario y contraseña de nuevo.

        Las respuestas con datos se guardan en memoria durante la ejecución:
        varias matrículas comparten matrices (p. ej. tres de El Retiro suben
        por 017-81855, 017-63800 y 017-7991), y un reintento no vuelve a gastar
        cupo en los folios ya consultados. Las respuestas vacías no se guardan,
        para que un fallo temporal de VUR pueda reintentarse.
        """
        if ruta in self._cache:
            self.consultas_cache += 1
            return self._cache[ruta]
        try:
            datos = self._api_get(ruta)
        except SesionExpirada:
            self.log("Sesión de VUR inactiva; recargando la página de inicio...")
            self.abrir_inicio()
            datos = self._api_get(ruta)
        if datos:
            self._cache[ruta] = datos
        return datos

    def _api_get(self, ruta):
        self._asegurar_origen()
        try:
            r = self.driver.execute_async_script(JS_FETCH, URL_API + ruta)
        except JavascriptException as e:
            raise ErrorVUR(f"Error de JavaScript: {mensaje_corto(e)}") from e
        self.consultas += 1
        if not r or r.get("error"):
            raise ErrorVUR(f"Fallo de red: {(r or {}).get('error')}")
        destino = (r.get("url") or "").lower()
        if "siteminder" in destino or "login" in destino or "json" not in r.get("tipo", ""):
            self.notificar("vencida")
            raise SesionExpirada("La sesión de VUR no está activa (respuesta de inicio de sesión).")
        if r["status"] != 200:
            raise ErrorVUR(f"VUR respondió HTTP {r['status']}")
        datos = json.loads(r["texto"] or "null")
        primero = datos[0] if isinstance(datos, list) and datos else datos
        if isinstance(primero, dict) and primero.get("estado") == -9999:
            self.notificar("limite")
            raise LimiteDiario("VUR: el usuario excedió el número de peticiones por día.")
        self.notificar("activa")
        return datos

    def datos_basicos(self, circulo, numero):
        return self.api_get(f"datosBasicos/{TIPO_CONSULTA}/{circulo}-{numero}") or None  # {} = no existe

    def estado_juridico(self, circulo, numero):
        return self.api_get(f"estadoJuridico/{TIPO_CONSULTA}/{circulo}-{numero}") or None


ERRORES_RECUPERABLES = (ErrorVUR, TimeoutException, WebDriverException, json.JSONDecodeError)

# ---------------------------------------------------------------------------
# Identidad visual (Manual de Identidad Visual Gobernación de Antioquia 2026)
# ---------------------------------------------------------------------------
# Colores primarios del manual: verde #018D38, verde oscuro #0B5640, blanco y
# verde menta #3AF9A2 (solo piezas digitales). Secundarios: azul #3561AB,
# naranja #F28E18, morado #8B4A97. Tipografía institucional: Poppins (si no
# está instalada en el equipo se usa Segoe UI).
# El logo NO se dibuja en el código: es el PNG oficial del portal antioquia.gov.co
# (sin su margen transparente), guardado en RECURSOS al final de este archivo.
# El manual pide mínimo 150 x 67 px en pantalla, sobre fondo blanco sin
# contenedor (o contenedor blanco sobre color) y con área de protección libre.
# Por eso el encabezado es blanco y el logo va directamente sobre él.

COLOR = {
    "verde": "#018D38",
    "verde_oscuro": "#0B5640",
    "menta": "#3AF9A2",
    "blanco": "#FFFFFF",
    "azul": "#3561AB",
    "naranja": "#F28E18",
    "morado": "#8B4A97",
    "fondo": "#F3F6F4",       # neutro de trabajo con leve tinte verde
    "linea": "#D5DED9",
    "texto": "#1C2723",
    "texto_suave": "#5B6B64",
    "rojo": "#B3261E",        # solo para errores (color semántico, no institucional)
}

TITULO_APLICATIVO = "Consulta de valorización en VUR"
SECRETARIA = "Secretaría de Infraestructura"
DEPENDENCIA = "Dirección Administrativa y Financiera"
CREDITO = "Desarrollado por Alexander Gómez A."

CARPETA_APP = Path(__file__).resolve().parent
# Instalación: VUR_valorizacion\vur_valorizacion.py con consultas\, historico\, datos\ y
# respaldos\ al lado. (Versiones anteriores lo tenían dentro de una carpeta "programa".)
CARPETA_RAIZ = CARPETA_APP.parent if CARPETA_APP.name.lower() == "programa" else CARPETA_APP
CARPETA_CONSULTAS = CARPETA_RAIZ / "consultas" if (CARPETA_RAIZ / "consultas").is_dir() else CARPETA_APP
# Histórico de consultas. Hoy es una carpeta local; para compartirlo entre funcionarios
# (carpeta de red o nube) basta con definir la variable de entorno VUR_HISTORICO con
# la nueva ruta, sin cambiar el código.
CARPETA_HISTORICO = Path(os.environ["VUR_HISTORICO"]) if os.environ.get("VUR_HISTORICO") else CARPETA_RAIZ / "historico"

# Carpeta de trabajo de cada usuario (consultas, histórico, datos, respaldos). Se elige
# la primera vez que se abre el programa (o la elige el instalador) y se guarda en
# %APPDATA%\VUR_Valorizacion\config.json; desde entonces es la carpeta del aplicativo.
ARCHIVO_CONFIG = Path(os.environ.get("APPDATA") or Path.home()) / "VUR_Valorizacion" / "config.json"
SUBCARPETAS_TRABAJO = ("consultas", "datos", "historico", "respaldos")


def leer_carpeta_trabajo():
    try:
        carpeta = json.loads(ARCHIVO_CONFIG.read_text(encoding="utf-8-sig")).get("carpeta")  # el instalador escribe con BOM
        return Path(carpeta) if carpeta else None
    except (OSError, ValueError):
        return None


def aplicar_carpeta_trabajo(carpeta, guardar=False):
    """Crea las subcarpetas, (opcional) guarda la elección y la usa desde ese momento."""
    global CARPETA_CONSULTAS, CARPETA_HISTORICO
    carpeta = Path(carpeta)
    for sub in SUBCARPETAS_TRABAJO:
        (carpeta / sub).mkdir(parents=True, exist_ok=True)
    if guardar:
        ARCHIVO_CONFIG.parent.mkdir(parents=True, exist_ok=True)
        ARCHIVO_CONFIG.write_text(json.dumps({"carpeta": str(carpeta), "elegida": f"{datetime.now():%Y-%m-%d %H:%M}",
                                              "version": VERSION_APLICATIVO}, ensure_ascii=False, indent=2),
                                  encoding="utf-8")
    CARPETA_CONSULTAS = carpeta / "consultas"
    if not os.environ.get("VUR_HISTORICO"):
        CARPETA_HISTORICO = carpeta / "historico"
    return carpeta


def primer_inicio(raiz):
    """Primera vez en este usuario de Windows: pide la carpeta de trabajo y la crea."""
    carpeta = leer_carpeta_trabajo()
    if carpeta and carpeta.is_dir():
        return aplicar_carpeta_trabajo(carpeta)
    messagebox.showinfo(
        "Bienvenido", "Es la primera vez que abre el programa en este usuario.\n\n"
        "Elija la carpeta donde quiere trabajar. Allí se crearán:\n"
        "  consultas  (Excel de matrículas y resultados)\n  historico  (consultas anteriores)\n"
        "  datos  y  respaldos\n\nSi no elige ninguna, se usará la carpeta del programa:\n"
        f"{CARPETA_RAIZ}", parent=raiz)
    elegida = filedialog.askdirectory(title="Carpeta de trabajo del aplicativo VUR", initialdir=str(CARPETA_RAIZ),
                                      mustexist=False, parent=raiz)
    try:
        return aplicar_carpeta_trabajo(Path(elegida) if elegida else CARPETA_RAIZ, guardar=True)
    except OSError as e:
        messagebox.showwarning("Carpeta de trabajo", f"No se pudo usar esa carpeta ({type(e).__name__}: {e}).\n"
                               f"Se usará {CARPETA_RAIZ}.", parent=raiz)
        return aplicar_carpeta_trabajo(CARPETA_RAIZ, guardar=True)
CARPETA_RECURSOS = CARPETA_USUARIO / "recursos"  # logo, icono y Poppins (licencia SIL OFL 1.1) recreados aquí


def cargar_fuentes_privadas():
    """Registra Poppins solo para este programa (Windows), sin instalarla en el equipo.

    Debe llamarse antes de crear la ventana de Tk. En otros sistemas, o si
    faltan los archivos, la interfaz usa Segoe UI / Helvetica.
    """
    import sys
    if sys.platform != "win32" or not CARPETA_RECURSOS.is_dir():
        return 0
    import ctypes
    FR_PRIVATE = 0x10
    cargadas = 0
    for ttf in CARPETA_RECURSOS.glob("*.ttf"):
        cargadas += bool(ctypes.windll.gdi32.AddFontResourceExW(str(ttf), FR_PRIVATE, 0))
    return cargadas


def preparar_recursos():
    """Recrea logo, icono y fuentes (guardados al final de este archivo) si faltan o cambiaron."""
    import base64
    import hashlib
    try:
        CARPETA_RECURSOS.mkdir(parents=True, exist_ok=True)
        for nombre, texto in RECURSOS.items():
            datos = base64.b64decode("".join(texto.split()))
            destino = CARPETA_RECURSOS / nombre
            if not destino.exists() or hashlib.sha256(destino.read_bytes()).digest() != hashlib.sha256(datos).digest():
                destino.write_bytes(datos)
        return CARPETA_RECURSOS
    except OSError:
        return None  # sin recursos: la interfaz usa texto y Segoe UI


# Estado -> (etiqueta corta del contador, color del texto, color de fondo de fila)
ESTILO_ESTADO = {
    "GRAVADA": ("Gravadas", "#A3570A", "#FDEBD3"),  # naranja institucional oscurecido para texto legible
    "GRAVAMEN CANCELADO": ("Canceladas", COLOR["azul"], "#E1E9F6"),
    "NO GRAVADA": ("No gravadas", COLOR["verde"], "#E3F4EA"),
    ESTADO_INCOMPLETO: ("Incompletas", COLOR["morado"], "#EFE4F2"),
    "SIN DATOS EN VUR": ("Sin datos", COLOR["texto_suave"], "#ECEFED"),
    "ERROR": ("Error", COLOR["rojo"], "#F8E1DF"),
}


CONCLUSIONES = {
    "GRAVADA": "TIENE VALORIZACIÓN VIGENTE",
    "GRAVAMEN CANCELADO": "TUVO VALORIZACIÓN, YA CANCELADA",
    "NO GRAVADA": "NO TIENE VALORIZACIÓN",
    ESTADO_INCOMPLETO: "NO SE HALLÓ VALORIZACIÓN, PERO LA REVISIÓN DE MATRICES QUEDÓ INCOMPLETA",
    "SIN DATOS EN VUR": "NO SE PUDO CONSULTAR (VUR SIN DATOS)",
    "ERROR": "NO SE PUDO CONSULTAR (ERROR)",
}


def conclusion(reg):
    """Frase concluyente de la primera columna del Excel."""
    estado = reg.get("Estado", "")
    texto = CONCLUSIONES.get(estado, estado or "SIN RESULTADO")
    folio, nivel = reg.get("Matricula_Con_Anotacion"), str(reg.get("Nivel_Matriz", "")).strip()
    if folio and estado in ("GRAVADA", "GRAVAMEN CANCELADO"):
        texto += (" · anotación en la misma matrícula" if nivel in ("", "0")
                  else f" · anotación en la matriz {folio} (nivel {nivel})")
    if reg.get("Confiabilidad") == CONFIABILIDAD_REVISAR:
        texto += " · REQUIERE REVISIÓN MANUAL"
    return texto


TEXTOS_LARGOS = {"Observaciones", "Cadena_Revisada", "Detalle_Valorizacion", "Notas_Proceso",
                 "Proyectos_Municipio", "Afectaciones_Valorizacion"}


def exportar_excel(registros, ruta):
    """Escribe el Excel de resultados con encabezado institucional.

    Todas las celdas se escriben como texto o número, nunca como fórmula: un
    texto de VUR que empiece por "=" no se ejecuta al abrir el archivo.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(COLUMNAS_EXCEL)
    for reg in registros:
        reg = dict(reg, Resultado=conclusion(reg))
        ws.append([reg.get(c, "") for c in COLUMNAS_EXCEL])
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.data_type = "s"
    encabezado = PatternFill("solid", fgColor=COLOR["verde_oscuro"][1:])
    for celda in ws[1]:
        celda.fill = encabezado
        celda.font = Font(bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    anchos = {"Observaciones": 80, "Cadena_Revisada": 60, "Detalle_Valorizacion": 50,
              "Propietarios_Actuales": 45, "Propietario_Gravado": 40, "Notas_Proceso": 50,
              "Proyectos_Municipio": 45, "Motivos_Revision": 55, "Afectaciones_Valorizacion": 60,
              "Resultado": 46, "Estado": 22, "Propietario": 28, "Municipio": 18}
    for i, col in enumerate(COLUMNAS_EXCEL, 1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(col, max(12, len(col) + 2))
        # Los textos largos (anotaciones completas, cadena...) no se ajustan: harían filas
        # gigantes. Se leen al seleccionar la celda o en la ventana de detalle del programa.
        ajustar = col in anchos and col not in TEXTOS_LARGOS
        for celda in ws.iter_cols(min_col=i, max_col=i, min_row=2):
            for c in celda:
                c.alignment = Alignment(wrap_text=ajustar, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    col_estado = COLUMNAS_EXCEL.index("Estado") + 1
    col_resultado = COLUMNAS_EXCEL.index("Resultado") + 1
    for fila in range(2, ws.max_row + 1):
        estilo = ESTILO_ESTADO.get(ws.cell(row=fila, column=col_estado).value)
        if estilo:
            for col in (col_estado, col_resultado):  # el resultado lleva el color de su estado
                celda = ws.cell(row=fila, column=col)
                celda.fill = PatternFill("solid", fgColor=estilo[2][1:])
                celda.font = Font(bold=True, color=estilo[1][1:])
    # Columnas de detalle: agrupadas y ocultas (se despliegan con el botón [+]).
    primera = get_column_letter(len(COLUMNAS_VISIBLES) + 1)
    ultima = get_column_letter(len(COLUMNAS_EXCEL))
    ws.column_dimensions.group(primera, ultima, hidden=True, outline_level=1)
    col_conf = COLUMNAS_EXCEL.index("Confiabilidad") + 1
    for fila in range(2, ws.max_row + 1):
        celda = ws.cell(row=fila, column=col_conf)
        if celda.value == CONFIABILIDAD_REVISAR:
            celda.fill = PatternFill("solid", fgColor="FDEBD3")
            celda.font = Font(bold=True, color="A3570A")
        elif celda.value == CONFIABILIDAD_ALTA:
            celda.font = Font(bold=True, color=COLOR["verde"][1:])
    _hoja_resumen(wb, registros)
    wb.save(ruta)


def _hoja_resumen(wb, registros):
    """Hoja 'Resumen': totales por estado y lista de lo que debe revisar el profesional."""
    from collections import Counter
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Resumen")  # 2.ª hoja: "Resultados" sigue siendo la primera para otros programas
    wb.active = wb.index(ws)          # pero el archivo abre en el resumen
    titulo = Font(bold=True, size=13, color=COLOR["verde_oscuro"][1:])
    negrita = Font(bold=True)
    encabezado = PatternFill("solid", fgColor=COLOR["verde_oscuro"][1:])
    ws.append([f"{TITULO_APLICATIVO} - {SECRETARIA}"])
    ws["A1"].font = titulo
    ws.append([DEPENDENCIA])
    ws.append(["Generado", f"{datetime.now():%d/%m/%Y %H:%M}"])
    ws.append(["Versión del programa", VERSION_APLICATIVO])
    ws.append(["Matrículas en el archivo", len(registros)])
    ws.append([])
    ws.append(["Estado", "Cantidad"])
    for celda in ws[ws.max_row]:
        celda.fill, celda.font = encabezado, Font(bold=True, color="FFFFFF")
    conteo = Counter(r.get("Estado", "") for r in registros)
    for estado in list(ESTILO_ESTADO) + sorted(e for e in conteo if e not in ESTILO_ESTADO):
        if conteo.get(estado):
            ws.append([estado or "(sin estado)", conteo[estado]])
            estilo = ESTILO_ESTADO.get(estado)
            if estilo:
                ws.cell(ws.max_row, 1).font = Font(bold=True, color=estilo[1][1:])
    revisar = [r for r in registros if r.get("Confiabilidad") == CONFIABILIDAD_REVISAR]
    ws.append(["Requieren revisión del profesional", len(revisar)])
    ws.cell(ws.max_row, 1).font = negrita
    if revisar:
        ws.append([])
        ws.append(["Matrícula", "Municipio", "Estado", "Por qué revisar"])
        for celda in ws[ws.max_row]:
            celda.fill, celda.font = encabezado, Font(bold=True, color="FFFFFF")
        for r in revisar:
            ws.append([r.get("Matricula_Limpia", ""), r.get("Municipio", ""), r.get("Estado", ""),
                       r.get("Motivos_Revision", "")])
            ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.data_type = "s"
    for col, ancho in zip("ABCD", (34, 22, 30, 90)):
        ws.column_dimensions[col].width = ancho


# ---------------------------------------------------------------------------
# Comparación de dos Excel de resultados
# ---------------------------------------------------------------------------

CAMPOS_COMPARADOS = [  # (columna, nombre para el usuario)
    ("Estado", "Estado"), ("Confiabilidad", "Confiabilidad"), ("Municipio", "Municipio"),
    ("Matricula_Con_Anotacion", "Folio con la anotación"), ("Nivel_Matriz", "Nivel de matriz"),
    ("Propietario_Gravado", "Propietario gravado"), ("Matricula_Cancelacion", "Folio de la cancelación"),
    ("Afectaciones_Valorizacion", "Afectaciones"), ("Folios_Revisados", "Folios revisados"),
    ("Propietarios_Actuales", "Propietarios actuales"),
]


def _leer_resultados(ruta):
    """Lee la hoja 'Resultados' (o la primera) de un Excel generado por el programa."""
    # "with" cierra el archivo al terminar: en Windows un Excel abierto queda bloqueado para otros programas.
    with pd.ExcelFile(ruta) as libro:
        hoja = "Resultados" if "Resultados" in libro.sheet_names else 0
        df = pd.read_excel(libro, sheet_name=hoja, dtype=str).fillna("")
    if "Matricula_Limpia" not in df.columns:
        raise ValueError(f"{Path(ruta).name} no parece un Excel de resultados del programa "
                         "(no tiene la columna Matricula_Limpia).")
    return {fila["Matricula_Limpia"].strip(): fila for fila in df.to_dict("records") if fila["Matricula_Limpia"].strip()}


def _anotaciones_en_cadena(texto):
    """Total de anotaciones que dice la columna Cadena_Revisada ('..., 13 anotaciones: ...')."""
    return sum(int(n) for n in re.findall(r"(\d+) anotaciones", texto or ""))


def _explicar_cambio(antes, ahora):
    """Causa probable de que una matrícula cambie entre dos consultas (es una pista, no una certeza)."""
    causas = []
    if antes.get("Estado") in ("SIN DATOS EN VUR", "ERROR", ESTADO_INCOMPLETO):
        causas.append("la consulta anterior no fue completa (sin datos, error o matrices sin respuesta)")
    if ahora.get("Estado") in ("SIN DATOS EN VUR", "ERROR", ESTADO_INCOMPLETO):
        causas.append("la consulta nueva no fue completa: repetirla")
    n_antes, n_ahora = _anotaciones_en_cadena(antes.get("Cadena_Revisada")), _anotaciones_en_cadena(ahora.get("Cadena_Revisada"))
    if n_antes and n_ahora and n_antes != n_ahora:
        causas.append(f"VUR muestra otra cantidad de anotaciones ({n_antes} -> {n_ahora}): hubo inscripciones nuevas")
    v_antes, v_ahora = antes.get("Version_Aplicativo", ""), ahora.get("Version_Aplicativo", "")
    if v_antes != v_ahora:
        causas.append(f"cambiaron las reglas del programa (versión {v_antes or 'anterior a 2026.10.23'} -> {v_ahora or '?'})")
    return "; ".join(causas) or "sin causa evidente: revisar manualmente en VUR"


def comparar_resultados(ruta_antes, ruta_ahora, ruta_salida):
    """Compara dos Excel de resultados y escribe un Excel con las diferencias.

    Devuelve un resumen (dict). Si los archivos no comparten ninguna matrícula,
    lanza ValueError: casi siempre significa que son de lotes distintos.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    antes, ahora = _leer_resultados(ruta_antes), _leer_resultados(ruta_ahora)
    comunes = [m for m in ahora if m in antes]
    if not comunes:
        raise ValueError("Los dos archivos no tienen ninguna matrícula en común: probablemente son de lotes "
                         "distintos. Elija el resultado anterior del mismo Excel de matrículas.")
    cambios, detalle = [], []
    for m in comunes:
        a, b = antes[m], ahora[m]
        campos = [(col, nombre) for col, nombre in CAMPOS_COMPARADOS
                  if col in a and col in b  # una columna nueva en una sola versión no es un cambio
                  and str(a[col]).strip() != str(b[col]).strip()]
        if not campos:
            continue
        cambios.append([m, b.get("Municipio") or a.get("Municipio", ""), a.get("Estado", ""), b.get("Estado", ""),
                        a.get("Confiabilidad", ""), b.get("Confiabilidad", ""), ", ".join(n for _, n in campos),
                        _explicar_cambio(a, b), b.get("Motivos_Revision", "")])
        for col, nombre in campos:
            detalle.append([m, nombre, a.get(col) or "(vacío)", b.get(col) or "(vacío)"])
    solo_antes = [m for m in antes if m not in ahora]
    solo_ahora = [m for m in ahora if m not in antes]
    cambio_estado = sum(1 for c in cambios if c[2] != c[3])
    resumen = {"comunes": len(comunes), "iguales": len(comunes) - len(cambios), "con_cambios": len(cambios),
               "cambio_estado": cambio_estado, "solo_antes": len(solo_antes), "solo_ahora": len(solo_ahora)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    enc = PatternFill("solid", fgColor=COLOR["verde_oscuro"][1:])
    blanco = Font(bold=True, color="FFFFFF")
    ws.append(["Comparación de resultados VUR"])
    ws["A1"].font = Font(bold=True, size=13, color=COLOR["verde_oscuro"][1:])
    for fila in (["Resultado anterior", Path(ruta_antes).name], ["Resultado nuevo", Path(ruta_ahora).name],
                 ["Generado", f"{datetime.now():%d/%m/%Y %H:%M}  (versión {VERSION_APLICATIVO})"], [],
                 ["Matrículas en ambos archivos", resumen["comunes"]], ["Sin ningún cambio", resumen["iguales"]],
                 ["Con algún cambio", resumen["con_cambios"]], ["   de ellas, cambió el estado", cambio_estado],
                 ["Solo en el anterior", resumen["solo_antes"]], ["Solo en el nuevo", resumen["solo_ahora"]]):
        ws.append(fila)
    if solo_antes or solo_ahora:
        ws.append([])
        ws.append(["Solo en el anterior", ", ".join(solo_antes) or "—"])
        ws.append(["Solo en el nuevo", ", ".join(solo_ahora) or "—"])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 70

    def hoja(titulo, columnas, filas, anchos):
        h = wb.create_sheet(titulo)
        h.append(columnas)
        for celda in h[1]:
            celda.fill, celda.font = enc, blanco
        for f in filas:
            h.append([recortar(str(x)) for x in f])
        for i, ancho in enumerate(anchos):
            h.column_dimensions[chr(65 + i)].width = ancho
            for celda in h.iter_cols(min_col=i + 1, max_col=i + 1, min_row=2):
                for c in celda:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        h.freeze_panes = "B2"
        return h

    hc = hoja("Cambios", ["Matrícula", "Municipio", "Estado antes", "Estado ahora", "Confianza antes",
                          "Confianza ahora", "Qué cambió", "Causa probable", "Motivos de revisión (ahora)"],
              cambios, (14, 18, 22, 22, 12, 12, 34, 55, 55))
    for fila in hc.iter_rows(min_row=2):
        if fila[2].value != fila[3].value:
            for c in fila[2:4]:
                c.font = Font(bold=True, color="A3570A")
    hoja("Detalle", ["Matrícula", "Campo", "Antes", "Ahora"], detalle, (14, 22, 70, 70))
    for h in wb.worksheets:
        for fila in h.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    celda.data_type = "s"
    wb.save(ruta_salida)
    return resumen


# ---------------------------------------------------------------------------
# Histórico de consultas (evita repetir en VUR matrículas ya consultadas)
# ---------------------------------------------------------------------------
# Estructura de la carpeta:
#   registros\AAAA-MM-DD_HHMMSS_usuario_equipo.jsonl  una línea JSON por consulta a VUR.
#        Cada sesión escribe SOLO su propio archivo y nunca modifica los de otros: así
#        varios funcionarios (o varias ventanas) pueden trabajar a la vez sin dañar nada.
#   Historico_consultas.xlsx  última consulta válida de cada matrícula, para leer y
#        filtrar. Se regenera al terminar cada proceso; no se edita a mano.

VIGENCIA_HISTORICO_DIAS = 30   # un resultado más antiguo se vuelve a consultar en VUR
ESTADOS_REUTILIZABLES = {"NO GRAVADA", "GRAVADA", "GRAVAMEN CANCELADO"}
ESTADOS_VALIDOS = ESTADOS_REUTILIZABLES | {ESTADO_INCOMPLETO}


def _fecha_registro(reg):
    try:
        return datetime.strptime(str(reg.get("Fecha_Consulta", ""))[:19], "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def _prioridad(reg):
    """Orden para elegir el registro que representa a una matrícula: primero los
    resultados válidos (un ERROR de hoy no tapa un resultado bueno de ayer), luego el más reciente."""
    return (reg.get("Estado") in ESTADOS_VALIDOS, _fecha_registro(reg) or datetime.min)


def evaluar_historico(reg, ahora=None):
    """¿Se puede reutilizar este registro del histórico sin consultar VUR?

    Devuelve (True/False, explicación para el usuario). Solo se reutiliza si:
    se analizó con las reglas vigentes, el estado es definitivo, la confiabilidad
    es ALTA y la consulta tiene menos de VIGENCIA_HISTORICO_DIAS días.
    """
    if not reg:
        return False, ""
    fecha = _fecha_registro(reg)
    cuando = f"{fecha:%d/%m/%Y}" if fecha else "fecha desconocida"
    quien = reg.get("Funcionario") or "otro usuario"
    if reg.get("Version_Reglas") != VERSION_REGLAS:
        return False, f"consultada el {cuando} con reglas anteriores del programa; se consulta de nuevo"
    if reg.get("Estado") not in ESTADOS_REUTILIZABLES or reg.get("Confiabilidad") != CONFIABILIDAD_ALTA:
        return False, f"la consulta del {cuando} quedó en {reg.get('Estado')} / {reg.get('Confiabilidad')}; se consulta de nuevo"
    dias = ((ahora or datetime.now()) - fecha).days if fecha else 10**6
    if dias > VIGENCIA_HISTORICO_DIAS:
        return False, f"consultada hace {dias} días (vigencia {VIGENCIA_HISTORICO_DIAS}); se consulta de nuevo"
    hace = "hoy" if dias == 0 else f"hace {dias} día{'s' if dias != 1 else ''}"
    return True, f"consultada el {cuando} por {quien} ({hace})"


def _usuario_equipo():
    import getpass
    import platform
    try:
        usuario = getpass.getuser()
    except Exception:  # noqa: BLE001 - nunca debe impedir consultar
        usuario = "usuario"
    return usuario, os.environ.get("COMPUTERNAME") or platform.node() or "equipo"


class Historico:
    def __init__(self, carpeta=None, log=lambda m: None):
        self.carpeta = Path(carpeta or CARPETA_HISTORICO)  # se lee al crear: respeta la carpeta de trabajo
        self.dir_registros = self.carpeta / "registros"
        self.archivo_excel = self.carpeta / "Historico_consultas.xlsx"
        self.log = log
        self.ultimo = {}       # matrícula -> registro que la representa
        self.total_lineas = 0
        self._archivo_sesion = None
        self._lock = threading.Lock()

    def cargar(self):
        """Lee todos los registros. Las líneas dañadas se ignoran (y se avisa)."""
        self.ultimo, self.total_lineas, dañadas = {}, 0, 0
        if not self.dir_registros.is_dir():
            return 0
        for archivo in sorted(self.dir_registros.glob("*.jsonl")):
            try:
                lineas = archivo.read_text(encoding="utf-8").splitlines()
            except OSError as e:
                self.log(f"Histórico: no se pudo leer {archivo.name} ({type(e).__name__})")
                continue
            for linea in lineas:
                if not linea.strip():
                    continue
                try:
                    reg = json.loads(linea)
                except json.JSONDecodeError:
                    dañadas += 1
                    continue
                self._indexar(reg)
                self.total_lineas += 1
        if dañadas:
            self.log(f"Histórico: se ignoraron {dañadas} línea(s) dañadas")
        return len(self.ultimo)

    def _indexar(self, reg):
        m = reg.get("Matricula_Limpia")
        if m and (m not in self.ultimo or _prioridad(reg) >= _prioridad(self.ultimo[m])):
            self.ultimo[m] = reg

    def buscar(self, matricula):
        return self.ultimo.get(matricula)

    def agregar(self, registro):
        """Guarda una consulta hecha en VUR (se escribe en disco de inmediato)."""
        reg = {k: v for k, v in registro.items() if k != "Fuente"}
        reg["Version_Reglas"] = VERSION_REGLAS
        with self._lock:
            if self._archivo_sesion is None:
                usuario, equipo = _usuario_equipo()
                limpio = re.sub(r"[^A-Za-z0-9.-]+", "_", f"{usuario}_{equipo}")
                self.dir_registros.mkdir(parents=True, exist_ok=True)
                self._archivo_sesion = self.dir_registros / f"{datetime.now():%Y-%m-%d_%H%M%S}_{limpio}.jsonl"
            with open(self._archivo_sesion, "a", encoding="utf-8") as f:
                f.write(json.dumps(reg, ensure_ascii=False, default=str) + "\n")
            self._indexar(reg)
            self.total_lineas += 1

    def exportar_excel(self):
        """Regenera Historico_consultas.xlsx. Si está abierto en Excel, se avisa y se deja para después."""
        if not self.ultimo:
            return False
        regs = sorted(self.ultimo.values(), key=lambda r: r.get("Matricula_Limpia", ""))
        temporal = self.archivo_excel.with_name("~Historico_consultas_nuevo.xlsx")
        try:
            self.carpeta.mkdir(parents=True, exist_ok=True)
            exportar_excel([dict(r, Fuente="Histórico") for r in regs], temporal)
            os.replace(temporal, self.archivo_excel)
            return True
        except OSError as e:
            self.log(f"Histórico: no se pudo actualizar {self.archivo_excel.name} (¿está abierto en Excel?): "
                     f"{type(e).__name__}. Se actualizará en el próximo proceso.")
            try:
                temporal.unlink()
            except OSError:
                pass
            return False


# ---------------------------------------------------------------------------
# Interfaz Tkinter
# ---------------------------------------------------------------------------

COLUMNAS_TABLA = [  # (columna del registro, título, ancho)
    ("Matricula_Limpia", "Matrícula", 110),
    ("Municipio", "Municipio", 120),
    ("Estado", "Estado", 255),
    ("Confiabilidad", "Confianza", 100),
    ("Matricula_Con_Anotacion", "Folio hallado", 140),
    ("Nivel_Matriz", "Nivel", 60),
    ("Propietario_Gravado", "Propietario gravado", 200),
    ("Folios_Revisados", "Folios", 62),
    ("Fuente", "Fuente", 92),
]


def valores_tabla(reg):
    """Valores de una fila de la tabla; la fuente se muestra corta (el detalle va en el Excel)."""
    valores = []
    for col, _, _ in COLUMNAS_TABLA:
        valor = reg.get(col, "")
        if col == "Fuente" and str(valor).startswith("Histórico"):
            valor = "Histórico"
        valores.append(valor)
    return valores


# ---------------------------------------------------------------------------
# Componentes visuales con relieve (Tk puro, sin librerías adicionales)
# ---------------------------------------------------------------------------
# Tkinter no trae botones redondeados, sombras ni transparencias. Aquí se
# "pintan" píxel a píxel en una imagen (tk.PhotoImage): bordes suavizados
# (antialiasing), sombra difusa debajo y un brillo de 1 px arriba que da la
# sensación de volumen. Como Tk no mezcla transparencias entre widgets, cada
# componente recibe el color del fondo sobre el que va y la mezcla se calcula
# aquí. Las imágenes se guardan en caché: cada combinación se pinta una vez.

SOMBRA_RGB = (11, 40, 28)  # verde casi negro: la sombra conserva el tono institucional


def _rgb(color):
    color = color.lstrip("#")
    return tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))


def _hex(rgb):
    return "#%02x%02x%02x" % tuple(max(0, min(255, int(v + 0.5))) for v in rgb)


def _mezcla(a, b, t):
    return (a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t, a[2] + (b[2] - a[2]) * t)


def mezclar(c1, c2, t):
    """Color intermedio entre dos colores '#rrggbb' (t=0 -> c1, t=1 -> c2). Simula transparencia."""
    return _hex(_mezcla(_rgb(c1), _rgb(c2), t))


def _distancia(x, y, ancho, alto, radio):
    """Distancia con signo de (x, y) al borde de un rectángulo redondeado: negativa = dentro."""
    qx = abs(x - ancho / 2) - (ancho / 2 - radio)
    qy = abs(y - alto / 2) - (alto / 2 - radio)
    return math.hypot(max(qx, 0.0), max(qy, 0.0)) + min(max(qx, qy), 0.0) - radio


def _cobertura(d):
    return 0.0 if d >= 0.5 else 1.0 if d <= -0.5 else 0.5 - d


def pintar_en(filas, x0, y0, ancho, alto, radio, arriba, abajo=None, borde=None,
              sombra=(0.0, 0.0, 1.0), brillo=0.0, anillo=None):
    """Pinta una forma redondeada sobre una matriz de colores RGB ya existente.

    arriba/abajo: colores del degradado vertical. sombra: (opacidad, desplazamiento
    hacia abajo, difuminado en px). brillo: opacidad del reflejo blanco superior.
    anillo: color del anillo de foco de teclado (accesibilidad), o None.
    """
    arr, abj = _rgb(arriba), _rgb(abajo or arriba)
    bor = _rgb(borde) if borde else None
    ani = _rgb(anillo) if anillo else None
    alfa_s, dy, difu = sombra
    blanco = (255, 255, 255)
    for py, fila in enumerate(filas):
        sy = py + 0.5 - y0
        t = min(max(sy / alto, 0.0), 1.0)
        relleno = _mezcla(arr, abj, t)
        for px in range(len(fila)):
            sx = px + 0.5 - x0
            col = fila[px]
            if alfa_s > 0:
                ds = _distancia(sx, sy - dy, ancho, alto, radio)
                if ds < difu:
                    col = _mezcla(col, SOMBRA_RGB, alfa_s * (1 - max(ds, 0.0) / difu) ** 2)
            d = _distancia(sx, sy, ancho, alto, radio)
            if ani is not None:
                a = min(max(1.5 - abs(d - 3.0), 0.0), 1.0)
                if a:
                    col = _mezcla(col, ani, a)
            cob = _cobertura(d)
            if cob > 0:
                c = relleno
                if brillo:
                    h = cob - _cobertura(_distancia(sx, sy - 1.2, ancho, alto, radio))
                    if h > 0:
                        c = _mezcla(c, blanco, brillo * h)
                if bor is not None:
                    b = cob - _cobertura(d + 1.0)
                    if b > 0:
                        c = _mezcla(c, bor, b)
                col = _mezcla(col, c, cob)
            fila[px] = col
    return filas


def lienzo(ancho, alto, fondo):
    f = _rgb(fondo)
    return [[f] * ancho for _ in range(alto)]


def a_imagen(maestro, filas):
    img = tk.PhotoImage(master=maestro, width=len(filas[0]), height=len(filas))
    img.put(" ".join("{" + " ".join(_hex(c) for c in fila) + "}" for fila in filas), to=(0, 0))
    return img


_CACHE_IMAGENES = {}


def imagen_cache(maestro, clave, fabricar):
    """Devuelve la imagen guardada para 'clave' o la fabrica una sola vez."""
    if clave not in _CACHE_IMAGENES:
        _CACHE_IMAGENES[clave] = a_imagen(maestro, fabricar())
    return _CACHE_IMAGENES[clave]


# Apariencia de los botones: (arriba, abajo, borde, texto, sombra, brillo) por estado.
_VERDE_TXT = "#FFFFFF"
APARIENCIA_BOTON = {
    "primario": {
        "normal": ("#078F3E", "#01722E", "#01652A", _VERDE_TXT, (0.30, 2.0, 5.0), 0.30),
        "hover": ("#0A9A44", "#017A31", "#01692B", _VERDE_TXT, (0.34, 3.0, 5.0), 0.34),
        "presionado": ("#015F26", "#016F2D", "#015524", _VERDE_TXT, (0.14, 0.5, 2.0), 0.0),
        "deshabilitado": ("#E6EBE8", "#E6EBE8", "#D9E0DC", "#97A69F", (0.04, 1.0, 2.0), 0.0),
    },
    "secundario": {
        "normal": ("#FFFFFF", "#F2F6F3", "#CBD5CF", "#0B5640", (0.12, 1.5, 4.0), 0.0),
        "hover": ("#FFFFFF", "#E9F4ED", "#9CC9AE", "#0B5640", (0.17, 2.5, 5.0), 0.0),
        "presionado": ("#DFEDE4", "#EAF3ED", "#9CC9AE", "#0B5640", (0.06, 0.5, 2.0), 0.0),
        "deshabilitado": ("#F5F7F6", "#F5F7F6", "#E1E7E3", "#A7B3AD", (0.0, 0.0, 1.0), 0.0),
    },
    "peligro": {
        "normal": ("#FFFFFF", "#F8F4F3", "#DCC6C3", "#B3261E", (0.12, 1.5, 4.0), 0.0),
        "hover": ("#FFFFFF", "#FBEAE8", "#E2A9A4", "#B3261E", (0.17, 2.5, 5.0), 0.0),
        "presionado": ("#F4D9D6", "#F9E7E5", "#E2A9A4", "#9C1F18", (0.06, 0.5, 2.0), 0.0),
        "deshabilitado": ("#F5F7F6", "#F5F7F6", "#E1E7E3", "#A7B3AD", (0.0, 0.0, 1.0), 0.0),
    },
}


class BotonPro(tk.Canvas):
    """Botón redondeado con degradado, sombra, efecto al pasar el mouse y al presionar.

    Se usa igual que un botón normal: configure(state="disabled"/"normal"),
    cget("state"), invoke(). Responde a Tab, Enter y barra espaciadora.
    """
    ALTO, RADIO, MARGEN = 36, 10, 8

    def __init__(self, padre, texto, comando=None, variante="secundario", fondo="#FFFFFF",
                 fuente=None, ancho=None, state="normal"):
        from tkinter import font as tkfont
        self._fuente = tkfont.Font(root=padre, font=fuente) if fuente else tkfont.nametofont("TkDefaultFont")
        self._ancho = ancho or max(92, self._fuente.measure(texto) + 40)
        m = self.MARGEN
        super().__init__(padre, width=self._ancho + 2 * m, height=self.ALTO + 2 * m, bg=fondo,
                         highlightthickness=0, bd=0, takefocus=1)
        self._fondo, self._variante, self._comando = fondo, variante, comando
        self._estado = state
        self._sobre = self._presionado = self._foco = False
        self._img = self.create_image(0, 0, anchor="nw")
        self._txt = self.create_text((self._ancho + 2 * m) / 2, (self.ALTO + 2 * m) / 2, text=texto,
                                     font=self._fuente)
        for evento, accion in (("<Enter>", self._entrar), ("<Leave>", self._salir),
                               ("<ButtonPress-1>", self._presionar), ("<ButtonRelease-1>", self._soltar),
                               ("<FocusIn>", lambda e: self._enfocar(True)),
                               ("<FocusOut>", lambda e: self._enfocar(False)),
                               ("<KeyPress-space>", lambda e: self.invoke()),
                               ("<KeyPress-Return>", lambda e: self.invoke())):
            self.bind(evento, accion, add="+")
        self._pintar()

    @property
    def activo(self):
        return self._estado != "disabled"

    def _modo(self):
        if not self.activo:
            return "deshabilitado"
        if self._presionado and self._sobre:
            return "presionado"
        return "hover" if self._sobre else "normal"

    def _pintar(self):
        modo = self._modo()
        arriba, abajo, borde, texto, sombra, brillo = APARIENCIA_BOTON[self._variante][modo]
        anillo = COLOR["verde"] if self._foco and self.activo else None
        m, w, h = self.MARGEN, self._ancho, self.ALTO
        clave = ("boton", w, self._variante, modo, bool(anillo), self._fondo)
        img = imagen_cache(self, clave, lambda: pintar_en(
            lienzo(w + 2 * m, h + 2 * m, self._fondo), m, m, w, h, self.RADIO,
            arriba, abajo, borde, sombra, brillo, anillo))
        self.itemconfigure(self._img, image=img)
        self.itemconfigure(self._txt, fill=texto)
        desplazamiento = 1 if modo == "presionado" else 0
        self.coords(self._txt, (w + 2 * m) / 2, (h + 2 * m) / 2 + desplazamiento)
        super().configure(cursor="hand2" if self.activo else "arrow", takefocus=1 if self.activo else 0)

    def _entrar(self, _e):
        self._sobre = True
        self._pintar()

    def _salir(self, _e):
        self._sobre = False
        self._pintar()

    def _presionar(self, _e):
        if self.activo:
            self._presionado = True
            self._pintar()

    def _soltar(self, _e):
        ejecutar = self._presionado and self._sobre and self.activo
        self._presionado = False
        self._pintar()
        if ejecutar:
            self.invoke()

    def _enfocar(self, valor):
        self._foco = valor
        self._pintar()

    def invoke(self):
        if self.activo and self._comando:
            self._comando()

    def configure(self, cnf=None, **kw):
        if cnf:
            kw.update(cnf)
        repintar = False
        if "state" in kw:
            self._estado = str(kw.pop("state"))
            self._presionado = self._presionado and self.activo
            repintar = True
        if "text" in kw:
            self.itemconfigure(self._txt, text=kw.pop("text"))
        if "command" in kw:
            self._comando = kw.pop("command")
        resultado = super().configure(**kw) if kw else None
        if repintar:
            self._pintar()
        return resultado

    config = configure

    def cget(self, clave):
        if clave == "state":
            return self._estado
        if clave == "text":
            return self.itemcget(self._txt, "text")
        return super().cget(clave)

    __getitem__ = cget


class Panel(tk.Canvas):
    """Tarjeta blanca con esquinas redondeadas, borde fino y sombra suave.

    Los widgets se ponen dentro de panel.interior. Para que cambiar el tamaño de
    la ventana sea fluido, la tarjeta se arma en 9 piezas: 4 esquinas pintadas y
    bordes/centro como líneas de color (técnica "9-slice").
    """
    RADIO, MARGEN = 14, 10
    SOMBRA = (0.07, 2.0, 8.0)

    def __init__(self, padre, fondo, relleno=(16, 10), ajustar_alto=True):
        super().__init__(padre, bg=fondo, highlightthickness=0, bd=0)
        self._fondo, self._ajustar = fondo, ajustar_alto
        self._px, self._py = relleno
        r, m = self.RADIO, self.MARGEN
        self._k = k = r + m
        self._base = pintar_en(lienzo(2 * k + 1, 2 * k + 1, fondo), m, m, 2 * r + 1, 2 * r + 1, r,
                               COLOR["blanco"], borde="#E1E7E3", sombra=self.SOMBRA)
        b = self._base
        self._esquinas = [a_imagen(self, [fila[x0:x0 + k] for fila in b[y0:y0 + k]])
                          for x0, y0 in ((0, 0), (k + 1, 0), (0, k + 1), (k + 1, k + 1))]
        self.interior = tk.Frame(self, bg=COLOR["blanco"])
        self._ventana = self.create_window(k + self._px - r, k + self._py - r, window=self.interior, anchor="nw")
        self.bind("<Configure>", self._redibujar)
        if ajustar_alto:
            self.interior.bind("<Configure>", self._ajustar_altura, add="+")

    def _desde_borde(self):
        return self._k - self.RADIO  # = MARGEN: espacio de la sombra alrededor

    def _ajustar_altura(self, evento):
        alto = evento.height + 2 * (self._desde_borde() + self._py)
        if int(self.cget("height")) != alto:
            super().configure(height=alto)

    def _redibujar(self, evento):
        w, h, k, b = evento.width, evento.height, self._k, self._base
        if w < 2 * k + 2 or h < 2 * k + 2:
            return
        self.delete("fondo")
        for img, (x, y) in zip(self._esquinas, ((0, 0), (w - k, 0), (0, h - k), (w - k, h - k))):
            self.create_image(x, y, image=img, anchor="nw", tags="fondo")
        for i in range(k):
            self.create_rectangle(k, i, w - k, i + 1, fill=_hex(b[i][k]), width=0, tags="fondo")
            self.create_rectangle(k, h - k + i, w - k, h - k + i + 1, fill=_hex(b[k + 1 + i][k]), width=0,
                                  tags="fondo")
            self.create_rectangle(i, k, i + 1, h - k, fill=_hex(b[k][i]), width=0, tags="fondo")
            self.create_rectangle(w - k + i, k, w - k + i + 1, h - k, fill=_hex(b[k][k + 1 + i]), width=0,
                                  tags="fondo")
        self.create_rectangle(k, k, w - k, h - k, fill=_hex(b[k][k]), width=0, tags="fondo")
        self.tag_lower("fondo")
        e = self._desde_borde()
        self.coords(self._ventana, e + self._px, e + self._py)
        opciones = {"width": max(w - 2 * (e + self._px), 1)}
        if not self._ajustar:
            opciones["height"] = max(h - 2 * (e + self._py), 1)
        self.itemconfigure(self._ventana, **opciones)


class Segmentado(tk.Canvas):
    """Control segmentado estilo Apple (p. ej. Resultados | Registro del proceso)."""
    ALTO, RADIO, MARGEN, HOLGURA = 30, 9, 5, 3

    def __init__(self, padre, opciones, al_cambiar, fondo, fuente, fuente_sel):
        from tkinter import font as tkfont
        self._f, self._fs = tkfont.Font(root=padre, font=fuente), tkfont.Font(root=padre, font=fuente_sel)
        self._seg = max(self._fs.measure(o) for o in opciones) + 32
        self._ancho = self._seg * len(opciones) + 2 * self.HOLGURA
        m = self.MARGEN
        super().__init__(padre, width=self._ancho + 2 * m, height=self.ALTO + 2 * m, bg=fondo,
                         highlightthickness=0, bd=0, takefocus=1, cursor="hand2")
        self._opciones, self._al_cambiar, self._fondo = opciones, al_cambiar, fondo
        self.seleccion, self._foco = 0, False
        self._img = self.create_image(0, 0, anchor="nw")
        self._textos = [self.create_text(m + self.HOLGURA + self._seg * (i + 0.5), m + self.ALTO / 2, text=o)
                        for i, o in enumerate(opciones)]
        self.bind("<ButtonPress-1>", self._clic)
        self.bind("<KeyPress-Left>", lambda e: self.seleccionar(self.seleccion - 1))
        self.bind("<KeyPress-Right>", lambda e: self.seleccionar(self.seleccion + 1))
        self.bind("<FocusIn>", lambda e: self._enfocar(True))
        self.bind("<FocusOut>", lambda e: self._enfocar(False))
        self._pintar()

    def _pintar(self):
        m, hg, sel = self.MARGEN, self.HOLGURA, self.seleccion
        pista = mezclar(self._fondo, COLOR["verde_oscuro"], 0.07)  # gris verdoso translúcido

        def fabricar():
            filas = lienzo(self._ancho + 2 * m, self.ALTO + 2 * m, self._fondo)
            pintar_en(filas, m, m, self._ancho, self.ALTO, self.RADIO, pista,
                      anillo=COLOR["verde"] if self._foco else None)
            return pintar_en(filas, m + hg + self._seg * sel, m + hg, self._seg, self.ALTO - 2 * hg,
                             self.RADIO - 2, "#FFFFFF", "#FBFCFB", "#DCE3DF", (0.14, 1.0, 3.0))
        self.itemconfigure(self._img, image=imagen_cache(
            self, ("seg", self._ancho, sel, self._foco, self._fondo), fabricar))
        for i, t in enumerate(self._textos):
            self.itemconfigure(t, font=self._fs if i == sel else self._f,
                               fill=COLOR["verde_oscuro"] if i == sel else COLOR["texto_suave"])

    def _clic(self, evento):
        i = int((evento.x - self.MARGEN - self.HOLGURA) // self._seg)
        if 0 <= i < len(self._opciones):
            self.seleccionar(i)

    def _enfocar(self, valor):
        self._foco = valor
        self._pintar()

    def seleccionar(self, i):
        i = max(0, min(len(self._opciones) - 1, i))
        if i != self.seleccion:
            self.seleccion = i
            self._pintar()
            self._al_cambiar(i)


class Chip(tk.Canvas):
    """Contador en forma de píldora de color: número en negrita y nombre del estado."""
    ALTO = 26

    def __init__(self, padre, titulo, color, fondo_chip, fondo, fuente, fuente_num):
        from tkinter import font as tkfont
        self._fn = tkfont.Font(root=padre, font=fuente_num)
        f = tkfont.Font(root=padre, font=fuente)
        self._ancho = 12 + self._fn.measure("0000") + 6 + f.measure(titulo) + 12
        super().__init__(padre, width=self._ancho, height=self.ALTO, bg=fondo, highlightthickness=0, bd=0)
        img = imagen_cache(self, ("chip", self._ancho, fondo_chip, fondo), lambda: pintar_en(
            lienzo(self._ancho, self.ALTO, fondo), 0, 0, self._ancho, self.ALTO, self.ALTO / 2,
            fondo_chip, borde=mezclar(fondo_chip, color, 0.16)))
        self.create_image(0, 0, image=img, anchor="nw")
        self._num = self.create_text(12, self.ALTO / 2, text="0", anchor="w", font=self._fn, fill=color)
        self._tit = self.create_text(0, self.ALTO / 2, text=titulo, anchor="w", font=f, fill=color)
        self.valor = 0
        self.fijar(0)

    def fijar(self, n):
        self.valor = n
        self.itemconfigure(self._num, text=str(n))
        self.coords(self._tit, 12 + self._fn.measure(str(n)) + 6, self.ALTO / 2)


class BarraProgreso(tk.Canvas):
    """Barra de progreso delgada con extremos redondeados."""
    GROSOR = 6

    def __init__(self, padre, fondo):
        super().__init__(padre, height=self.GROSOR + 4, bg=fondo, highlightthickness=0, bd=0)
        self.valor, self.maximo = 0, 1
        self.bind("<Configure>", lambda e: self._pintar())

    def fijar(self, valor=None, maximo=None):
        if maximo is not None:
            self.maximo = max(maximo, 1)
        if valor is not None:
            self.valor = valor
        self._pintar()

    def _pintar(self):
        self.delete("all")
        w, y, g = self.winfo_width(), (self.GROSOR + 4) / 2, self.GROSOR
        if w < 2 * g:
            return
        self.create_line(g / 2, y, w - g / 2, y, width=g, capstyle="round", fill="#E3E9E5")
        fraccion = min(max(self.valor / self.maximo, 0), 1)
        if fraccion > 0:
            self.create_line(g / 2, y, g / 2 + (w - g) * fraccion, y, width=g, capstyle="round",
                             fill=COLOR["verde"])


class Insignia(tk.Canvas):
    """Círculo verde con el número del paso (1, 2, 3)."""
    TAM = 24

    def __init__(self, padre, numero, fondo, fuente):
        t = self.TAM
        super().__init__(padre, width=t, height=t, bg=fondo, highlightthickness=0, bd=0)
        img = imagen_cache(self, ("insignia", fondo), lambda: pintar_en(
            lienzo(t, t, fondo), 0, 0, t, t, t / 2, "#0A9341", "#01722E", brillo=0.25))
        self.create_image(0, 0, image=img, anchor="nw")
        self.create_text(t / 2, t / 2, text=str(numero), fill="#FFFFFF", font=fuente)


class Ayuda:
    """Globo de ayuda que aparece al dejar el mouse sobre un control."""

    def __init__(self, widget, texto, fuente):
        self.widget, self.texto, self.fuente = widget, texto, fuente
        self._ventana = self._tarea = None
        widget.bind("<Enter>", self._programar, add="+")
        widget.bind("<Leave>", self._ocultar, add="+")
        widget.bind("<ButtonPress>", self._ocultar, add="+")

    def _programar(self, _e):
        self._tarea = self.widget.after(650, self._mostrar)

    def _mostrar(self):
        if self._ventana:
            return
        x = self.widget.winfo_rootx() + 8
        y = self.widget.winfo_rooty() + self.widget.winfo_height() - 4
        self._ventana = v = tk.Toplevel(self.widget)
        v.wm_overrideredirect(True)
        v.wm_geometry(f"+{x}+{y}")
        marco = tk.Frame(v, bg=COLOR["linea"], padx=1, pady=1)
        marco.pack()
        tk.Label(marco, text=self.texto, bg="#FFFFFF", fg=COLOR["texto"], font=self.fuente, justify="left",
                 wraplength=320, padx=10, pady=6).pack()

    def _ocultar(self, _e=None):
        if self._tarea:
            self.widget.after_cancel(self._tarea)
            self._tarea = None
        if self._ventana:
            self._ventana.destroy()
            self._ventana = None


# Estado de la sesión -> (color del punto, título, explicación)
AVISO_INACTIVIDAD_MIN = 10  # VUR cierra la sesión por inactividad; se avisa antes
ESTADOS_SESION = {
    "cerrada": ("#9AA8A1", "Sin sesión", "Chrome está cerrado. Se abrirá al iniciar consultas."),
    "esperando": ("#E0A100", "Esperando ingreso", "Inicie sesión y resuelva el captcha en Chrome."),
    "activa": ("#018D38", "Sesión activa", "VUR está respondiendo."),
    "inactiva": ("#E0A100", "Sesión inactiva", "Sin actividad reciente: VUR puede pedir ingresar de nuevo."),
    "vencida": ("#B3261E", "Sesión vencida", "VUR cerró la sesión: ingrese de nuevo en Chrome."),
    "limite": ("#B3261E", "Cupo agotado", "VUR no acepta más consultas de este usuario hoy."),
}


class Semaforo(tk.Frame):
    """Indicador del estado de la sesión de VUR (punto de color + texto)."""

    def __init__(self, padre, fondo, fuente, fuente_detalle):
        super().__init__(padre, bg=fondo)
        self._fondo = fondo
        self._punto = tk.Canvas(self, width=20, height=20, bg=fondo, highlightthickness=0, bd=0)
        self._punto.pack(side="left", padx=(0, 8))
        self._img = self._punto.create_image(0, 0, anchor="nw")
        textos = tk.Frame(self, bg=fondo)
        textos.pack(side="left")
        self._titulo = tk.Label(textos, bg=fondo, fg=COLOR["texto"], font=fuente, anchor="w")
        self._titulo.pack(anchor="w")
        self._detalle = tk.Label(textos, bg=fondo, fg=COLOR["texto_suave"], font=fuente_detalle, anchor="w")
        self._detalle.pack(anchor="w")
        self.estado = None
        self.fijar("cerrada")

    def fijar(self, estado, detalle=None):
        color, titulo, explicacion = ESTADOS_SESION[estado]
        img = imagen_cache(self, ("semaforo", color, self._fondo), lambda: pintar_en(
            lienzo(20, 20, self._fondo), 2, 2, 16, 16, 8, mezclar(color, "#FFFFFF", 0.25), color,
            sombra=(0.18, 0.5, 2.0), brillo=0.35))
        self._punto.itemconfigure(self._img, image=img)
        self.estado = estado
        self._titulo.configure(text=f"VUR · {titulo}")
        self._detalle.configure(text=detalle or explicacion)


def colorear_barra_titulo(ventana, fondo="#FFFFFF", texto="#1C2723"):
    """Windows 11: pinta la barra de título del color del encabezado (en Windows 10 no hace nada)."""
    import sys
    if sys.platform != "win32":
        return
    try:
        import ctypes
        ventana.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(ventana.winfo_id())

        def colorref(color):
            r, g, b = _rgb(color)
            return ctypes.c_int(r | (g << 8) | (b << 16))
        for atributo, valor in ((34, colorref(COLOR["linea"])), (35, colorref(fondo)), (36, colorref(texto))):
            ctypes.windll.dwmapi.DwmSetWindowAttribute(hwnd, atributo, ctypes.byref(valor), ctypes.sizeof(valor))
    except Exception:  # noqa: BLE001 - estético; nunca debe impedir abrir el programa
        pass


class App:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{TITULO_APLICATIVO} · {SECRETARIA} · {DEPENDENCIA} · Gobernación de Antioquia")
        self.root.geometry("1180x860")
        self.root.minsize(1120, 640)
        self.root.configure(bg=COLOR["fondo"])

        self.cola = queue.Queue()
        self.detener = threading.Event()
        self.lock = threading.Lock()
        self.df = None
        self.resultados = []
        self.ruta_excel = None
        self.hilo = None
        self.cliente = None
        self._logo = None

        self._fuentes()
        self._estilos()
        self.forzar_vur = tk.BooleanVar(value=False)
        self._ultima_respuesta = None   # hora de la última respuesta de VUR (para el aviso de inactividad)
        self._forzar = False
        self.n_historico = 0
        self._construir_ui()
        colorear_barra_titulo(self.root)
        icono = CARPETA_RECURSOS / "icono_vur.ico"
        if icono.exists():
            try:
                self.root.iconbitmap(default=str(icono))
            except tk.TclError:
                pass
        self.usuario, self.equipo = _usuario_equipo()
        self.historico = Historico(log=self.log_hilo)
        try:
            n = self.historico.cargar()
            self._log(f"Histórico: {n} matrícula(s) consultadas antes, en {self.historico.carpeta}")
        except Exception as e:  # noqa: BLE001 - el histórico nunca debe impedir abrir el programa
            self._log(f"Histórico no disponible: {type(e).__name__}: {e}")
        self.root.protocol("WM_DELETE_WINDOW", self._al_cerrar)
        self.root.after(100, self._procesar_cola)
        self.root.after(15000, self._vigilar_sesion)

    # --- Apariencia ------------------------------------------------------------

    def _fuentes(self):
        from tkinter import font as tkfont
        familias = set(tkfont.families(self.root))
        base = next((f for f in ("Poppins", "Segoe UI", "Helvetica") if f in familias), "TkDefaultFont")
        self.f_titulo = (base, 15, "bold")
        self.f_subtitulo = (base, 10)
        self.f_subtitulo_negrita = (base, 10, "bold")
        self.f_pequena_negrita = (base, 9, "bold")
        self.f_seccion = (base, 11, "bold")
        self.f_boton = (base, 10, "bold")
        self.f_texto = (base, 10)
        self.f_pequena = (base, 9)
        self.f_mini = (base, 8)
        self.f_numero = (base, 10, "bold")
        self.f_mono = ("Consolas", 9) if "Consolas" in familias else ("Courier New", 9)

    def _estilos(self):
        st = ttk.Style(self.root)
        st.theme_use("clam")
        c = COLOR
        st.configure(".", background=c["fondo"], foreground=c["texto"], font=self.f_texto)
        st.configure("Treeview", background=c["blanco"], fieldbackground=c["blanco"], rowheight=28,
                     font=self.f_texto, borderwidth=0, relief="flat")
        st.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])  # sin marco alrededor
        st.configure("Treeview.Heading", background="#F7F9F8", foreground=c["texto_suave"],
                     font=self.f_pequena_negrita, relief="flat", borderwidth=0, padding=(8, 7))
        st.map("Treeview.Heading", background=[("active", "#EDF3EF")])
        st.map("Treeview", background=[("selected", c["verde"])], foreground=[("selected", c["blanco"])])
        st.configure("Vertical.TScrollbar", gripcount=0, background="#D5DDD8", troughcolor=c["blanco"],
                     bordercolor=c["blanco"], lightcolor="#D5DDD8", darkcolor="#D5DDD8",
                     arrowcolor=c["texto_suave"], relief="flat", arrowsize=11)
        st.map("Vertical.TScrollbar", background=[("active", "#B9C6BE")])
        st.configure("Opcion.TCheckbutton", background=c["blanco"], foreground=c["texto"], font=self.f_pequena,
                     indicatorcolor=c["blanco"], indicatorbackground=c["blanco"])
        st.map("Opcion.TCheckbutton", indicatorcolor=[("selected", c["verde"]), ("pressed", "#E3F4EA")],
               background=[("active", c["blanco"])])

    def _cargar_logo(self, contenedor):
        """Muestra el logo oficial sin caja: el PNG tiene fondo transparente y va sobre blanco."""
        # Logo oficial sin margen, 157 x 70 px (mínimo del manual: 150 x 67), guardado al final del archivo.
        try:
            self._logo = tk.PhotoImage(data="".join(RECURSOS["logo_gobernacion_app.png"].split()))
            tk.Label(contenedor, image=self._logo, bg=COLOR["blanco"], bd=0).pack(side="left", padx=(24, 18), pady=10)
            return
        except (tk.TclError, KeyError):
            pass
        tk.Label(contenedor, text="GOBERNACIÓN\nDE ANTIOQUIA", bg=COLOR["blanco"], fg=COLOR["verde_oscuro"],
                 font=self.f_seccion, justify="left").pack(side="left", padx=(24, 18), pady=14)

    # --- Construcción de la interfaz -------------------------------------------

    def _boton(self, padre, texto, comando, variante="secundario", ayuda=None, **kw):
        b = BotonPro(padre, texto, comando, variante=variante, fondo=COLOR["blanco"], fuente=self.f_boton, **kw)
        if ayuda:
            Ayuda(b, ayuda, self.f_pequena)
        return b

    def _paso(self, padre, columna, numero):
        grupo = tk.Frame(padre, bg=COLOR["blanco"])
        grupo.grid(row=0, column=columna, sticky="w")
        Insignia(grupo, numero, COLOR["blanco"], self.f_pequena_negrita).pack(side="left", padx=(0, 4))
        return grupo

    def _divisor(self, padre, columna):
        tk.Frame(padre, bg=COLOR["linea"], width=1).grid(row=0, column=columna, sticky="ns", padx=10, pady=10)

    def _construir_ui(self):
        c = COLOR
        # Encabezado: blanco, como pide el manual para el logo (fondo blanco sin contenedor).
        cab = tk.Frame(self.root, bg=c["blanco"])
        cab.pack(fill="x")
        self._cargar_logo(cab)
        tk.Frame(cab, bg=c["linea"], width=1).pack(side="left", fill="y", pady=20, padx=(0, 18))
        textos = tk.Frame(cab, bg=c["blanco"])
        textos.pack(side="left")
        tk.Label(textos, text=TITULO_APLICATIVO, bg=c["blanco"], fg=c["verde_oscuro"],
                 font=self.f_titulo).pack(anchor="w")
        tk.Label(textos, text=f"{SECRETARIA}  ·  {DEPENDENCIA}", bg=c["blanco"],
                 fg=c["verde"], font=self.f_subtitulo_negrita).pack(anchor="w")
        self.semaforo = Semaforo(cab, c["blanco"], self.f_subtitulo_negrita, self.f_pequena)
        self.semaforo.pack(side="right", padx=(0, 28))
        Ayuda(self.semaforo, "Estado de la sesión en VUR.\nVerde: activa. Amarillo: esperando ingreso o sin "
                             f"actividad por más de {AVISO_INACTIVIDAD_MIN} min. Rojo: vencida o cupo diario "
                             "agotado. Gris: Chrome cerrado.\nAl terminar un proceso, Chrome queda abierto para "
                             "reutilizar la sesión en el siguiente sin volver a resolver el captcha.", self.f_pequena)
        tk.Frame(self.root, bg=c["verde"], height=2).pack(fill="x")  # filete institucional

        # Pie
        pie = tk.Frame(self.root, bg=c["fondo"])
        pie.pack(side="bottom", fill="x", padx=24, pady=(0, 8))
        tk.Label(pie, text=f"Proyectos: {FUENTE_DATOS}", bg=c["fondo"], fg=c["texto_suave"],
                 font=self.f_mini).pack(side="left")
        tk.Label(pie, text=CREDITO, bg=c["fondo"], fg=c["verde_oscuro"],
                 font=self.f_pequena_negrita).pack(side="right")
        tk.Label(pie, text=f"Versión {VERSION_APLICATIVO}   ·", bg=c["fondo"], fg=c["texto_suave"],
                 font=self.f_mini).pack(side="right", padx=(0, 6))

        cuerpo = tk.Frame(self.root, bg=c["fondo"])
        cuerpo.pack(fill="both", expand=True, padx=12, pady=(6, 0))

        # 1) Barra de acciones: los tres pasos en una sola fila compacta.
        barra = Panel(cuerpo, c["fondo"], relleno=(14, 6))
        barra.pack(fill="x")
        fila = barra.interior
        fila.columnconfigure(0, weight=1)

        g1 = self._paso(fila, 0, 1)
        self.btn_seleccionar = self._boton(
            g1, "Seleccionar Excel", self.seleccionar_excel,
            ayuda="Elija el Excel con la columna de matrículas (por ejemplo 001-123456).")
        self.btn_seleccionar.pack(side="left")
        self._divisor(fila, 1)

        g2 = self._paso(fila, 2, 2)
        self.btn_iniciar = self._boton(
            g2, "Iniciar consultas", self.iniciar, "primario",
            ayuda="Abre Chrome en VUR. Usted inicia sesión y resuelve el captcha; luego el programa consulta "
                  "todas las matrículas y sus matrices.")
        self.btn_iniciar.pack(side="left")
        self.btn_detener = self._boton(g2, "Detener", self.solicitar_detener, "peligro", state="disabled",
                                       ayuda="Se detiene al terminar la matrícula actual y guarda el parcial.")
        self.btn_detener.pack(side="left")
        self.btn_diagnostico = self._boton(
            g2, "Diagnosticar", self.diagnosticar,
            ayuda="Descarga la respuesta completa de VUR de una matrícula a un archivo JSON, para soporte técnico.")
        self.btn_diagnostico.pack(side="left")
        self._divisor(fila, 3)

        g3 = self._paso(fila, 4, 3)
        self.btn_guardar = self._boton(g3, "Guardar Excel", self.guardar,
                                       ayuda="Guarda todos los resultados con el formato institucional. Además, durante la "
                                             f"consulta se guarda un parcial automático cada {GUARDADO_PARCIAL_CADA} "
                                             "matrículas.")
        self.btn_guardar.pack(side="left")
        self.btn_comparar = self._boton(
            g3, "Comparar", self.comparar,
            ayuda="Compara dos Excel de resultados del mismo lote (por ejemplo el de hoy con uno anterior) y "
                  "genera un Excel con lo que cambió y la causa probable.")
        self.btn_comparar.pack(side="left")

        # 2) Estado de la consulta: progreso, avisos y contadores en poco espacio.
        estado = Panel(cuerpo, c["fondo"], relleno=(18, 8))
        estado.pack(fill="x", pady=(0, 0))
        e = estado.interior
        sup = tk.Frame(e, bg=c["blanco"])
        sup.pack(fill="x")
        self.lbl_archivo = tk.Label(sup, text="Sin archivo cargado", bg=c["blanco"], fg=c["texto_suave"],
                                    font=self.f_pequena)
        self.lbl_archivo.pack(side="left", padx=(0, 14))
        self.lbl_progreso = tk.Label(sup, text="Sin consulta en curso", bg=c["blanco"], fg=c["texto"],
                                     font=self.f_subtitulo_negrita)
        self.lbl_progreso.pack(side="left")
        self.lbl_revisar = tk.Label(sup, text="", bg=c["blanco"], fg="#A3570A", font=self.f_pequena_negrita)
        self.lbl_revisar.pack(side="left", padx=(16, 0))
        self.n_revisar = 0
        self.lbl_consultas = tk.Label(sup, text="", bg=c["blanco"], fg=c["texto_suave"], font=self.f_pequena)
        self.lbl_consultas.pack(side="right")
        self.progreso = BarraProgreso(e, c["blanco"])
        self.progreso.pack(fill="x", pady=(6, 8))
        chips = tk.Frame(e, bg=c["blanco"])
        chips.pack(fill="x")
        self.contadores = {}
        for estado_reg, (titulo, color, fondo) in ESTILO_ESTADO.items():
            chip = Chip(chips, titulo, color, fondo, c["blanco"], self.f_pequena, self.f_numero)
            chip.pack(side="left", padx=(0, 8))
            self.contadores[estado_reg] = chip
        self.chk_forzar = ttk.Checkbutton(chips, text="Forzar consulta en VUR", variable=self.forzar_vur,
                                          style="Opcion.TCheckbutton", takefocus=1)
        self.chk_forzar.pack(side="right")
        Ayuda(self.chk_forzar, "Marcado: consulta todas las matrículas en VUR aunque estén en el histórico. "
                               f"Sin marcar: las consultadas hace menos de {VIGENCIA_HISTORICO_DIAS} días con "
                               "resultado ALTA se toman del histórico, sin gastar cupo de VUR.", self.f_pequena)

        # 3) Resultados: ocupan todo el espacio restante.
        resultados = Panel(cuerpo, c["fondo"], relleno=(2, 2), ajustar_alto=False)
        resultados.pack(fill="both", expand=True)
        r = resultados.interior
        r.rowconfigure(2, weight=1)
        r.columnconfigure(0, weight=1)
        cab_r = tk.Frame(r, bg=c["blanco"])
        cab_r.grid(row=0, column=0, sticky="ew", padx=(8, 14), pady=(2, 2))
        self.segmentos = Segmentado(cab_r, ["Resultados", "Registro del proceso"], self._mostrar_vista,
                                    c["blanco"], self.f_texto, self.f_subtitulo_negrita)
        self.segmentos.pack(side="left")
        tk.Label(cab_r, text="Doble clic o Enter sobre una fila para ver el detalle completo",
                 bg=c["blanco"], fg=c["texto_suave"], font=self.f_pequena).pack(side="right")
        tk.Frame(r, bg=c["linea"], height=1).grid(row=1, column=0, sticky="ew")

        self.vista_tabla = tk.Frame(r, bg=c["blanco"])
        self.vista_tabla.grid(row=2, column=0, sticky="nsew", padx=(1, 1), pady=(0, 10))
        self.tabla = ttk.Treeview(self.vista_tabla, columns=[col[0] for col in COLUMNAS_TABLA], show="headings",
                                  selectmode="browse")
        for col, titulo, ancho in COLUMNAS_TABLA:
            self.tabla.heading(col, text=titulo.upper(), anchor="w")
            self.tabla.column(col, width=ancho, anchor="w", stretch=col in ("Propietario_Gravado", "Estado"))
        for estado_reg, (_, color, fondo) in ESTILO_ESTADO.items():
            self.tabla.tag_configure(estado_reg, background=fondo, foreground=COLOR["texto"])
        barra_v = ttk.Scrollbar(self.vista_tabla, orient="vertical", command=self.tabla.yview)
        self.tabla.configure(yscrollcommand=barra_v.set)
        barra_v.pack(side="right", fill="y")  # primero la barra: así nunca queda fuera de la ventana
        self.tabla.pack(side="left", fill="both", expand=True)
        self.tabla.bind("<Double-1>", self._ver_detalle)
        self.tabla.bind("<Return>", self._ver_detalle)
        # En Tk 8.6 el color de estado de la fila tapa el resaltado de selección; por eso
        # la fila elegida recibe temporalmente la etiqueta "seleccion" (verde, texto blanco).
        self.tabla.tag_configure("seleccion", background=c["verde"], foreground=c["blanco"])
        self._fila_marcada = None  # (iid, etiquetas originales)
        self.tabla.bind("<<TreeviewSelect>>", self._marcar_seleccion)
        self.lbl_vacio = tk.Label(self.vista_tabla, bg=c["blanco"], fg=c["texto_suave"], font=self.f_texto,
                                  justify="center", text="Aún no hay resultados.\n"
                                  "Seleccione el Excel de matrículas (paso 1) y pulse «Iniciar consultas» (paso 2).")
        self.lbl_vacio.place(relx=0.5, rely=0.45, anchor="center")

        self.vista_log = tk.Frame(r, bg=c["blanco"])
        self.vista_log.grid(row=2, column=0, sticky="nsew", padx=(1, 1), pady=(0, 10))
        self.txt_log = tk.Text(self.vista_log, height=12, state="disabled", font=self.f_mono, relief="flat",
                               bd=0, bg=COLOR["blanco"], fg=COLOR["texto"], padx=14, pady=10, wrap="word",
                               highlightthickness=0)
        barra_log = ttk.Scrollbar(self.vista_log, orient="vertical", command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=barra_log.set)
        self.txt_log.pack(side="left", fill="both", expand=True)
        barra_log.pack(side="right", fill="y")
        self.vista_tabla.tkraise()

    def _marcar_seleccion(self, _evento=None):
        if self._fila_marcada and self.tabla.exists(self._fila_marcada[0]):
            self.tabla.item(self._fila_marcada[0], tags=self._fila_marcada[1])
        self._fila_marcada = None
        sel = self.tabla.selection()
        if sel:
            self._fila_marcada = (sel[0], self.tabla.item(sel[0], "tags"))
            self.tabla.item(sel[0], tags=("seleccion",))

    def _notificar_sesion(self, estado):
        """Llamado desde el hilo de consultas: el cambio se pinta en el hilo de la ventana."""
        self.cola.put(("sesion", estado))

    def _estado_sesion(self, estado):
        if estado == "activa":
            self._ultima_respuesta = datetime.now()
            self.semaforo.fijar("activa", f"Última respuesta de VUR: {self._ultima_respuesta:%H:%M}")
        else:
            if estado in ("cerrada", "vencida", "limite"):
                self._ultima_respuesta = None
            self.semaforo.fijar(estado)

    def _vigilar_sesion(self):
        """Cada 15 s: avisa la inactividad y detecta si el usuario cerró Chrome (solo sin proceso en curso)."""
        try:
            ocupado = bool(self.hilo and self.hilo.is_alive())
            if not ocupado and self.cliente and self.semaforo.estado != "cerrada" and not self.cliente.vivo():
                self.cliente.driver = None
                self._estado_sesion("cerrada")
            elif self.semaforo.estado in ("activa", "inactiva") and self._ultima_respuesta:
                minutos = int((datetime.now() - self._ultima_respuesta).total_seconds() // 60)
                if minutos >= AVISO_INACTIVIDAD_MIN:
                    self.semaforo.fijar("inactiva", f"Sin actividad hace {minutos} min: puede pedir ingresar de nuevo")
        finally:
            self.root.after(15000, self._vigilar_sesion)

    def _obtener_cliente(self):
        """Reutiliza el Chrome abierto (y su sesión) o prepara uno nuevo."""
        if self.cliente and self.cliente.vivo():
            self.cliente.nueva_ejecucion()
            return self.cliente, self.semaforo.estado in ("activa", "inactiva")
        self.cliente = VURClient(self.log_hilo, self._notificar_sesion)
        return self.cliente, False

    def _mostrar_vista(self, indice):
        (self.vista_tabla if indice == 0 else self.vista_log).tkraise()

    # --- Registro, progreso y tabla ---------------------------------------------

    def _log(self, mensaje):
        self.txt_log.configure(state="normal")
        self.txt_log.insert("end", f"[{datetime.now():%H:%M:%S}] {mensaje}\n")
        self.txt_log.see("end")
        self.txt_log.configure(state="disabled")

    def log_hilo(self, mensaje):
        self.cola.put(("log", mensaje))

    def _en_ejecucion(self, activo):
        estado = "disabled" if activo else "normal"
        for b in (self.btn_seleccionar, self.btn_iniciar, self.btn_diagnostico, self.btn_guardar, self.btn_comparar):
            b.configure(state=estado)
        self.chk_forzar.configure(state=estado)
        self.btn_detener.configure(state="normal" if activo else "disabled")

    def _sumar_contador(self, estado, delta):
        if estado in self.contadores:
            chip = self.contadores[estado]
            chip.fijar(chip.valor + delta)

    def _texto_revisar(self):
        self.lbl_revisar.configure(
            text=f"⚑ {self.n_revisar} {'requiere' if self.n_revisar == 1 else 'requieren'} revisión del profesional"
            if self.n_revisar else "")

    def _reiniciar_vista(self):
        self.tabla.delete(*self.tabla.get_children())
        self._fila_marcada = None
        for chip in self.contadores.values():
            chip.fijar(0)
        self.progreso.fijar(0, len(self.df))
        self.lbl_progreso.configure(text=f"0 de {len(self.df)} matrículas")
        self.lbl_consultas.configure(text="")
        self.n_revisar = 0
        self._texto_revisar()
        self.lbl_vacio.configure(text=f"Archivo listo: {len(self.df)} matrículas.\n"
                                      "Pulse «Iniciar consultas» (paso 2) para comenzar.")
        self.lbl_vacio.place(relx=0.5, rely=0.45, anchor="center")

    def _agregar_fila(self, reg):
        self.lbl_vacio.place_forget()
        estado = reg.get("Estado", "")
        valores = valores_tabla(reg)
        iid = self.tabla.insert("", "end", values=valores, tags=(estado,) if estado in ESTILO_ESTADO else ())
        self.tabla.see(iid)
        self._sumar_contador(estado, 1)
        if reg.get("Confiabilidad") == CONFIABILIDAD_REVISAR:
            self.n_revisar += 1
            self._texto_revisar()

    def _actualizar_fila(self, indice, anterior, nuevo):
        """Reemplaza una fila de la tabla tras la segunda pasada y ajusta los contadores."""
        hijos = self.tabla.get_children()
        if indice >= len(hijos):
            return
        estado = nuevo.get("Estado", "")
        etiquetas = (estado,) if estado in ESTILO_ESTADO else ()
        if self._fila_marcada and self._fila_marcada[0] == hijos[indice]:
            self._fila_marcada = (hijos[indice], etiquetas)  # conserva el resaltado de selección
            etiquetas = ("seleccion",)
        self.tabla.item(hijos[indice], values=valores_tabla(nuevo),
                        tags=etiquetas)
        for reg, delta in ((anterior, -1), (nuevo, 1)):
            self._sumar_contador(reg.get("Estado", ""), delta)
            if reg.get("Confiabilidad") == CONFIABILIDAD_REVISAR:
                self.n_revisar += delta
        self._texto_revisar()

    def _ver_detalle(self, _evento=None):
        sel = self.tabla.selection()
        if not sel:
            return
        indice = self.tabla.index(sel[0])
        with self.lock:
            if indice >= len(self.resultados):
                return
            reg = dict(self.resultados[indice])
        c = COLOR
        ven = tk.Toplevel(self.root)
        ven.title(f"Detalle {reg.get('Matricula_Limpia', '')}")
        ven.geometry("860x600")
        ven.configure(bg=c["blanco"])
        cab = tk.Frame(ven, bg=c["blanco"])
        cab.pack(fill="x", padx=24, pady=(16, 12))
        tk.Label(cab, text=f"{reg.get('Matricula_Limpia', '')}  ·  {reg.get('Municipio', '')}", bg=c["blanco"],
                 fg=c["verde_oscuro"], font=self.f_titulo).pack(anchor="w")
        linea = tk.Frame(cab, bg=c["blanco"])
        linea.pack(anchor="w", pady=(6, 0))
        estado = reg.get("Estado", "")
        _, color, fondo = ESTILO_ESTADO.get(estado, ("", c["texto"], "#ECEFED"))
        tk.Label(linea, text=f"  {estado or '—'}  ", bg=fondo, fg=color, font=self.f_pequena_negrita,
                 pady=3).pack(side="left")
        confiab = reg.get("Confiabilidad", "")
        tk.Label(linea, text=f"Confiabilidad: {confiab or '—'}", bg=c["blanco"],
                 fg="#A3570A" if confiab == CONFIABILIDAD_REVISAR else c["texto_suave"],
                 font=self.f_pequena_negrita).pack(side="left", padx=(12, 0))
        tk.Frame(ven, bg=c["verde"], height=2).pack(fill="x")
        marco = tk.Frame(ven, bg=c["blanco"])
        marco.pack(fill="both", expand=True)
        txt = tk.Text(marco, wrap="word", font=self.f_texto, relief="flat", bd=0, padx=24, pady=10,
                      bg=c["blanco"], fg=c["texto"], highlightthickness=0, spacing1=1, spacing3=1)
        barra = ttk.Scrollbar(marco, orient="vertical", command=txt.yview)
        txt.configure(yscrollcommand=barra.set)
        txt.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        txt.tag_configure("t", font=self.f_pequena_negrita, foreground=c["verde"], spacing1=14, spacing3=4)
        for campo, titulo in (("Fuente", "Fuente del resultado"),
                              ("Motivos_Revision", "Por qué requiere revisión"),
                              ("Segunda_Pasada", "Segunda pasada"),
                              ("Propietarios_Actuales", "Propietarios actuales"),
                              ("Propietario_Gravado", "Propietario gravado"),
                              ("Detalle_Valorizacion", "Anotación de valorización"),
                              ("Observaciones", "Anotaciones completas"),
                              ("Afectaciones_Valorizacion", "Afectaciones por valorización (no definen el paz y salvo)"),
                              ("Cadena_Revisada", "Cadena de matrices revisada"),
                              ("Proyectos_Municipio", "Proyectos del municipio"),
                              ("Notas_Proceso", "Notas del proceso")):
            valor = str(reg.get(campo, "") or "—")
            if campo == "Cadena_Revisada":
                valor = valor.replace(" -> ", "\n→ ")
            txt.insert("end", titulo.upper() + "\n", "t")
            txt.insert("end", valor + "\n")
        txt.configure(state="disabled")
        colorear_barra_titulo(ven)

    def _procesar_cola(self):
        try:
            while True:
                tipo, dato = self.cola.get_nowait()
                if tipo == "log":
                    self._log(dato)
                elif tipo == "sesion":
                    self._estado_sesion(dato)
                elif tipo == "progreso":
                    hechas, consultas, reutilizadas, del_historico = dato
                    total = len(self.df) if self.df is not None else 0
                    self.progreso.fijar(hechas, total)
                    self.lbl_progreso.configure(text=f"{hechas} de {total} matrículas")
                    self.lbl_consultas.configure(
                        text=f"{del_historico} del histórico · {consultas} consultas a VUR · "
                             f"{reutilizadas} reutilizadas")
                elif tipo == "resultado":
                    self._agregar_fila(dato)
                elif tipo == "actualizar":
                    self._actualizar_fila(*dato)
                elif tipo == "estado":
                    self.lbl_progreso.configure(text=dato)
                elif tipo == "confirmar":
                    mensaje, evento, respuesta = dato
                    respuesta["ok"] = messagebox.askokcancel("VUR", mensaje, parent=self.root)
                    evento.set()
                elif tipo == "info":
                    messagebox.showinfo("VUR", dato, parent=self.root)
                elif tipo == "fin":
                    self._en_ejecucion(False)
        except queue.Empty:
            pass
        self.root.after(100, self._procesar_cola)

    def _esperar_confirmacion(self, mensaje):
        """Llamado desde el hilo de trabajo: muestra el diálogo en el hilo principal y espera."""
        evento, respuesta = threading.Event(), {}
        self.cola.put(("confirmar", (mensaje, evento, respuesta)))
        while not evento.wait(0.5):
            if self.detener.is_set():
                return False
        return respuesta.get("ok", False)

    # --- Acciones de botones ------------------------------------------------

    def seleccionar_excel(self):
        ruta = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xls")], initialdir=str(CARPETA_CONSULTAS),
                                          title="Excel de matrículas")
        if not ruta:
            return
        try:
            df = leer_excel(ruta)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el archivo:\n{type(e).__name__}: {e}")
            return
        self.df, self.ruta_excel = df, Path(ruta)
        with self.lock:
            self.resultados = []
        repetidas = int(df["Matricula_Limpia"].duplicated().sum())
        extra = f" · {repetidas} repetidas" if repetidas else ""
        unicas = list(dict.fromkeys(df["Matricula_Limpia"]))
        en_historico = [m for m in unicas if self.historico.buscar(m)]
        vigentes = [m for m in en_historico if evaluar_historico(self.historico.buscar(m))[0]]
        hist = f" · {len(en_historico)} en histórico ({len(vigentes)} vigentes)" if en_historico else ""
        self.lbl_archivo.configure(text=f"{self.ruta_excel.name} · {len(df)} matrículas{extra}{hist}")
        self._reiniciar_vista()
        self._log(f"Archivo cargado: {self.ruta_excel.name} ({len(df)} matrículas{extra}{hist})")
        if en_historico:
            self._avisar_historico(en_historico, vigentes)

    def _avisar_historico(self, en_historico, vigentes):
        lista = ", ".join(en_historico[:12]) + (" …" if len(en_historico) > 12 else "")
        texto = (f"{len(en_historico)} matrícula(s) de este archivo ya fueron consultadas antes:\n{lista}\n\n"
                 f"• {len(vigentes)} tienen un resultado vigente (menos de {VIGENCIA_HISTORICO_DIAS} días, "
                 "confiabilidad ALTA): se tomarán del histórico sin consultar VUR.\n"
                 f"• {len(en_historico) - len(vigentes)} se consultarán de nuevo (resultado antiguo o que "
                 "requería revisión).\n\n"
                 "Para consultarlas todas en VUR, marque «Forzar consulta en VUR».\n\n"
                 f"¿Abrir ahora el Excel del histórico ({self.historico.archivo_excel.name})?")
        if messagebox.askyesno("Matrículas ya consultadas", texto, parent=self.root):
            if not self.historico.archivo_excel.exists():
                self.historico.exportar_excel()
            try:
                os.startfile(self.historico.archivo_excel)  # solo Windows
            except (AttributeError, OSError) as e:
                self._log(f"No se pudo abrir el histórico: {type(e).__name__}")

    def _validar_carga(self):
        if self.df is None or self.df.empty:
            messagebox.showwarning("VUR", "Primero cargue un Excel con matrículas.")
            return False
        return not (self.hilo and self.hilo.is_alive())

    def iniciar(self):
        if not self._validar_carga():
            return
        self.detener.clear()
        self._forzar = self.forzar_vur.get()  # se lee aquí: las variables de Tk no se tocan desde el hilo
        with self.lock:
            self.resultados = []
        self._reiniciar_vista()
        self._en_ejecucion(True)
        self.hilo = threading.Thread(target=self._flujo_consultas, daemon=True)
        self.hilo.start()

    def diagnosticar(self):
        if not self._validar_carga():
            return
        sugerida = next((m for m in self.df["Matricula_Limpia"] if separar_matricula(m)), "")
        texto = simpledialog.askstring("Diagnóstico VUR", "Matrícula a diagnosticar:",
                                       initialvalue=sugerida, parent=self.root)
        if not texto:
            return
        partes = separar_matricula(limpiar_matricula(texto))
        if not partes:
            messagebox.showwarning("VUR", f"Formato de matrícula inválido: {texto}")
            return
        self.detener.clear()
        self._en_ejecucion(True)
        self.hilo = threading.Thread(target=self._flujo_diagnostico, args=(partes,), daemon=True)
        self.hilo.start()

    def solicitar_detener(self):
        self.detener.set()
        self._log("Detención solicitada: se detendrá al terminar la matrícula actual.")

    def guardar(self):
        with self.lock:
            datos = list(self.resultados)
        if not datos:
            messagebox.showinfo("VUR", "No hay resultados para guardar.")
            return
        sugerido = f"{self.ruta_excel.stem}_resultados.xlsx" if self.ruta_excel else "resultados_vur.xlsx"
        ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=sugerido,
                                            filetypes=[("Excel", "*.xlsx")])
        if not ruta:
            return
        try:
            exportar_excel(datos, ruta)
            self._log(f"Resultados guardados en {ruta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar (¿está abierto en Excel?):\n{type(e).__name__}: {e}")

    def comparar(self):
        tipos = [("Excel", "*.xlsx")]
        ruta_ahora = filedialog.askopenfilename(title="1 de 2: resultado NUEVO", filetypes=tipos,
                                                initialdir=str(CARPETA_CONSULTAS), parent=self.root)
        if not ruta_ahora:
            return
        ruta_antes = filedialog.askopenfilename(title="2 de 2: resultado ANTERIOR (mismo lote)", filetypes=tipos,
                                                initialdir=str(Path(ruta_ahora).parent), parent=self.root)
        if not ruta_antes:
            return
        if Path(ruta_antes).resolve() == Path(ruta_ahora).resolve():
            messagebox.showwarning("Comparar", "Eligió el mismo archivo dos veces.", parent=self.root)
            return
        salida = Path(ruta_ahora).with_name(f"comparacion_{Path(ruta_ahora).stem}_vs_{Path(ruta_antes).stem}.xlsx")
        try:
            r = comparar_resultados(ruta_antes, ruta_ahora, salida)
        except ValueError as e:
            messagebox.showwarning("Comparar", str(e), parent=self.root)
            return
        except Exception as e:  # archivo abierto en Excel, dañado, etc.
            messagebox.showerror("Comparar", f"No se pudo comparar:\n{type(e).__name__}: {e}", parent=self.root)
            return
        self._log(f"Comparación guardada en {salida}")
        texto = (f"Matrículas en ambos: {r['comunes']}\nSin cambios: {r['iguales']}\n"
                 f"Con cambios: {r['con_cambios']} (cambió el estado en {r['cambio_estado']})\n"
                 f"Solo en el anterior: {r['solo_antes']} · solo en el nuevo: {r['solo_ahora']}\n\n"
                 f"Archivo: {salida.name}\n\n¿Abrirlo ahora?")
        if messagebox.askyesno("Comparación lista", texto, parent=self.root):
            try:
                import os
                os.startfile(salida)  # solo Windows
            except (AttributeError, OSError):
                pass

    def _al_cerrar(self):
        if self.hilo and self.hilo.is_alive():
            if not messagebox.askokcancel("VUR", "Hay un proceso en curso. ¿Cerrar de todos modos?\n"
                                                 "Se guardará el parcial con lo consultado hasta ahora."):
                return
            self.detener.set()
            self.hilo.join(timeout=5)
        if self.cliente:
            self.cliente.cerrar_navegador()  # evita dejar Chrome abierto al cerrar la ventana
        self.root.destroy()

    # --- Hilo de trabajo ----------------------------------------------------

    def _iniciar_sesion(self, cliente):
        if not cliente.vivo():
            self.log_hilo("Iniciando Chrome...")
            cliente.iniciar_navegador()
        else:
            cliente.notificar("esperando")
        cliente.abrir_inicio()
        return self._esperar_confirmacion(
            "Inicie sesión en VUR en la ventana de Chrome (incluido el captcha).\n\n"
            "Pulse Aceptar cuando haya ingresado."
        )

    def _guardar_parcial(self):
        with self.lock:
            datos = list(self.resultados)
        if not datos or not self.ruta_excel:
            return
        ruta = self.ruta_excel.with_name(f"{self.ruta_excel.stem}_resultados_parcial.xlsx")
        try:
            exportar_excel(datos, ruta)
            self.log_hilo(f"Guardado parcial: {ruta.name} ({len(datos)} registros)")
        except Exception as e:
            self.log_hilo(f"No se pudo guardar el parcial (¿está abierto en Excel?): {type(e).__name__}: {e}")

    def _flujo_consultas(self):
        cliente, sesion_vur = self._obtener_cliente()
        if sesion_vur:
            self.log_hilo("Se reutiliza la sesión de VUR abierta en Chrome.")
        self.n_historico = 0
        try:
            self.historico.cargar()  # puede haber consultas nuevas de otra ventana o funcionario
        except Exception as e:  # noqa: BLE001
            self.log_hilo(f"Histórico no disponible: {type(e).__name__}: {e}")
        self._previo = dict(self.historico.ultimo)  # foto del histórico antes de este proceso
        if self._forzar:
            self.log_hilo("Forzar consulta en VUR: no se usa el histórico en este proceso.")
        try:
            for i, fila in self.df.iterrows():
                if self.detener.is_set():
                    raise ProcesoDetenido()
                resultado = None if self._forzar else self._desde_historico(fila)
                if resultado is None:
                    if not sesion_vur:  # Chrome se abre solo si de verdad hay que consultar VUR
                        if not self._iniciar_sesion(cliente):
                            self.log_hilo("Proceso cancelado por el usuario.")
                            return
                        sesion_vur = True
                    resultado = self._consultar_y_registrar(cliente, fila)
                else:
                    self.n_historico += 1
                with self.lock:
                    self.resultados.append(resultado)
                self.cola.put(("resultado", resultado))
                self.cola.put(("progreso", (i + 1, cliente.consultas, cliente.consultas_cache, self.n_historico)))
                if (i + 1) % GUARDADO_PARCIAL_CADA == 0:
                    self._guardar_parcial()
            if sesion_vur:
                self._segunda_pasada(cliente)
            self.log_hilo(f"Consultas finalizadas: {self.n_historico} tomadas del histórico, {cliente.consultas} "
                          f"peticiones a VUR, {cliente.consultas_cache} respuestas reutilizadas.")
            with self.lock:
                por_revisar = sum(1 for x in self.resultados if x.get("Confiabilidad") == CONFIABILIDAD_REVISAR)
            del_hist = (f"{self.n_historico} matrícula(s) se tomaron del histórico sin consultar VUR "
                        "(columna Fuente).\n\n") if self.n_historico else ""
            self.cola.put(("info", f"Consultas finalizadas.\n\n{del_hist}{por_revisar} matrícula(s) requieren "
                                   "revisión del profesional (columna Confiabilidad = REVISAR).\n\n"
                                   "Use 'Guardar Excel' para guardar los resultados."))
        except ProcesoDetenido:
            self.log_hilo("Proceso detenido.")
        except LimiteDiario as e:
            self.log_hilo(f"Proceso detenido: {e}")
            self.cola.put(("info", f"{e}\nLos resultados obtenidos se guardaron en el archivo parcial."))
        except NavegadorCerrado as e:
            self.log_hilo(f"Proceso interrumpido: {e}")
        except Exception as e:
            self.log_hilo(f"Error general: {type(e).__name__}: {mensaje_corto(e)}")
        finally:
            self._guardar_parcial()
            # Chrome queda abierto para reutilizar la sesión en el siguiente proceso;
            # se cierra al cerrar la ventana del programa.
            if not cliente.vivo():
                cliente.cerrar_navegador()
            if self.historico.exportar_excel():
                self.log_hilo(f"Histórico actualizado: {self.historico.archivo_excel}")
            self.cola.put(("fin", None))

    def _desde_historico(self, fila):
        """Resultado tomado del histórico si es reutilizable; None si hay que consultar VUR."""
        previo = self._previo.get(fila["Matricula_Limpia"])
        reutilizable, explicacion = evaluar_historico(previo)
        if not reutilizable:
            if explicacion:
                self.log_hilo(f"{fila['Matricula_Limpia']}: {explicacion}")
            return None
        registro = {c: previo.get(c, "") for c in COLUMNAS_SALIDA}
        registro.update(Matricula_Original=fila["Matricula_Original"], Fuente=f"Histórico: {explicacion}")
        self.log_hilo(f"{fila['Matricula_Limpia']}: tomada del histórico ({explicacion}) -> {registro['Estado']}")
        return registro

    def _consultar_y_registrar(self, cliente, fila):
        """Consulta VUR, compara con la consulta anterior (si la hay) y guarda en el histórico."""
        nuevo = self._procesar_matricula(cliente, fila)
        nuevo.update(Fuente="VUR", Funcionario=self.usuario, Equipo=self.equipo)
        previo = self._previo.get(nuevo["Matricula_Limpia"])
        if previo and previo.get("Estado") in ESTADOS_VALIDOS and nuevo.get("Estado") in ESTADOS_VALIDOS \
                and previo["Estado"] != nuevo["Estado"]:
            fecha = _fecha_registro(previo)
            cambio = (f"Cambió respecto a la consulta del {fecha:%d/%m/%Y} " if fecha else
                      "Cambió respecto a la consulta anterior ") + f"({previo['Estado']} -> {nuevo['Estado']})"
            if previo.get("Version_Reglas") == VERSION_REGLAS:
                # Mismas reglas y otro resultado: VUR cambió (inscripción nueva). Lo debe ver el profesional.
                nuevo["Motivos_Revision"] = "; ".join(x for x in (nuevo.get("Motivos_Revision"), cambio) if x)
                nuevo["Confiabilidad"] = CONFIABILIDAD_REVISAR
            else:
                nuevo["Notas_Proceso"] = "; ".join(x for x in (nuevo.get("Notas_Proceso"),
                                                               cambio + " por cambio de reglas del programa") if x)
        try:
            self.historico.agregar(nuevo)
        except OSError as e:
            self.log_hilo(f"No se pudo guardar en el histórico: {type(e).__name__}: {e}")
        return nuevo

    def _segunda_pasada(self, cliente):
        """Vuelve a consultar las matrículas con resultado dudoso por fallas de VUR."""
        import time
        with self.lock:
            pendientes = [i for i, r in enumerate(self.resultados) if necesita_segunda_pasada(r)]
        if not pendientes:
            return
        self.log_hilo(f"Segunda pasada: {len(pendientes)} matrícula(s) con respuesta incompleta de VUR. "
                      f"Se reintentan en {ESPERA_SEGUNDA_PASADA} s.")
        self.cola.put(("estado", f"Segunda pasada en {ESPERA_SEGUNDA_PASADA} s: {len(pendientes)} matrícula(s)…"))
        for _ in range(ESPERA_SEGUNDA_PASADA):
            if self.detener.is_set():
                raise ProcesoDetenido()
            time.sleep(1)
        for n, i in enumerate(pendientes, 1):
            if self.detener.is_set():
                raise ProcesoDetenido()
            with self.lock:
                anterior = dict(self.resultados[i])
            self.cola.put(("estado", f"Segunda pasada: {n} de {len(pendientes)}"))
            nuevo = self._consultar_y_registrar(cliente, self.df.iloc[i])
            if nuevo["Estado"] == anterior["Estado"]:
                nuevo["Segunda_Pasada"] = f"Consultada de nuevo: se confirmó {nuevo['Estado']}"
            else:
                nuevo["Segunda_Pasada"] = f"Consultada de nuevo: cambió de {anterior['Estado']} a {nuevo['Estado']}"
            self.log_hilo(f"{nuevo['Matricula_Limpia']}: {nuevo['Segunda_Pasada']}")
            with self.lock:
                self.resultados[i] = nuevo
            self.cola.put(("actualizar", (i, anterior, nuevo)))
        self._guardar_parcial()

    def _procesar_matricula(self, cliente, fila):
        original, limpia = fila["Matricula_Original"], fila["Matricula_Limpia"]
        registro = {col: "" for col in COLUMNAS_SALIDA}
        registro.update(Matricula_Original=original, Matricula_Limpia=limpia, Intentos=0,
                        Version_Aplicativo=VERSION_APLICATIVO, Confiabilidad=CONFIABILIDAD_REVISAR)

        partes = separar_matricula(limpia)
        if not partes:
            registro.update(Estado="ERROR", Notas_Proceso="Formato de matrícula inválido",
                            Motivos_Revision="Formato de matrícula inválido",
                            Fecha_Consulta=f"{datetime.now():%Y-%m-%d %H:%M:%S}")
            self.log_hilo(f"{original}: formato inválido, se omite.")
            return registro
        circulo, numero = partes

        intentos, ultimo_error = 0, ""
        while intentos < MAX_INTENTOS:
            if self.detener.is_set():
                raise ProcesoDetenido()
            intentos += 1
            registro["Intentos"] = intentos
            self.log_hilo(f"{limpia}: intento {intentos}/{MAX_INTENTOS}")
            try:
                return self._consultar(cliente, circulo, numero, registro)
            except SesionExpirada:
                intentos -= 1  # la expiración de sesión no consume intentos
                cliente.abrir_inicio()
                cliente.notificar("vencida")
                if not self._esperar_confirmacion("La sesión de VUR expiró. Inicie sesión de nuevo y pulse Aceptar."):
                    raise ProcesoDetenido()
            except ERRORES_RECUPERABLES as e:
                ultimo_error = f"{type(e).__name__}: {mensaje_corto(e)}"
                self.log_hilo(f"{limpia}: {ultimo_error}")
                cliente.recuperar()

        registro.update(Estado="ERROR", Fecha_Consulta=f"{datetime.now():%Y-%m-%d %H:%M:%S}",
                        Motivos_Revision="La consulta a VUR falló; repetirla",
                        Notas_Proceso=f"Falló tras {MAX_INTENTOS} intentos. Último error: {ultimo_error}")
        return registro

    def _consultar(self, cliente, circulo, numero, registro):
        registro["Fecha_Consulta"] = f"{datetime.now():%Y-%m-%d %H:%M:%S}"
        registro["Version_Aplicativo"] = VERSION_APLICATIVO
        db = cliente.datos_basicos(circulo, numero)
        if not db:
            # VUR a veces responde vacío por una falla momentánea (024-0024613 respondió vacío
            # el 22/10/2026 y con datos el 23/10/2026). Se espera y se reintenta antes de concluir.
            import time
            for espera in (3, 10):
                self.log_hilo(f"{circulo}-{numero}: VUR respondió sin datos; reintento en {espera} s")
                time.sleep(espera)
                db = cliente.datos_basicos(circulo, numero)
                if db:
                    break
        if not db:
            registro.update(Estado="SIN DATOS EN VUR", Notas_Proceso="VUR no devolvió datos para la matrícula",
                            Confiabilidad=CONFIABILIDAD_REVISAR, Version_Aplicativo=VERSION_APLICATIVO,
                            Motivos_Revision="VUR no devolvió datos en 3 intentos: verificar la matrícula y "
                                             "consultarla manualmente (puede ser una falla temporal de VUR)")
            self.log_hilo(f"{circulo}-{numero}: sin datos en VUR")
            return dict(registro)

        campos = extraer_campos(db)
        municipio = campos["Municipio"]
        estado_mpio, etiqueta, info = clasificar_municipio(municipio)
        limite = fecha_limite_municipio(info)
        notas = [f"Hoja Proy x Mpio: {info['nota']}"] if info and info.get("nota") else []

        registro.update(campos)
        registro.update(
            Proyectos_Municipio=describir_proyectos(info),
            Fecha_Limite_Revision=f"{limite:%d/%m/%Y}" if limite else "",
        )

        if estado_mpio in ESTADOS_SIN_CONSULTA:
            registro.update(
                Clasificacion_Municipio=etiqueta,
                Estado="NO GRAVADA" if estado_mpio == "000" else etiqueta,
                Confiabilidad=CONFIABILIDAD_ALTA if estado_mpio == "000" else CONFIABILIDAD_REVISAR,
                Motivos_Revision="" if estado_mpio == "000" else f"Municipio {etiqueta} en la hoja Proy x Mpio",
                Notas_Proceso="; ".join(notas + [f"Municipio {etiqueta} en {FUENTE_DATOS}; no se consultó estado jurídico"]),
            )
            self.log_hilo(f"{circulo}-{numero}: {registro['Estado']} | {municipio}")
            return dict(registro)

        motivos = []
        if estado_mpio is None:
            motivos.append("Municipio no está en la hoja Proy x Mpio")
        if estado_mpio == "CONDONADA":
            motivos.append("Municipio con proyecto condonado")
        if estado_mpio is None:
            etiqueta = "SIN CLASIFICAR"
            notas.append(f"Municipio '{municipio or 'no leído'}' no está en la hoja Proy x Mpio; se consultó por precaución")

        r = revisar_cadena(cliente, circulo, numero, db, limite, self.log_hilo)
        gravados, textos, detalle = [], [], []
        for a in r["valorizacion"]:
            if a["a"]:
                gravados.extend(limpiar_persona(p) for p in a["a"])
            else:
                # La anotación no nombra al propietario ("A:"). 1) Se toma el dueño a esa fecha:
                # la última adquisición anterior en el mismo folio. 2) Si el folio no tiene
                # compra anterior, el dueño de su matriz cuando el folio se abrió.
                personas, nro_adq = propietarios_en_fecha(r["anotaciones_folio"], a["nro"])
                desde_matriz = [] if personas else propietario_desde_matrices(
                    cliente, r["matricula"], r.get("datos_folio"), a["fecha"])
                if personas:
                    gravados.extend(f"{limpiar_persona(p)} (propietario según anotación {nro_adq})" for p in personas)
                    motivos.append("Propietario gravado deducido de la adquisición anterior (VUR no lo indica)")
                elif desde_matriz:
                    gravados.extend(desde_matriz)
                    motivos.append("Propietario gravado deducido de la matriz del folio (VUR no lo indica y el "
                                   "folio no tiene compra anterior): verificar")
                else:
                    motivos.append("Propietario gravado no identificado")
                    notas.append(f"Anotación {a['nro']} de {r['matricula']}: VUR no indica el propietario; "
                                 "no hay adquisición anterior en ese folio ni en su matriz")
            marca = f" [cancelada por la anotación {a['cancelada_por']}]" if a["cancelada_por"] else ""
            if a.get("vinculo") == "texto":
                marca += " (cancelación deducida del texto: la anotación no cita el número; verificar)"
                motivos.append(f"Cancelación de la anotación {a['nro']} deducida del texto (sin número)")
            if tipo_valorizacion(a) == "dudosa":
                motivos.append(f"Anotación {a['nro']}: código de valorización no reconocido ({_codigo(a['especificacion']) or 'sin código'})")
            textos.append(f"{r['matricula']} - {a['texto']}{marca}")
            detalle.append(f"{r['matricula']} anotación {a['nro']} del {a['fecha']}: {a['especificacion']}{marca}")
        folio_canc = r["matricula_cancelacion"] or r["matricula"]
        for a in r["cancelaciones"]:
            textos.append(f"{folio_canc} - {a['texto']}")
            detalle.append(f"{folio_canc} anotación {a['nro']} del {a['fecha']}: {a['especificacion']}")
        if not registro.get("Propietarios_Actuales"):
            personas, nro_adq = propietarios_en_fecha(r["anotaciones_inicial"])
            if personas:
                registro["Propietarios_Actuales"] = "; ".join(limpiar_persona(p) for p in personas) + \
                    f" (según anotación {nro_adq} del estado jurídico; datos básicos no trae propietarios)"
                registro["Propietario"] = limpiar_persona(personas[0])
                registro["Total_Propietarios"] = len(personas)
                motivos.append("Propietarios actuales tomados del estado jurídico (datos básicos no los trae)")
        if r["estado"] == "NO GRAVADA":
            notas.append("NO ESTÁ GRAVADA EN VUR")
        elif r["estado"] == ESTADO_INCOMPLETO:
            notas.append("No se halló valorización, pero la revisión de matrices quedó incompleta: verificar manualmente")
            motivos.append("Revisión de matrices incompleta")
        if r["estado"] == "GRAVAMEN CANCELADO" and not r["valorizacion"]:
            motivos.append("Se halló la cancelación pero no el gravamen original")
        afect = []
        for folio, a in r["afectaciones"]:
            marca = f" [cancelada por la anotación {a['cancelada_por']}]" if a["cancelada_por"] else " [VIGENTE]"
            if a.get("vinculo") == "texto":
                marca += " (cancelación deducida del texto; verificar)"
            bloque = f"{folio} anotación {a['nro']} del {a['fecha']}: {a['especificacion']}{marca}\n{a['texto']}"
            canc = a.get("cancelacion")
            if canc:
                bloque += (f"\n--- Cancelación: {folio} anotación {canc['nro']} del {canc['fecha']}: "
                           f"{canc['especificacion']}\n{canc['texto']}")
            afect.append(bloque)
        if afect:
            # Afectación vigente, o cuya cancelación se dedujo del texto -> revisión del profesional.
            # Afectación cancelada con número -> no exige revisión, pero queda a la vista.
            dudosas = [a for _, a in r["afectaciones"] if not a["cancelada_por"] or a.get("vinculo") == "texto"]
            if dudosas:
                motivos.append("Tiene afectación por valorización VIGENTE (o su cancelación no es clara): no define "
                               "el paz y salvo, pero debe tenerse en cuenta")
            else:
                notas.append("Tuvo afectación por valorización ya cancelada (ver Afectaciones_Valorizacion); "
                             "no define el paz y salvo")
        if r["no_reconocidas"]:
            motivos.append("Anotaciones que no se pudieron leer: " + "; ".join(r["no_reconocidas"]))
        motivos = list(dict.fromkeys(motivos))

        registro.update(
            Clasificacion_Municipio=etiqueta,
            Estado=r["estado"],
            Matricula_Con_Anotacion=r["matricula"],
            Nivel_Matriz="" if r["nivel"] is None else r["nivel"],
            Detalle_Valorizacion=recortar(" | ".join(detalle)),
            Propietario_Gravado="; ".join(dict.fromkeys(gravados)),
            Matricula_Cancelacion=r["matricula_cancelacion"],
            Observaciones=recortar("\n\n".join(textos)),
            Afectaciones_Valorizacion=recortar("\n\n".join(afect)),
            Confiabilidad=CONFIABILIDAD_REVISAR if motivos else CONFIABILIDAD_ALTA,
            Motivos_Revision="; ".join(motivos),
            Cadena_Revisada=recortar(" -> ".join(r["cadena"])),
            Folios_Revisados=len(r["cadena"]),
            Notas_Proceso="; ".join(notas + r["notas"]),
        )
        self.log_hilo(f"{circulo}-{numero}: {r['estado']} | {municipio} | {len(r['cadena'])} folio(s) revisado(s)")
        return dict(registro)

    def _flujo_diagnostico(self, partes):
        """Guarda las respuestas JSON de VUR de una matrícula para revisión."""
        cliente, sesion_vur = self._obtener_cliente()
        try:
            if not sesion_vur and not self._iniciar_sesion(cliente):
                return
            etiqueta = f"{partes[0]}-{partes[1]}"
            ruta = self.ruta_excel.with_name(f"diagnostico_{etiqueta}.json")
            datos = {"datosBasicos": cliente.datos_basicos(*partes),
                     "estadoJuridico": cliente.estado_juridico(*partes)}
            with open(ruta, "w", encoding="utf-8") as f:
                json.dump(datos, f, ensure_ascii=False, indent=2)
            self.log_hilo(f"Diagnóstico guardado en {ruta}")
            self.cola.put(("info", f"Diagnóstico guardado en:\n{ruta}"))
        except Exception as e:
            self.log_hilo(f"Error en diagnóstico: {type(e).__name__}: {mensaje_corto(e)}")
        finally:
            if not cliente.vivo():
                cliente.cerrar_navegador()
            self.cola.put(("fin", None))


# ---------------------------------------------------------------------------
# Pruebas automáticas de las reglas (python vur_valorizacion.py --probar)
# ---------------------------------------------------------------------------
# Ejecútelas después de CUALQUIER cambio. Si alguna falla, no use el programa.
# Cada caso reproduce el formato real de las anotaciones de VUR con nombres y
# documentos ficticios. Cuando aparezca un error nuevo, agregue aquí un caso
# que lo reproduzca antes de corregirlo.


def ejecutar_autopruebas():
    import sys
    import unittest
    from datetime import date  # noqa: F401 - usada por los casos
    v = sys.modules[__name__]  # este mismo archivo
    P = "PERSONAS QUE INTERVIENEN EN EL ACTO (X-Titular de derecho real de dominio,I-Titular de dominio incompleto)"


    def an(nro, fecha, espec, lineas, cancela=None):
        """Arma una anotación con el mismo formato HTML que entrega VUR."""
        c = f"<br>    Se cancela anotación No: {cancela}" if cancela else ""
        return (f"ANOTACION: Nro {nro} Fecha: {fecha} Radicación: 1<br>Doc: OFICIO 1 VALOR ACTO: ${c}<br/>"
                f"ESPECIFICACION: {espec}<br/>{P}<br/>" + "<br/>".join(lineas))


    COMPRA = an(1, "01-01-1980", "101 COMPRAVENTA (MODO DE ADQUISICION)", ["DE: VENDEDOR FICTICIO", "A: TITULAR FICTICIO   CC 1  X"])
    HIPOTECA = an(2, "12-04-1996", "210 HIPOTECA ABIERTA CUANTIA INDETERMINADA NATURALEZA Y N: ESC. 446\r\nNATURALEZA JURIDICA DEL ACTO: GRAVAMEN (GRAVAMEN)",
                  ["DE: TITULAR FICTICIO", "A: BANCO FICTICIO"])
    VAL_430 = an(5, "13-05-1999", "430 GRAVAMEN POR CONTRIBUCION DE VALORIZACION (GRAVAMEN)", ["DE: DEPARTAMENTO DE ANTIOQUIA", "A: TITULAR FICTICIO"])
    CAN_5 = an(6, "16-06-2009", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA (CANCELACION)", ["A: TITULAR FICTICIO"], cancela=5)
    VAL_BELLO = an(5, "28-12-1999", "430 VALORIZACION OBRA 808 RESOLUCION 314 DE 01-12-99 (GRAVAMEN)", ["DE: IMPES DE BELLO"])
    CAN_BELLO = an(6, "10-11-2006", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA DE VALORIZACION OBRA 808. (CANCELACION)", ["DE: MUNICIPIO DE BELLO"])
    AFECT = an(8, "08-08-2002", "380 AFECTACION DE INENAJENABILIDAD POR VALORIZACION (MEDIDA CAUTELAR)", ["DE: PLANEACION Y VALORIZACION", "A: TITULAR FICTICIO"])
    CAN_AFECT = an(11, "24-08-2026", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA DE AFECTACION DE INENAJENABILIDAD POR VALORIZACION (CANCELACION)",
                   ["DE: ALCALDIA"], cancela=8)


    class API:
        """Sustituye a VUR: devuelve respuestas fijas por matrícula."""
        def __init__(self, db, ej):
            self.db, self.ej = db, ej

        def datos_basicos(self, c, n):
            return self.db.get(f"{c}-{n}")

        def estado_juridico(self, c, n):
            return self.ej.get(f"{c}-{n}")


    def folio(anotaciones):
        return v.analizar_folio([v.parsear_anotacion(x) for x in anotaciones])


    def consultar(db, ej, circulo, numero):
        app = v.App.__new__(v.App)
        app.cola = queue.Queue()
        reg = {c: "" for c in v.COLUMNAS_SALIDA}
        return app._consultar(API(db, ej), circulo, numero, reg)


    def db_basico(mpio, apertura="20/04/1981", matrices=(), propietarios=True):
        return {"circulo": "008", "numeroMatricula": "1", "municipio": {"nombre": mpio}, "fechaAperturaFolio": apertura,
                "matriculasMatriz": list(matrices),
                "propietarios": [{"nombre": "TITULAR FICTICIO", "numeroDocumento": "1", "titular": "X"}] if propietarios else []}


    class Clasificacion(unittest.TestCase):
        def test_hipoteca_no_es_valorizacion(self):
            self.assertEqual(folio([COMPRA, HIPOTECA])["estado"], "NO GRAVADA")

        def test_gravamen_vigente(self):
            self.assertEqual(folio([COMPRA, VAL_430])["estado"], "GRAVADA")

        def test_cancelacion_con_numero(self):
            r = folio([COMPRA, VAL_430, CAN_5])
            self.assertEqual(r["estado"], "GRAVAMEN CANCELADO")
            self.assertEqual(r["valorizacion"][0]["vinculo"], "numero")

        def test_cancelacion_sin_numero_bello(self):
            """Error del 23/10/2026: la cancelación sin número se tomaba como otro gravamen."""
            r = folio([COMPRA, VAL_BELLO, CAN_BELLO])
            self.assertEqual(r["estado"], "GRAVAMEN CANCELADO")
            self.assertEqual(r["valorizacion"][0]["vinculo"], "texto")

        def test_afectacion_no_define_estado(self):
            r = folio([COMPRA, AFECT, CAN_AFECT])
            self.assertEqual(r["estado"], "NO GRAVADA")
            self.assertEqual(len(r["afectaciones"]), 1)
            self.assertEqual(r["afectaciones"][0]["cancelada_por"], 11)

        def test_cancelacion_codigo_0723(self):
            """Caso real 024-12609 (23/10/2026): cancelación con código 0723, no 0842."""
            val = an(2, "23-07-2001", "0212 VALORIZACION (GRAVAMEN)", ["DE: MUNICIPIO"])
            can = an(3, "29-05-2003", "0723 CANCELACION VALORIZACION (CANCELACION)", ["DE: MUNICIPIO"], cancela=2)
            self.assertEqual(folio([COMPRA, val, can])["estado"], "GRAVAMEN CANCELADO")
            can_sin_marca = an(3, "29-05-2003", "0723 CANCELACION VALORIZACION", ["DE: MUNICIPIO"])
            self.assertEqual(folio([COMPRA, val, can_sin_marca])["estado"], "GRAVAMEN CANCELADO")

        def test_codigo_desconocido_es_dudosa(self):
            a = v.parsear_anotacion(an(7, "01-01-2010", "0999 OTRA ANOTACION DE VALORIZACION (OTRO)", ["DE: MUNICIPIO"]))
            self.assertEqual(v.tipo_valorizacion(a), "dudosa")


    class Resultados(unittest.TestCase):
        def test_confiabilidad_alta(self):
            r = consultar({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430, CAN_5]}}, "008", "1")
            self.assertEqual((r["Estado"], r["Confiabilidad"]), ("GRAVAMEN CANCELADO", "ALTA"))

        def test_afectacion_vigente_queda_para_revision(self):
            r = consultar({"008-1": db_basico("APARTADO")}, {"008-1": {"textoAnotaciones": [COMPRA, AFECT]}}, "008", "1")
            self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "REVISAR"))
            self.assertIn("afectación por valorización VIGENTE", r["Motivos_Revision"])
            self.assertIn("380 AFECTACION", r["Afectaciones_Valorizacion"])

        def test_bello_en_matriz_con_propietario_deducido(self):
            hijo = db_basico("BELLO", "19/12/2024", ["008-2"], propietarios=False)
            madre = db_basico("BELLO", "01/01/1990")
            r = consultar({"008-1": hijo, "008-2": madre},
                          {"008-1": {"textoAnotaciones": [an(2, "20-12-2024", "0125 COMPRAVENTA (MODO DE ADQUISICION)", ["A: COMPRADOR FICTICIO  CC 2  X"])]},
                           "008-2": {"textoAnotaciones": [COMPRA, VAL_BELLO, CAN_BELLO]}}, "008", "1")
            self.assertEqual(r["Estado"], "GRAVAMEN CANCELADO")
            self.assertEqual(r["Matricula_Con_Anotacion"], "008-2")
            self.assertEqual(r["Confiabilidad"], "REVISAR")
            self.assertIn("anotación 1", r["Propietario_Gravado"])
            self.assertIn("anotación 2", r["Propietarios_Actuales"])

        def test_municipio_000_no_consulta(self):
            r = consultar({"008-1": db_basico("GUARNE")}, {}, "008", "1")
            self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "ALTA"))

        def test_reintento_si_vur_responde_vacio(self):
            """Caso real 024-0024613: el 22/10 VUR respondió vacío y el 23/10 con datos."""
            import time
            class APIFalla(API):
                def __init__(self, *a):
                    super().__init__(*a)
                    self.llamadas = 0
                def datos_basicos(self, c, n):
                    self.llamadas += 1
                    return None if self.llamadas == 1 else super().datos_basicos(c, n)
            dormir, time.sleep = time.sleep, lambda s: None
            try:
                app = v.App.__new__(v.App)
                app.cola = queue.Queue()
                reg = {c: "" for c in v.COLUMNAS_SALIDA}
                api = APIFalla({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430]}})
                r = app._consultar(api, "008", "1", reg)
            finally:
                time.sleep = dormir
            self.assertEqual(r["Estado"], "GRAVADA")

        def test_sin_datos(self):
            import time
            dormir, time.sleep = time.sleep, lambda s: None
            try:
                r = consultar({}, {}, "008", "1")
            finally:
                time.sleep = dormir
            self.assertEqual((r["Estado"], r["Confiabilidad"]), ("SIN DATOS EN VUR", "REVISAR"))


    class SegundaPasada(unittest.TestCase):
        def test_seleccion(self):
            f = v.necesita_segunda_pasada
            self.assertTrue(f({"Estado": "SIN DATOS EN VUR"}))
            self.assertTrue(f({"Estado": "ERROR", "Notas_Proceso": "Falló tras 3 intentos"}))
            self.assertFalse(f({"Estado": "ERROR", "Notas_Proceso": "Formato de matrícula inválido"}))
            self.assertTrue(f({"Estado": v.ESTADO_INCOMPLETO,
                               "Notas_Proceso": "Matriz 017-7991: VUR no devolvió datos (dos intentos); no se pudo revisar"}))
            self.assertFalse(f({"Estado": v.ESTADO_INCOMPLETO, "Notas_Proceso": "Tope de seguridad de 300 folios alcanzado"}))
            self.assertFalse(f({"Estado": "NO GRAVADA", "Confiabilidad": "REVISAR",
                                "Motivos_Revision": "Tiene afectación por valorización"}))
            self.assertFalse(f({"Estado": "GRAVADA"}))

        def test_reconsulta_y_reemplaza(self):
            """Caso 024-0024613: primera pasada sin datos, segunda con valorización."""
            import threading
            import time
            import pandas as pd

            class APIRecupera(API):
                def __init__(self, *a):
                    super().__init__(*a)
                    self.activa = False
                def datos_basicos(self, c, n):
                    return super().datos_basicos(c, n) if self.activa else None

            api = APIRecupera({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430]}})
            app = v.App.__new__(v.App)
            app.cola, app.lock, app.detener, app.ruta_excel = queue.Queue(), threading.Lock(), threading.Event(), None
            app.df = pd.DataFrame([{"Matricula_Original": "MI 008-0000001", "Matricula_Limpia": "008-0000001"}])
            import tempfile
            temporal = tempfile.TemporaryDirectory()
            self.addCleanup(temporal.cleanup)
            app.historico, app._previo, app.usuario, app.equipo = v.Historico(temporal.name), {}, "prueba", "PC"
            dormir, espera, time.sleep = time.sleep, v.ESPERA_SEGUNDA_PASADA, lambda s: None
            v.ESPERA_SEGUNDA_PASADA = 0
            try:
                app.resultados = [app._procesar_matricula(api, app.df.iloc[0])]
                self.assertEqual(app.resultados[0]["Estado"], "SIN DATOS EN VUR")
                api.activa = True                      # VUR se recupera
                app._segunda_pasada(api)
            finally:
                time.sleep, v.ESPERA_SEGUNDA_PASADA = dormir, espera
            self.assertEqual(app.resultados[0]["Estado"], "GRAVADA")
            self.assertIn("cambió de SIN DATOS EN VUR a GRAVADA", app.resultados[0]["Segunda_Pasada"])


    class Matrices(unittest.TestCase):
        def test_ocho_niveles(self):
            db, ej = {}, {}
            for i in range(9):
                db[f"020-{100 + i}"] = {"fechaAperturaFolio": f"01/01/{2020 - i * 5}",
                                        "matriculasMatriz": [f"020-{101 + i}"] if i < 8 else []}
            ej["020-107"] = {"textoAnotaciones": [COMPRA, VAL_430]}
            r = v.revisar_cadena(API(db, ej), "020", "100", db["020-100"], date(1985, 6, 16))
            self.assertEqual((r["estado"], r["matricula"], r["nivel"]), ("GRAVADA", "020-107", 7))

        def test_corte_por_fecha(self):
            db = {"025-1": {"fechaAperturaFolio": "01/01/1970", "matriculasMatriz": ["025-9"]}}
            r = v.revisar_cadena(API(db, {}), "025", "1", db["025-1"], date(1973, 3, 8))
            self.assertEqual(len(r["cadena"]), 1)

        def test_matriz_sin_respuesta_es_incompleta(self):
            db = {"017-1": {"fechaAperturaFolio": "19/12/2023", "matriculasMatriz": ["017-2"]}}
            r = v.revisar_cadena(API(db, {}), "017", "1", db["017-1"], date(1988, 10, 19))
            self.assertEqual(r["estado"], v.ESTADO_INCOMPLETO)


    class Mejoras261026(unittest.TestCase):
        """Cambios del 26/10/2026 tras comparar el lote 'matriculas 2 monica' (23/10 vs 26/10)."""

        def test_anotacion_sin_especificacion_usa_naturaleza(self):
            antigua = ("ANOTACION: Nro 4 Fecha: 01-01-1985 Radicación: 1<br>Doc: RESOLUCION 1<br/>"
                       "NATURALEZA JURIDICA DEL ACTO: 430 GRAVAMEN POR CONTRIBUCION DE VALORIZACION<br/>DE: MUNICIPIO")
            self.assertEqual(folio([COMPRA, antigua])["estado"], "GRAVADA")

        def test_no_reconocida_dice_cual_y_si_menciona_valorizacion(self):
            rota = "Anotación ilegible de VALORIZACION sin número<br/>DE: MUNICIPIO"
            r = consultar({"017-1": db_basico("LA CEJA")}, {"017-1": {"textoAnotaciones": [COMPRA, rota, "  "]}},
                          "017", "1")
            self.assertEqual(r["Confiabilidad"], "REVISAR")
            self.assertIn("entrada 2 sin número de anotación y MENCIONA VALORIZACIÓN", r["Motivos_Revision"])
            self.assertNotIn("entrada 3", r["Motivos_Revision"])  # la entrada vacía se ignora
            self.assertIn("Anotación ilegible", r["Notas_Proceso"])

        def test_afectacion_cancelada_trae_texto_de_la_cancelacion(self):
            r = consultar({"008-1": db_basico("APARTADO")}, {"008-1": {"textoAnotaciones": [COMPRA, AFECT, CAN_AFECT]}},
                          "008", "1")
            self.assertIn("--- Cancelación: 008-1 anotación 11", r["Afectaciones_Valorizacion"])
            # Decisión del usuario (26/10/2026): afectación cancelada con número -> ALTA, pero visible.
            self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "ALTA"))
            self.assertIn("afectación por valorización ya cancelada", r["Notas_Proceso"])

        def test_propietario_gravado_desde_la_matriz(self):
            """Caso Bello 01N-5052127: el gravamen no nombra dueño y el folio no tiene compra anterior."""
            hijo = db_basico("BELLO", "19/12/2024", ["008-2"])
            con_gravamen = db_basico("BELLO", "24/06/1992", ["008-3"])
            abuela = db_basico("BELLO", "01/01/1970")
            venta_posterior = an(9, "01-01-1995", "0125 COMPRAVENTA (MODO DE ADQUISICION)", ["A: OTRO DUEÑO FICTICIO"])
            r = consultar({"008-1": hijo, "008-2": con_gravamen, "008-3": abuela},
                          {"008-1": {"textoAnotaciones": []},
                           "008-2": {"textoAnotaciones": [VAL_BELLO, CAN_BELLO]},
                           "008-3": {"textoAnotaciones": [COMPRA, venta_posterior]}}, "008", "1")
            self.assertEqual(r["Matricula_Con_Anotacion"], "008-2")
            self.assertIn("TITULAR FICTICIO", r["Propietario_Gravado"])
            self.assertIn("(dueño de la matriz 008-3 según su anotación 1, deducido)", r["Propietario_Gravado"])
            self.assertNotIn("OTRO DUEÑO", r["Propietario_Gravado"])  # venta posterior a la apertura del folio
            self.assertEqual(r["Confiabilidad"], "REVISAR")
            self.assertIn("deducido de la matriz", r["Motivos_Revision"])

        def test_excel_con_hoja_resumen(self):
            import tempfile
            import openpyxl
            regs = [dict(Matricula_Limpia="008-1", Estado="NO GRAVADA", Confiabilidad="ALTA"),
                    dict(Matricula_Limpia="008-2", Estado="GRAVADA", Confiabilidad="REVISAR", Motivos_Revision="=prueba")]
            with tempfile.TemporaryDirectory() as d:
                ruta = Path(d) / "r.xlsx"
                v.exportar_excel(regs, ruta)
                wb = openpyxl.load_workbook(ruta)
                self.assertEqual(wb.sheetnames, ["Resultados", "Resumen"])
                self.assertEqual(wb.active.title, "Resumen")
                valores = [c for fila in wb["Resumen"].iter_rows(values_only=True) for c in fila if c is not None]
                self.assertIn("Requieren revisión del profesional", valores)
                self.assertIn("=prueba", valores)  # texto, no fórmula


    class Comparacion(unittest.TestCase):
        def _excel(self, d, nombre, regs):
            ruta = Path(d) / nombre
            v.exportar_excel(regs, ruta)
            return ruta

        def test_detecta_cambios_y_causa(self):
            import tempfile
            import openpyxl
            base = dict(Matricula_Limpia="01N-1", Municipio="BELLO", Confiabilidad="ALTA",
                        Cadena_Revisada="01N-1 (nivel 0, apertura ?, 5 anotaciones: GRAVADA)", Version_Aplicativo="A")
            with tempfile.TemporaryDirectory() as d:
                antes = self._excel(d, "a.xlsx", [dict(base, Estado="GRAVADA"), dict(base, Matricula_Limpia="01N-2", Estado="NO GRAVADA")])
                ahora = self._excel(d, "b.xlsx", [dict(base, Estado="GRAVAMEN CANCELADO",
                                                       Cadena_Revisada="01N-1 (nivel 0, apertura ?, 6 anotaciones: X)"),
                                                  dict(base, Matricula_Limpia="01N-2", Estado="NO GRAVADA"),
                                                  dict(base, Matricula_Limpia="01N-3", Estado="NO GRAVADA")])
                r = v.comparar_resultados(antes, ahora, Path(d) / "c.xlsx")
                self.assertEqual((r["comunes"], r["iguales"], r["cambio_estado"], r["solo_ahora"]), (2, 1, 1, 1))
                fila = list(openpyxl.load_workbook(Path(d) / "c.xlsx")["Cambios"].iter_rows(min_row=2, values_only=True))[0]
                self.assertEqual(fila[0], "01N-1")
                self.assertIn("inscripciones nuevas", fila[7])

        def test_lotes_distintos(self):
            import tempfile
            with tempfile.TemporaryDirectory() as d:
                a = self._excel(d, "a.xlsx", [dict(Matricula_Limpia="001-1", Estado="NO GRAVADA")])
                b = self._excel(d, "b.xlsx", [dict(Matricula_Limpia="002-1", Estado="NO GRAVADA")])
                with self.assertRaises(ValueError):
                    v.comparar_resultados(a, b, Path(d) / "c.xlsx")


    class HistoricoConsultas(unittest.TestCase):
        def setUp(self):
            import tempfile
            self.tmp = tempfile.TemporaryDirectory()
            self.addCleanup(self.tmp.cleanup)

        def reg(self, m="008-1", estado="NO GRAVADA", conf="ALTA", dias=1, **extra):
            from datetime import datetime, timedelta
            r = {"Matricula_Limpia": m, "Estado": estado, "Confiabilidad": conf,
                 "Fecha_Consulta": f"{datetime.now() - timedelta(days=dias):%Y-%m-%d %H:%M:%S}", "Funcionario": "jperez",
                 "Version_Reglas": v.VERSION_REGLAS}
            r.update(extra)
            return r

        def test_reglas_de_reutilizacion(self):
            e = v.evaluar_historico
            self.assertTrue(e(self.reg())[0])
            self.assertTrue(e(self.reg(estado="GRAVADA"))[0])
            self.assertFalse(e(self.reg(dias=v.VIGENCIA_HISTORICO_DIAS + 1))[0])       # vencida
            self.assertFalse(e(self.reg(conf="REVISAR"))[0])                           # requería revisión
            self.assertFalse(e(self.reg(estado="SIN DATOS EN VUR"))[0])
            self.assertFalse(e(self.reg(estado=v.ESTADO_INCOMPLETO))[0])
            self.assertFalse(e(self.reg(Version_Reglas="2026.01.01-1"))[0])            # reglas anteriores
            self.assertIn("jperez", e(self.reg())[1])

        def test_guardar_y_recargar_con_varios_usuarios(self):
            a = v.Historico(self.tmp.name)
            a.agregar(self.reg("008-1", dias=3))
            b = v.Historico(self.tmp.name)        # otra ventana / otro funcionario
            b._archivo_sesion = Path(self.tmp.name) / "registros" / "otro.jsonl"
            b.agregar(self.reg("008-1", estado="GRAVADA", dias=0))
            b.agregar(self.reg("008-2", estado="ERROR", conf="REVISAR", dias=0))
            (Path(self.tmp.name) / "registros" / "danado.jsonl").write_text('{"roto": \n', encoding="utf-8")
            h = v.Historico(self.tmp.name)
            self.assertEqual(h.cargar(), 2)
            self.assertEqual(h.buscar("008-1")["Estado"], "GRAVADA")            # el más reciente
            self.assertTrue(h.exportar_excel())
            self.assertTrue(h.archivo_excel.exists())

        def test_un_error_no_tapa_un_resultado_valido(self):
            h = v.Historico(self.tmp.name)
            h.agregar(self.reg("008-1", estado="NO GRAVADA", dias=2))
            h.agregar(self.reg("008-1", estado="ERROR", conf="REVISAR", dias=0))
            self.assertEqual(h.buscar("008-1")["Estado"], "NO GRAVADA")


    class CarpetaTrabajo(unittest.TestCase):
        def test_crea_subcarpetas_y_guarda_eleccion(self):
            import tempfile
            with tempfile.TemporaryDirectory() as d:
                config_original = v.ARCHIVO_CONFIG
                consultas, historico = v.CARPETA_CONSULTAS, v.CARPETA_HISTORICO
                try:
                    v.ARCHIVO_CONFIG = Path(d) / "appdata" / "VUR_Valorizacion" / "config.json"
                    destino = Path(d) / "Mi carpeta VUR"
                    v.aplicar_carpeta_trabajo(destino, guardar=True)
                    for sub in v.SUBCARPETAS_TRABAJO:
                        self.assertTrue((destino / sub).is_dir())
                    self.assertEqual(v.leer_carpeta_trabajo(), destino)
                    self.assertEqual(v.CARPETA_CONSULTAS, destino / "consultas")
                finally:
                    v.ARCHIVO_CONFIG, v.CARPETA_CONSULTAS, v.CARPETA_HISTORICO = config_original, consultas, historico


    class ExcelConcluyente(unittest.TestCase):
        """Pedido de los profesionales (27/10/2026): las primeras columnas deben ser concluyentes."""

        def test_conclusion(self):
            self.assertEqual(v.conclusion({"Estado": "NO GRAVADA", "Confiabilidad": "ALTA"}), "NO TIENE VALORIZACIÓN")
            t = v.conclusion({"Estado": "GRAVAMEN CANCELADO", "Matricula_Con_Anotacion": "01N-5052127",
                              "Nivel_Matriz": 3, "Confiabilidad": "REVISAR"})
            self.assertEqual(t, "TUVO VALORIZACIÓN, YA CANCELADA · anotación en la matriz 01N-5052127 (nivel 3)"
                                " · REQUIERE REVISIÓN MANUAL")
            self.assertIn("misma matrícula", v.conclusion({"Estado": "GRAVADA", "Matricula_Con_Anotacion": "008-1",
                                                            "Nivel_Matriz": 0}))

        def test_orden_y_columnas_agrupadas(self):
            import tempfile
            import openpyxl
            with tempfile.TemporaryDirectory() as d:
                ruta = Path(d) / "r.xlsx"
                v.exportar_excel([dict(Matricula_Original="MI 008-1", Matricula_Limpia="008-1", Estado="GRAVADA",
                                       Confiabilidad="ALTA")], ruta)
                wb = openpyxl.load_workbook(ruta)
                ws = wb["Resultados"]
                encabezado = [c.value for c in ws[1]]
                self.assertEqual(encabezado[:3], ["Matricula_Original", "Resultado", "Confiabilidad"])
                self.assertEqual(ws["B2"].value, "TIENE VALORIZACIÓN VIGENTE")
                self.assertEqual(set(encabezado) - {"Resultado"}, set(v.COLUMNAS_SALIDA))  # no se pierde ninguna
                oculta = openpyxl.utils.get_column_letter(len(v.COLUMNAS_VISIBLES) + 1)
                self.assertTrue(any(dim.hidden and dim.min <= openpyxl.utils.column_index_from_string(oculta) <= dim.max
                                    for dim in ws.column_dimensions.values()))
                wb.close()

    casos = [c for c in list(locals().values())
             if isinstance(c, type) and issubclass(c, unittest.TestCase)]
    cargador = unittest.TestLoader()
    suite = unittest.TestSuite(cargador.loadTestsFromTestCase(c) for c in casos)
    resultado = unittest.TextTestRunner(verbosity=2).run(suite)
    print("\n" + ("TODAS LAS PRUEBAS PASARON" if resultado.wasSuccessful()
                  else "HAY PRUEBAS FALLIDAS: no use el programa hasta corregirlo"))
    return resultado.wasSuccessful()


# ---------------------------------------------------------------------------
# Recursos incorporados (base64): logo oficial, icono y tipografía Poppins
# ---------------------------------------------------------------------------
# Poppins: Indian Type Foundry, licencia SIL Open Font License 1.1 (texto completo
# en OFL-Poppins.txt, que se recrea junto a las fuentes).

RECURSOS = {
    "logo_gobernacion_app.png": """
iVBORw0KGgoAAAANSUhEUgAAAJ0AAABGCAYAAADW35Y6AAAtIUlEQVR42u19eZhcR3Xv75yqe3uZnn3R7D1Lz6LR7jGObWxkG7yw
2WZRbJL4BcIWEuCFhwPEEMYKEGMgJsQQB8IXAmQxFg48drDBVmy8Im+ytbZGmtGMpNFo9t5u36o674+ekcbCBhO+F95jdL6ve6bv
vV3rr06d86tT1YQVKJu3bEms2XB+ykecbdQKQh/wAKAIhADgLz5ZXPxb+uwsU+lzuPhehHKOAA/MVhACrJx43lJOHoJCgRD3YAMl
XukSnGUVLswe/8LNHxhdie2vV1Jlh4aGeOvWrS7Zvv6s8vZN/+rFK32ICEgAAQCCnHzH4v84+b70Z0kEAgjo5BsERApCCiIMIocy
OBHnln2TLLHE548euh3AO0FkIUKL2Z0B3W+rJBIVcY5VNHrlNVpsCCgCHIGIlqAEEEBCi+BygNAy7BFAyzAiDiCGgw9TyEAXjyEi
8wgpgWKkDX4kAjgDMCDOiReNU2ZmsnIlAW3Fg86LxYTgjHVFVgKwYwiA0ARQrEAgiAissxAIPBUFASBC6ZoTgAAHAUPgWMMU8qjK
3oe1sYfRUTWKmM4jb6J4erYPTxWuBidawKGBMFsXBhrOGaxQWZGgc65IABE50EIwT7lwFjG/HC5bwL4jOzCQugC+F8XPdn8fAy3n
wkUIzxx9AOWqGqubzoMfiZamXhE49mEyR7DGfAMXtz+K6oQBsUJECZwEqC1/AGp0Ho/n3w4Vr4W4UIiYCERnQLeCJAiKrlyxzM0d
x4PP3CnHaJwSVTUoN+UYG9+FjJ5DVaweew89CIsAB+weOGtRW6xHa3kPGuJJmLAIqAh05iAu8r+CC7r2wYv6+NedhCfHixi6OIKp
PPCZn+bxznOewdzED7G/+Ab4WiAicIIVK7xSK04iiMbjsL6FQREnMuN4JvMo3CoP+zJP4qGx78OWA7vzj4OVRkW8FvmyAnZPPYJs
dhbKiyFcOIa2/FfRVLUXRYoCRLioneEXGY8fdnh8vIinRgqor/BxQd1PEcs+A6XjJUeF4c5ouhUkYuAgkD0HH8CxqWGYsAAvGsPq
5DkIEeLIbBrzmIXn+4ipClgbYtaegCOHfOZxzKWncGH7K1E9911sP/E4vj0agc9ZXJJUeOugj6s3RPH0UYsHJx16k2VIlCn06Hm0
H/02HjpeidbGdlDJX1lZbutKBl3OWFMs5iQS+miNdWMmMYWyyho4FmjRaCnrRrBQRLYwjTK/CtXxRoRhHpncDJJV/ciSYOboNvzx
+l3gdZUoGsLonMVXfpbDkakQ16wtx/gs44Exi3e+OIYnjlp87nGHA9M7kQ0+ioPHX4qWSMoBEDmj6VaGZLI5G41UWFfr0awUJFk3
iILJwycfPkdwsLgLlC8gCh8WBuMLB1CmE4hwFJMnjuLs5kG8MTWM+cICvv44cCIXorc+gt9fU4Z79hfx3d1FHJ1nnJi12Hnc4va0
RVhkpKoJydgMEtE9mA1i7iQZeMam++0Xn/fI/hNfl0eOfQ21XoCZuTSK+SdQrUIcmtuF3OHd2Lh3BOVhCB8hYKYwmTsMqwNkaBwL
c/+C5tgYtu0somgd1tR5GD5hcOuDsxASHJ8BsgFQUeXhK3tDjM9YVKCITD7EAzMax80o1pXdwwBAJGc03W+3bAUAzBd/RsH8Xm6P
MGYmj8KVM7Tv4ccjw8gjgW4uoKOYw3GXwzwRKrVDlQIW7AG0x4G+aIBIpBx/8dJKnAgI/3nEocFzyHs+vjOcw6bqEG1VHl7eEsE3
hwM0RgUXtPoYaFBorSBYFhw7ukAEhog7A7qVIPmCxqq4lqlsCOszGuIe9s0TquIJmGIZairzUO1xRFuimDeArwhVUYcODzgrU8Ql
T+RQviqK4ZoIPvJwDkSMpipGZaVCZV0Uo0GIyakQqTIPUR8oMOPxWYftEwZiDN79ogRSNZ4RODBhxdEnKxJ0sbxBc2jE04L6KsFC
LoP+RBRHwhgSfoCmYgE2olFeRpicEzSXAVaAjpjFFTaC+P1z+OHDB/HkJdW45vx6vHatByUWIRRu3004mlP4/mGHhjJg6GyFxnIP
PZUMAeHxCYOnJoooeKVFtZXIm6wo0G3FEICtmM02y+HxfrAXQYWaRYIz0JxHJi/YVJFF20IBhahDnW+wXzSiniACiw0JQndtG3af
pbAvO4UHHhYc3J/G5Vv7UFHOGE7PYWHfFN5xaTtGAkZrBfCGDRGgaAHnAAHaUoxX98WwY18JdLwCgbcyebq6BZqrmaKjuXWYmt+E
CDl4CxkI5lBWcKhZeBIVeho1s3E0uVlUh0XUxwUX1pQhny3D2reWo2GuHfEHovirL9yBb987jeu2NGLkaBbHx6dxfLYeFS6Ccge4
osCGCooFAoILHTwCogorVlaU9zq06EhEOE9NlWla234XKsuOYMHFMCeVmJMkHgv6sHu+AffMnY3vTlyF7PwlWBg/B5eXV6LKeJie
O4Z4Iof6qiJSqypRqerw5W+MIjsbIhYlRDyCsw79oUGNEzARFJUiUxiAIoJA4FbugsQKXQYzQN5qNNgCkjQPRwIFQAkQQQAVZpFT
MRSoEqP51VDl9TirnUG8CiItGDvciPTBJHbfq3BZy4Uoi3Thkafm0NWSQGdnNepr4qB8CBNYPNe6PhHAxCuVpluhC/4MaTRW1gRx
PBlULc65BIDAzsCzBZBSCELGhsqDuPG8HahNFJHhSXiexg/v8/Dd7xbgpxmbX9qDvsu7EYnci4YGQl1XM/yoBpRCIZTnXnMQgAV0
RtOtIOnLWrkoa2Rsvht7wib4VJrqHDGUCxGFwbyLoSt+DB9/8c+wpmEBxijEI0WsasjjyFQGw6MzgBIUMY76mv1Y1xvB9FQIlQ0R
YUJggcCcikCW5eAjQGjlxpmsMNANAQBaOaTduRR/I/siWGjwYrQww0GRwJcQGyrH8OkL78eFyeMwpmSXOcsAEa6+grAqpWFrLdpX
x5CbywESoLhg0CICrYBs3oKEsRhnDKJSYOhiiDvsIghXorpbYZRJSX6YWY/5hQswo6uhycIBUHAIJYIgtKg0U3jX6km8uDMOUyAo
RRAnUFrgLDA3WofXXHgxOjvHEMvn8J93pFD5R0+hsw9whhEWHZyR57DlaIUGqJ+x6TBqqsG2zCktEAGYgYLVqHJH8Ye9T6DJTWNm
jGBNHEQM5wRKCeZmPdx9Rz2OjLYg1T+NjgaF4acY+58uYG7aA3tFwBIcMxQrLBTk51wFXlRvi6FNZ6bXlVXxUq8TA0EIrPV247YL
78aNl42ip8fDgV1ZHD0SQHmlCTC0Cnf/WyPuub0F41ML0Ik8lCpDc7IB8TKFmUmFpQg5gkATwdolB+WULbfswxnQrTQREIgEoWWs
j6fxuUsfwhVr5yEgJHsTsEWHR386jTAUKN9i/5MxPP6TRgQ6i5bOKLrXVQAcYn5hHzxl4PKJkqdKJY81CCxmM+aUu7q4SdGd9CRW
rqZbkdMrwrAEOyKYYhHnNuzB6voQxkSg2SHZHsWl1zQhGuESUBxh/2M1KMzGkUppXH5lAwSzGD+4E9PHpxHa7pMaTQQgJlQnvZKW
FDm5tXE5ZeKcPQO6lSTWFoWdAwnASmNiPo+du4cRrahHRZmPyrIo1m2qAKyDtUBYJJyYdsgUs6jI+ghygFMTyMyPwORqEeQ9cCwA
qES+MQn+x4VlIADiHEQEzLyMNhHkgtAAgHWLmDwDut9uMcX1zrPGAgTtKezIb8LTJ3Zik9qLqVwEh8cSqK9vQndbHCICzYKXbTkG
qvQwvDOHmckcmrtq4ekESBmcdfFRpNblAMuLphrBE7dkNZ6iS6TE2EGAmXlbXDb1riitt7JsuhtvFADwqoMZgswIAQpGjlIfbj7w
Jnxx70uwe6oKWs3g+PFDODCWBysFZkJVvY+XXN6ES19jULQ7YYwBKR+xqnmsPv8o6hrzEHvqFIBFeJ1OmoAZVCwK4Hj6zPS6EmRx
Hpved/t4Xf1bh5m4SwBRbHFUUvjyZBLfOjGF/she/EHH3VAqjWyuCYlYFeASiEbGkVqdgQkLyOSzaOwoR13tAiK+j6hf0mRE/Gxe
Dli8TpDS+RNqPg+TCezBxZGw4jaErUDvVQjp9LxPxafZGiwdl+ORRcxXyOgm3F+8BH+z73UYz8ShzTAmJvdgLrcLlVUH0dJWQHl1
BVgB1dWM2lqNRIJATHi+9YVnGW1MyOZxbPyY3lXC3MrrgZUX1bVll8KuXRLvXluJSNWV8CKaSj4nlc4msfCVw5Q04/HJJrDJo6V8
Fr5MYnIaiEbLARtFRfkUKhMzpQOXXInwWw6u0+Eni/tcWTEPj+M/P3nzodvGF+bNjfduP8PT/dbLtjscABQmhu9CIfsYWJEsWl+l
VVKGE8BnizFeg8+MvQl/9rO34lsjm2DMcexLDyMTTCISzYK1BrEGaQZzSdsxE5hKa6zLF/lFAMWg6Tnrho+G33hobCzvPjzEWIEL
YyszvGZoiLF1q2t8/Z//iTT2fE50zAGWSpYXnaQ1CAIhRigaCBZwdd2P8ZY12xG3WRhJIBZLIBrx4Uc8EJUOTlSKURaPIBHj0olO
DiXvVJzjqFKP73E7bvj8+Mt/8IP0pJSszBUHupUZNL39XgA3kqr5q7Qqa3sRlVd1iRNLpUPoTp2AuBgOp8iCvSj2FHqwP9cGrRWU
KiC0WYTFeWRzGcxm5rGQmUMuO4m56Rks5BmkI4hGPBGCUx7R0Ukb3PtQ4QOf/uLOh2UITNtX5qrEig0kXNJ2La9870W2Nfk1V1bT
INY4IuKShfdzVhmIgKL1oClETWQeNf4CylUOER0AKEUfN0RmsalyN/pjw0iQRVmiXtqay20kEtXbd+S/eOmbH/4TEZglx/YM6FYe
8hjY6pqv+eA7bG3r37louYIzruRjEi2ZZEvkLqhkBAsAAw3rAFm84hYbkxwQ11lsqhmWl9U9Kuuiu6S9FurQ8ar7br+bfv/Ldzxz
2LlFBmWFilrZoNsuEKGFLS99rKL7d4rw/ZfAj3giYktYo9LbIvAWwzFBKMXfabLw2MEjhwhbeGzhKysWHo1kVsmDM+vlQNCtTozN
Pnho+963fPYb48NDQ+CLL17ZUXUrXNOdbAMBgJYtH3qPrW74S1deWw1rXenoQiJCifGVk4v3Sx7p8sV8Kk2X4gQkwuwpMUUUT0z8
aObA9Dvx8Cf3D2GIt2KrW+kNrs5g7iTwaGHXfz5Ynux9gpi7of126AiTiAg5BxLBScwtf4cIRIjECUPAShFrlsz0jD2e/nTN3m+/
b2bHfxweGhrirdvPAO6Mpju9LUoehKy69DUNqFv3h5So+n3xytYhlmApnXINiBMBZMnOIyJaJOlAJoAU5qcll/lBePzgF2Z/8I/b
T9emZ+QM6J7XqwWA5guvarONay+jsvILDOtNpP0W5UerHEifdD6dzdl8flLB7HPF/AMyNXbPiR/8430AHEQWaT86A7gzoHsB7TI0
REvgA8DVl13VosL6hnhNXVPAFAcADeu0zU9N5aYnvMkHx2d2DM+dbNahD/Oy75+RM/IraL2hIX7BY1OESs+fGcxnNN2v3h7ynNeH
hp67vbbeKCe91zNyRs4MwP/PKIRf0BG/6P4Lvbf8GfoF36NfEyC/rC4n0+vsbFnf1dZ29v+lvOiX1POX1ZFeYL6/Tt/hVyjf6WV6
weXl57kmy16nJ/DL7p9+j5/n3vJnTr9Gp011P78Q+sIBsDzd56qvAoBUqrHeOX5X4Fz54nP8K+T169T9F6X1QussLyD952qP5xr4
z3f/l5XxBZf3ucLV3ebNA4mJiXyzUpx/5pkDh09LwA12dVVmRVaJ72f27t17ZOn+wMCAn0gkokEQsIiETz31VBaLWz0HAF+vWuUF
VVXcqJRMApifn7djY2OFc1tbo/MVFQqYBFBf3LVr19IPrXIqlfLS6XSwrEE4mUz6IyMjwbJyLZHc9vRK9/X1NRNRYhXzke27dmUA
0MDAgBeLxWTHjh3h0ndEys7z2O2qNLlHxk8djkkDAwPeYnnotPzcae2iNg8MxGY9T4dhaHaV8jq5JJtMJv3KykqVCAJGLWCOV9Ij
6UfyWPrx2NMG/pYtW2jbtm0n6zMA+LtKP0B7shyDg4NeIrFDZmbWR4IgYKWUAECiWGQu1IUPjYUG2IHFPAQA1q5duyqTyVTV10dP
PPronqnT26u/v79WpFAn4k/t27fvxDJALtX5pEeeSqUimzZtMgDwzDPPqGX9hs2bN+t7773X0nMcH/9z6jGZTL4GcG9joogA2jn3
iDHur48cOTJV6vDWa4XwFiXMQqSZ1SPMfFM6nZ5sa2v7MyJ3tSI16+CYhLKa1T+kDx3a3t7a+ufE9CqlKAeAxCHmQA97nvd5Uww+
CqZVBBTEWSJSTyYqKm8OgqAqCHKfFHH/MDp65G4AqCsv741UlL/LGHvLxMTEQQC6ra3lY8xMIyOH/2IJRK2tAzVaZ66H0EuUUgJn
C468zxw8mP5OT0/3u0wQysHR0c+mUo31Yei/l0AXKqKChSMW+urwyMiX29vbK621f6MUPTA6OvZFANTU1BTTmv/CWrnzyJEjTyx2
hu3u7rjGGvtWEPIljplno5q/sid98K6enp6uMCzeDEgVU+mIKHGIksFnhsdH/gOnnQLb3t76IYCb6+vrr9+xY0euoqKipqqq8lMi
7uHDh8c/vwSCZGfy/b72x0ncZcaZZmetAbEwc0SxfjQ0Zp9SUnfgwMjHUqlURbFY+J9EdDkzWSKQCN1+MFb2RezaFZ7b2ho9SvR2
R7haK4Zzohj0TeV5n02n09Tc3PjhYtF8/8SJE/ctAa5YyP81WN2pta60YXhJx+Guv9iO7Wbz5s1Vhw4N3worD46Mjf39aQP2pPpn
AJJMJq8WMVuNcXcWQ/u2vAneS4INUd+7GQB1dra/xglutE5uD0LzNsDeYE3xbGfMp7dsuUMRST+Ao1bCj4Sh+zggswb2rzckN1SB
qB9ER33l30LMf8+aPq21/loY5sodZKNS3jc0qVuY9Fcd3Evz+ezbnXMx5+TlROpjGzdurAdAyvOqROQcESkHgK6urtVEeJ2IXJtM
JtctaVylMjc5584D2Q86kTeHzv5IJLxldWr1Js3cJSR9AGDDyMdE5Dxj7Q2O6B3WypeN2A8lk63XAAATXiwiH0smkxsBiOd5EQJd
ICINywerczIAUjBGPiJCH2fIWL5ob+vp6fkd51xMRDZoon9nxZ8Cqc8o5k9ZjceWaRkGgO7u7jYRvAYkW2ZnZzcCoPJyHQdwPoCP
p1KpTQAklUopWHsOAM9j9Tnf9z/DpOqV4gWt+SYVk38VkVaABwmAtfZDAK6CCT9GZN9iDG61xr27PZP5YwAyoeg9AvwhBLeEoX2L
teHfOLFvCoL8ewBAEZ/vedy6BJzKbFY5yDkuDBuIXBuYzh1PjSsAOHTo0MXO2auEcW0qlao43bTRi43mNm/erEcOHXwTK/0fo6Mj
/7j0QHd36x8T8as3btxYNz01+VaA/2388NgXFm/v62hpeZeBuf2JJz54NsBZRTRycGRkBwB0t3fDwnw6LJuNco4KzuGJPen0XctV
bUdH/QZQZNrzqr6+Z8+OowDQ3tGaEut6WVufCD8TQWx2dvo9AG7wy8ocnCsYlHbIi9grmdV2ImgiugrAEwsLC+uJ3PlK0ZsPHhx7
ZDGrTyaTSePEVZG4kD2V6+zsXG9MeIEI/nBsbOzRpTq1tja3g/mPPKV2WGd3QyRO4t6fSqXeCMCGYRAwswGAzQC2AxARq1kdPDR2
Mr+ftrY214ZFeavnR29lpsl4eeXXFk2O57KNFsEbXsmMZ5zDA8YUrwHwgMsJq3IcAslsGAYfbG1tvS6dTheTbW0hs0w/s3f/T4kI
qVT3G8XiiX370t8vacyOVwI83796dU8ul3mltXL9+Pix7y/mub8t2VarCdesWdOzPbNQuNLY4k1Hjkx8a/F+uq2jOSIG72tMNv5v
WGTI0SmyuxmgCQpAsM6JASSoTFeWdouQXOlAX2DQgIhcBOBbz2nTpdPpRq25wfeiDwHglpaWZmNMTyYTini429q5aitSr5juWkSt
AmDzxqSjKnLYFovdROKsuFf3prqsAyqLJnyFhv7url3Dx9ra2iIAXtvR0RFjDz4solr73w+CYNpaE19YOPKqllWr9ivP6xaLyylC
t8ViMQqCzDRAX7TWfaSjo+N/5/P5rOdpba0Nm5qa4mLlfNb6FudCBbj3tra2xgB0gVUmCMyBxbISADcyMvJpAOjp6byKhAqhMauJ
KENE6eV1IpJ7SeRlZG0rtMqTuH9wIu+ztnhdTU39l09MTtBzU8PiDQG8tZSOAfCgdXIdGRNxzlZlMvMfSaW6JpxDRCmVnZ+f/+rE
xMTxUwMf+tAhbBbhbVqrvRD5TGtra00ikQgLhaxzTn9SxL3d8/jNAD7ryLExogDw5Zdf7qXT+z0HGzmpVdgRoGyxWGwFkBGRR5Y5
SU5becwpeu3CQnCOs8aFYf6xZe3gPI4+ZpSxvrFNwgiFhZ89UogdRNgBIlA7sMOkUqmuYjHfGS5kb9SVla8XZ14HyLdLkYanplUC
AKOMB5BmZgsASqFXKX6jUurdEZIvE7luAgXWWrvc8GRmYUboyOnFVUZxzpGI5EiwV8gNrlvX0yXiQsD5SkmZAiWYVTmASCkdiYrI
K0XxtVbc65k4KcLDItowU6K+vv5+Znenc/bDZWVlVQBZrXWgtT4L7Aa05j7P8zqttWvicf8sZg4ZDOaq5d4xBgYGEq961WB8yZJ1
znnMMFrrpQYp2VuiAgHZxVPEomVedBTgW5zD2xcWplc7kaxzTmFRy53KQ2TrMs1FBHcy+F3ARKqMiMqUorhmjkWjjpdP0SMjHWsA
d5bv+12+7w86cV2a6MLx8fEioGJKqcMA3ywib+7u7l4DoaxaBOzk5KQrnc/j5HSDn9kpZrKe59lne5aeACTOuah1QLF4EhgGgI3F
Yo4AGBE+zSOnI6VUTjluXCLHrQ2vBKguVl19NmtV50Re1NXV3La8nnopsfV9q44fODCXCcOwD8CPRkfHtwP4SUdHcx/A/+x5Oh+G
NmCP1wJ4aKmTtNZVTqRFa3/C2nADFH0/PXzo5hKBP8T/9E9fvGd+3lyktVewVrYdODDyt8+eXjs2MPNMNBq/IZ1O71rUVF8WsS8F
8E1mFgDegle4NVGMn2PD8G0AAqWUx8wvExErgotK27nIBIG9gpm/ba2JAPM1AE7upA9t+L7htDvApANLVpPnjYrYSmN0DYC5JQ2g
FNZaCw2tjkNEW6Xio6Oj323vaHt5ENobmNnzPOVKXhqwfftS57Jd3uEi6PM8dYRZhZZk2iQS7x958snZ56ce3JUAi4ic7ZwBM/Kk
6PLmquZHAxRCIpcYHR39cUdH+wNhWLyBiKKkl4AyCMJPiFgv47ZKm3E9T58oFm0FESUBPLlk2IuWJiZEleJdzrEvEmubm5sbrq7u
qmxsjFQtLCwknbMM4ARAEa09b6luakIJa0QF5KAgZGE2b96cGBk5dDER8kqp3yXiqA3DGNi/BMA/L+W7hGC+666nssWi/a4xxbd0
dHT8TiqV0uvXdzWA/D8ToV6lyo8D8k1y8s7+/u7Bwbe9TfX399cy04eJOBOPxx9x4JiAkps2bUr29Kzr+upX//kqIqol0mPOuYjn
qd5169b1rV6/elN/f//g4OBgu1JKaa1jkUjEA0BjY2N5gTvCzLEwhAeQb4yJTKen5wG+2cG9GpBupVSSSF7C7H9g48ZNb0hvHHyD
7/k3gHB5NBrNKkXHtK+H1q/vakilroj09nZfbcNwS2jtMMCaRMrCfH4HQGNKFT7Q29tbNzg4qFpbW9cK8Cda6x85546DKOac41Jj
qU+Idb0icik0CgAokxlcPCEbLIL6DRs2dKRSqe7Ozs4/0Eq/nkj/iw8EAKoi+cy5qzesXjOwYWDj+rPP3tQ/ONi0NPoHBgZqILiM
Wd104KwD11ZUVP0ee96fifC5tsyuYYJzjjUAikRin2TmXoG82lGJDZirH2Uw+0LkLdE0IsJgVBTbu3cT0R7Afbi/v7/pbYODuru7
OyXirnfAk23XvWk7ET0djUY+2NbW11xdHWnIZrN/RyS3kcKDR46MPwmocSK+pr29vWn9+vVlWus/IOb6RFTvAlwUgD189PA5ItJs
Ld6U9v3rNm7csAWa/llr/7VDdwz5SzPAswjQIAhuA+NHDvYfrDPfXFhwXyNW7PmRu8iHx6w/T6TuKRTtF+Z+cvfXi2HxTlbc6yLR
9+3cuXMGzk2LuHNn5mZuNTbzOQe6Xnnev+v+/u1CsmBtcXMms/C5QiZ/S1AMPjczM3M926glUUcAf4njIVLqsHNIEJGAaFxEQgFo
dHT0QWb9GTCNEVEvsZqMRKrv2bZtm8W2bVYnO3+otHccQHsslrgegtZMTr4p2PuN0NoPae3fmt6//z4HLIB5dmJiIsusPiSCVmPC
O2fmZu7wPO+LxOpnWuu/E8kSE00QUWHRKxuxpD7qBMNKuyIASezYUdpEDT0ZmjA1Nz/3OWvDvwfhraa28m8PHDjwg5BInHXZYmj/
spDJ/10+m7slMzt9G7L51y0Rsblcbp2AZkNd8UNsg92xY0cY0ZH7AdpvimY1sR4hsnkAsnfv3kNa8yeIeISFcwCksn3SEcsxEpo+
Se4qzJDQifQPfhD4Pj5ErLhYDO68e256mxP3L2A6KDWxm7Zv3WryEgwR8ZT2itvAwSeIpUtEotrT/vr1G+IU8f6WNcWI5Ttz83Pf
AuHdmtWte/eOHHIWeYDGnMEgM983Njb2DHbtKm7btq1oJPq9XBBEP3/9l1KnEf30LPK8valuMNna9HupjtRmABo1qMAQokvmSX97
+2Cyufna7o6OyzegsmrZdF3jAz0+0OcD/dXA0lyOcqAWQLcP9C6++tGCVgi8+GvQhHuglxHg5T2tsZbU9/ZFkEA9nn38vQ+gGU2D
dQmg/llkOjGAVQ34vXdUA0A7Kqt7k8krutvafre3Kdl/6tn2aqCiZqn+fUB5Z3v7y7rb2q7pamw7+1TDbFFlwCpgi7+UxxDAQLwJ
qSsizybyayoipfr1+UAfgPqTY3rwbV4UaFtsm6VXfyLRWH8yqz//YnnXeesbnl0fAH9aXost76kBUA9sji5rI47VoBV3ILaURhnQ
gKtwqj/e9v5KNA3WLTXcIBDv6+6+uKOt+ZpkdfN5ALzldO0VSEV6Onpekmxuvra3s/NFnXWdvanOzj86B6kKAGhEor5zVfurUu0t
W5K1Tf2nonHuSGxMXVC/EYl6ABWnraSpc3t6WoY+//n4SaXS/o6LB2PNq14O+PBZtBeEQkoFQrDka7ILhajRDroiiliglDMGpHXB
kDXsmGzR+qy0R77vpBg4scYS2AGAdYa18jzxFLnAWIvQOCERcQIhcb7VVCaeFMlYo7QUlGFfw7eWIr44W1apTGbO2QgrV/Dhq5C0
tWKFrfN9YD4D60e0JSJFJNAMKgaGEuXsFXIa4kKlVFC0IhIaJbC+0xEmZw2DYSO+8m1AZJ0lzwsgIgzignURsGLHykkh58jTyhGL
WJFIGFJBTGjYh3PWMStWWrEmIWVCJ2AHWFgnrMhXjgErRCosWoFzxOQAhiYh40UUmBnOANGEU2FGyDpVJK9kovqOle+cQqVgbgZG
+YpIiSMryoGZCtaVsTKZKCCOKChYriQiG1UiRORHnc0tOFGKfGdZGXG6LFa0NnTWgqyYCGmlLdhpYwBnnYpEA3LWEbGyzrCIs1oo
4sgpY13o+dEiKwUXFHUoYURFYmKicRcLCuLCUMLQqVh5mSo6A+MALYFExA+5zI/MzGSm8sPD/07tl1/YpNq8FABoVqSUInGnDsol
LjUAjIElFvaIsXS/SGLPskLCoh8DVIRJWMiYpe+SGGNAxKIUkxSF0AHYPiuYBGSfLnJGW+jSmCusMn60I0c6HSUZZQIKgI6CrBMg
RKiEFEcIulQeYk9MoYBIfW1YVhtT82PH2AWw2GSFjygtRxRBG9giiWslg1ohHoHxNrAvniMatWL2e1ARR+IUAQa4QIN2WaEJK9BA
GIkQbQhKbfCEFRSjZOGcbTJGWo0SJxQlRVLlsz6sYfYAuGjRATxhILNSlON+KE4R90LDhBEaJlFFRbbXCpmiw15FYCsuqa10OVbW
ldrXeFCPEqNggCoNWmPFlLEgBHRckexlQtrARsQ5huH1ylN7mSjrRFSRnCODdg1JGoVJD2qYSeAYAYDVVqibxO5gQWeIUn8oUr4Q
ohpmozlJqDEpJ9aRjgEyLoSdiuCRWGfFIxaRIhloaA1orWGMKfm/y1xhIcdskc1MTD39Gw1l2dTYmFzT0rEh1ZzctCaZXP1fTac/
2X/uxo0bLwWAVUBZEktTzClZv6qr4QogMtA00P7rlrsJiA9UtNa8kGeTlcmqnmTXK7qT3Venalpaf9GzXWWrGv7LbVDX33T64nx/
eXPtAF5YOf87pcTTyX/jawsUBNTX19c8Xxa7LevJR11M/WnBU1/o6uq6dWBgIPGcEbmny5bSIn97e3s1YnhDYIMCAESTrdcUW5re
vvzRvg0bOgrV+uPjZ511Dmrxxt7e3rrnIC3oOa8JCEPPjk6Jpjpfmavz37sYJfzzzy+Vrbv7LKqlO5zGG+DhClMd+beO7o6rniuv
rq6uHlsX/VhjY2P9CyrbMule090W1sonetev7zhlZ4Fda/UrXH/i8uet12/idZKn++8NXxQQJBxQWjlWSnl/vW/fvvvXrl3bn8vN
f8Va259KnbPPYnILi8Sj0bLvPUN0qLe3d02hUKjXmtcq5T22f9v++5qbO/qiUZ+UUp8qFubc6tWrk9lsVpGYls7O1OsA2+R5hdtt
NmuY0RDmcjNKqR+Xl5fPdXV1tQP2FURsjHHfGaGRY6lU6jIrdh0cRnO53PcmaCK7tFqQSqUqjDFXKSBw1gw4oTi2bnVdXV0vdmTO
U6B060Ud39m+fbsBYJPJZJSsucE5ua+2vvbjO3bsCDs7O9+sWV0+dKN8+47VAxvDMNwMYM459x9huOBEIlXHjh0zmzZtap5bmLsS
DhSl6F27aNdwamBgtQRBg3FuwFPRB7WWaGCCc33l320DOyNSrDVhuDnZmXy5p3h/euvBu9AbPslam9bW1lgkErlSRJLM/ESa0nfj
N/gzs7+Jfa8EQGoqKiqIcKUmHq1ubs6YQmEjRLoR4bvJ5f6UgFVELrTWvbaiom5YJPyU1qrbOcwR4U319Q2HnQuurqmv9Z98/Mn7
4/GKLURyibV02PP4f2itjoDcgIh3udbebhE5Wyl10NrwDYXCwiHn6APM5APoVIpfWV/foFnkD5xgLyBXRaN+ZHp69vGliI6qqooP
K6LzCDQngteLc3vq6poKIPMehqSd40symYVV09MzOwCgsbF2tTX2WmPdx/bu3TsJgGdnZ59qaW29f/fO284ummDIOZoQcmcT6EJm
/0lrzdkVFZVPi7V/RaAyYhczsNdVN1TvRGhuIOIXEeBE7PUENBHEh6PfVTqy05jCy4mkkUlNEak31dfXz1hrN4QuTEZ9f6OIvJTZ
SxO5P6qvbzg8NTV1aFls5H/79PqbmdeZnTGmLAjD6ySb/YAx4U0CTKhAYtaa8621zzinnrLWdhAFFzrnTjDLZw8ePHiTteZ7YRC8
jpl5YXbW27JlixIRJSKaWbS17qF0Ov3xYmDeJYIyEfMSQHLOOc85x87xiwGRAwcOvj+TyV3PTN8h8odDR19xjveIs3nnZMn+ks7O
zm4Rd5Zy8pfpgwc/7cTdSWAVBvlXWSOWnXpKxB0IQ/OKrq6uCgAgQwkQF6y1xcXBzQDczp07Z/Jh9iIiPDQ8vP+jpmjezyztRPI7
AM1oTS+1YnPpAwfem04fugGQo2TkJVbcXGCCLw0PD39IxE0L8OV0evi9QmI0oV2ET2gduXV4ePgTWus7w7BwmbU2CotoJKLuB+jr
IuaQc05r59p/0zbdb0REhJTS80rrT/oL0f/lsXo3nKwqGNMDR8Ja92jNL2JWDzPzYecwMz5+fBQArJURAeIiosLQFbZt22ZFbNFa
G1prXRgUDwPA2NhYHiwT1rpKIjKlbBESqTLAngCAiYmJbFVV7Q9E8n0i9jql5MUgxca4k9Eg8Xi8UoSzew8dGlvM/6A4iIOLMHM5
+/ol2tdRpegn2WzWAMAC0WER8WKx2KrFGD/T3d29uqene6uIlBnjRkoh8p2TTlwmDMNqOFMwRsoBHsNJ/49HRVw5gbLkKNPU1BRl
5mmt9VRy8+YoEQWOQmEtmbIyNQYAJghGnBMFQIjE5vP2QhF5vae8Ac0qgFLFFQs6gDxmzu2a3JWJlpU9JYQIFOaFMEZC92jt/ysA
X+vIPBG629qar+rr6+vQWl1FLE8xo6AZ529cvbrH8/yXMyPGLMbz/VesWbNmdW9v7wWwlPJY73ZAQsQwIHFjgkNisba/v3uwr6/v
8unp6Vudk2s9hWHP874EuIJS1Da06CjkcrlRBvze7u7XbtiwoUMzbxER1lofIqEj8XjiS6ZIewHl9/f3BwBoPJ0eA+hB59yf9/b2
rh8YSA0YY95nQ/FFaA/gLl3X09N1dGzsVWKQYNZPOlC1UpGniWhDKpU6t7+/f61z9gJmTkNcObPTSikhSBmVlrsAQsxaiDO2bW4u
85pNm1YnoehKzfoo4AISWwPYyxTwXS/ifdOKixFJw28SdL+xs0yi0agfjUZTRPTo1NTUZCKRIKW4P0Lq/qK1B5Tit4u4S0WwK5vN
3qc1vYxZrXLOvsI5LHie/7fWyjhILgWpFxGTAfCIc3aUSTUR87nO2Vc4kTv9SOQuJrQx66eU0lUAfV2siwj4zSJu0Dm+A5CfWLGX
OSsvAtMME8f2P5F+9Pjs8ezs7Gy2uqZmDpDrQmM2gaQIoscWFhb+3Y/4Lzam+LuA6xTB15544onDS7ZSZSV2KPLXgOha59xLRTBm
nb0pCIKnfO2tMSTXArSBgH9irX8mYjt939smYhQzv9E5+zJY92OhuTuVincze0/E4/EpEKes9R6OFjJzTOgl4qdFJEHMXcbIFc5Z
a6z7ZDTq1wr0MWNcGkTXAK5XRE44wK+urrl/enraYKXtghsYGPBlGR0yODjoDQ4OektUSDKZbASAurp4U3t765caGxuTi9fU8jQW
o1MBgLZs2aIWw+6rllMjmzdv1kNDQ7yUPgCs6upqaG9vrz4ZDZtMVvU399eWwt1bY6fvIamtrS1vamqqW0rvVHh5e1MqlYo8L1eX
TDYu1WW5tLe3N3V1dVUulf20NKtXrTrF2y2Wm5blTcvrtaiVVVtbW/NSuTdv3qyXaJ3Ozs5VAwP1CQBoamqK48yWy1+8ta2ioqKm
paXtf1ZWVlY/x+6mXyWtX7Q/hP4L6dCv+Pn58qJfs2wvpH7/z8j/AQK3X7wXtnUpAAAAAElFTkSuQmCC
""",
    "icono_vur.ico": """
AAABAAUAEBAAAAAAIABUAwAAVgAAABgYAAAAACAA3wUAAKoDAAAgIAAAAAAgABQJAACJCQAAMDAAAAAAIACVEAAAnRIAAEBAAAAA
ACAAuRoAADIjAACJUE5HDQoaCgAAAA1JSERSAAAAEAAAABAIBgAAAB/z/2EAAAMbSURBVHicZZNfbFN1FMc/5/b+abt27aztBuoQ
NIHBFvBPiIJAs2QJBg1CspmYaEwmLyT64KMP1hljfODFB2P2NOMSH4D4oA+agTonGAxbCMEhS0VlsM12ls7u2t7e3nt/PpR2MX4f
zvk+nPM9f3KOoJQgosbPzfa6qiMjhub5nicGURyvIgb/UGULC/XNqgPYlbgWcitB6Y2je34HJXpuejo0Bt6qE3nLiCdPOHbFEQy9
EVQxQmEJtCThxm1eiEyqdT/l/+o9Zxl66TPg5dOn0XTuwWu4oe/mJ7VC5bdwJBTXHGXT3z1EXyLNUOf7hFWBctUJzFham3Oe1lp5
+vT0NACF9duNv2sldDPi54uzYnYkCJcvMGTeZKa4hOfeR75q+N0949qasxAAfDSfFQ2aAsX1cRWJzeLWb0hfrSQha1WetC/KNpal
oDrlp0JNrpZF1pwK+xJf0e6gRYxqF6ZvkamWSCfDdCV9Rus657+1OX7E47JA35YYjz1ssrIcAuCdLLRn6U0uUXT3cqv0LJcWj3Ek
soNN21Os/rmbOzdD9McDHjEgGRO2JQWALKA1HdSWd1KpJ/nLdhjdcZ2TBxe5a1ex/BT7B3qIR8Os1wJUAI6v2iNo398jZyvPoOsB
Hzx1ieH0ZewVxbWpRxnMdrN43SIeFeyqQkIQEgnaAi3jeKbsjV7l5OMzbOpLsHClk7vF+1HqZ2zbJZM2eKhbhwb/gQZZAiATrq3+
UO5XI18OMn+nwvb+FQ4M/0FXzzoDB8p0mgabuzTwUQq9tCGwa1UBdFmNqSDQ5YvKcUYvvsrHsyaZzBxP7CmSyrgQuCAaeCK2G51q
bjGjmuvM5TTt3bFg65sTnztW/FjVdt26J/pQ+kc5mMlj6Zrqja55hwc806knz3996rXDwztHlIwRtO8gUEr2pT49caFcy6hYdL/l
unyzluVc+RBoIh2WMp9fmpl7e/eVV0bOjAS5HAI0DUDrKycmJsKn8vrrdoMXA9/bqkBMXbtl6frZX/IvfcgZsVux/A9KyQZVMvje
Jw8cyk0+qJTSWtVyuY3jA/gXkwtQNFHurJoAAAAASUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAGAAAABgIBgAAAOB3PfgAAAWm
SURBVHictZRrbJZnGcd/1/087/O+b0tp+/YwoC2HQgDpTMAEqMNgmW4iM2OZFpN9IFmyfZk6jdGp0aS8fBFNnPHDYowJusVtoV3m
B4bxMEwJEXZih44i2A0KpdDS0tPb9/Q8z31ffmiBVjR+MF7fruTO9bsO//8N/+cQgM7Obq+nZ5/92dF399hgydOxjYtiPINTxAjq
FM/4xDYUjMHDElJJKVQ1toAXVJL2S1Y8U2GKhd985+HNRzq7u72effusD7DpqQahB8pO1idrln3B5os4GyOBIbYh4nuUNCSZrAUc
BZeiptzH5+v+Sn3iBu/nttFnd1NbU0WhcOlNgE0NDQLgLxwnnawsDw712dODR60NQ68qWYOWLYWgiMSw+xP7cYkMq0ovs3/1S7xz
rUiw1Kd68iTrkhoPF77qB2KLC2vOAXrn92USOjo24P3j0ikyTau9a4VBcoUJ6pc2YYCTQ8e5P1ONxi9wbMhndaqG59+Laa2H9sZX
9bnRDV6mok7vBsxHHBXjkpRYt2ob9bXNXB2/wGw8hoqicRITjtDm/YGxkqVvUBhPh5wdhYs2wdGrNyRIHqKq+gEF6O3tvQPo5QAA
5288q9fDYUKFj3LXqLxyhQ1uhuFaxVmf9cFHtC0TXvl7ggszjrcul6ldkuC1i5bPLvN017o8q5MnBKCDXk4AZuEEtfaGrmWMBr1I
GA3SnMlRs0KpScbsas7zyFvjvPfqGCvvCXhia4q1LT6yRPnGFp9HNviy1ANx1gF0dCxaUQdwgly+mZu5e4lmY5Ync6wMB5hI54kk
4qGaJPfurOPQazdZ88YIj+6tYdqfZGT5cp7cCH4AJH0uXxe5VXEBYH5fjeekIrGcN0YfRkqCN7WEy7km2lsusyPVT7Smis9ta+WX
R95m++YiqQBWFiNKcUAaMAFUJxcbzSyk3Zcv69pcgBFLhSlSjiP8wOfpT11j9dqQMx8W+fXPHU891E7T8mpa1zQQhW7OiJ6oAM7N
mXcRoHce8fHoVnl98jNghOlyQANTvPzgn9jZMox4ERvWWNa3V1JVVaDSQsaHfBlEBJwKChj0LsCtOFJs1xGtIY5i9rac5bGmM6yy
w2g6wdDZCo4fXsMTj7UweL6RidEUDSsSlAqKdcx9OgJG5W7Aifmk2osoRcK+pnd58dG/sX1LknMDEc4qPb9opRxlWFJ7k7A8SWk2
BThCwDqFudtijHELAXNHbmubo8axqHrU6SUmrwxQ39RC66Z6SpNKrhCzZVuaqfEzzM5EpKvmOv7KzgqWpgGnoFCO3N0TdN5KjD9e
4Zd5aeQB6Tq1md4PrzN6fQQJDF/+bpmlyz4gLKW570vjtKybgdCjoQo8AQfgIHSJqTu6nJ+gp79fAVozwQeXrxZn816m8vDY4/r7
sRH5duEwu5v7aLynSHNzHbmcR9MKH6zgHKiCETBgygXiojS+D9DR0eHgxEJJdRnVA7rxmReOTZH8ookLtqxpL+HyfLrqFPtbjrOj
NUV1bS0J30eMAfEQMThnrUmKmZlJv/1s/tiOAx1i50+id1TU2SYioisq9FDCRRqR0ARlVRPwx+k9fH/gWxw+38yf+6c5dWGc0+fH
6f94iOmJETVinWpCJkqZH2d3SUxPp4E5uS4yBV1dRrJZt/kHzx8YobIrLhZiUSeewSu7gFgN1YkcKVMEVer8KXd//Tu2c+Nwor6y
/rmWXX/5unapkSy3lbQYAEJXl3gHs+6T3/vtT8Zc8EzkBI0iZ8QpiMZqUERURayXMgnfsK2i/1ePn/vR1/o3dWk2m9Vb3f87wO17
GLJuZ/Z3e6/m3Q+LVrZaP3X7uRGHsWXSRvuWJcOfnjz45Iuok3kvLJLpfwDMrYts1ql2ew8edO3jYbQ9F9pVxiGVCYYzSe/N17dP
nJY93yyrIgjIvxT/79HZ7f2vb/4JeNWHhfCcYEwAAAAASUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAIAAAACAIBgAAAHN6evQA
AAjbSURBVHicxZdbbB1XFYa/tffMnKuPjy9JnJtdN01TkrYJDVUbILSpuKYEEO1xhUQkXgAheAUJhEiMQAIkBOIiBEVcSluQXYQg
AnFVE0gDJVCqJHWTNJfmZqd27OPLsY/PzOy9eLAdkrpukXhgPY32zJ7/3+tf65818H8OWbhQVRERfezJf7TXGsE9qlaNUXWvsMl7
L1iLeCfXrofGqTE5jYMiFofEExI7CMR4ry7KpfWnPrTrjZcWsACChc09/f0GcJcn0tviXPmJFAPqEbGICF4d1gQ477DWoigqOn8E
xQjEmsVNT5H1L+LJkmQ7yBiH15RMNsfM5OUeoL9/Hus6AldPp6qNJFEJst5oIHHcIPENcpki1dowTdlWxpMxQhsRmgweh0pAI3Z0
xvvY3naQFdmXqKeWg2M7eM68j1xkXJzENitR8nK8RQSMtYpXjrzwJzPCoOR9nvHqEOu77+LMqX9SXrWKU6NH2NLyFjavu5eGF4K0
zvbgR+zoPMj+i558GNGc9SxLH6fblDhr79eizBqROcn6X5WAserSmDNXnuXs9HNEhQKBDTnzwvMELkBGT5ENCzw/dZi2sZW0Fdch
gw/Dyr9QM82UQ8/TZzxXGp6sNXxw86/5+vnbJcjfiIlHdRHeyxdKhbKvVgc1mZqmrK2sbd7AyuYbSQKlaqvkM80YMczoOIcGD5If
f4TN7U9zdCLPF/fXUZNSzob87kXP7V1NPDt6hedPf5uBFw9R96FfUoKNlYrOMTJpGnifXbksaC3epLkgI7NJTG5WcF6YSaeIiMjE
RTY0ZdjWfICB4YQ0dqzNWn5zzOOd49yY5+GBhKPDQkf2GOtmGpTLDy1NAPYCMHTlCX1+9JA2GjMoGWaDEpcm62w4cYLBtS1M+SFm
UqUuBZqDI7QULbmqpaaGgxcSfAxrc8Ly1pB9JxN2rTHc3FFkWWmc7nCfAajQf7UOrkow0D8gADX3rLSG5yTwF2n4C3K6dpqYSVZ3
WPJlR94kdJVi7ioO884TVboaniSf4VJsWH9DnktiOTEzS2e7sKZkyOUDBqcceU1Zt9z6OQKvJMFzGxWgZaJdl8+s1qZolFQamstE
UtAazc2ObItBPbRnHT3lAp2/SPjCgbMUdi3jq5V2atMxd+aFA+M5yhn41Bsi2jKGqdhyvArnqymwRBcMDMxl4HIAZ+rr5fL4DopB
XRqmyg0yQqmudMQZqmaaLXnHHa1l7MdKvO7kjXzzp8/w0NYsRlJqJwfZtukmVGDTSovWPR0FWL885PLEHNa1GfhPF8yvZgrj0rH2
EFoc42h9Pefirfyjegd/HdvK4eGHmBx6G+9oWoGfsDStVnZs7aIQd/G9J4awFjQQytMJbbHHp+DU4J3g40Ud+DIC8zFWD7lpNqXb
zWIlIcKTdZPECIONDrat9dzVndDSZjn8bIlvfdXwtpVvYvNNdxNlLJvvXEvOGhqzKUYEmQcxAk5kEd7Vhcq8MD1+WNvHVjIwvZ5Q
PB4hrw3qLuTBrmN8afvfiCRh5ZorjNSG+dNfRym3z7DrPiUyAZnEosZQXzBdBc/86VO/tBEtFMbx2i3y6MQOrmiBSFJSCZicVh5c
dZTv3HeQYsaR+hCMcvemZt7V003X7SHH/27IBkpn2VKvO/Byzbd26bhKYOPGPgX4xfQWPe+WacY4ptJQOxhi981HudWcpiVq4NKA
MOc48mQL+x/p5sHtLbS3ljnwyyZy+Qy5FRHqhTidf7GAWWBi5bWtOEdKiKfhLK9vOkPf2/fxmfeMcHHQ88zhKWxJuXQqzyNf6uLi
SwmlNqVYijAizExa0Lniq01fY3oyJ4WqeW0CAN4IxDN84uZD3N7ZwGSL3LNzGZmsAI6nf7OMeLLE/T2rMOFlLpx/ijRW1AegUGgx
lNssXMNBAXRxDSyy4jSpIz6jYiOOD07K+ZOncNFyOte30lQM0HqMKc4gxQk0zTM1fo7p6hTF1jYy+QQayvaNEQhoPDeoAOo8zMYs
GrAWfY7XFIoTxyeSugmj6LHhXeoHfiXb2k8weuUya7s2sKIt4N73p0Tt09TjZ2g27TS1n2fD1rPkC4J6QVDwCiLonAISO9ysi2pw
vRNelaB3714F+O4Db3xRLBcC6xmhU7926WN8/OhnePzsZoYuHuPk6RFGxuDN26vcsmUcQsOd90RsvE1RvyA4IFdbwGMwjnB4Itp2
FuC5edu/PgMiCnuM7e6e3fLpHx6upcGtoTR8FGJGdQ0/eGk3Dc3ywJpDtOWHaDSto6m8mlxhmLYWATWoyrzW1+ArSmg0nsn+683b
Pzqoioj0Xq2O64uwskk8sLqUeTR0sfcgXiHDLJnI8tjIB/j40U/y+OktnDlzmsnxk0TRILMNIXEGpxaPxQRmDlkV1Kv6QCYbLT9D
HfRXrsNcbBU6N+ve+skf7BuRwk5JGokiISgGaBCRxI73rtrPQ2uepN1O0ZwTojBE1RJapVhooancShjY1ERpMDxW/OfTV/a+Zdeu
N9TnM6NLE9izx9Db6z/yjSdu+cOFmT/XUrPMqEsVCVTnqlpFmHUh7dEEnfkhWsMJImmgCCuiMe4uH+O2ctWtXtEiQaEjPjd561tv
u+/rT+mePUZ6e6+bil7RLPfsUdPbK35n74/fe3RS+mbVROJdghIgIqAIilNLrCFeDcrcP4J4aMtOJHeVjwS7Op6RTS35D9/xwB+/
31ep2J7+/kVtuLRb71FDr/j7P/+T9xyfcg9PSW65xg0v4FE1CDLn9oqgCqKAR1RSW7CBc7VbmtNP7N+7+8eVSsX2vwI4LOGEAPSK
p9Jnf/253b96Z2dm+yob92etOhtFgQahUTEylweLk1B8EBqNskFgjF1hZ3777tW1ew+8BvirZ2AhKn2W/h4XAO//yqNveqHqeqZT
7k683hg7LQESGqlFlnNZy987S/bnv/vs7t8nHqhULK8C/t8RgIXCVEAF8KrBh7/8aMdlZ4repdJWyk0/srP4krl5Z0Ov37NoDP+f
olLps1T67NIP9NnK3P3/7mDAvwGc4xZRqnITQQAAAABJRU5ErkJggolQTkcNChoKAAAADUlIRFIAAAAwAAAAMAgGAAAAVwL5hwAA
EFxJREFUeJzFmmmMXmd1x3/nufe++7yzezyeGa9x7CROnI0sLoGEkIYEshCI1QraoKYNKqIgQEWVUJkxSNAPVEVtEVLVFkhTloiQ
QgoJaZBD0ux2QuyME3vsWWyPPfv27vfe5zn98M44zjKTOND2vB/eK92r5/7/Z3vOOc8VVhBVlb6+Plnpmf8L2bVrl/v/xvC/Jstq
V1XNv/7ymY/HJt2DEysGXbrnVtCHMQYVJ/Xn3GnrG8CBAyOoYNUYA8YDLC52KJ7ggVqjIqoq+BJVFm5odv+ydseOypu9z18WSX+/
P1bx/4KGlkutrSFOUD3tvuir/JcuFdQqIku4BREWmSuiCqKIeDh8osiiNgIxBEECT2JcVGeIcxg/IK7Z8akUPwIqqioicjqK5QkM
HD0qKp2hjSNnnY0BgwoI+OLjnOLU4hmDYIhtjBEw4qPq0CW2p2ygoIozAXGlSrpykG5ziGZvgtClORpvpZS9mEQiiVoHgqq1nohf
mnehLgNzBQucBdrvfBExvpfwDcYAOBczs3CCbKaJVJClUlugWJqhrWktDmW+Mkk6mSMwSZwqgqDqMEaI1COYe4XzzUNc2XWIzoYy
scSgwnT5aX5+9DDDfJREIoVzVsV4gqppWBbkSgQODyB0aGxjRicOcaJ8BIzQ3bCJoYE9lP0K1170xzw//AjF6Sm2b72GFyYfY2x2
mLOzF3LJWdcTpFKoOkSEWqTkF3ZzbdNPubh7mqGSx1OTwuWrE4DjWLnMpc0PU5trYty/GU8sS85XeCcEwrWdyhFxsY0ZOfoC/fPP
kmjNsW/qcUwtxoaWxw7/mONjL5OMEjw0fA+RhDRlWynEM5TCeVrTWazGWGfoLj/MLWvvZ01bGVySV8Yj7tlTJfPeJFs6fL72cInP
vSfF+zt2c/fJrdB0NoHGOCwZm1nWhcxyN8470I8oeL6HzYH1LdXCAtVqgbhBoCHBnolfMWummc7OEJoaSS9FRUqcYIS9xx5mduEE
4iWJJ5+iNn8fv54o8OyoT80qHzzH57YtafYdgyeHYyYLMR0NKba0T3Np4iH8uEoiSCHG2HQ6XjbvLUuAjWvUE88ODj3H8MG95Cop
ZCZiY9M2zmreTqvfQSABFalQ0iKgYIRaXMUaZTQ+wi+HfsT82FOscT9joDDNI4PKV54s8amHSozMOq7c5FGLA+4fsLSvybGh1aPm
DB4vsv/QfRyZPIK1tVOQ3iznLx8Dl9zl4uN7bXOmg01dF3HUDdHdsBHrK3PVGXLJJlpyPRyfO4hYJeVlyfh5jBpcOWJzx0XUBBqr
D/CJS05Sppla7Ahj+P4LFb70iwKf/r085chj90DEjecH7Jl0/O2zISdLQqv5GbMHnqPccZtbe+0dy8JcngCiyUR/OOcPMRgcYWP7
dpxzoJZV2R6OzL5MdW6CTClE8i3MhRPMVMdoTLSS9AMOHt/PLRu28rHNBf6jv8wPD1iqYUxXU8C1GzIkXcAD+yJsBNNFx4Fpx7NP
hNjYsCEvNKUCNuYqNGWG3PQ0dReSN9pgBQLKyGQvz42/QKOxLMy8QpGjbGjYxMDEMQZnBjnv8AjOlRlozJIhpuiKTJRLdGU7iMNp
XK2fKAo4MG758LlJ2lMZ9o/H3PNCgY5sQFc2xXxF2dKV5L/HI/IBbM1aEupRtIan50q8P7HH29i63wCovpHDCgRgqvSMWZOaYWEi
ZMYbpakpyVPjU4Q2QeDlaE/WIIw56pUBZVVSsFKlpsOck7Vs8hydbSl2Xd/IbyYdDx11jBufru4GnjxcZkOpxHltPld1p6kcV6pV
x6VrEmxp81mfV7qaA2aLKsND9h1sZEBzOtBqRRlVoaclyciC0phOUw4bSPgx2WafsCGBZDxSThHjWJtxZALh2uM1LtkfQmOSh1M+
//ibkDU5n85GaMwYejpTzBVC9sxY1qUsmRTMhzBYhZcHQyaLlovaQ+66MBOv78kum4VWJLA6FqfWkWgVJCzjAoNNJXipGnN2uoyn
Fi+bIJtweCH05MA55YKc8h4/y/x/zfOTl4bof3cLn31vGx84JwC1KIb7GuFkOcUjJxyZlOEvN0M68Dm/LSBhYKoCvxissfdETbZf
0wK8aQisTOD4xDadKW8mwwKBKXJuqszwQsQ2iXiXqxC7mDhtaRTHhBOaUorWHDtakjRtXMfwFTUGvIgnHqsxNniYK3u30pAVDg8X
GfnNGHdev5Z5fAqh40NnJ8FacDFOhaYAPnNZkvG5BHsH5wXeQQyUGse9YVo5MXkF4JEuVnFmliaJOF6cpqf2AtlChvZkSKBFsmHM
WY1wbjpPpSHN5Z9poWtqPcHuCt+8535ufXKOj97cxuhYiemT85yYLpGIG2hWh4scsTV4pl7aRg78UEmIW7GhWomAtGaOy3nZ/aAh
w5OXEkoWXIpZkiyUGikVx3hlcgdxukqHmSQdjfKhjccoTztq9jjd+Yi2sMzm1S0ELs/dPz3KB69uJpkUkkmIrGNTbClbhxjwXb34
OwVAlo3dU7L8TgxUNSBBzBYzjYfD4PDU4KMk4wKRg4KXo2B7eLm0jY6OZs5qDVBvPdVqK6Oj3Ywc72T02TTXdV9FtbqK5/bPs64z
w7qePKvbcwRhjK1YFOE1e+2rLQWeWVjWCitZQFNW2VA19BebiDH4ahebEgjiCgbFGEMUOz68tp8/v3iAILB0dJwgkVAefTrHo7+K
kUOW3//A+ay5cjO4x+lY5bN6czepbBLxHNU4Rp1i3ixKQbxC8cwJ9AHvroY2LDXxXGUTiNRrewEnQspWEc+nEAe8r+MIvVf205Ku
EcVCS0ORpA9PvRgxMpxhi5/DmSnWd1XpXJ2hUKpiIkfCgFWhasHpa5R++oVamzvzarQPkFqjeaBwGcdsOwFu0cz11jBJiIfl2o5D
fOM9z7ChuUhsDb4xxJFHNmf40HVKLV2mmKrS3F0jrA6TTobMT4Q0zYakA6FQcbgYZNH7BVDVekenwltFwYpZ6PuFy+XZageBWFxd
/4utuaEaKjtyA9z27nk2t5SJY4Pn1Xtiz1dAWZ3v5PZbr2ZD6xBrmxM8/2CWpuw+OnuStLcnUVWi0GEM6GJPDSCyFMqKKNh8/swt
ADAcr7YW79RiGCjHHt2McO2aQVpqU+TjBRRTz88OxFdKJZ9H7l3Fi//ZxYUNERdtDzg5XGDPoxGzUymSGSGd9RAVfGMo1eqTjtc4
+tKg4C2mUisS8E83oBhcGHFNfi//dM2jfOq6Oco15eX9JcRbnLkYIQwND36nk59/t4uijTHBLEYb6VzbhQlgbsaAA3X1WPKMEEev
HdWILG1YgoCN7fK10IoE7OK/ikBc4yOrn+bbN+zlyrNKNLcFdHQm2PfENCMjVTxPEN/y4hONPP9IJ/PVMrFWWbMxg58QxJ8h5Qvl
uXRduwqC4jsIKxanr8Oor/lbVlYkAPURTeiE9WaIL1wxRGe7IawJubRw3uUtdKxPU6vEiFEqJeHAk22UiwEbz0lz086NZBpjpsYO
Mn7yeZCQQJOALvq8UvQdCx6oCojWuZ1CLQhmRQ4rBjG4RY6GpFtg/MQJcn4LubSHJ8L5F+TYel4Gg6Aos9MBo0cySGToaW8nlw+Y
GHuJuYn91Mp5SmVDTAiy6NoK112eZUeopPw6CaEexG6xh1FVF0/UXeiMizmJrUN9Ag83Fq42Bycq5LyXGPUaCfwMTflGelanUAtq
IZm2dF18grGFMgPDPpfOriMRBIgINgpYtabGqnXVRd3WZW0jIAa19dRZn+opIA4RrxZjZ1L5U978eozLulDgGZfyZU4RjMbMeD18
b/Rmnp9ahbNzeHaI46OHOXy0iAnqY8fW1ohb7ihz26ezbL50iMLCAMl0M6opUvl5rrjlCOu3zqOxYckznAUbu9PAL1adRhRRrAaV
vYNr4uVwLmMBFedE85nUS1oDNaIBqv32MvnrV9ZyQeYgl7W8zJUdh9GpEWK3jvVrskSlFDPTbfSsTZDPhETxCHHcTLbZo3l1mYa8
kkl7oOZU0SbCqeslEkuzVASc5A7t2umHZ0agt0/cLrQxMC8GWotCFaMovrFEqQ6eiTt5anQHD57czxfO+THZxAAvvtyC57XTuSpF
U3aMNZ0x1VqS0ZPCurOa6WivIfiguljXv9GhX7WAKig29om9fD9YtBcju3hDZ7ZiFtra4r3S6OnJ+l4pDgRRS0JiUkmPQ+YSvnnw
Vg7NNtOQmCCQA5ycOECsJVIpn0IxTzYzRTZdQDCoq+9OIm+V/EQRvIVyUC5VW55Z6ck3X6mvTwHp+1OOZhPmMd8YERFXV47gEFSV
lBdzkHfx5QN38g/97+fATCfEkxw8coxjEyGtLUXW9QyRbyjXM8zi5qSnfm8UVVQQ6/mGUBv2Hqt87nmAPnrPgICI3n77vUZkZ7il
xf9BUuOaUzUC+mqdIqgKCeOYCTbyo5mP8KX+P+OeIzuIbYHxE4c4Nn6SybmIYtUjjA2xCk4Fp/XUbDwQcYvzJpZG8oo4IuvLTKnh
Jzfe2L6gvb1muWOmZdPovffe7kSQuz+//ZFL/urJX1Xi3I04G6FWTq/cFfDUYnylaHr4t4mdTNPFnef8mkRhiGNzhsDPkAgSBMkA
SKKAMUIymaW1OUMmJahVnBM14ExAcHImd2CwesMPVb8jfX3LoXyrUqm317Brl/vCt+571/1D8w8W4qDZiFNUjWKW9qJT3bbgQDyc
Kpsyx9jRdoDNuaO0BnPkTBlPHLZ+SEPWq5GmQiJoINfYTkdbXtMJcSJWKmEiOjy77WMXXvd3970mMZ0xAaCXXvNVdrmbv/7vn90z
GX2zhomMqkFFnKp5fTbRxa4hcj6KkvJCsl6NlKmR8KPFlyotQZELGge4omkf3d4YCT/tOtpbbEd7PhiZ3fDtze/74adUo1OJ6R0T
UFWRvj7Rvr7g6i/f/fcHy8Fd1sWRQQ2qRrXeep+WAkHqNRQKjnrQ1/3+1cBVZxAsXZlJvar9gNvRuNdubZxIONn0y4nkJz9+000f
mwbl9WdiZ0zgtOdUdSh11Zcf/9ZgmT+JrDoBJ+BB/VXLqOBNXyR11eCcZ2vq052Z9W9a9cov/qi7685tO7849lausyRvWY0uoaif
EG6oPv6Vqz+9LadfT0scO+P5KiZWdLHCX9yD3gAVFFmsokUVcYpYFGuM+k1e2W8wqe+VG7/6iW07vzjW29tr3g74+upnIvVkrgkD
H/2bH9yyb6qyayb2t8cI6pwVg11E6snigeoSBUUUUSeCs4onYjxPDI1SHdrc5H3tgS/d8V0RiXtXSJm/PYE6C4E+gV1u9+7dbd94
evSO4bnoD0qR2V41QWCdIs5SryYXAwIFxVPj1dOnC8n47mBn1v/xzWe3//Pn//CGYT1NQWeC5p1/B7GYYgEee+yx9r97fPC9J6vB
FdU4uqgWu42VmNWx1UQ9oMWmfDOZ9nQ4GXj72hLuiVu3Nj/xyZ23DOnSWn19eqbgfzsCAKjQ2ydLRAIg3L0799lDhZaJQq21HEaJ
pPHVSexaUzJ717ne1I7rd86HSw7yWwD/nYqqCrff60Hv20kKwu33er29b+vZt7HY71pUheU+0alrG966V3/b8j8u+xjM0FQaNwAA
AABJRU5ErkJggolQTkcNChoKAAAADUlIRFIAAABAAAAAQAgGAAAAqmlx3gAAGoBJREFUeJztm3mMXfd13z/nd5e3zrx5s3JmNNwk
iiK10LJomZJskzJrW/Iat2GatmqsuK2Rxm2AdIlhtCjFFghSt0HhwEBbpIkRR4ZTEW7jfU1ESpYsa7Nkivs25HA2zvbmvXnrvb/f
6R/3DTlkJHIoKV2AHoBzL2fuu7/f+f7Ofs4TVk9yA8/+30D6f3oD/0/Qqk71ySef9FOpmWBwsO+NUR1dxYvWw+TkjABc8a5RmEzN
XLGXS38fTT73eu9f/sxgc8W71oM3mZKLYRhv3749ut6WrgnAnj17vP3799uvfP+nH2ho8HFrAjUCIuJWPucu/bjOYskHBUDVXl7b
gYoTMJjlZxVVI+qpOvGS3xljAIeLwSqoEZM8m+xH1Yp4nqqLMqGLDvzDh3d8zamKiLzhwfnX2vDWrb8psJ+JSryjle78J+qnL2uW
aPt+BYZXrCMkD8iV2piw376R9kUvP6+a3CngFCMezgrWLf8SjBGMKBq7q5Zz4HBB2Gm0Mpn24GvyGKKqvBEI1wRgmYwQgTZQTRYB
o5cYUxBwmgBi2kw5HEalzZzi3GVGjQoCqjjEQMJymxkHSox4AaoBzUZNaM0RxnOEuoTVNHWvB80OEoYBai1IGzwRFZHYRXHKl6Cq
AEf2C+x5cxJwoH1Vhw8ECftOcAmflwHy8MTDqkXVgQqeCKqKOociGGPwxMfhcOpQUURJgFkhIoKC8YlaTbyl4wy5QwylTzPQMUvO
bxBrivFaN8dKd1LK3Y+f7QaXvEuNAzyjIr5FvdUc7jUB2AUcBFSMoIgg6nmhGM8IqiAGxbG4OEktrtJXGEGMD+qYK48Tter092wk
8NNUGyXqUZlsukDaz6GqlzRAl0VbBCdCvDTNQP1Z7uz4GXf0z9Lb4QCHEwGWuMvOsbU0zvdHZ5mST2IyhUQAFG1LwiVb8pYAONC+
iqAqYJ1lqbrAXGOSalTCk4CB7DD1hRIvnvohw2tv456NH6TcnOPpo1+nPxiikO/j5OxLHJ95iYXKRdZlNnHn2vdRLAwnAIi011Cc
Cm7xBFv0B+wYPMTNA00qzue1BWEg49MVOgRltuGxaCs8OPQc358osBh+FDEhTl1bHVZlk68PwC4SCXA2ciLG1RpLHD1xkGPzL1IP
6mQ7u8j6nXTEeRbqk8xNTpPPFak2FxmbOUqhu5Nnz3+b46WXCMMMuVQnrUaNVrOBMdK2Gw7EYBVk4TDb/G+wc90punKCSsDL4xF/
+FKDe4dCHn1HQGfe4+DRFt89WuF3H/bZ0fssP1zYAD33gMYIpi0Eb2z5Vw3ApYc8T5vqCIMUhVwPqVKGMJtFfWGidhbnIrJ9OcIg
4KWJv8DFMelMlvH4LJWlGp35XvJeAQQWzQLnq8fJ1YvkM0VQRQGvfIbtqW+w6+az5FKGOFIi69jYZdjRl+KVMy2OrhHeeXPIwRN1
JuctgmHbYIkzpac4tLSedK6AagTiobrsM94GAGhbPTWw6JUopeZplqu4RfDDkGw2j++nqdk6S61JPCekc1lsUCUd5nDWMh9NEYZZ
nBczWnoNzxluHdxOR+cAtjyBN/9NToTHOfmaMtLhsa3XZ10BRrosj9wTUq0Zzsw4shnlF5Mt7r4ppCsthCbm3YPHmBr7KbXcw+AZ
nFOM2FVpwaoAUDyrWHfizHPez1/9PkEuQ9jIgCesGbiNTEeBZlxnqbFAo7ZI5FkacS0RbVFSfhrP+JTrs6gIkcnwSulpZppTPDDy
fvqrzzHbfIWfjtcRz6PumliFD28MefTOFGvyyqahgOmS5cB5y6wJ2bIuT8ZTxPcY7GoxNPkMB8e7Ka7ZQmfYgRqzqmTg2gDsAvZB
OgjsfH0JajG399xLlLVMNkZxvkfkW3yNcOrI+p2sya7nYmuMuqsRxCGFfC+d6R5EoeaVaUVNhgsbqWmdxfoCk5Pf5ePrjzO81aPc
7CD0hVLd8dRoi28dqnFhJuJ3duboz8GxCY9fTEXEvrCux6NkhR+ciPjeOcvh6VEuVv+E7Kk1cvvI/Wwd2OyS6GA/sOetSUCsLg5T
ObduwzZvMSxxbukQnQMDrCtuZmZpjHJ9gYzJkg4zuMgnqitiHcYINbtEVGvhqcEREy3Vqbbm2Dyyg0xBuDP6Dhu758h6IaWGMm8t
XWnDo3dkuKcY8qVnyvznZ+p88NY8i3XHS6NN8nmFGP7dcy2+caJFR+hzR6/PRzdUmK8s6VRtGBdtfPuMILjYA3do7lk5Un2JTQPv
ZLBjHecWjmNMwI6Rh6i2Khy++Dz1+SmKpQqVtCGVzRHZKktujqzXQYffTaGrl6nyGKWTJR697RZ+eUuJqZLwn368wMszjjiKSfse
t/Sm+Nitaf7+PV0cONXi4Mkmzdin4YTA9/nSq3VO1j3uGggZznrUrGUi9tjaH8rOjjk0dco64Ik3PvxVAHDgAABx9DLHy9McXXiZ
YqhkoxKvXjiJkVm2dt/GTOlVDs8fp1S+yNrpBQanJji0sZcaTboIKGuTcn2JMnMM50fozfksVM9gozNY9fjzQ1WMGD67PUvWOGbr
wrPnmnzhqUXu6A9512Ce87NKLRLWdPicbMaMTcBIznFmKuKMKNm0IfYNz81E+tDwaT4x0gzAh/37r6UB1wuEEgAWGy+aycqodIc1
epxw7PyP8Tt9Bjo8Xr74NLM1QDL46tPVWKS/WqO7Wifuq2CdI+tBMTDMRhEVd4ZAHJtzMWHsyAYFHrknT2/eQ31hvCrMLClrelMc
mYz50fEqY5Uy7x1ME6Rjdq9N0TMLz0/ENCPY0Sd8eFOem3sN6QAm63B2ASarVQPmOhZgtTbAQSGEDl8Zu1DHZNOMdHmM15SZlqEz
m6bcyuCnHJ39HtLKkhtM4XwlEEOkls6Uo5CCaiyMpJV7SxFbrEc+hJ6BgMOzlq+82uSnF5UF6zBAb0cKU0wzOtXgjmZMuWVZmzOs
607x81nH5h6fXJfPX0zG/OWUY20W3n9zyCO3Z5hdCmPQtgq8cda/KgBsBMW0MlOxlPG4qctAK6JS9RjMC2CYb0C33yTnK7YYkhtO
4xbAGKUQGKrOsalDIVK2hXD/0RpFZ5jBo35Tmj8YjTh10XL/UIr1PR4pT5lpCM9dMFSaAWXfcHTe0ag5hosWI5ZzS2CNIyOWpno8
dT7iWydb8hv3GN67znOrqYqtCoC8Z53aWOs4bu4x9AQRU/MthnIp8tmAEwuGTGC5KRMTOIsLDDlfyfgQIwzmoNxKssM1KceOVMia
Skj5hRLPvVpiblcXazYX+ezuAlv7QNViUDzf52B/wPdOQ+ALJxqQS3u8p1e59W6fO/rT3NUXkDYOq8J4VfnW8RpHZ1sM5VsCvDUV
2EWSC/iLAxrWlZ5okT6/SSpqkpMIL1RKEZyxjo25OsNeFa/ZwmV9fF/IGqjHkPEcxZzSiJVbs7C52EF5QzdTx1/j5YJw6lhM+dA5
bmn2svnhfpx1+J5hZrbGkWMlHhzIs3NrkdF6FRFh94aQgUwAanG2iWvnQMUuYdOOjF6sGcrVdsHlrUjAgTYEo65LTi91UqpnyKoj
H9QITZ18Y4myV6M/qrPFqzMYC624ioTQImIUj5LzaAEZX+mzyv0DKVLSDZsDejt6GY7SxHP9fO27z/HFx8fZtrWbO2/xaDUcJ89U
mDg2xbpUH3HcgY8Seg51EVFkUBVMm0+HYC14xspw3id0SUJ8HS94PRU4AEBHYcwrhCVZLA1yemELtdYwAYpPBKZMWiNejDzW6RQb
Wy+h1RT+bJpiq4RPg7WxJdKYe3tD1qUzlGZydG+qsvn+boqn+tlWvYVT58Z58oVf8N0DM9x+8xDWWsq1FlFkWVyqU2taCgpZ6wgl
wDfgnEkAINF2Z8AX0ShWYnWrKviuygZkTJOe9ATZgfOkvJjRi9uwmsaJj9JJBY+SDag2FL/WyWlzG+XZtRgu0OuXaNkqmwvneV9/
i/pixPzcGF3dFkOLXLZMqbzISE8/1aryw2em2fPBbkYGk62lAsFZRa1ls7U4QBAQh5HLPCamuF1yW201BFZXOHGAVQ9jmgyHE6RN
hFWDoBgX4CMEWMK4io0tZa+TGVtkIdrIidp2fla+gw19vQylA6JoiFSqn/lZmJ7tw5MeorKhOL+ezZ2bOTcufO8nc3ieR1eHR0+X
MNCdJRv6ZFsWiRSHIvp6B7xCHFZJq5IAMagTpdgSwqbHL1Rx4jDtoqg4UANiW0mR0gvJSJyUx2yTB9eNc/+tFfxASaVTNKMeSgsl
bDzCoWNZDj/dZOmUxye278DfsMi5iUNUlmJuGSkwtujo6u8gk/ZoxpZa06FWeeOz07Y4qFw2gstV5xsEYBeJFwBHfwS3NAxHa0Vq
LoMgrKy8GVXCuIGHw+KjYrBRi78xeJrf2HaCwUyN2EFfzwXyuYtMmphv/WiCn71QoHAxS87F3L1lPZve43Hi/BLWLlLszlAcMkhG
8A0EqQBbt1ir12SqXR9clQ1YlQoMNjy7vRa5TK2DlxvrqRHir5AzZ5Iid8o2CHyD9X3i2PKu7jF+650n2dxTwzpBgEzo6Ouu03KW
Ay9UmJu1pDyPwPcxwRKZcIF7tqTp7PCoNZOSWdoXnFOsCg2bRKZtPpPr1bUfAVVZEQi9MRbXBGBX+7ohik0mypinl7ZysjWId7X+
qSLqCFwLh9CM4d3F03z+3ld599AC6hztjhIignPChmHDjnug5VWZiavUzSKxd4xK/Sfkc7MEYcD0aBV7qkpfu3fQaFpsS6G9/vIu
ZLkEr6/TiLkOXROAA+1rS9Pyo9rtPF3fSiwBKztjyrKsCYGJ6ZIq7+8/yee3H+KBkTls0pVJquiaSIpaoafbsuMeobM3jSkabr27
k01bt9Fc2ERtUQBL6MO6gYBCh0+j5WhGihrhku97XUpqjG9TWXwXcJAfL93GTxp55jSDSWzwciOmrYqKA+JI2ehP8MAWw33rwUUO
TzzECCrtJokAxlArG3JmA5/65V0UchUGMi8Q1wJe+l4/lXfOM3JLmcGRHDICQQi1ikLs8DDwOiyKXDZ4K/pMbw2Ag+3rmOt3S85p
smRSaEr+41ACXNxirRnlrsIF0vM1wsoCRgpYJ8kKbR1VSbz15PkUL/y4yNnXstx0ywU2PgAmI/ziZwsceblJ90gG26oQBg6HIJo0
RkPfo1ZzRDFc0wjewCTDKitCy3jLJXEXBaceQbzAffkj/IO7xxhoLXHgzx3nzlTZclcH+WzS1PSWW3eecuzlLD96fIC5co477uug
vz8C9ejM3kLvYAvPn2JxwWFbBj+j4JJqtErSFG1ESmSvsVEVURQjuiohuIGYaVm7kqtVKNpxHl37U/7Dh47x0O1V1gwIhULAxbN1
Rk9XEd8kLVAVjO8xdT7Dk0/0M3FskNm5GseOT1GrNMlkDWIM3QMh2ZyhXAqoN7xLKiaaGLlQBHHSFqgrj/mSBkjy40ov8DYAoO12
dhKHGXp1ks/c+jyf23WBDX0WdZZC0WdwfYZqqcVrLy0yPR3heYJzjihSXvzLTiaO91LTKtWozvDaPLfd20kqb4jjJo3mBKEfY5p5
mnXTZqS9uoKnQqNmaUXJ4S57v2Uje/WOV0M3JAGIYjyDbdbZ3X2UT99XoTMrxHFyUpmUYcOWAr3rMqgq1aUWIHi+cnHCcPZQkbgV
EtPgvl2D7HpoGD901Gtl5qaPMj93nJgGKUklxk7c8oHigHmNmXMxkZNr67muLhWGVdsAy0r/amyTLrlIpVQj2xMS+l6SiAhsvDlF
sXuIIBA6O30iqwQBTJ/PUZvN4+EzWOjmrju7CEJHdWmWuemjtBozxFGWlo2JtEmsNmGk7eOzaeHWrWmywyF9hUS1li3/FYGQtiVD
5JKlkGtgsSoAXLKCtgVGxUtxsdLk7Pkx5ksFOnMp8rmAXCogk/EYHPRxVnAKziVNuoUFmCs3iRo+eEJpNmZYAxqNearVKQLPp1lN
06ykYSjGC1wicZLkfylP+didmcQHWZvMaiSd8EuBEKAYh7PQajl96zbgQPvaiixWnYhBnCphmpPRRuYbEbZxjumZ85wdu8DRM5Mc
O7NIpargFHXtnr8V1t66xNbdk2Q2zXLRTjA2Ok7UjMlkioRBDnWKoAxvqLHxzkXyHS4ZRlmZ0sQxLo5XbPAyg+1IUMEQWaUeEYHj
QN+lWZ0bB2DnruSaS5mqZ6QFahRBjHA8fgd/du59jFa6yWUadKbn8c000zPjHD8zT8OC5yuqDgU2bYl4+G/XefjvZdnxkSZda45T
r00RhFnCIIuq0tm7xLYHx9h63xTpjE2GElZ0ua8W+b+qAkkIZFU01rAFyi52XovFa6tA/5HbFWDrmp7Rw4uVeXXSkwR+Vp3fKU/V
dnPq8DD3dh3jruI5tvRM01+sMVee5shJYeumHrIhxBYqZY9mo8jITf1k/UXmZ0rMz58lne1CjU9MTJB3pDKWfCHEiKE9RHIFyQqF
1hV2QBJdUMBDvXJMcRyAXf3X1IPrTIkdVoB/+dFbTz1x7PlpjL8JG5PMSDlxQZYz9m7GFjZxoDzPyIXz7Og+zK6Rk1Trk/z8aMzI
QAfZVJbyUoFmMyAVjNM9sEBnMcDaJVq2TL47zZp1PrmMR+AH5HOGy2HM6hyVWxYDX8QSzrfM+tOQNIbeNACPPfaY7tuHyfXfPpMP
nj9ZbsTviZOoBMUkqaqnqNfBjBaYtus4OnELc40f8MiWlwnNNGcvLOJcjp5Ci5uG6qSDCM9z+F5ApWqYmDDkO3oZGpokkwFxieFM
zPm1XdkV0tCOtNUZIudPVlLvPQewZ8/Wa0rANeEVEWUnxojYNWl3PIVtODUeIm1PK6gaRC2BWlK+o54e4jsLu/n6iduJVejtqFLI
XmShcoLRC2M4bREESq3hMTvfg+/XSKenMAIuVmK3PDwliNxAmKKCiJN6y9OlRu7ownvedS4ZOXzszQMAsOezt6sCO4byz4e+GVsO
MdqxCZpM++EkuQ/EUQtv4utzn+D3f/4gz44N0nQBvYUWtcpFDh2fpVJNoslMqkJ//0n6+6bx/fakmAJq0HaJc+W/N+RdcSJJthG5
YLHcGnruV0TsgQM7vevNCl03XFJVEXlMVB+T7Z//8uNTTf9XW4ozSSXArDREsFwfEKyC5+p06xQbU6fZ2fMKu9aeQxWarW42rl/D
mt4Q37hkVlATREUul9qWc4/LWegb7DFJGlU88eaqXc8eaX7ukZ07d5xV3WtE9l0zKbquBIiI7tyLERF733D2f6SNm5dkIFQFnFwV
Zi1v3sfhvAwz3jpeiHfxXy/8Kn90+N1UojT59Bxnzp3h8MkZxqYbVBpKyxmsM0nwpMlkaRLstAMBcUkeshwBsez7SUItcVJrGntu
tuvpnTvvO/vk3p3+9Zi/NqwrKJECQVVTd//z//KNi5r7oEVt2+zgwMgVr1ouTQntOWbU85DqPLu7X+BTd/yMkfQUtSVlqekReDlS
qZBU4JPKBPheCmc9xDh8HzLpFPlsSDpMpEtUUTEkY6mCoiq+8SbnO48dj37rl3bv3n18NaefHNQqSER07141ItL4/a9+83e/9OLF
u6uk+kTUqooY1OkV0iRtCNzlextDusjB+k5mjw3zgf5X2FY8Q3exhG+riCyBWupLBqsezmliBFWJWgYv7KKvt0hfd5Zs2gNV18Yf
4+EtVP366OzaP9z9N3cff+KJPZ7IvmtVDVbs9IZor/HNv3Uf2vfffu+Vefntlvr+ckfGgbn6dbocrK9cSBJlT3tNBtKzjKRnWJOd
ozdcIu/XSJmY5YlyT6GQqjKUniGvi9gmWJNjTV8/fd05lwnB8yGOjRkr9f5PffC//52vikSPJSK7qnz4hgBI3gu6ON79wBd+/OUz
VT6mQqxqLpUpnaq52i4k7CwveLmUlUwAC0YThfdNjGcsyyrk4ciZJjdlptnee4p3dZ9kjUygsXVIBzet6dT+7qw3Xe0/PB790q/f
/4FPv/jY3r2yb9/1Rf9NAdBGQRDRx7/95Du+cPDMl6fi1DvAxaKYxANo0rXTtke4Kl5XWS6stR3dpVF3SPqZgqIspwGqhtgaUl6N
jflJfaDvuL6reIwh/4LLB5GPv3ZyVj70Tz/wyX//9eUC9I2wc+MAcMko6hce/8YDX/753B/NOW8zEKF4l16rbSv4Bsn4srvU9mlf
lowr928ufZkCjZPSMiOZGfuu7hPB+4qnJwvh2n/9sU/9yR/r3n9j5AZO/tL7b/QDsGwU95rfeeQTz/zjB9b+2lDKvhSoDZxgE8un
iKBXf7VmJW8ryqtXqMWy91gOt9QZ59Q4RQmx6hs/vtAcCJ6ZevfowfKj/+zjn/rKHz+xR703wzy8SQlYpr1795p9+/a5L3712/c8
cWT2X4xV7a9UJTQBNFTVT4rny0UtUHWSqMTrGEuWY/uksJGE2YnFFKMqmNg6DUPjZDDdemrnmuA/fvG3P/2tT/4tvP37WZXFfz16
SwAAsHevYd8+N37sxd5Pf+21z4zX3D+atd76SMX6IrGoM5o4CnGoCKZdv7lSN6Rd4EmcPJrUOVUFcbEaY1C/6EfzfSn7+G/uWP+l
Rz72wZO6F8O+VfdAXpfeOgBc/nZZyjN8+g/+bOeRuebfnalGH1lo6FCML+JJJGBRRBRPRUXdsn28ZARVVDRxCTgrRtS6lEHp9Nxi
f9b/0fqcfu2rn/u1b4tIiz17PPbvf9Mnv0xvCwBwSR0UUJ2ezv/6Vw/eO1VrfWSmGn9gtm431yQdKg5xzokYq4JrN7narSZE1fki
YlQMadfUYqhne3PeT3qz4TcfvWPkqY9++H0zLlnM8CZ1/mp62wBISIW9j8ny5nR8PPvZ7zy7YbYsm6otd/ditbqtarm5FuvaemQ7
l797JWoJPdPIhv5Y3rOjnZnMkWxKnu/NxMce3dQ3+t6HHppPjnqvUX1s1UHOauhtBiCh5QwSEiA8IFb1//RPf9h9cGa6sxbku5v1
Wk5NKMTgE+OCoNGV8RY2+Fr6V58ZLAXm/nq8ohy49wYDnNXSXwsAl0i17e3+6old7QtenzNte8i378Rfbx//u6i91jW6ulfSXxvT
/59W0P8C59hJipQJoosAAAAASUVORK5CYII=
""",
    "Poppins-Regular.ttf": """
AAEAAAAMAIAAAwBAR1BPU0R0THUAAD64AAAAHkdTVUIfQCdhAAA+2AAAAC5PUy8yWeqfAAAAAUgAAABgY21hcLKy0wgAAAUQAAAA
1GdseWYyXireAAAHnAAAMzJoZWFkLrMAywAAAMwAAAA2aGhlYQx1AtgAAAEEAAAAJGhtdHjVjiiZAAABqAAAA2hsb2Nh8UXlZgAA
BeQAAAG2bWF4cAFcASwAAAEoAAAAIG5hbWUlg0NzAAA60AAAAdxwb3N0qV1zlgAAPKwAAAILAAEAAAAEAQYtIAebXw889QgDA+gA
AAAA2KSpvgAAAADnABM2/e39xAoHBCkAAAAHAAIAAAAAAAAAAQAABBr+ogBkCe397fluCgcAAQAAAAAAAAAAAAAAAAAAANoAAQAA
ANoAlQAMAHYABgABAAIAHgAGAAAAZAAAAAMAAgAEA1MBkAAFAAACigJYAAAASwKKAlgAAAFeADIBSAAAAAAFAAAAAAAAAAAAAAcA
AAAAAAAAAAAAAABJVEZPAMAAACIVBBr+ogBkBG8CcyAAAJMAAAAAAiQCugAAACAABAH0AAAAAAAAAQsAAAELAAABKgBYASQAIwNI
ACECbgA8AvcAKgLjACUAnwAjAcYAbAHGACEB5gA9AqsAVwDG//4CJwA6ANIALAHcADUCdAA+AUAAJgI/ADMCTQAxAnUAKAJ0AEkC
ewBJAiIAIQJ3ADwCdgBJANUALAEIABcCKwBUAtMAZQIbAF4CDAAlA/UASAKiACECZQBNAwQAKwLDAE0CAQBNAfgATQMKACsCtABN
APYATQISACkCVwBNAbAATQNdAE0CvwBNAxIAKwJDAE0DFAArAmAATQJLADkCHQAiAqMASwKkABYD0AAXAm0ALQJIABMCHQAuAacA
hQKSAIMBpwBuAnUAJQLdAGkBAQARAqQAKwKkAE0CXwArAqQAKwJsACsBSQAXAqQAKwKAAE0A9gA+APj/5QIDAE0A9gBNBAYATQKA
AE0CgAArAqQATQKkACsBdQBNAgoALwFsABoCgABIAjEADAM0AAwB3wANAjMADAHHACkBzgBpASMAZAHOAEQCBwAlAQsAAAEqAFgC
mQBBAmwALAIZAD8CUwAZASMAZAI/AC4BOgAXAxgAMQHCACEBzQAtAooAKQInADoB/QA0AYMAEQGaABwCrgBYAUoAHgFMABcA9wAR
AoUATQJUAB8A1AAvARAAEQDEABkBtAAhAc0AMgKAACYCowAkAsUAIwIHAC4CogAhAqIAIQKiACECogAhAqIAIQKiACEDhAAMAwQA
KwIBAE0CAQBNAgEATQIBAE0A9gAHAPYABwD2//AA9v/1AtUACQK/AE0DEgArAxIAKwMSACsDEgArAxIAKwKDAGcDEgAjAqMASwKj
AEsCowBLAqMASwJIABMCQwBNAqkAPAKkACsCpAArAqQAKwKkACsCpAArAqQAKwRJACsCXwArAmwAKwJsACsCbAArAmwAKwD2AAcA
9gAHAPb/8AD2//UCfgArAoAATQKAACsCgAArAoAAKwKAACsCgAArApEATgKAACYCgABIAoAASAKAAEgCgABIAjMADAKkAE0CMwAM
APYATQQ+ACsEOAArANsAIAE3ABEA8wAPAVkAEQKlAEwDdQBMANsAGADbACAAuv/+AXwAGAF8ACABXv/+AZwAUwJFACwBGAAtARgA
MgHuAAkDKwAZAw0AGQInAEwB3AA1AAAAAgAAAAMAAAAUAAMAAQAAABQABADAAAAALAAgAAQADAAAAA0AfgD/ATEBUwK8AsYC2gLc
IBQgGiAeICIgJiA6IEQgrCEiIhIiFf//AAAAAAANACAAoAExAVICvALGAtoC3CATIBggHCAiICYgOSBEIKwhIiISIhX//wAB//X/
4//C/5H/cf4J/gD97f3s4Lbgs+Cy4K/grOCa4JHgKt+13sbexAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA
AAAAAAAAAAAAAAAmACYAJgAmAEQAVwCGANQBIwFuAXsBmAG2AdgB7QH7AgMCGQIpAlUCZAKLAskC5QMXA1ADZAOvA+gD9AQABBAE
JAQ0BGkEzQTnBR8FTQVwBYYFmgXOBecF9AYQBioGOQZWBmwGnQa9BvUHHQdeB3AHkQemB8QH4Af3CAwIHggtCD4IUAhdCGoIoAjW
CQEJNwlrCYgJzgnxCg4KNApMClkKjwqwCuELFwtNC2MLowvBC+QL9gwUDC4MRAxZDJsMpwzpDRANEA0tDV0NlQ3TDfoODQ5qDo8O
4w8SDysPOg9CD40Pmg/CD9sQAxA+EEsQbxCIEJEQrhC9EOkRAhEtEWcRvRHyEf4SChIWEiISLhI6El4SoRKtErkSxRLREt0S6RL1
EwETLBM4E0QTUBNcE2gTdBOOE9YT4hPuE/oUBhQSFDUUfxSLFJcUoxSvFLsUxxUnFWcVcxV/FYsVlxWiFa0VuBXDFgcWExYfFisW
NxZDFk8WfBbBFs0W2RblFvEW/RczFz4XSxeLF+cX7xf/GCUYTxhcGGkYdhiDGI8Yohi1GMcY3RjtGP0ZDRkdGV4ZhBmRGZkAAAAF
AAAAAAH0ArwAAwAGAAkADAAPAAARIREhASEXAzcnAREHEycHAfT+DAGk/qyqyKqqAZCqjKqqArz9RAKK//7U///+AgH+//7U//8A
AgBY//oA0gK5AAMADwAAEwMjAxImNTQ2MzIWFRQGI8MLTAsbJCQaGSMjGQK5/gwB9P1BJBoaJCQaGiQAAgAjAlUBAAMcAAMABwAA
EwcjJzMHIyd5Cz8M3Qs/DAMcx8fHxwACACEAAAMRAuQAGwAfAAABBzMVIwcjNyMHIzcjNTM3IzUzNzMHMzczBzMVIyMHMwJ1Joqb
K1cr0itXK52uJpusK1cr0itXK4vz0ibSAcy0UMjIyMhQtFDIyMjIULQAAwA8/6kCDwMUACQAKwAyAAAkBgYHFSM1JiYnMxYWFzUu
AjU0Njc1MxUWFhcjJiYnFR4CFQAWFzUGBhUSNjU0JicVAg8wXUE2W3IBYQQ3MkNSOnFeNlVoCGEFNCtCUzn+ijo4ND7ePTw3j1c6
BFFRCGVOKj4I+hEkTEBRagZTUwdfTCM5CfQRI0xBARU0EecFOTP+RUIuMTUQ7AAABQAq//QCzALFAAsADwAbACcAMwAAEjYzMhYV
FAYjIiY1JQEjAQQGFRQWMzI2NTQmIxI2MzIWFRQGIyImNTYGFRQWMzI2NTQmIypSQEBSUkBAUgJI/mtaAZX+gicnIiIoKCLtUkBA
UVFAQFJvKCgiIigoIgJ0UVFERVFRRYr9RgK6LS8uLjAxLS4v/kJRUUVEUlFFXS8uLS8vLS4vAAACACX/9ALbAscAKAAxAAAhJwYG
IyImJjU0NjcmJjU0NjYzMhYWByM2JiMiBhUUFhcXNzczBwYHFyQ2NycGFRQWMwJkWzN0SUhvPVNRIBwtVDc2UCgCWwEwJigxISfQ
EkRiUhEXnf50VyfNgFVEXDUzNGBASW8gJkAmLEcoKksuKS8tIx44KdEcco4eIZ1CKizOMWo7TgABACMCVQB7AxwAAwAAEwcjJ3sM
Pw0DHMfHAAEAbP9GAaUDpgAPAAAWAjU0EjczFQYCFRQSFxUj12txaV9wcWtnX1wBIqisASthCWn+2Z+b/uZqCQABACH/RgFaA6YA
DwAAFzU2EjU0Aic1MxYSFRQCBzBna3FwX2lxa2C6CWoBGpufASdpCWH+1ayo/t5eAAABAD0BaAGoAt4AEQAAARcHFwcnFyM3Byc3
JzcXJzMHAYMki4wmeRBLD3kni4slfBBMEQKeQTo5Q1iXl1lEOjhDWZiYAAEAVwBvAlUCbQALAAABIxUjNSM1MzUzFTMCVdVV1NRV
1QFH2NhN2dkAAAH//v9xAKsAfgADAAA3AyMTq3A9SX7+8wENAP//ADoBSAHJAZUAAgDY7gAAAQAs//oApgB2AAsAABYmNTQ2MzIW
FRQGI1AkJBoZIyMZBiQaGiQkGhokAAABADX/VgGhA64AAwAAAQEjAQGh/u5aAREDrvuoBFgAAAIAPgAAAjUC3gALABsAABI2MzIW
FRQGIyImNSQmJiMiBgYVFBYWMzI2NjU+cIyLcHCLjHABnRdHQ0RHFxdHRENHFwIdwcGsr8LCr1d5Skp5V1p6Skp6WgABACYAAADi
AtMABQAAEzUzESMRJrxcAoBT/S0CgAABADMACAIUAuAAGAAAPgI1NCYjIgYHIzY2MzIWFRQGBgchFSE1spBjP0ZESwNYBIBlZ3lj
fmIBWP4fr4OPRUFLVUdwenJkU513T0xBAAABADH/+wIQAuEAKwAAEjYzMhYWFRQGBxUWFhUUBgYjIiYnMxYWMzI2NTQmIyM1MzY2
NTQmIyIGByNDgGZEYzNHOUFMNGhJaogIWAdURkZJZGUXGFxeRUA+TAdZAnloMVQ1PloNBRBeTDlbNG9mPExJOktCTAE7PjVAQDYA
AgAoAAACWALMAAoADQAANzUBMxEzFSMVIzUTASEoAV9tZGRaBP7zAQ2hRQHm/iRPoaEBzP6DAAABAEn//gI5AtQAIAAAASEVNjYz
MhYWFRQGBiMiJiczFhYzMjY1NCYjIgYHIxEhAgv+oRdaNFNoLTZvUmmCDlkNVEBPUFBONU8SVgGzAoTkIClFbT1IckJsWjlAX05O
VTMsAZgAAgBJ//4CPgLiABgAJQAAACMiBhc2NjMyFhUUBgYjIiYmNRAhMhYXIw4CFRQWMzI2NTQmIwG7c1lYARdrQmd7NWtMZ3Qu
AQBicApUs0wvWFJGU1BNApaJnTQ7gHFEbEBconcBb2pRlSNHM0xfV0pOWQAAAQAhAAAB+wLRAAYAAAEBIwEhNSEB+/7qXAEa/n4B
2gKM/XQCg04AAAMAPP/sAjwC3gAbACcAMwAAEiY1NDY2MzIWFhUUBgcWFhUUBgYjIiYmNTQ2NyQmIyIGFRQWMzI2NQYGFRQWMzI2
NTQmI4o8NmtMTGs3PjY+SUB1TExzQEg+AQ9QRkVQU0JDU+BeW01NWlxLAY5UPDZXMzNXNjpXFRNfQ0FhNTVhQUNgEthCQj44REU4
w0VFQE5OQENHAAIASf//AjQC4QAaACYAADYWMzI2JwYGIyImJjU0NjMyFhUUBgYjIiYnMzY2NTQmIyIGFRQWM71JPlNPAhViPENp
PHxyjHEta11mcghU0FZUTEZTUk6JPoKaLTM3aktrg7a0fqRWblWSU0dLXFhISVgA//8ALP/6AKYCIgAiABEAAAAHABEAAAGs//8A
F/9xAMQCIwAnABEAGAGtAAIADxkAAAEAVACIAbgCUwAFAAAlJzczBxcBRfHxc/LyiObl5eYAAgBlANwCbQIAAAMABwAAARUhNQUV
ITUCbf34Agj9+AIATU3XTU0AAAEAXgCIAcMCUwAFAAATMxcHIzdec/Lyc/ICU+Xm5gAAAgAl//oB3ALCABcAIwAAABYVFAYjByMn
MzI2NTQmIyIGFSM0NjYzAiY1NDYzMhYVFAYjAWJ6gm8DTwQdY3BIPT5HVjZjQVkkJBoZIyMZAsJwXmllX548UzxGQzo8WzL9OCQa
GiQkGhokAAIASP85A6wCgQA4AEYAAAAWFhUUBgYjIiYnBgYjIiY1NDY2MzIXNzMDBhUUMzI2NjU0JiMiBgYVBhYzMjcXBiMiJiY1
NDY2MwI2NjU0JiMiBgYVFBYzArCkWDRjQjI2BB9fNUxXP3BHXyELTzIDNSg7Hp2Pd79sAaCOXkgMWWtro1mC55BFTCw5NDBJKDYz
AoFUl2NPkVoyKysyXU9Lgk9IPv7fExA8THM4g5BzyHmEjyI+JVGXZI3phv2/NVs2NEE5XTI1PgAAAgAhAAACgQK1AAcACgAAJSEH
IxMzEyMnAwMB6f7QOGD8aftgUn5+m5sCtf1L5QFg/qAAAAMATQAAAjQCuQARABoAIwAAABYWFRQGBiMhETMyFhYVFAYHJTMyNjU0
JiMjEjY1NCYjIxUzAdA/JTZnRv78+kdkMkM3/v6ZQEZGQpflTFBDnqEBYTJKKjVVMQK5L1AyPlESJTw1NTz93UA5OkP2AAABACv/
+gLKAsIAHQAAEjY2MzIWFyMmJiMiBgYVFBYWMzI2NzMGBiMiJiY1K1yfYXKqJ20dbUxJdEJCdElMbR1tJ6pyYZ9cAcOjXG5lP0RE
fVRTfURDP2RtW6JmAAIATQAAApgCuQAKABMAAAAWFhUUBgYjIxEzEjY1NCYjIxEzAZinWVmnctnZh46Phn5+ArlVnmtrnVMCuf2S
j4GCkv3cAAABAE0AAAHEAroACwAAExUzFSMVIRUhESEVqP7+ARz+iQF3Am/pS/BLArpLAAEATQAAAdcCuQAJAAABFSEVMxUjESMR
Adf+0fb2WwK5SutK/sYCuQABACv/+gLfAsEAIQAAASYmIyIGBhUUFhYzMjY3ITUhFQ4CIyImJjU0NjYzMhYXAl0dbklJdUNDdUlm
hAv+6gF3CFyWXGGgXV2gYW+rJwHvPUNDfVJSfEN6aEpGVo9TW6JmZqNbbmQAAAEATQAAAmYCuQALAAABESMRIREjETMRIRECZlv+
nVtbAWMCuf1HATz+xAK5/s4BMgAAAQBNAAAAqAK5AAMAABMRIxGoWwK5/UcCuQABACn/+QGsArkADwAAAREUBiMiJjUzFhYzMjY1
EQGsaldYalsBMzMzMgK5/f1WZ2lbM0A9LwIDAAEATQAAAj4CuQAKAAAhAREjETMRATMBAQHH/uFbWwEgc/7EAT8BPv7CArn+vQFD
/qP+pAABAE0AAAGcArkABQAANzMVIREzqPT+sVtKSgK5AAABAE0AAAMQArQADAAAAREjEQMjAxEjETMBAQMQW+ZA51tiAQABAAK0
/UwCBP38AgX9+wK0/cQCPAAAAQBNAAACcQK6AAkAACEjAREjETMBETMCcVv+kltbAW5bAiv91QK6/dYCKgACACv/+QLnAsEADwAf
AAAEJiY1NDY2MzIWFhUUBgYjPgI1NCYmIyIGBhUUFhYzASigXV2gYWKgXFygYkl1Q0N0Skp0Q0N1SQdbo2Zmo1tbomdnoltPRH5T
VH1ERH1UU35EAAIATQAAAh8CuQAKABIAAAAGIyMRIxEzMhYVBjY1NCMjETMCH3d6hlvhdnulSJSGhgGWc/7dArlyWn9CPYH/AAAC
ACv/fALyAsEAEwAjAAAFJwYjIiYmNTQ2NjMyFhYVFAYHFwAWFjMyNjY1NCYmIyIGBhUCfIszNWGgXV2gYWKgXFhNsP2WQ3VJSXVD
Q3RKSnRDhIsOW6NmZqNbW6JnZKAurwGOfkREflNUfUREfVQAAgBNAAACKQK5AA4AFwAAIQMjESMRMzIWFhUUBgcTATMyNjU0JiMj
AbymblvhT202U1Ov/n+GSkpJS4YBHf7jArk2XDtIbhL+3AFmST0+RAABADn/+QISAsIALAAAFiYmJzMWFjMyNjU0JiYnLgI1NDY2
MzIWFyMmJiMiBhUUFhYXHgIVFAYGI+VtPgFhBUdEQUssQjhFUzs6aURifQhkBUo9OUgrPzlFVDw2akgHMVc5MUNBMygyGg8SJE1B
OVgwYlEoPTs1JjAZEBMlTUIzWjgAAQAiAAAB+gK5AAcAAAEVIxEjESM1Afq+W78CuUr9kQJvSgAAAQBL//kCWAK5ABMAABMRFBYz
MjY1ETMRFAYGIyImJjURpltRUFtbRndKSndFArn+R11aWl0Buf5IV3c6OndXAbgAAQAWAAACjQK5AAYAAAEBIwEzExMCjf75af75
YdvbArn9RwK5/acCWQAAAQAX//8DuQK5AAwAAAEDIwMDBwMzExMzExMDucpmoqhlw2GYqGagmQK5/UcCMf3PAQK6/bECT/2zAk0A
AQAtAAACQQK5AAsAAAETIwMDIxMDMxMTMwFq1menn2XV1maooGYBXP6kARD+8AFcAV3+7wERAAABABMAAAI1ArkACAAAAQMRIxED
MxMTAjXjW+RlrKwCuf5O/vkBBwGy/p8BYQAAAQAuAAAB7gK5AAkAADchFSE1ASE1IRWaAVT+QAFS/rIBuk9PSAIiT0gAAQCF/0YB
OQOmAAcAAAEVIxEzFSMRATlfX7QDpkv8N0wEYAAAAQCD/1YCFgOuAAMAAAUBMwEBvP7HWgE5qgRY+6gAAQBu/0YBIgOmAAcAAAUj
NTMRIzUzASK0X1+0ukwDyUsAAAEAJQCpAlECtQAGAAA3IxMzEyMDg17rVexftqkCDP30AZwAAAEAaf+GAn3/2gADAAAFFSE1An39
7CZUVAAAAQARAk0A5gMOAAMAABMVJzXm1QKRRHJPAAACACv/9wJXAi0AEgAiAAASNjYzMhYXNTMRIzUGBiMiJiY1JCYmIyIGBhUU
FhYzMjY2NStEd0lIahpcXBtrR0l2RAHQMlUzM1QyMlQzM1UyAWh/Rj4vZP3cZjA/SIJTPVwxMFw+P10xMV0+AAACAE3/9wJ5AuQA
EgAiAAASNjMyFhYVFAYGIyImJxUjETMRBCYmIyIGBhUUFhYzMjY2NcRsRUp2RER3SUdrG1tbAXQyVTMyVTMzVTIzVTIB7z5Gf1RT
gkg+MGUC5P7abFwwMV09Pl0xMV0/AAABACv/9wIzAi0AGwAAEjY2MzIWFyMmJiMiBhUUFjMyNjczBgYjIiYmNStEeU5lgxViDlE8
TmBgTjxQD2IWhGNOeUQBaH9GYlcyOmtiY2w4NFRlRoBWAAIAK//3AlcC5AASACIAABI2NjMyFhcRMxEjNQYGIyImJjUkJiYjIgYG
FRQWFjMyNjY1K0R3SkBuHVxcG2pHSXdEAdAyVTMzVDIyVDMzVTIBaH9GOzABIv0cZzE/SIJTPVwxMFw+P10xMV0+AAIAK//3AkEC
LQAYACAAAAAHIRYWMzI2NzMGBiMiJiY1NDY2MzIWFhUuAiMiBgchAkED/koFZUg7TxBiFoRiTntGRHtQTnhBXi5PMEVhBwFaAQ0d
UVs3Lk9jRoFVVYBFRHdLR0snWE4AAQAXAAABLAL8ABEAAAEjESMRIzUzNTQ2MxUiBhUVMwEsc1tHR19pPDFzAdn+JwHZSydcVUwv
NicAAAIAK/7yAlcCLQAfAC8AAAAWFzUzERQGBiMiJiczFhYzMjY1NQYGIyImJjU0NjYzFiYmIyIGBhUUFhYzMjY2NQF2axpcQHdP
bJANWg9eQktfG2pHSXdERHdJzDJVMzNUMjJUMzNVMgItPi9k/dBLdUJmWDI9XlVzMEBIglNUf0bcXDEwXD4/XTExXT4AAQBNAAAC
OALkABQAAAAWFhURIxE0JiMiBhURIxEzETY2MwGbZDlaUkdIVVtbG187Ai41akz+vQE2UldaVv7RAuT+8iouAAIAPgAAALgC+QAL
AA8AABImNTQ2MzIWFRQGIxcRIxFiJCQaGSMjGSxbAn0kGhokJBoaJFn93AIkAAAC/+X+/AC7AvkACwAXAAASJjU0NjMyFhUUBiMT
FAYjIzUzMjY1ETNkIyMaGiMjGi1KRzUmJh9bAn0kGhokJBoaJP0MSURNHiQCmQAAAQBNAAAB+wLkAAoAACEnFSMRMxE3MwEBAX/X
W1vTf/7+AQPy8gLk/k3z/u/+7QAAAQBNAAAAqALkAAMAABMRIxGoWwLk/RwC5AABAE0AAAO+Ai4AIwAAABYWFREjETQmIyIGFREj
ETQmIyIGFREjETMVNjYzMhYXNjYzAyBkOlpRRUdUWlFFR1RbWxtbN0VqGhdqQQIuNWpM/r0BNlJXW1b+0gE2UldbVv7SAiRPKy4+
PDpAAAABAE0AAAI4Ai4AEwAAABYVESMRNCYjIgYVESMRMxU2NjMBvHxaUkdIVVtbG104Ai55cv69ATZSV1pW/tECJE4qLgACACv/
9wJVAi0ADwAfAAAWJiY1NDY2MzIWFhUUBgYjPgI1NCYmIyIGBhUUFhYz731HSX5OTn5JS4BOMVY1NFQxMlMyMVIxCUaBVVSBRUWA
VVWBRlAuXEJCXC0tXEJDXC0AAAIATf78AnkCLQASACIAABI2MzIWFhUUBgYjIiYnESMRMxUEJiYjIgYGFRQWFjMyNjY1w2tHSXdE
RHdJRmscW1sBdDJVMzJVMzNVMjNVMgHuP0Z/VFOCSD8v/pcDKGVtXDAxXT0+XTExXT8AAAIAK/78AlcCLQASACIAABI2NjMyFhc1
MxEjEQYGIyImJjUkJiYjIgYGFRQWFjMyNjY1K0R3SkdrGVxcGmxISHZEAdAyVTMzVDIyVDMzVTIBaH9GPy5k/NgBaS5ASIJTPVwx
MFw+P10xMV0+AAEATQAAAVkCLgALAAASNjMVIyIVESMRMxXAWUAYmVtbAfo0Xqb+1gIkWQAAAQAv//cB1wItACsAABYmJiczFhYz
MjY1NCYnLgI1NDY2MzIWFyMmJiMiBhUUFhYXHgIXFAYGI9BkOgNeBEU4NDxAQz1NNzJcO1twBFsDPzUxOiM1LztKNQEyWzsJK00y
KTQuIyQjERAhQTUqRilcUCs0KiIbIxUNEB8+Mi1IKQABABoAAAFLAq4AEwAAExEUFjMzFSMiJjURIzUzNTMVMxW8IipDUkxMR0db
jwHZ/r0oIU1GUAFDS4qKSwABAEj/+AIzAiQAFAAAAREjNQYGIyImJjURMxEUFjMyNjURAjNbGl04QGY7WlJHSVQCJP3cUSovNWpM
AUH+y1FXWlYBLQAAAQAMAAACJAIkAAYAACUTMwMjAzMBGaph12rXYlQB0P3cAiQAAQAMAAADJwIkAAwAAAEDIwMDIwMzExMzExMD
J6tehIRerF1+iF2FfAIk/dwBs/5NAiT+NAHM/jMBzQAAAQANAAAB0gIkAAsAACEnByMTAzMXNzMDEwFrgn1fsbFngnxfsLHMzAEQ
ARTLy/7x/usAAAEADP7+AiUCJAAHAAABASMTAzMTEwIl/rZebN1lrKoCJPzaAQgCHv5EAbwAAQApAAABngIkAAkAADchFSE1ASE1
IRWRAQ3+iwEK/vgBcUtLSwGPSkoAAQBp/0UBigOmAC8AABM2NTQnJjU0NjMzFSMiBhUUFhcWFRQGBxUWFhUUBwYGFRQWMzMVIyIm
NTQ3NjU0J2lmCgpRRDopJyYIAQo1Nzc1CgEIJicpOkRRCgpmAZkTViZcYixITE8jKB1VD2AsNUcMAgxINSxgD1UdKCNPTEgsYlwm
VhMAAQBk/5YAvwMbAAMAABcjETO/W1tqA4UAAQBE/0UBZQOmAC8AABIVFBcWFRQGIyM1MzI2NTQmJyY1NDY3NSYmNTQ3NjY1NCYj
IzUzMhYVFAcGFRQXFf8KClFEOiknJggBCjU3NzUKAQgmJyk6RFEKCmYBP1YmXGIsSExPIygdVQ9gLDVIDAIMRzUsYA9VHSgjT0xI
LGJcJlYTRwAAAQAlAPcB4gGKABcAABIzMhYXFhYzMjY3MwYjIiYnJiYjIgYHIz9zGysbGB8RGiUERBl1GigdGB4TGiYEQwGKExIP
DiMekhISDw8jHgAAAgBY/3UA0gI0AAsADwAAEhYVFAYjIiY1NDYzBzMTI68jIxkaJCQaKkwLYgI0JBoaJCQaGiTL/gwAAAEAQf+n
AiwCYwAfAAAkNjczBgYHFSM1JiY1NDY3NTMVFhYXIyYmIyIGFRQWMwF/TQtVCXFXNmx4eGw2V3EJVQxMPUxfX0xQNzFKYgdeXgmN
a2uNCVxcB2JJMDhhVlZgAAEALP/+AiECzwAmAAAlFSEnNjY1NCcjNTMmJjU0NjYzMhYXIyYmIyIGFRQWFzMVIxYVFAcCIf5CECop
CXFbDQ43Y0BmbgRUAkM/OkkODsKuCUdKTDQvVj8aJEAjPyZAXzR2WjtIREYhOyhAJBpfTQACAD8AqgHZAkoAGwAnAAAABxcHJwYj
IicHJzcmNTQ3JzcXNjMyFzcXBxYVBjY1NCYjIgYVFBYzAb4gOyk8Lzo8KzsqOx4gPSo9Kzs7LD0pPCGJPDsrKjk4KwE9LjsqPB4d
Oyo7Ljw+LTwqPR0dPSo8LD92Pzc3QD84OD4AAQAZAAACOwK5ABgAAAEzFQcHFTMVIxUjNSM1MzUnJzUzAzMTEzMBl4mtG8jIW8jI
HKyJpWarq2YBf0ICNCdDnZ1DJzQCQgE6/qABYAAAAgBk/5YAvwMbAAMABwAAEzMRIxMjETNkW1tbW1sDG/6G/fUBewACAC7/PAIR
AsgANQBBAAAkFhUUBgYjIiYnMxYWMzI2NTQmJicmJjU0NjcmJjU0NjYzMhYXIyYmIyIGFRQWFhcWFhUUBgcmFjMyNjU0JiMiBhUB
ujkxXT5cdQJbBDs5NjsdREBmY0dDMzkxXT9cdAJbAzs5NjwdREBlZEZE/FJFQ09SRERPSEAvK0gqV04tOCsrGSUhFSJURjhXExk/
LytIKlhOLDkrKRolIhUhVEY4VhVyQkMwMkBCMAACABcCbgEjAtcACwAXAAASJjU0NjMyFhUUBiMyJjU0NjMyFhUUBiM2Hx8WFR8f
FYweHhYWHx8WAm4eFhYfHxYWHh4WFh8fFhYeAAMAMf/6AucCwAAPAB8AOgAAABYWFRQGBiMiJiY1NDY2Mw4CFRQWFjMyNjY1NCYm
Iw4CFRQWFjMyNjcjBiMiJjU0NjMyFhczJiYjAfKeV1eeZWaeWFieZl2KSkqKXVyJSkqJXD5jOjpjPEtuFFQeWzxMTDwsQA1UE25M
AsBboWdnoVtboWdnoVsoT45eXY9QUI9dXo5PUzhpRkZpOVFEU1ZPUFYrKEdOAAIAIQFbAZMCxQASAB4AABI2NjMyFhc1MxEjNQYG
IyImJjUWFjMyNjU0JiMiBhUhLEwvL0MRSEgQQy4wTSxLQC8wQEAwL0ACRlItJR49/qI8HSUvUzU6REM5OUNCOQACAC0AeAGbAdYA
BQALAAA3JzczBxczJzczBxeNYGBZY2NbX19aY2N4r6+vr6+vr68AAAEAKQDcAl4BsAAFAAABFSM1ITUCXlv+JgGw1IxI//8AOgFI
AckBlQACABAAAAAEADQBKQHJAsIADwAbACkAMgAAABYWFRQGBiMiJiY1NDY2MxI2NTQmIyIGFRQWMzYGBxcHJyMVIzUzMhYVBzMy
NjU0JiMjATlcNDRcOjtcNDRcO0hZWUhJV1dJWhkWOD8wETRgIil3KQsPDwspAsI1XTo6XjU1Xjo6XTX+j1xJSVxcSUlctyEHUwFQ
UNckHxkMCwsLAAABABECgQFxAsUAAwAAARUhNQFx/qACxUREAAIAHAFbAX0CuwAMABgAABI2MzIWFRQGBiMiJjUkJiMiBhUUFjMy
NjUcY05OYi5RM05hARg7LSw7OS0tPAJbYGBPNVAsYVA5QUE5Oz9AOgAAAQBYAHYCVQJwAA8AAAEVMxUhNTM1IzUzNTMVMxUBgtP+
A9XV1VXTAXm2TU22TaqqTQABAB4BYAEhAsYAGAAAARUhNTc2NjU0JiMiBgcjNjYzMhYVFAYHBwEh/v9tISMdFxgbA0kCRDs5RSwx
TwGbOzRYGzEdGx4bGC88QC4mOyQ4AAEAFwFbASoCxgApAAASNjMyFhUUBgcVFhYVFAYjIiYnMxYWMzI2NTQmIyM1MzI2NTQmIyIG
ByMgSTo4TCIcHCVNPTxJBEgEIRscISMcKykdIiAcGh8FRQKPNzwrHigIAgorHik4ODMbGhwYGBwwGhgXHRgUAAABABECTQDmAw4A
AwAAEwc1N+bV1QK/ckR9AAEATf78AjgCJAAVAAABESM1BgYjIiYnESMRMxEUFjMyNjURAjhbGl89KEMVWlpVR0dTAiT93GQ2Nx4h
/sYDKP7UV1xaVgEvAAEAHwAAAgYCtAAOAAAhIxEjESMRIyImNTQ2MzMCBlFNUQt1eHl0+gJt/ZMBHnBaWnL//wAvARcAqQGTAAcA
EQADAR0AAQAR/vwA/gAHABIAABYWFRQGIyM1MzI2NTQmIyM1MxW4RkQ2c2EhHx8hMj05NDAvODsUGBcUeUEAAAEAGQFgAJcCwAAF
AAATNTMRIxEZfkgCiDj+oAEoAAIAIQFaAZMCxQAPABsAABI2NjMyFhYVFAYGIyImJjUkJiMiBhUUFjMyNjUhMFU0NVQwMVU0NVMw
AShBLi8/Pi8uQgJGUi0tUjY1Uy4uUjY6QkI7OkFCOQACADIAeAGgAdYABQALAAATJzMXByMlJzMXByOVY1lgYFkBF2JZYGBZASev
r6+vr6+vAAQAJgAAAlwCwAAFAAkAFAAXAAATNTMRIxElASMBEyMVIzUjNTczFTMnBzMmfkgBrv6KVgF2qDBIt55hMHZvbwKIOP6g
ASgy/UYCuv2LRUUq8u2vrwAAAwAk//8CcwLAAAUACQAiAAATIzUzESMBASMBExUhNTc2NjU0JiMiBgcjNjYzMhYVFAYHB1o2fkgB
o/6TVwFuzP7/bSEjHRcYGwNJAkQ7OUUsMU8CiDj+oAFa/UYCuv2AOzRYGzEdGx4bGC88QC4mOyQ4AAAEACMAAAKqAsYAKQAtADgA
OwAAEjYzMhYVFAYHFRYWFRQGIyImJzMWFjMyNjU0JiMjNTMyNjU0JiMiBgcjJQEjARMjFSM1IzU3MxUzJwczLEk6OEwiHBwlTT08
SQRIBCEbHCEjHCspHSIgHBofBUUCQv57VQGEly9It55hL3ZvbwKPNzwrHigIAgorHik4ODMbGhwYGBwwGhgXHRgUV/1GArr9i0VF
KvLtr68AAgAu/2QB5QIsAAsAIwAAABYVFAYjIiY1NDYzAiY1NDYzNzMXIyIGFRQWMzI2NTMUBgYjAWQkJBoZIyMZonqCbwNPBB1j
cEg9PkdWNmNBAiwkGhokJBoaJP04cF5pZV+ePFM8RkM6PFsy//8AIQAAAoEDnwAiACQAAAAHAEMAzQCR//8AIQAAAoEDnwAiACQA
AAAHAHYAzQCR//8AIQAAAoEDgQAiACQAAAAHAMYAtgCR//8AIQAAAoEDWQAiACQAAAAHAMgArQCR//8AIQAAAoEDaAAiACQAAAAH
AGoAtQCR//8AIQAAAoEDwgAiACQAAAAHAMcA2ACRAAIADAAAA0YCugAPABIAAAEVMxUjFSEVITUhByMBIRUBEQMCLPz8ARr+i/71
V2MBiwGv/ovjAnLvRvVIm5sCukj+cAGT/m0AAAEAK/78AsoCwgAvAAAkNjczBgYHFTYWFRQGIyM1MzI2NTQmIyM1LgI1NDY2MzIW
FyMmJiMiBgYVFBYWMwHTbR1tJJdlO0ZENnNhIR8fITJbk1Rcn2FyqidtHW1MSXRCQnRJSUM/XGwINQE0MC84OxQYFxRtBl2dYmaj
XG5lP0REfVRTfUT//wBNAAABxAOfACIAKAAAAAcAQwB6AJH//wBNAAABxAOfACIAKAAAAAcAdgB6AJH//wBNAAABxAOBACIAKAAA
AAcAxgBjAJH//wBNAAABxANoACIAKAAAAAcAagBiAJH//wAHAAAA3AOfACIALAAAAAcAQ//2AJH//wAHAAAA3AOfACIALAAAAAcA
dv/2AJH////wAAABBQOBACIALAAAAAcAxv/fAJH////1AAABAQNoACIALAAAAAcAav/eAJEAAgAJAAACqgK5AA4AGwAAABYWFRQG
BiMjESM1MxEzEjY1NCYjIxUzFSMVMwGqp1lZp3HZV1fZho6Ohn7AwH4CuVWea2udUwExVAE0/Y2TgoKS6lTrAP//AE0AAAJxA1kA
IgAxAAAABwDIALwAkf//ACv/+QLnA58AIgAyAAAABwBDAQcAkf//ACv/+QLnA58AIgAyAAAABwB2AQcAkf//ACv/+QLnA4EAIgAy
AAAABwDGAPAAkf//ACv/+QLnA1kAIgAyAAAABwDIAOgAkf//ACv/+QLnA2gAIgAyAAAABwBqAO8AkQABAGcAlgIbAkgACwAAJScH
JzcnNxc3FwcXAeGfnzyfnDqdnTydnpafnzyfnDudnTydnwAAAwAj//kC7wLBABkAIgArAAABFhYVFAYGIyImJwcjNyYmNTQ2NjMy
Fhc3MwAXASYjIgYGFSQnARYzMjY2NQKOKy5coGI/cS0/SmEqL12gYT9xLT9K/Zk3AW9EYUp0QwICNv6QRGFJdUMCUS99SGeiWyck
RGowfEdmo1snJET+N0gBkDlEfVRrSf5wOUR+UwD//wBL//kCWAOfACIAOAAAAAcAQwDNAJH//wBL//kCWAOfACIAOAAAAAcAdgDN
AJH//wBL//kCWAOBACIAOAAAAAcAxgC2AJH//wBL//kCWANoACIAOAAAAAcAagC1AJH//wATAAACNQOfACIAPAAAAAcAdgCgAJEA
AgBNAAACHwK6AAwAFQAAAAYjIxUjETMVMzIWFQY2NTQmIyMRMwIfd3qGW1uGdnumSUlLhoYBBXOSArqRclqCRT0+Rv76AAEAPP/2
Am8DAwAzAAASNjMyFhYVFAYHBgYVFBYXFhYVFAYjIiYnMxYWMzI2NTQmJyYmNTQ2NzY2NTQmIyIVESMRPHd0RWMzMCobGC5BSUJl
WVduB1wDOzIvMiw3UUUmJSIhRTyPWwKOdS5MLDBDIxcdEBYgFhhNO0hZYFQxOionJCwSGzwuIzQhHywcLTOQ/doCMf//ACv/9wJX
Aw4AIgBEAAAAAwBDAM4AAP//ACv/9wJXAw4AIgBEAAAAAwB2AM4AAP//ACv/9wJXAvAAIgBEAAAAAwDGALcAAP//ACv/9wJXAsgA
IgBEAAAAAwDIAK8AAP//ACv/9wJXAtcAIgBEAAAAAwBqALYAAP//ACv/9wJXAzEAIgBEAAAAAwDHANkAAAADACv/9wQfAi0AKgAy
AEIAAAAHIR4CMzI2NzMGBiMiJicVIzUGBiMiJiY1NDY2MzIWFzUzFTYzMhYWFS4CIyIGByEENjY1NCYmIyIGBhUUFhYzBB8E/j0F
NlAsQFQOYhaAX0NxHU8ebEVJd0RFeExDaSFMR4pMdUBeL1ExSGkHAWn9s1UyMlUzM1QyMlQzAQwZOE8oNy5PYz41amkzP0iCU1R/
Rjk1ZWVuRXdIQU0pXE/uMV0+PlwxMFw+P10xAAABACv+/AIzAi0ALQAAJDY3MwYGBxU2FhUUBiMjNTMyNjU0JiMjNS4CNTQ2NjMy
FhcjJiYjIgYVFBYzAXJQD2IUdVg7RkQ2c2EhHx8hMkVqO0R5TmWDFWIOUTxOYGBORDg0TmMHMgE0MC84OxQYFxRrB0l7T1V/RmJX
MjprYmNs//8AK//3AkEDDgAiAEgAAAADAEMAsgAA//8AK//3AkEDDgAiAEgAAAADAHYAsgAA//8AK//3AkEC8AAiAEgAAAADAMYA
mwAA//8AK//3AkEC1wAiAEgAAAADAGoAmQAA//8ABwAAANwDDgAiAMIAAAACAEP2AP//AAcAAADcAw4AIgDCAAAAAgB29gD////w
AAABBQLwACIAwgAAAAIAxt8A////9QAAAQEC1wAiAMIAAAACAGreAAACACv/9wJUAucAHAAsAAAAFRQGBiMiJiY1NDY2MzIXJicH
NTcmJzMWFzcVBwI2NjU0JiYjIgYGFRQWFjMCVEh/UU99RUZ6THc+IEWCXiApXSgGg15cVTMyVDIyUzEwUjIB+8BpkUpHg1VTf0Zc
X00sMiAfISMGLDMg/bAwXkFBXi8vW0FCXzH//wBNAAACOALIACIAUQAAAAMAyACeAAD//wAr//cCVQMVACIAUgAAAAcAQwC8AAf/
/wAr//cCVQMVACIAUgAAAAcAdgC8AAf//wAr//cCVQL3ACIAUgAAAAcAxgClAAf//wAr//cCVQLPACIAUgAAAAcAyACdAAf//wAr
//cCVQLeACIAUgAAAAcAagCkAAcAAwBOAFsCTAKEAAsADwAbAAAAJjU0NjMyFhUUBiMXFSE1EiY1NDYzMhYVFAYjATUkJBoZIyMZ
/f4C5yQkGhkjIxkCCCQaGiQkGhokc01N/sYkGhokJBoaJAAAAwAm//cCWwItABcAIAApAAABFhYVFAYGIyInByM3JiY1NDY2MzIX
NzMAFwEmIyIGBhUkJwEWMzI2NjUCEiAjS4BOYUoqQUggI0l+TmVKK0H+LSMBCjBGMlMyAXAj/vYvQjFWNQHUJWI6VYFGNy5PJmM7
VIFFOTD+ojUBJCktXEJNMf7cJi5cQgD//wBI//gCMwMOACIAWAAAAAMAQwC8AAD//wBI//gCMwMOACIAWAAAAAMAdgC8AAD//wBI
//gCMwLwACIAWAAAAAMAxgClAAD//wBI//gCMwLXACIAWAAAAAMAagCkAAD//wAM/v4CJQMOACIAXAAAAAMAdgCVAAAAAgBN/vwC
eQLkABIAIgAAEjYzMhYWFRQGBiMiJicRIxEzEQQmJiMiBgYVFBYWMzI2NjXCa0hJd0REd0lIahtbWwF0MlUzMlUzM1UyM1UyAehF
Rn9UU4JIRzr+hAPo/sxeXDAxXT0+XTExXT///wAM/v4CJQLXACIAXAAAAAIAan0AAAEATQAAAKgCJAADAAATESMRqFsCJP3cAiQA
AgAr//kEAALBABoAKgAAATMVIxUhFSE1BgYjIiYmNTQ2NjMyFhc1IRUhABYWMzI2NjU0JiYjIgYGFQLl/f0BG/6KJ4lYXp1cXJ1d
WIonAXb+5f2jQ3VJSXVDQ3VJSXVDAYNH9EiPSE5bo2ZnoltMR4tH/pl8RER8UlJ9Q0N9UgADACv/9wQNAi0AJAA0ADwAAAAHIRYW
MzI2NzMGBiMiJicGBiMiJiY1NDY2MzIWFzY2MzIWFhUENjY1NCYmIyIGBhUUFhYzACYmIyIGByEEDQT+TARgSTxPEGIWhGNKdR8i
e0tOfEdJfk5KeSEgdkxOd0L9YFY0M1UxMVMyMVIxAnMuTzBHYQUBWgEMGVJdNy5PY0dAQEdGgVVUgUVGPz9GRHZK5S9eQkJdLi5d
QkNdLwEkTylcUgD//wAgAdYAwAK5AAIAzAAAAAEAEQJVASYC8AAFAAATBzU3FxWci4uKAqpVSFNTSAAAAgAPAl8A5QMxAAsAFwAA
EgYjIiY1NDYzMhYVJiYjIgYVFBYzMjY15T0uLT4+LS49NR8XFx8fFxcfApg5Oi8vOjkwGh8fGRkgIBkAAAEAEQJdAUcCyAAZAAAS
NjMyFhcWFjMyNjczBgYjIiYnJiYjIgYHIxg1KBMbEw4WDREYAzQGNigTHRIQEw0RFwM1ApI2DQwKCxYVMjYNDQoKFhYAAAEATAFI
AloBlQADAAABFSE1Alr98gGVTU0AAQBMAUgDKQGVAAMAAAEVITUDKf0jAZVNTQABABgB1gC5ArkAAwAAEyM3M35mYEEB1uMAAAEA
IAHWAMACuQADAAATMwcjW2VgQAK54wAAAf/+/5AAnwBzAAMAADczByM6ZWBBc+MAAgAYAdYBWwK5AAMABwAAEyM3MxcjNzN+ZmBB
Z2ZgQQHW4+PjAAACACAB1gFiArkAAwAHAAATMwcjNzMHI1tlYEDdZWBBArnj4+MAAAL//v+QAUMAcwADAAcAADczByM3MwcjOWZg
QeBlYEFz4+PjAAEAUwDTAUkByAALAAA2JjU0NjMyFhUUBiOaR0czNEhINNNHMzRHRzQzRwD//wAs//oCHQB2ACMAEQC7AAAAIwAR
AXcAAAACABEAAAABAC0AeADmAdYABQAANyc3MwcXjWBgWWNjeK+vr68AAAEAMgB4AOsB1gAFAAATJzMXByOVY1lgYFkBJ6+vrwAA
AQAJAAAB+QK6AAMAAAEBIwEB+f5rWwGVArr9RgK6AAABABn/+gMHAsIALgAAJDY3MwYGIyImJyM1MyY1NDcjNTM+AjMyFhcjJiYj
IgYHMxUjBhUUFzMVIxYWMwIRbR1sJqpxdrQiYVMDA1NgFmGHT3GqJmwdbExUfRvs/AME++sdfVFMQj1kbYJtQBoaGRg/SW89bmU+
QlZMPxgZHRdASlMAAgAZAWAC1QLAAAcAFAAAEzMVIxEjEyMlESMnByMDESMRMxMTGfJVSAFWArxDAW0zbEJdbXACwDj+2AEoOP6g
//8BA/79AWD+9QELAAEATAFIAdsBlQADAAABFSE1Adv+cQGVTU3//wA1/1YBoQOuAAIAEgAAAAAAAAAIAGYAAwABBAkAAACiANQA
AwABBAkAAQAOAMYAAwABBAkAAgAOALgAAwABBAkAAwA8AHwAAwABBAkABAAeAF4AAwABBAkABQAKAFQAAwABBAkABgAeADYAAwAB
BAkADgA2AAAAaAB0AHQAcABzADoALwAvAHMAYwByAGkAcAB0AHMALgBzAGkAbAAuAG8AcgBnAC8ATwBGAEwAUABvAHAAcABpAG4A
cwAtAFIAZQBnAHUAbABhAHIANAAuADAAMAA0AFAAbwBwAHAAaQBuAHMAIABSAGUAZwB1AGwAYQByAEkAVABGAE8AOwAgAFAAbwBw
AHAAaQBuAHMAIABSAGUAZwB1AGwAYQByADsAIAA0AC4AMAAwADQAYgA4AFIAZQBnAHUAbABhAHIAUABvAHAAcABpAG4AcwBDAG8A
cAB5AHIAaQBnAGgAdAAgADIAMAAyADAAIABUAGgAZQAgAFAAbwBwAHAAaQBuAHMAIABQAHIAbwBqAGUAYwB0ACAAQQB1AHQAaABv
AHIAcwAgACgAaAB0AHQAcABzADoALwAvAGcAaQB0AGgAdQBiAC4AYwBvAG0ALwBpAHQAZgBvAHUAbgBkAHIAeQAvAFAAbwBwAHAA
aQBuAHMAKQACAAAAAAAA/7UAMgAAAAAAAAAAAAAAAAAAAAAAAAAAANoAAAECAQMAAwAEAAUABgAHAAgACQAKAAsADAANAA4ADwAQ
ABEAEgATABQAFQAWABcAGAAZABoAGwAcAB0AHgAfACAAIQAiACMAJAAlACYAJwAoACkAKgArACwALQAuAC8AMAAxADIAMwA0ADUA
NgA3ADgAOQA6ADsAPAA9AD4APwBAAEEAQgBDAEQARQBGAEcASABJAEoASwBMAE0ATgBPAFAAUQBSAFMAVABVAFYAVwBYAFkAWgBb
AFwAXQBeAF8AYABhAQQAowCEAIUAvQCWAOgAhgCOAIsAnQCpAKQBBQCKANoAgwCTAPIA8wCNAQYAiADDAN4A8QCeAKoA9QD0APYA
ogCtAMkAxwCuAGIAYwCQAGQAywBlAMgAygDPAMwAzQDOAOkAZgDTANAA0QCvAGcA8ACRANYA1ADVAGgA6wDtAIkAagBpAGsAbQBs
AG4AoABvAHEAcAByAHMAdQB0AHYAdwDqAHgAegB5AHsAfQB8ALgAoQB/AH4AgACBAOwA7gC6ANcAsACxAQcA2ADdANkAsgCzALYA
twDEALQAtQDFAIcAqwC+AL8AvAEIAIwA7wEJBE5VTEwCQ1IHdW5pMDBBMAd1bmkwMEFEB3VuaTAwQjUHdW5pMDJCQwRFdXJvB3Vu
aTIyMTUAAAEAAAAKABwAHAABREZMVAAIAAQAAAAA//8AAAAAAAAAAQAAAAoALAAsAANERkxUABhkZXYyABRkZXZhABQAAAAAAAQA
AAAA//8AAAAAAAA=
""",
    "Poppins-Bold.ttf": """
AAEAAAAMAIAAAwBAR1BPU0R0THUAAD1sAAAAHkdTVUIfQCdhAAA9jAAAAC5PUy8yWyOiAAAAAUgAAABgY21hcLKy0wgAAAUQAAAA
1GdseWbsn7cbAAAHnAAAMf5oZWFkLjsA3gAAAMwAAAA2aGhlYQv8AjIAAAEEAAAAJGhtdHj2wSFGAAABqAAAA2hsb2NhzuTDOgAA
BeQAAAG2bWF4cAFcASUAAAEoAAAAIG5hbWUiZ0DZAAA5nAAAAcRwb3N0qV1zlgAAO2AAAAILAAEAAAAEAQbAoUjrXw889QgDA+gA
AAAA2KSpyAAAAADnABM2/c79ogmtBFQAAQAHAAIAAAAAAAAAAQAABBr+ogBkCZ39zvkYCa0AAQAAAAAAAAAAAAAAAAAAANoAAQAA
ANoAkAAMAHQABgABAAIAHgAGAAAAZAAAAAMAAgAEA2kCvAAFAAACigJYAAAASwKKAlgAAAFeADIBTgAAAAAIAAAAAAAAAAAAAAcA
AAAAAAAAAAAAAABJVEZPAKAAACIVBBr+ogBkBG8CcyAAAJMAAAAAAi4CwQAAACAABAH0AAAAAAAAANQAAADUAAABiABfAZsAKAOI
ABQCkgAsA2YAHwMeACQA3QAoAd0AVwHdAAcCFgBKAnQAOwEfABoCRAA0ARoAKAHFAAoCjAAtAXgAGwI7ACUCXQAqAqUAMQKKAFQC
fQA+AhcAHAKIADYCZwA0ARwAKAFbADICJwA5ArgARAIdAD8CGgAaBDgASQLhABACkwA+AvoAIQLXAD4CHQA+AiMAPgL6ACEC2wA+
AScAPgJCABoCuQA+Ad0APgOWAD4C8AA+AxIAIQJwAD4DFAAhAowAPgJnACoCTwAYAsEAOwLaAAkEHAAZAssAGAKfAAcCVAAyAf4A
eAMiALcB/gBaAsUAHgMMAFUBIwAKAqcAHAKnAD4CXQAcAqcAHAJoABwBaAARAqcAHAKiAD4BJwAuASb/4QJqAD4BJwA+BCMAPgKi
AD4CfQAcAqcAPgKnABwBrAA+Ai4AIAGWABUCogA5AnIACQNiAAQCTAAFAnj//wHwAB8B8wBbASMAPAHzAEcCfAAcANQAAAGIAF8C
dwAqApcAJAJBADwCnwAHAX0AaQJLAB4BbwAKAwsAJgHVABkCLQAtAs0ALwJEADQCAgAnAaQACgHPABkCdQA8AX8AIgF2AA8BBQAK
AsEASwLEABwBKwAyATkACgD5ABQBygAZAi0AMgMSACgDNQAoA3sAMgIaACIC4QAQAuEAEALhABAC4QAQAuEAEALhABADuwAGAvoA
IQIdAD4CHQA+Ah0APgIdAD4BJwAXAScAFwEn//kBJ//mAtwABALwAD4DEgAhAxIAIQMSACEDEgAhAxIAIQJ9AEUDEgAWAsEAOwLB
ADsCwQA7AsEAOwKfAAcCcAA+AvgAMQKnABwCpwAcAqcAHAKnABwCpwAcAqcAHAQtABwCXQAcAmgAHAJoABwCaAAcAmgAHAFBACQB
QQAkAUEABgFB//MCgAAcAqIAPgJ9ABwCfQAcAn0AHAJ9ABwCfQAcAmMANQJ1AA0CogA5AqIAOQKiADkCogA5Anj//wKnAD4CeP//
AUEASwQnACEEEQAcAToAMQFKAAoBBgAJAZoABQK3AEYDogBGAToAJgE6ADEBJAAZAh0AJgIdADECCAAaAfUAZAL4ACgBUQAtAVEA
MgI6AAoDHAAUA3MAGwIfADQBxQAKAAAAAgAAAAMAAAAUAAMAAQAAABQABADAAAAALAAgAAQADAAAAA0AfgD/ATEBUwK8AsYC2gLc
IBQgGiAeICIgJiA6IEQgrCEiIhIiFf//AAAAAAANACAAoAExAVICvALGAtoC3CATIBggHCAiICYgOSBEIKwhIiISIhX//wAB//X/
4//C/5H/cf4J/gD97f3s4Lbgs+Cy4K/grOCa4JHgKt+13sbexAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA
AAAAAAAAAAAAAAAmACYAJgAmAEUAWQCJANQBHQFoAXUBkgGwAdIB5wH1Af0CEwIjAkwCXAKFAsMC3QMPA0kDXAOlA98D6wP3BAcE
GwQrBGEEwQTbBQ4FOQVdBXMFhwW3BdAF3QX5BhIGIQY9BlMGfwajBtUG/gc8B04HbweDB6EHuwfRB+YH+AgHCBkIKwg4CEUIdQil
CM0I/QkvCU8JjwmxCc4J9AoLChgKTAptCpgKyAr+CxYLVQtyC5ULpwvFC98L9QwJDE4MWgyfDMYMxgzjDRINRw2FDaoNvQ4RDjYO
ig62Ds8O3w7nDy4POw9jD3wPoA/VD+IQBRAiECsQSBBXEIMQnBDHEP0RThGFEZERnRGpEbURwRHNEfASMBI8EkgSVBJgEmwSeBKE
EpASvBLIEtQS4BLsEvgTBBMeE2ETbRN5E4UTkROdE8MUExQfFCsUNxRDFE8UWxSyFO4U+hUGFRIVHRUoFTMVPhVJFYcVkxWfFasV
txXCFc4V+xY3FkMWTxZbFmcWcxamFrIWvxb4F0oXUhdiF4gXrBe5F8YX0xfgF+0YARgWGCoYRhhWGGYYdhiGGMYY6hj3GP8AAAAF
AAAAAAH0ArwAAwAGAAkADAAPAAARIREhASEXAzcnAREHEycHAfT+DAGk/qyqyKqqAZCqjKqqArz9RAKK//7U///+AgH+//7U//8A
AgBf//gBKgLXAAMADwAAAQMjAxImNTQ2MzIWFRQGIwEfFJEUMjk5LSw5OSwC1/4hAd/9ITUnKDY2KCc1AAACACgCQwFzAyAAAwAH
AAATByMnIQcjJ7UNcw0BSw1zDQMg3d3d3QAAAgAUAAADZwLkABsAHwAAAQczFSMHIzcjByM3IzUzNyM1MzczBzM3MwczFSEjBzMC
yxx4mSGgIaYhoCGUtRyRsSCgIKYgoCB8/sSmHKYBtoOamZmZmZqDmpSUlJSagwAAAwAs/6kCUQMYACIAKQAwAAAkBgYHFSM1JiYn
MxYXNS4CNTQ2NzUzFRYWFyMmJxUeAhUAFhc1BgYVEjY1NCYnFQJROW9MQGqCBLYGNE5dRodqQGl6B7cGLVNaR/6QHx0cIJojISCV
XTsDUVIJbV48EZ4UJlRJW24HUVEIal42EJsXJFNIARoiDYsFIR3+hSYcGiEMjgAABQAf//UDRwLLAAsADwAYACQALwAAEjYzMhYV
FAYjIiY1JQEjAQQVFDMyNjU0IwA2MzIWFRQGIyImNTYGFRQzMjY1NCYjH11KSlxcSkpdAqf+f6UBgf50MBcaMQE2XEpKXFxKSlyO
GTAXGhoXAnNYWFFRWFhRnv1AAsBZRUQiIkX+iFhYUVFYWFFFIiNEIiIiIwACACT/9QMeAswAKQAxAAAhJwYjIiYmNTQ2NyYmNTQ2
NjMyFhYHIzYmIyIGFRQWFxc2Njc3MwcGBxckNycGFRQWMwJXQWGHUXhBR0caFjRjREVgMAKhAR0YGCAbHasCBwQlqzAbI6n+WzSc
SDsyP0o1YUFCayIgOyMyUi8yVDIbHRwWFS4epwUOCT9XOzWmfyaXJT8mMwABACgCQwC1AyAAAwAAEwcjJ7UNcw0DIN3dAAEAV/8v
AdYDkQAPAAAWAjU0EjczFQYCFRQSFxUjw2xwZ6hpcm5kqXQBJqqsASteEGj+3ZqZ/uRoEAABAAf/LwGGA5EADwAAFzU2EjU0Aic1
MxYSFRQCBxBkbnJpqGdwbGHREGgBHJmaASNoEF7+1ayq/tpdAAABAEoBaAHPAuQAEQAAARcHFwcnFyM3Byc3JzcXJzMHAZc4e3s8
ZBl1GGM+e3o5aBp3GgK3ZisqaliCgllsKyhnV4SEAAEAOwBfAjkCXQALAAABIxUjNSM1MzUzFTMCOa6irq6irgERsrKZs7MAAAEA
Gv9/APYAngADAAA3AyMT9m9tOp7+4QEfAP//ADQBGAHsAaYAAgDYAAAAAQAo//gA8wCyAAsAABYmNTQ2MzIWFRQGI2E5OS0sOTks
CDUnKDY2KCc1AAABAAr/RQG7A7YAAwAAAQEjAQG7/vSlAQwDtvuPBHEAAAIALQADAl8C6QALABkAABI2MzIWFRQGIyImNSQmIyIG
FRQWFjMyNjY1LYWUlIWFlJSFAYosRUUsEDEwMDEQAiPGxqyuxsauZW1tZURZNzdZRAABABsAAAEwAtoABQAAEzUhESMRGwEVsgI7
n/0mAjsAAAEAJQAKAh0C6gAaAAA2Bz4CNTQmIyIGFSM+AjMyFhUUBgchFSE1SQNicU4mJiYrpQJFcUV3eZJxAQ7+CqQCUWhyNiku
OTRVcjZ6Ymu3V4t/AAABACr/+gImAusAKwAAEjYzMhYWFRQGBxUWFhUUBgYjIiYnMxYWMzI2NTQmIyM1MzI2NTQmIyIGByM3hHFL
azdCLDlCOW1LeI8EpgEwLicrPUMgIDM/KCMmJQOnAnd0NFo4QlMPBBNYRT5fNnZ3LDMtJTAriyMtJCguIgAAAgAxAAACfgLQAAoA
DQAANzUBMxEzFSMVIzUTBzMxATy/UlKrDKiogooBxP5GlIKCAZL+AAEAVAAAAlkC2wAgAAABIRU2NjMyFhYVFAYjIiYmJzMWFjMy
NjU0JiMiBgcjESECL/7IFEgqS2MuhHhQdkADpwYuKC8uMC4iLAelAdMCRYoWHERsP3WJN2JAHyk8MjE0IRsBswACAD7//AJJAukA
GgAmAAAAJiMiBgc2NjMyFhUUBgYjIiY1NDYzMhYWFyMGBhUUFjMyNjU0JiMBjiglOTIBFlMyYnc7cU2aeH+MS2c1BZ96OTUxLDMy
LwI6JWNrIyh8b0luPsG1vrk7YDnAMy8wNjUuLzYAAAEAHAAAAf8C2QAGAAABAyMTITUhAf/7r/7+yQHjAlr9pgJHkgAAAwA2//IC
UgLpABoAJgAyAAASNTQ2NjMyFhYVFAYHFhYVFAYGIyImJjU0Njc2JiMiBhUUFjMyNjUGBhUUFjMyNjU0JiNHOHJTU3E4Mys2Okh7
S0t7SDo19S8nJy8wJiYwhTo5MDA3OS4Bums2WTU1WTY2UBcaWz1HZzY2Z0c+WxmsLS4pJy0uJ80zLSo2NiosNAACADT//QI5AusA
GwAnAAA2MzI2NQYGIyImJjU0NjYzMhYVFAYGIyImJiczNjY1NCYjIgYVFBYz9E01LBZML0BkOT1xTZR2L29gTGs5BJ59MjQsLDMz
MIdeah8iNWlLSm48uLOCqVg8Yju/MywwNDUuKzUA//8AKP/4APMCSQAiABEAAAAHABEAAAGX//8AMv9/ARACQQAnABEAHQGPAAIA
DxgAAAEAOQCBAdsCOwAFAAAlJzczBxcBBMvL18zMgd3d3d0AAgBEAJMCdAIpAAMABwAAARUhNQUVITUCdP3QAjD90AIpl5f/l5cA
AAEAPwCBAeECOwAFAAATMxcHIzc/18vL18wCO93d3QAAAgAa//gB9wLlABcAIwAAABYVFAYHByMnMzI2NTQmIyIGFSMmNjYzAiY1
NDYzMhYVFAYjAXaBe2UFlgU8TVEmIiQooQI3bU1vOTktLDk5LALlcGZeaQFLuiYyIygpIz1iOf0TNScoNjYoJzUAAgBJ/x8D7wKc
ADcARAAAABYWFRQGBiMiJicGBiMiJjU0NjYzMhc3MwcGFRQzMjY1NCYjIgYGFRQWMzI3FwYjIiYmNTQ2NjMCNjY1NCYjIgYVFBYz
AuCwXztzTzVFDBxYNEpVPGpCVB4JiiwDJzAzkYNwt2iNf2NOG2J6cq1ejvmaUjEdJiEsOiIhApxZoWlVklcqJicrWU1JgU48NfsS
Dix7UHOCbLtweYYkbStYomyX9Yv92CI5ISIrTDQhKAAAAgAQAAAC0gK+AAcACgAAJSEHIxMzEyMLAgHz/voqs/7G/rVWV1Z8fAK+
/UIBAAEB/v8AAAMAPgAAAm4CvgAOABcAIAAAABYVFAYjIREhMhYVFAYHJzMyNjU0JiMjEjY1NCYjIxUzAiRKe27+uQE8a3lBNv5w
Ki0tKnCpLzErfH4BW189WGcCvmJUPlIQOiUkJCb+VSclJSqbAAEAIf/7As8CxgAbAAASNjYzMhYXIyYmIyIGFRQWMzI2NzMGBiMi
JiY1IVqhZn2yHrwVTTFPYmJPMU0VvB6yfWahWgHIo1uEciwublxcbi4scoNbomgAAgA+AAACtQK+AAoAEwAAABYWFRQGBiMhESES
NjU0JiMjETMBtKZbW6du/vkBB1ZsbGFRUQK+WJ9oZ6BYAr791mphYWz+aAAAAQA+AAAB7AK+AAsAABMVMxUjFSEVIREhFenl5QED
/lIBrgI1joSaiQK+iQABAD4AAAIHAr4ACQAAARUhFTMVIxEjEQIH/uLW1qsCvomUhf7kAr4AAQAh//sC2QLGAB8AAAEmJiMiBhUU
FjMyNjcjNSEVDgIjIiYmNTQ2NjMyFhcCEBNHMFNkaVw/VxTZAXQTW4xYaKNbW6Jofq0cAeAjJW1bYW1APH6fQG5EW6JoaKNbemwA
AQA+AAACngK+AAsAAAERIxEhESMRMxEhEQKeq/72q6sBCgK+/UIBIf7fAr7+7QETAAABAD4AAADpAr4AAwAAExEjEemrAr79QgK+
AAEAGv/5AfYCvgAPAAABERQGIyImNTMUFjMyNjURAfZ9am+GqiQiHyICvv4ib3h+dCwtKCYB3gAAAQA+AAACpgK+AAoAACEDESMR
MxETMwEBAdXsq6vqyf7wARoBNv7KAr7+zAE0/qj+mgABAD4AAAHJAr4ABQAANzMVIREz6eD+dauEhAK+AAABAD4AAANZAr4ADAAA
AREjEQMjAxEjETMTEwNZq52KnqvKxcMCvv1CAaX+WwGm/loCvv4aAeYAAAEAPgAAArICvgAJAAAhIwERIxEzAREzArKr/uKrqwEe
qwGx/k8Cvv5NAbMAAgAh//kC8QLIAA8AGwAABCYmNTQ2NjMyFhYVFAYGIzY2NTQmIyIGFRQWMwEnpWFhpWNjpV9gpGNUZWVUVWVl
VQdcpWdnpFxcpGdnpVyccFxdb25eXW8AAAIAPgAAAlUCvgAMABUAAAAGBiMjFSMRITIWFhUGNjU0JiMjFTMCVTh0VmqrARVUdDre
MDAxXV0Bn2U+/AK+OmZCWC4qKi6wAAIAIf+IAwECyAATAB8AAAUnBiMiJiY1NDY2MzIWFhUUBgcXABYzMjY1NCYjIgYVAjJaKCZj
pWFhpWNjpV9JQZr9z2VVVGVlVFVleHkIXKVnZ6RcXKRnWpUxuQF8b3BcXW9uXgACAD4AAAJlAr4ADgAXAAAhAyMRIxEhMhYWFRQG
BxMBMzI2NTQmIyMBpJIpqwEfU3U6T02i/oRqLy8vL2oBCf73Ar46ZT5Gbhf+6gGCLiooLgAAAQAq//kCPALIACoAABYmJiczFhYz
MjY1NCYmJy4CNTQ2MzIWFyMmJiMiBhUUFhceAhUUBgYj73pJArYELiUmLCMzL0RWPo5ydI4FuQIuJB8mQEREVT49dE8HMmJFJykj
HxoiFg4VKlJCYm9vYyInIR8iJhYXKlA/PGI6AAEAGAAAAjcCvgAHAAABFSMRIxEjNQI3uqu6Ar6J/csCNYkAAAEAO//5AocCvgAT
AAATERQWMzI2NREzERQGBiMiJiY1EeY+PDxAq1CHU1ODTAK+/lw/REQ/AaT+XV6CQkGDXgGjAAEACQAAAtECvgAGAAABAyMDMxMT
AtH51vm2rq8Cvv1CAr797gISAAABABkAAAQEAr4ADAAAAQMjAwMjAzMTEzMTEwQEt89wdM+yt2V9vHhmAr79QgHO/jICvv4BAf/+
AQH/AAABABgAAAKzAr4ACwAAIScHIxMDMxc3MwMTAeyPfsLh5seNfMLf6NfXAWUBWdTU/p7+pAAAAQAHAAACmAK+AAgAAAEDFSM1
AzMTEwKY86vzwoiHAr7+KujoAdb+2gEmAAABADIAAAIiAr4ACQAANyEVITUBITUhFfQBLv4QASz+1AHwjIyCAbCMggABAHj/LwGk
A5EABwAAARUjETMVIREBpIyM/tQDkYr8s4sEYgABALf/RQJoA7YAAwAABQEzAQHD/vSlAQy7BHH7jwABAFr/LwGGA5EABwAABSE1
MxEjNSEBhv7UjIwBLNGLA02KAAABAB4ApQKnAsAABgAANyMTMxMjA8qs8qbxrJilAhv95QFVAAABAFX/MwLB/9cAAwAABRUhNQLB
/ZQppKQAAAEACgJUAPsDPgADAAATFSc1+/ECynZlhQAAAgAc//gCaQI2ABIAHgAAEjY2MzIWFzUzESM1BgYjIiYmNSQmIyIGFRQW
MzI2NRxBcEU7WRirqxlZO0RwQQGiRzMzR0czM0cBboJGMCdP/dJPJzBHg1Y/SklAQEtKQAAAAgA+//gCiwLkABIAHgAAADYzMhYW
FRQGBiMiJicVIxEzERYmIyIGFRQWMzI2NQEBWjpFcEFBcEU7WBmrq/RHNDNHRzMzSAIGMEaCVlaDRy8nTgLk/vuHSUpAQEpLQAAA
AQAc//gCQAI2ABkAABI2NjMyFhcjJiMiBhUUFjMyNzMGBiMiJiY1HEd/UmmNFrYXQi84OC9CF7YWjmhSf0cBboJGbmRASURESUBi
cEaCVwACABz/+AJpAuQAEgAeAAASNjYzMhYXETMRIzUGBiMiJiY1JCYjIgYVFBYzMjY1HEFwRTdbGqurGFk7RXBBAaJHMzNHRzMz
RwFugkYuJwED/RxQKDBHg1Y/SklAQEtKQAACABz/+AJMAjYAGAAfAAAAByEWFjMyNzMOAiMiJiY1NDY2MzIWFhUnNCYjIgYHAkwD
/n0EOytAGbYOSW5EUoBIR4BTUX5HrzwtKzsHAQgaNDc2N1gyRoJXV4JGRH5ULSw0Mi4AAQARAAABVgMQABMAAAEjESMRIzUzNTQ2
MzIXFSYGFRUzAVZcqz4+dnITCTEoXAGg/mABoI4QZ2sBkQMiLAUAAAIAHP7tAmkCNgAfACsAAAAWFzUzERQGBiMiJiczFhYzMjY1
NQYGIyImJjU0NjYzFiYjIgYVFBYzMjY1AU1ZGKs9fV18mgupCDwsNT8ZWDtFcEFBcEWsRzMzR0czM0cCNjAnT/3TTX1KdWQgJT1D
TycxR4NWVoJG30pJQEBLSkAAAQA+AAACaQLkABMAAAAWFREjETQmIyIGFREjETMRNjYzAfV0qjoxMTqrqxpaOAI0f2/+ugEvOD4+
OP7RAuT+/yUsAAACAC4AAAD5AyEACwAPAAASJjU0NjMyFhUUBiMXESMRZzk5LSw5OSxVqwJoNScoNTUoJzU6/dICLgAAAv/h/vYA
+AMhAAsAFwAAEiY1NDYzMhYVFAYjExQGIyM1MzI2NREzZjk5LSw5OSxWaF1DKRwYqwJoNScoNTUoJzX9UGdbkRYZAngAAAEAPgAA
AmcC5AAKAAAhJxUjETMRNzMDEwGTqqurqdPo6urqAuT+Z+P+6P7qAAABAD4AAADpAuQAAwAAExEjEemrAuT9HALkAAEAPgAAA+kC
NAAiAAAAFhURIxE0JiMiBhURIxE0JiMiBhURIxEzFTY2MzIWFzY2MwNue6o5MjI5qjkyMjmrqxpUNT9jHB1kOwI0fnD+ugEvNjs7
Nv7RAS82Ozs2/tECLkYjKTYyLjoAAQA+AAACaQI0ABMAAAAWFREjETQmIyIGFREjETMVNjYzAfR1qjoxMTqrqxpYNwI0f2/+ugEv
OD4+OP7RAi5KJSsAAgAc//gCYAI2AA8AGwAAFiYmNTQ2NjMyFhYVFAYGIzY2NTQmIyIGFRQWM+qDS0yEUlKETE2FUjFFQzEyQkEx
CEaCV1aDRkaDVlaDRpRIQ0NIR0RDSAACAD7+9gKLAjYAEgAeAAAANjMyFhYVFAYGIyImJxEjETMVFiYjIgYVFBYzMjY1AQJYO0Vw
QUFwRTpZGaur9Ec0M0dHMzNIAgYwRoJWVoNHMCb+qAM4T4dJSkBASktAAAACABz+9gJpAjQAEgAiAAASNjYzMhYXNTMRIxEGBiMi
JiY1JCYmIyIGBhUUFhYzMjY2NRxFdUU1VBqrqx1TNUR0RQGiIzgfHjgkJDgeHjkjAWmCSSUhQPzIAVAjKkiCUy49Hx8+LS09Hx8+
LQABAD4AAAGYAjQADAAAADYzFSMiBhURIxEzFQEHWjcvQECrqwH/NbU3Rf79Ai5dAAABACD/+AIDAjYAKgAAFiYmJzMWFjMyNjU0
JicuAjU0NjYzMhYXIyYmIyIGFRQWFx4CFxQGBiPYckIEqQMwIyAjMjg8UDo1Z0dpeQmeBCsjHiAzNj5OOwE3Z0QIMlc2HSQZFBgX
Dg4fQzkwTy5oVh0iFxQYGAwQH0U6MU0sAAABABUAAAF1ArYAEwAAJRUjIiY1NSM1MzUzFTMVIxUUFjMBdVddaEREq3BwGByRkVtn
3o6IiI7gGRYAAQA5//oCZQIuABQAAAERIzUGBiMiJiY1ETMRFBYzMjY1EQJlqxpZNkBiNqo6MTI6Ai790kwlLTlsSQFG/tE4Pj44
AS8AAAEACQAAAmgCLgAGAAAlEzMDIwMzATl5tsbTxregAY790gIuAAEABAAAA18CLgAMAAABAyMDAyMDMxMTMxMTA1+XvVhbvJir
T161X04CLv3SAWr+lgIu/nEBj/5zAY0AAAEABQAAAkgCLgALAAAhJwcjEwMzFzczAxMBiGtaubi9wGtaubvAm5sBHQERmpr+5/7r
AAAB///+9wJ4Ai4ABwAAAQEjEwMzExMCeP6iuIDjv4GAAi78yQEcAhv+owFdAAEAHwAAAdACLgAJAAA3MxUhNRMjNSEV3vL+T+jm
AaqNjYgBGY2IAAABAFv/LQGsA5IAMQAAEzY2NTQnJiY1NDYzMxUjIgYVFBcWFRQGBxUWFhUUBwYVFBYzMxUjIiY1NDY3NjU0Jidb
MS4KAQpoV0gnIBwICTU2NjUJCBwgJ0hXaAoBCi4xAaQCJyQbUAhXI1VfjBwdFk5hHjdICgILRzgeYUwXHR2MX1YjVwhQGyIpAgAB
ADz/fwDnAy8AAwAAFyMRM+erq4EDsAABAEf/LQGYA5IAMQAAAAYVFBcWFhUUBiMjNTMyNjU0JyY1NDY3NSYmNTQ3NjU0JiMjNTMy
FhUUBgcGFRQWFxUBZy4KAQpoV0gnIBwICTU2NjUJCBwgJ0hXaAoBCi4xARopIhtQCFcjVl+MHR0XTGEeOEcLAgpINx5hThYdHIxf
VSNXCFAbJCcCiAABABwA0AJhAakAFwAAEjYzMhYXFhYzMjczBgYjIiYnJiYjIgcjLmdOIDMhGCISNgl/EmhOIDMhGSESNgp9AT9q
ExMODkJvahQSDg5CAAACAF//XQEqAjwACwAPAAASFhUUBiMiJjU0NjMDMxMj8Tk5LC05OS1LkRS5Ajw1Jyg2NignNf8A/iEAAQAq
/6sCMgJVAB4AACQ2NzMGBgcVIzUmJjU0Njc1MxUWFhcjJiMiBhUUFjMBXTAInQd6WkBwfX1wQFp6B50TSDM6OjOFHhtTZwhRUQqL
b2+LClFRCGdTOUE6OkEAAAEAJP/8AkwC6wAkAAAlFSEnNjY1NSM1MyY1NDY2MzIWFyMmJiMiBhUUFhczFSMVFAYHAkz+AhgvLG1R
Ez1uR3R8BJ0CKCkoLAsLtZwgIoKGVCpeRQJzRDlEZDR+ZSgvMTAXMCVzAydUHwACADwAlgIFAmAAGwAnAAAABxcHJwYjIicHJzcm
NTQ3JzcXNjMyFzcXBxYVBjY1NCYjIgYVFBYzAfIjNkJAKzc3LEBCNyEiOEJDLTQyLUJCNiOyKyogHyoqHwE9LjZDQBISQEM3Ljw+
LThDQxEQQkM2Lj9SKSgoKiooKCkAAQAHAAACmAK+ABYAAAEzFQcHMxUjFSM1IzUzJyc1MwMzExMzAglkoCDAyKvIwCCgZI/CiIfA
AalzAT1yhoZyPQFzARX+yQE3AAACAGn/fwEUAy8AAwAHAAATMxEjEyMRM2mrq6urqwMv/mb96gGaAAIAHv8gAi0CyAAvADsAACQV
FAYGIyImJzMWMzI1NCYmJyYmNTQ2NyY1NDY2MzIWFyMmIyIVFBYWFxYWFRQGByY2NTQmIyIGFRQWMwIJN2hHZXgEqQQ2OBs/P1JU
MzhHN2hHZXgEqQQ2OBs/P1JUMzh2LS0mJi0tJiNQM1IuZ1ZFMRUgIhsjWT0yTxU7UDNSLmdWRTEVICIbI1k9MUsaSC0hIS0tISEt
AAIACgJsAWUC/wALABcAABImNTQ2MzIWFRQGIzImNTQ2MzIWFRQGIzQqKh8gKisfqSoqHyAqKx8CbCofHysqIB8qKh8fKyogHyoA
AwAm//kC5gLIAA8AHwA6AAAAFhYVFAYGIyImJjU0NjYzDgIVFBYWMzI2NjU0JiYjDgIVFBYWMzI2NyMGIyImNTQ2MzIWFzMmJiMB
7aBZWaBnZ6BZWaBnWYRHR4RZWYRHR4RZP2U5OWVAT3IQkxEyJSoqJRkjB5MQcVACyFukaGikXFykaGikWzhKiVxbiktLiltciUpO
N2ZERGY3WkoxODY2OBoXTFgAAgAZAV0BtwLHABAAHAAAEjY2MzIXNTMRIzUGIyImJjUWFjMyNjU0JiMiBhUZLEwuTimBgSdNME0s
hSwgICwsICAsAklSLDEs/qAqLyxSNyYqKyUlKislAAIALQB4AgAB1gAFAAsAADcnNzMHFzMnNzMHF4FUVJ5XV0NUVJ5XV3ivr6+v
r6+vrwAAAQAvALQClgG6AAUAAAERIzUhNQKWq/5EAbr++oSCAP//ADQBGAHsAaYAAgAQAAAABAAnARkB2gLJAA8AGwApADAAAAAW
FhUUBgYjIiYmNTQ2NjMSNjU0JiMiBhUUFjM2BgcXIycjFSM1MzIWFQczMjU0IyMBQGM3N2M/P2M4OGM/SFZWSEhWVkhhFRIzUigG
SG8jKnUcEBAcAsk3Yj8+Yzc3Yz4/Yjf+iFhISFhYSEhYsyAIUkpK1CYfFA8OAAEACgJ0AZoC6gADAAABFSE1AZr+cALqdnYAAgAZ
AVsBtgLFAAwAGAAAEjYzMhYVFAYjIiYmNSQmIyIGFRQWMzI2NRlyXV1xcVw9XjUBGSweHywtHx4rAmRhYVNTYy1SNyouLisrLS0r
AAABADwAbwI6AnIADwAAARUzFSE1MzUjNTM1MxUzFQGMrv4Crq6uoq4BZV+Xl1+YdXWYAAEAIgFgAWECygAWAAATMxUhNTc2NTQm
IyIHIzY2MzIWFRQGB66u/st6OA4LGwOAA1tNSEw2QAHEZGBKIyUNDyQ6RkItJDYfAAEADwFcAV0CygAlAAASNjMyFhUUBgcVFhUU
BiMiJiczFjMyNjU0JiMjNTMyNTQmIyIHIxdUUkVVJCBKWEhVVQSFAx4OERUSIhwnEA8cBHoChEY7Lh8mBwMUPS43QDofEBARETgh
Dg8fAAABAAoCVAD7Az4AAwAAEwc1N/vx8QK5ZXZ0AAEAS/72AnYCLgAUAAABESM1BgYjIicRIxEzERQWMzI2NRECdqsZVzkXFqqq
OTIyOQIu/dJSKS8I/vQDOP7RNjs7NgEvAAABABwAAAJ5Ar4AEAAAISMRIxEjESMiJiY1NDY2MyECeY0+jRdRazIya1EBbwI8/cQB
ETthOjpiOwD//wAyAN4A/QGYAAcAEQAKAOYAAQAK/vUBLwAHABIAABYWFRQGIyM1MzI2NTQmIyM1MxXhTlI5moMUFBQUR2clOTk5
O14MDg4LgSwAAAEAFAFgANECwAAFAAATNTMRIzUUvYECXWP+oP0AAAIAGQFbAbECxQAPABsAABI2NjMyFhYVFAYGIyImJjUkJiMi
BhUUFjMyNjUZM108PF0zNF08PFwzARQqHx4pKB8fKgJGUS4uUTU1Ui8vUjUnKiooJyorJgACADIAeAIFAdYABQALAAATJzMXByMl
JzMXByOJV55UVJ4BOFeeVFSeASevr6+vr6+vAAQAKAAAAvQCwAAFAAkAFAAXAAATNTMRIzUlASMBASMVIzUjNTczFTMnBzMovYEC
Gv5/pQGBARsogcOSsiijTk4CXWP+oP1h/UICvv2APj5E3tmEgQAAAwAo//8C+gLAAAUACQAgAAATIzUzESMBASMBEzMVITU3NjU0
JiMiByM2NjMyFhUUBgdkPL2BAhr+f6UBgW6u/st6OA4LGwOAA1tNSEw2QAJdY/6gAWD9QALA/aNkYEojJQ0PJDpGQi0kNh8ABAAy
AAADZwLKACUAKQA0ADcAABI2MzIWFRQGBxUWFRQGIyImJzMWMzI2NTQmIyM1MzI1NCYjIgcjJQEjAQEjFSM1IzU3MxUzJwczOlRS
RVUkIEpYSFVVBIUDHg4RFRIiHCcQDxwEegK5/n+lAYEBGyiBw4e9KKNOTgKERjsuHyYHAxQ9LjdAOh8QEBEROCEODx9w/UACwP1+
Pj5E3tmOjgAAAgAi/08B/wI8AAsAIwAAABYVFAYjIiY1NDYzAiY1NDY3NzMXIyIGFxQWMzI2NTMWBgYjAX05OS0sOTksrYF7ZQWW
BTxOUQEnISQooQI3bU0CPDUnKDY2KCc1/RNwZl5pAUu6JzMiJykjPWI5AP//ABAAAALSA9AAIgAkAAAABwBDAOoAkv//ABAAAALS
A9AAIgAkAAAABwB2AOoAkv//ABAAAALSA6sAIgAkAAAABwDGAMwAkv//ABAAAALSA4MAIgAkAAAABwDIAKQAkv//ABAAAALSA5EA
IgAkAAAABwBqALkAkv//ABAAAALSA9sAIgAkAAAABwDHAO4AkgACAAYAAAOKAr4ADwASAAABFTMVIxUzFSE1IwcjASEVAREDAo3f
3/3+WOc8uQFnAh3+WKUCN5CCnod2dgK+h/7BAUf+uQAAAQAh/vECzwLGAC0AACQ2NzMGBgcVMhYVFAYjIzUzMjY1NCYjIzUuAjU0
NjYzMhYXIyYmIyIGFRQWMwGxTRW8G5NpNE5SOZqDFBQUFEdaikxaoWZ9sh68FU0xT2JiT5YuLGV/Dic5OTk7XgwODgt8Cl+aX2ij
W4RyLC5uXFxuAP//AD4AAAHsA9AAIgAoAAAABwBDAIoAkv//AD4AAAHsA9AAIgAoAAAABwB2AIoAkv//AD4AAAHsA6sAIgAoAAAA
BwDGAGsAkv//AD4AAAHsA5EAIgAoAAAABwBqAFkAkv//ABcAAAEIA9AAIgAsAAAABwBDAA0Akv//ABcAAAEIA9AAIgAsAAAABwB2
AA0Akv////kAAAEvA6sAIgAsAAAABwDG/+8Akv///+YAAAFBA5EAIgAsAAAABwBq/9wAkgACAAQAAAK6Ar4ADgAbAAAAFhYVFAYG
IyERIzUzESESNjU0JiMjFTMVIxUzAbmnWlqnb/75Pz8BB1ZsbGBRmJhRAr5Yn2hnoFgBGI4BGP3UbGFhbIaOhgD//wA+AAACsgOD
ACIAMQAAAAcAyACrAJL//wAh//kC8QPQACIAMgAAAAcAQwEDAJL//wAh//kC8QPQACIAMgAAAAcAdgEDAJL//wAh//kC8QOrACIA
MgAAAAcAxgDkAJL//wAh//kC8QODACIAMgAAAAcAyAC8AJL//wAh//kC8QORACIAMgAAAAcAagDSAJIAAQBFAGoCOAJYAAsAACUn
Byc3JzcXNxcHFwHIi4lviYVwhYNwhYxqi4pui4VvhYVuhYwAAAMAFv/5AvwCyAAYACAAKAAAARYWFRQGBiMiJicHIzcmJjU0NjYz
Mhc3MwAXASYjIgYVJCcBFjMyNjUCmCouYKRjPG0tNWhkKy9hpWN4XjVo/dMcAQguPFVlAXMb/vguPFRlAlIve0dnpVwiITtuL3xH
Z6RcQjr+XjIBIx5uXkMx/t4ecFz//wA7//kChwPQACIAOAAAAAcAQwDaAJL//wA7//kChwPQACIAOAAAAAcAdgDaAJL//wA7//kC
hwOrACIAOAAAAAcAxgC8AJL//wA7//kChwORACIAOAAAAAcAagCpAJL//wAHAAACmAPQACIAPAAAAAcAdgDJAJIAAgA+AAACVQK+
AA4AFwAAAAYGIyMVIxEzFTMyFhYVBjY1NCYjIxUzAlU4dFZqq6lsVHQ63jAwMV1dASVkPIUCvoI5Y0BaMCoqMLQAAAEAMf/2AsgD
AAA3AAASNjYzMhYWFRQGBw4CFRQWFxYWFRQGBiMiJiczFhYzMjY1NCYnJiY1NDY3NjY1NCYjIgYVESMRMUWAV0tzPygjAxYKIzlN
QzZjQWN6DJ0EJh8bHSIxVkkfHxkYMio1NKoCW2o7MVQxLT8iAxcRBw4YFR1OPDVQLWNcHyQWFhcfER89LR8wIRsjFB8jNjL99gIW
AP//ABz/+AJpAz4AIgBEAAAAAwBDANAAAP//ABz/+AJpAz4AIgBEAAAAAwB2ANAAAP//ABz/+AJpAxkAIgBEAAAAAwDGALIAAP//
ABz/+AJpAvEAIgBEAAAAAwDIAIoAAP//ABz/+AJpAv8AIgBEAAAAAwBqAJ8AAP//ABz/+AJpA0kAIgBEAAAAAwDHANQAAAADABz/
+AQRAjYAKQAwADwAAAAHIRYWMzI2NzMGBiMiJicVIzUGBiMiJiY1NDY2MzIWFzUzFTYzMhYWFScmJiMiBgcGNjU0JiMiBhUUFjME
EQX+YgdDLig0C7YVfFk8XxujG1o5RXBBQXJGOFgdoT15SHE/rwFEMy9ECfdHRzMzR0czAQwcNzgeGlVsLylQUCgwR4NWVoJGLCpO
TlZFflImLzc0MrpKQEBKSUBASwABABz+8QJAAjYAKgAAJDczBgYHFTIWFRQGIyM1MzI2NTQmIyM1JiY1NDY2MzIWFyMmIyIGFRQW
MwFzF7YTclU0TlI5moMUFBQUR2d8R39SaY0WthdCLzg4L4pAVWwNJTk5OTteDA4OC3oQl3RXgkZuZEBJRERJAP//ABz/+AJMA0IA
IgBIAAAABwBDAK4ABP//ABz/+AJMA0IAIgBIAAAABwB2AK4ABP//ABz/+AJMAx0AIgBIAAAABwDGAI8ABP//ABz/+AJMAwMAIgBI
AAAABgBqfAT//wAkAAABFQM+ACIAwgAAAAIAQxoA//8AJAAAARUDPgAiAMIAAAACAHYaAP//AAYAAAE8AxkAIgDCAAAAAgDG/AD/
///zAAABTgL/ACIAwgAAAAIAaukAAAIAHP/5AmQC5AAcACgAAAAWFRQGBiMiJiY1NDY2MzIXJicHNTcmJzMXNxUHAjY1NCYjIgYV
FBYzAiRATYRSUYZOSX5NPDkWHmk+ISSdHl8xeURGLzBFRjACRa1bZ5JLR4JWVoFFHy8oIEMTIx8eHUMP/fpIRERHR0RESP//AD4A
AAJpAvEAIgBRAAAAAwDIAJwAAP//ABz/+AJgA0IAIgBSAAAABwBDAMUABP//ABz/+AJgA0IAIgBSAAAABwB2AMUABP//ABz/+AJg
Ax0AIgBSAAAABwDGAKcABP//ABz/+AJgAvUAIgBSAAAABgDIfgT//wAc//gCYAMDACIAUgAAAAcAagCUAAQAAwA1AB8CNAKeAAsA
DwAbAAAAJjU0NjMyFhUUBiMXFSE1EiY1NDYzMhYVFAYjAQs5OS0sOTks/P4B1jk5LSw5OSwB5DUnKDY2KCc1OpeX/nU1Jyg2Nign
NQAAAwAN//gCaAI2ABUAHQAlAAABFhUUBgYjIicHIzcmNTQ2NjMyFzczABc3JiMiBhU2JwcWMzI2NQIZQ02FUmFJJVxOQ0yEUmNJ
Jlz+XgupGyUyQugLqRsjMUUB2E5zVoNGMChVTnRWg0YxKf7EHLgUR0QkHLkSSEP//wA5//oCZQM+ACIAWAAAAAMAQwDLAAD//wA5
//oCZQM+ACIAWAAAAAMAdgDLAAD//wA5//oCZQMZACIAWAAAAAMAxgCtAAD//wA5//oCZQL/ACIAWAAAAAMAagCaAAD//////vcC
eAM+ACIAXAAAAAMAdgC1AAAAAgA+/vYCiwLkABAAIAAAADMyFhYVFAYGIyInESMRMxUWJiYjIgYGFRQWFjMyNjY1ASJrRXVERHVF
aDyrq/QjOR4fOCMjOR4fOCMCNEmCU1OCSFj+pQPu/6I+Hx89LS0+Hx89LQD//////vcCeAL/ACIAXAAAAAMAagCEAAAAAQBLAAAA
9gIuAAMAABMRIxH2qwIu/dICLgACACH/9gP3AsgAGQAlAAABMxUjFSEVITUGByImJjU0NjYzMhYXNSEVIQAWMzI2NTQmIyIGFQLz
5uYBBP5RSYNhnlxcn2A+aCYBr/78/d1lVVRlZVRVZQGlgpyHSE8DXKZoaKRcKChGh/7McHBdXW5uXQADABz/+APnAjYAJAArADcA
AAAHIRYWMzI3Mw4CIyImJwYGIyImJjU0NjYzMhYXNjYzMhYWFScmJiMiBgcGNjU0JiMiBhUUFjMD5wT+fQQ4LEEZtg5Jb0Q8ZiMm
aj1Sg0tMhFI8aiUlZj1RfkevATwtLDsG9EVEMTFCQTEBBxc3NjY3WDIrKSkrRoJXVoNGLCgoLER+UyQxNzcxuklDREhIRENJ//8A
MQGyARQCvgACAMwAAAABAAoCVAFAAxkABQAAEwc1NxcVpZubmwKmUnZPT3YAAAIACQJXAP4DSQALABcAABIGIyImNTQ2MzIWFSYm
IyIGFRQWMzI2Nf5HNDRGRjQ0R0sbFRQbGxQVGwKZQkI3N0JCNxYbGxYVHBwVAAABAAUCXQGVAvEAFQAAEjMyFhcWFjMyNzMGIyIm
JyYmIyIHIx1yFiEYEhYMJQZYGHIWIRgSFgwlB1cC8Q0NCgktlA0NCgktAAABAEYBGAJxAaYAAwAAARUhNQJx/dUBpo6OAAEARgEY
A1wBpgADAAABFSE1A1z86gGmjo4AAQAmAbIBCQK+AAMAABMjEzPVr3BzAbIBDAABADEBsgEUAr4AAwAAEzMDI2WvcHMCvv70AAEA
Gf+XAPwAowADAAA3MwMjTa9wc6P+9AAAAgAmAbIB7AK+AAMABwAAEyMTMxMjEzPVr3Bzr65wcgGyAQz+9AEMAAIAMQGyAfcCvgAD
AAcAABMzAyMBMwMjZa9wcwEXr3BzAr7+9AEM/vQAAAIAGv+XAd8AowADAAcAADczAyMBMwMjTq9wcwEXrnByo/70AQz+9AABAGQA
twGRAeQADwAANiYmNTQ2NjMyFhYVFAYGI9FFKChFKSpFKChFKrcoRSkqRSgoRSopRSgA//8AKP/4AtQAsgAjABEA8AAAACMAEQHh
AAAAAgARAAAAAQAtAHgBHwHWAAUAADcnNzMHF4FUVJ5XV3ivr6+vAAABADIAeAEkAdYABQAAEyczFwcjiVeeVFSeASevr68AAAEA
CgAAAjACwAADAAABASMBAjD+f6UBgQLA/UACwAAAAQAU//sC/wLGAC0AACQ2NzMGBiMiJicjNTMmNTQ3IzUzNjYzMhYXIyYmIyIG
BzMVIwYVFBczFSMWFjMB4UwWvB6yfXyzIU4/AQE/TyGze32yHrwWTDEwTBenwwEBw6kXTTGXLityg4JvWQkSEwpaboGEcistKida
CRQSCVkpLAACABsBYwNGAsIABwAUAAATIRUjFyM3IyURIycHIycVIxEzFzcbASBTAXwBUwMrcQNSU1JtpkdSAsJj/Pxj/qHJydPT
AV/JyQABADQBGAHsAaYAAwAAARUhNQHs/kgBpo6O//8ACv9FAbsDtgACABIAAAAAAAAACABmAAMAAQQJAAAAogC8AAMAAQQJAAEA
DgCuAAMAAQQJAAIACACmAAMAAQQJAAMANgBwAAMAAQQJAAQAGABYAAMAAQQJAAUACgBOAAMAAQQJAAYAGAA2AAMAAQQJAA4ANgAA
AGgAdAB0AHAAcwA6AC8ALwBzAGMAcgBpAHAAdABzAC4AcwBpAGwALgBvAHIAZwAvAE8ARgBMAFAAbwBwAHAAaQBuAHMALQBCAG8A
bABkADQALgAwADAANABQAG8AcABwAGkAbgBzACAAQgBvAGwAZABJAFQARgBPADsAIABQAG8AcABwAGkAbgBzACAAQgBvAGwAZAA7
ACAANAAuADAAMAA0AGIAOABCAG8AbABkAFAAbwBwAHAAaQBuAHMAQwBvAHAAeQByAGkAZwBoAHQAIAAyADAAMgAwACAAVABoAGUA
IABQAG8AcABwAGkAbgBzACAAUAByAG8AagBlAGMAdAAgAEEAdQB0AGgAbwByAHMAIAAoAGgAdAB0AHAAcwA6AC8ALwBnAGkAdABo
AHUAYgAuAGMAbwBtAC8AaQB0AGYAbwB1AG4AZAByAHkALwBQAG8AcABwAGkAbgBzACkAAgAAAAAAAP+1ADIAAAAAAAAAAAAAAAAA
AAAAAAAAAADaAAABAgEDAAMABAAFAAYABwAIAAkACgALAAwADQAOAA8AEAARABIAEwAUABUAFgAXABgAGQAaABsAHAAdAB4AHwAg
ACEAIgAjACQAJQAmACcAKAApACoAKwAsAC0ALgAvADAAMQAyADMANAA1ADYANwA4ADkAOgA7ADwAPQA+AD8AQABBAEIAQwBEAEUA
RgBHAEgASQBKAEsATABNAE4ATwBQAFEAUgBTAFQAVQBWAFcAWABZAFoAWwBcAF0AXgBfAGAAYQEEAKMAhACFAL0AlgDoAIYAjgCL
AJ0AqQCkAQUAigDaAIMAkwDyAPMAjQEGAIgAwwDeAPEAngCqAPUA9AD2AKIArQDJAMcArgBiAGMAkABkAMsAZQDIAMoAzwDMAM0A
zgDpAGYA0wDQANEArwBnAPAAkQDWANQA1QBoAOsA7QCJAGoAaQBrAG0AbABuAKAAbwBxAHAAcgBzAHUAdAB2AHcA6gB4AHoAeQB7
AH0AfAC4AKEAfwB+AIAAgQDsAO4AugDXALAAsQEHANgA3QDZALIAswC2ALcAxAC0ALUAxQCHAKsAvgC/ALwBCACMAO8BCQROVUxM
AkNSB3VuaTAwQTAHdW5pMDBBRAd1bmkwMEI1B3VuaTAyQkMERXVybwd1bmkyMjE1AAABAAAACgAcABwAAURGTFQACAAEAAAAAP//
AAAAAAAAAAEAAAAKACwALAADREZMVAAYZGV2MgAUZGV2YQAUAAAAAAAEAAAAAP//AAAAAAAA
""",
    "OFL-Poppins.txt": """
Q29weXJpZ2h0IDIwMjAgVGhlIFBvcHBpbnMgUHJvamVjdCBBdXRob3JzIChodHRwczovL2dpdGh1Yi5jb20vaXRmb3VuZHJ5L1Bv
cHBpbnMpCgpUaGlzIEZvbnQgU29mdHdhcmUgaXMgbGljZW5zZWQgdW5kZXIgdGhlIFNJTCBPcGVuIEZvbnQgTGljZW5zZSwgVmVy
c2lvbiAxLjEuClRoaXMgbGljZW5zZSBpcyBjb3BpZWQgYmVsb3csIGFuZCBpcyBhbHNvIGF2YWlsYWJsZSB3aXRoIGEgRkFRIGF0
OgpodHRwOi8vc2NyaXB0cy5zaWwub3JnL09GTAoKCi0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0t
LS0tLS0tLS0tLS0tLS0tClNJTCBPUEVOIEZPTlQgTElDRU5TRSBWZXJzaW9uIDEuMSAtIDI2IEZlYnJ1YXJ5IDIwMDcKLS0tLS0t
LS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0tLS0KClBSRUFNQkxFClRoZSBnb2FscyBv
ZiB0aGUgT3BlbiBGb250IExpY2Vuc2UgKE9GTCkgYXJlIHRvIHN0aW11bGF0ZSB3b3JsZHdpZGUKZGV2ZWxvcG1lbnQgb2YgY29s
bGFib3JhdGl2ZSBmb250IHByb2plY3RzLCB0byBzdXBwb3J0IHRoZSBmb250IGNyZWF0aW9uCmVmZm9ydHMgb2YgYWNhZGVtaWMg
YW5kIGxpbmd1aXN0aWMgY29tbXVuaXRpZXMsIGFuZCB0byBwcm92aWRlIGEgZnJlZSBhbmQKb3BlbiBmcmFtZXdvcmsgaW4gd2hp
Y2ggZm9udHMgbWF5IGJlIHNoYXJlZCBhbmQgaW1wcm92ZWQgaW4gcGFydG5lcnNoaXAKd2l0aCBvdGhlcnMuCgpUaGUgT0ZMIGFs
bG93cyB0aGUgbGljZW5zZWQgZm9udHMgdG8gYmUgdXNlZCwgc3R1ZGllZCwgbW9kaWZpZWQgYW5kCnJlZGlzdHJpYnV0ZWQgZnJl
ZWx5IGFzIGxvbmcgYXMgdGhleSBhcmUgbm90IHNvbGQgYnkgdGhlbXNlbHZlcy4gVGhlCmZvbnRzLCBpbmNsdWRpbmcgYW55IGRl
cml2YXRpdmUgd29ya3MsIGNhbiBiZSBidW5kbGVkLCBlbWJlZGRlZCwKcmVkaXN0cmlidXRlZCBhbmQvb3Igc29sZCB3aXRoIGFu
eSBzb2Z0d2FyZSBwcm92aWRlZCB0aGF0IGFueSByZXNlcnZlZApuYW1lcyBhcmUgbm90IHVzZWQgYnkgZGVyaXZhdGl2ZSB3b3Jr
cy4gVGhlIGZvbnRzIGFuZCBkZXJpdmF0aXZlcywKaG93ZXZlciwgY2Fubm90IGJlIHJlbGVhc2VkIHVuZGVyIGFueSBvdGhlciB0
eXBlIG9mIGxpY2Vuc2UuIFRoZQpyZXF1aXJlbWVudCBmb3IgZm9udHMgdG8gcmVtYWluIHVuZGVyIHRoaXMgbGljZW5zZSBkb2Vz
IG5vdCBhcHBseQp0byBhbnkgZG9jdW1lbnQgY3JlYXRlZCB1c2luZyB0aGUgZm9udHMgb3IgdGhlaXIgZGVyaXZhdGl2ZXMuCgpE
RUZJTklUSU9OUwoiRm9udCBTb2Z0d2FyZSIgcmVmZXJzIHRvIHRoZSBzZXQgb2YgZmlsZXMgcmVsZWFzZWQgYnkgdGhlIENvcHly
aWdodApIb2xkZXIocykgdW5kZXIgdGhpcyBsaWNlbnNlIGFuZCBjbGVhcmx5IG1hcmtlZCBhcyBzdWNoLiBUaGlzIG1heQppbmNs
dWRlIHNvdXJjZSBmaWxlcywgYnVpbGQgc2NyaXB0cyBhbmQgZG9jdW1lbnRhdGlvbi4KCiJSZXNlcnZlZCBGb250IE5hbWUiIHJl
ZmVycyB0byBhbnkgbmFtZXMgc3BlY2lmaWVkIGFzIHN1Y2ggYWZ0ZXIgdGhlCmNvcHlyaWdodCBzdGF0ZW1lbnQocykuCgoiT3Jp
Z2luYWwgVmVyc2lvbiIgcmVmZXJzIHRvIHRoZSBjb2xsZWN0aW9uIG9mIEZvbnQgU29mdHdhcmUgY29tcG9uZW50cyBhcwpkaXN0
cmlidXRlZCBieSB0aGUgQ29weXJpZ2h0IEhvbGRlcihzKS4KCiJNb2RpZmllZCBWZXJzaW9uIiByZWZlcnMgdG8gYW55IGRlcml2
YXRpdmUgbWFkZSBieSBhZGRpbmcgdG8sIGRlbGV0aW5nLApvciBzdWJzdGl0dXRpbmcgLS0gaW4gcGFydCBvciBpbiB3aG9sZSAt
LSBhbnkgb2YgdGhlIGNvbXBvbmVudHMgb2YgdGhlCk9yaWdpbmFsIFZlcnNpb24sIGJ5IGNoYW5naW5nIGZvcm1hdHMgb3IgYnkg
cG9ydGluZyB0aGUgRm9udCBTb2Z0d2FyZSB0byBhCm5ldyBlbnZpcm9ubWVudC4KCiJBdXRob3IiIHJlZmVycyB0byBhbnkgZGVz
aWduZXIsIGVuZ2luZWVyLCBwcm9ncmFtbWVyLCB0ZWNobmljYWwKd3JpdGVyIG9yIG90aGVyIHBlcnNvbiB3aG8gY29udHJpYnV0
ZWQgdG8gdGhlIEZvbnQgU29mdHdhcmUuCgpQRVJNSVNTSU9OICYgQ09ORElUSU9OUwpQZXJtaXNzaW9uIGlzIGhlcmVieSBncmFu
dGVkLCBmcmVlIG9mIGNoYXJnZSwgdG8gYW55IHBlcnNvbiBvYnRhaW5pbmcKYSBjb3B5IG9mIHRoZSBGb250IFNvZnR3YXJlLCB0
byB1c2UsIHN0dWR5LCBjb3B5LCBtZXJnZSwgZW1iZWQsIG1vZGlmeSwKcmVkaXN0cmlidXRlLCBhbmQgc2VsbCBtb2RpZmllZCBh
bmQgdW5tb2RpZmllZCBjb3BpZXMgb2YgdGhlIEZvbnQKU29mdHdhcmUsIHN1YmplY3QgdG8gdGhlIGZvbGxvd2luZyBjb25kaXRp
b25zOgoKMSkgTmVpdGhlciB0aGUgRm9udCBTb2Z0d2FyZSBub3IgYW55IG9mIGl0cyBpbmRpdmlkdWFsIGNvbXBvbmVudHMsCmlu
IE9yaWdpbmFsIG9yIE1vZGlmaWVkIFZlcnNpb25zLCBtYXkgYmUgc29sZCBieSBpdHNlbGYuCgoyKSBPcmlnaW5hbCBvciBNb2Rp
ZmllZCBWZXJzaW9ucyBvZiB0aGUgRm9udCBTb2Z0d2FyZSBtYXkgYmUgYnVuZGxlZCwKcmVkaXN0cmlidXRlZCBhbmQvb3Igc29s
ZCB3aXRoIGFueSBzb2Z0d2FyZSwgcHJvdmlkZWQgdGhhdCBlYWNoIGNvcHkKY29udGFpbnMgdGhlIGFib3ZlIGNvcHlyaWdodCBu
b3RpY2UgYW5kIHRoaXMgbGljZW5zZS4gVGhlc2UgY2FuIGJlCmluY2x1ZGVkIGVpdGhlciBhcyBzdGFuZC1hbG9uZSB0ZXh0IGZp
bGVzLCBodW1hbi1yZWFkYWJsZSBoZWFkZXJzIG9yCmluIHRoZSBhcHByb3ByaWF0ZSBtYWNoaW5lLXJlYWRhYmxlIG1ldGFkYXRh
IGZpZWxkcyB3aXRoaW4gdGV4dCBvcgpiaW5hcnkgZmlsZXMgYXMgbG9uZyBhcyB0aG9zZSBmaWVsZHMgY2FuIGJlIGVhc2lseSB2
aWV3ZWQgYnkgdGhlIHVzZXIuCgozKSBObyBNb2RpZmllZCBWZXJzaW9uIG9mIHRoZSBGb250IFNvZnR3YXJlIG1heSB1c2UgdGhl
IFJlc2VydmVkIEZvbnQKTmFtZShzKSB1bmxlc3MgZXhwbGljaXQgd3JpdHRlbiBwZXJtaXNzaW9uIGlzIGdyYW50ZWQgYnkgdGhl
IGNvcnJlc3BvbmRpbmcKQ29weXJpZ2h0IEhvbGRlci4gVGhpcyByZXN0cmljdGlvbiBvbmx5IGFwcGxpZXMgdG8gdGhlIHByaW1h
cnkgZm9udCBuYW1lIGFzCnByZXNlbnRlZCB0byB0aGUgdXNlcnMuCgo0KSBUaGUgbmFtZShzKSBvZiB0aGUgQ29weXJpZ2h0IEhv
bGRlcihzKSBvciB0aGUgQXV0aG9yKHMpIG9mIHRoZSBGb250ClNvZnR3YXJlIHNoYWxsIG5vdCBiZSB1c2VkIHRvIHByb21vdGUs
IGVuZG9yc2Ugb3IgYWR2ZXJ0aXNlIGFueQpNb2RpZmllZCBWZXJzaW9uLCBleGNlcHQgdG8gYWNrbm93bGVkZ2UgdGhlIGNvbnRy
aWJ1dGlvbihzKSBvZiB0aGUKQ29weXJpZ2h0IEhvbGRlcihzKSBhbmQgdGhlIEF1dGhvcihzKSBvciB3aXRoIHRoZWlyIGV4cGxp
Y2l0IHdyaXR0ZW4KcGVybWlzc2lvbi4KCjUpIFRoZSBGb250IFNvZnR3YXJlLCBtb2RpZmllZCBvciB1bm1vZGlmaWVkLCBpbiBw
YXJ0IG9yIGluIHdob2xlLAptdXN0IGJlIGRpc3RyaWJ1dGVkIGVudGlyZWx5IHVuZGVyIHRoaXMgbGljZW5zZSwgYW5kIG11c3Qg
bm90IGJlCmRpc3RyaWJ1dGVkIHVuZGVyIGFueSBvdGhlciBsaWNlbnNlLiBUaGUgcmVxdWlyZW1lbnQgZm9yIGZvbnRzIHRvCnJl
bWFpbiB1bmRlciB0aGlzIGxpY2Vuc2UgZG9lcyBub3QgYXBwbHkgdG8gYW55IGRvY3VtZW50IGNyZWF0ZWQKdXNpbmcgdGhlIEZv
bnQgU29mdHdhcmUuCgpURVJNSU5BVElPTgpUaGlzIGxpY2Vuc2UgYmVjb21lcyBudWxsIGFuZCB2b2lkIGlmIGFueSBvZiB0aGUg
YWJvdmUgY29uZGl0aW9ucyBhcmUKbm90IG1ldC4KCkRJU0NMQUlNRVIKVEhFIEZPTlQgU09GVFdBUkUgSVMgUFJPVklERUQgIkFT
IElTIiwgV0lUSE9VVCBXQVJSQU5UWSBPRiBBTlkgS0lORCwKRVhQUkVTUyBPUiBJTVBMSUVELCBJTkNMVURJTkcgQlVUIE5PVCBM
SU1JVEVEIFRPIEFOWSBXQVJSQU5USUVTIE9GCk1FUkNIQU5UQUJJTElUWSwgRklUTkVTUyBGT1IgQSBQQVJUSUNVTEFSIFBVUlBP
U0UgQU5EIE5PTklORlJJTkdFTUVOVApPRiBDT1BZUklHSFQsIFBBVEVOVCwgVFJBREVNQVJLLCBPUiBPVEhFUiBSSUdIVC4gSU4g
Tk8gRVZFTlQgU0hBTEwgVEhFCkNPUFlSSUdIVCBIT0xERVIgQkUgTElBQkxFIEZPUiBBTlkgQ0xBSU0sIERBTUFHRVMgT1IgT1RI
RVIgTElBQklMSVRZLApJTkNMVURJTkcgQU5ZIEdFTkVSQUwsIFNQRUNJQUwsIElORElSRUNULCBJTkNJREVOVEFMLCBPUiBDT05T
RVFVRU5USUFMCkRBTUFHRVMsIFdIRVRIRVIgSU4gQU4gQUNUSU9OIE9GIENPTlRSQUNULCBUT1JUIE9SIE9USEVSV0lTRSwgQVJJ
U0lORwpGUk9NLCBPVVQgT0YgVEhFIFVTRSBPUiBJTkFCSUxJVFkgVE8gVVNFIFRIRSBGT05UIFNPRlRXQVJFIE9SIEZST00KT1RI
RVIgREVBTElOR1MgSU4gVEhFIEZPTlQgU09GVFdBUkUuCg==
""",
}


# ==== INICIO DATOS DE PROYECTOS ====
# Generado por desarrollo/generar_datos_valorizacion.py a partir del Excel de proyectos
# (hoja 'Proy x Mpio'). No editar a mano: regenerar con ese programa.
# estado: '000' no gravado | '001' tiene proyectos | 'NO_APLICA' | 'CONDONADA'

FUENTE_DATOS = 'INFORMACION PROYECTOS DE VALORIZACION (1).xlsx'

MUNICIPIOS_VALORIZACION = {'ABEJORRAL': {'cod_mpio': '002',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1990,
                              'codigo': '007',
                              'fecha_resolucion': '1990-02-20',
                              'obra': 'PUENTE EL ORO SOBRE EL RIO ARMA',
                              'resolucion': '1'}],
               'texto_hoja': '007 - PUENTE EL ORO SOBRE EL RIO ARMA'},
 'ABRIAQUI': {'cod_mpio': '004',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'ALEJANDRIA': {'cod_mpio': '021',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1995,
                               'codigo': '772',
                               'fecha_inicio': '1996-02-28',
                               'fecha_resolucion': '1995-11-29',
                               'fecha_vencimiento': '1999-01-01',
                               'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                               'oficinas_registro': ['018 - MARINILLA'],
                               'resolucion': '772'}],
                'texto_hoja': '772 - GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'AMAGA': {'cod_mpio': '030',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1981,
                          'codigo': '064',
                          'fecha_resolucion': '1981-12-10',
                          'obra': 'CAMILOCE - EL CINCO - FREDONIA',
                          'resolucion': '17'},
                         {'anio': 1991,
                          'codigo': '033',
                          'fecha_resolucion': '1991-04-25',
                          'obra': 'ALBANIA - TITIRIBI',
                          'resolucion': '20'}],
           'texto_hoja': '064 - CAMILOCE - EL CINCO - FREDONIA\n  033 - ALBANIA - TITIRIBI'},
 'AMALFI': {'cod_mpio': '031',
            'estado': '001',
            'nota': '',
            'proyectos': [{'anio': 2015,
                           'codigo': '200',
                           'fecha_inicio': '2016-05-01',
                           'fecha_resolucion': '2015-11-20',
                           'fecha_vencimiento': '2021-04-01',
                           'obra': 'EL LIMON - ANORI',
                           'oficinas_registro': ['003 - AMALFI'],
                           'resolucion': '305531'}],
            'texto_hoja': '200 - EL LIMON - ANORI'},
 'ANDES': {'cod_mpio': '034',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 2002,
                          'codigo': '017',
                          'fecha_inicio': '2002-03-31',
                          'fecha_resolucion': '2002-01-31',
                          'fecha_vencimiento': '2010-02-01',
                          'obra': 'ANDES - TAPARTO',
                          'oficinas_registro': ['004 - ANDES', '005 - CIUDAD BOLIVAR'],
                          'resolucion': '1008'}],
           'texto_hoja': '017 - ANDES - TAPARTO'},
 'ANGELOPOLIS': {'cod_mpio': '036',
                 'estado': '000',
                 'nota': '',
                 'proyectos': [],
                 'texto_hoja': '000 - NO GRAVADO'},
 'ANGOSTURA': {'cod_mpio': '038',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1977,
                              'codigo': '021',
                              'fecha_resolucion': '1977-09-06',
                              'obra': 'GUADALUPE - RIO SAN PABLO',
                              'resolucion': '21'},
                             {'anio': 1973,
                              'codigo': '031',
                              'fecha_resolucion': '1973-03-08',
                              'obra': 'EL ROBLE - MINA VIEJA - MIRAFLORES',
                              'resolucion': '31'}],
               'texto_hoja': '021 - GUADALUPE - RIO SAN PABLO\n  031 - EL ROBLE - MINA VIEJA - MIRAFLORES'},
 'ANORI': {'cod_mpio': '040',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 2015,
                          'codigo': '200',
                          'fecha_inicio': '2016-05-01',
                          'fecha_resolucion': '2015-11-20',
                          'fecha_vencimiento': '2021-04-01',
                          'obra': 'EL LIMON - ANORI',
                          'oficinas_registro': ['003 - AMALFI'],
                          'resolucion': '305531'}],
           'texto_hoja': '200 - EL LIMON - ANORI'},
 'ANZA': {'cod_mpio': '044',
          'estado': '001',
          'nota': '',
          'proyectos': [{'anio': 1996,
                         'codigo': '014',
                         'fecha_inicio': '1996-07-31',
                         'fecha_resolucion': '1996-05-28',
                         'fecha_vencimiento': '2001-06-01',
                         'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                         'oficinas_registro': ['004 - ANDES',
                                               '024 - ANTIOQUIA',
                                               '035 - URRAO',
                                               '036 - URRAO'],
                         'resolucion': '0422'}],
          'texto_hoja': '014 - BOLOMBOLO - SANTA FE DE ANT'},
 'APARTADO': {'cod_mpio': '045',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1986,
                             'codigo': '050',
                             'fecha_resolucion': '1986-07-09',
                             'obra': 'ELECTRIF CHIGORODO - BARRANQUILLITA',
                             'resolucion': '49'}],
              'texto_hoja': '050 - ELECTRIF CHIGORODO - BARRANQUILLITA'},
 'ARBOLETES': {'cod_mpio': '051',
               'estado': 'CONDONADA',
               'nota': 'CONDONADA ORDENANZA 001',
               'proyectos': [{'codigo': '025', 'obra': 'ARBOLETES SAN JUAN DE URABA'}],
               'texto_hoja': '025 ARBOLETES SAN JUAN DE URABA'},
 'ARGELIA': {'cod_mpio': '055',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'ARMENIA': {'cod_mpio': '059',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1996,
                            'codigo': '014',
                            'fecha_inicio': '1996-07-31',
                            'fecha_resolucion': '1996-05-28',
                            'fecha_vencimiento': '2001-06-01',
                            'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                            'oficinas_registro': ['004 - ANDES',
                                                  '024 - ANTIOQUIA',
                                                  '035 - URRAO',
                                                  '036 - URRAO'],
                            'resolucion': '0422'}],
             'texto_hoja': '014 - BOLOMBOLO - SANTA FE DE ANT'},
 'BARBOSA': {'cod_mpio': '079',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 2005,
                            'codigo': '020',
                            'fecha_inicio': '2005-03-31',
                            'fecha_resolucion': '2005-01-25',
                            'fecha_vencimiento': '2010-02-01',
                            'obra': 'DOBLE CALZADA BELLO - HATILLO',
                            'oficinas_registro': ['012 - GIRARDOTA', '01N - MEDELLIN Z.NORTE'],
                            'resolucion': '0707'},
                           {'anio': 2014,
                            'codigo': '100',
                            'fecha_inicio': '2014-09-01',
                            'fecha_resolucion': '2014-08-04',
                            'fecha_vencimiento': '2019-08-01',
                            'obra': 'HATILLO - BARBOSA - PRADERA',
                            'oficinas_registro': ['012 - GIRARDOTA',
                                                  '025 - SANTA ROSA DE O.',
                                                  '026 - SANTO DOMINGO',
                                                  '038 - YOLOMBO'],
                            'resolucion': '120405'},
                           {'anio': 1983,
                            'codigo': '052',
                            'fecha_resolucion': '1983-05-18',
                            'obra': 'GIRARDOTA - CABILDO - HATILLO',
                            'resolucion': '4'}],
             'texto_hoja': '020 - DOBLE CALZADA BELLO - HATILLO\n'
                           '  100 - HATILLO - BARBOSA - PRADERA\n'
                           '  052 - GIRARDOTA - CABILDO - HATILLO'},
 'BELLO': {'cod_mpio': '088',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 2002,
                          'codigo': '019',
                          'fecha_inicio': '2002-08-31',
                          'fecha_resolucion': '2002-07-05',
                          'fecha_vencimiento': '2010-07-01',
                          'obra': 'PAJARITO - SAN PEDRO',
                          'oficinas_registro': ['01N - MEDELLIN Z.NORTE'],
                          'resolucion': '10929'}],
           'texto_hoja': '019 - PAJARITO - SAN PEDRO'},
 'BELMIRA': {'cod_mpio': '086',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 2002,
                            'codigo': '000',
                            'fecha_resolucion': '2002-03-22',
                            'nota_cruce': 'año tomado por nombre de obra (código distinto en hoja Proyectos)',
                            'obra': 'LLANOS DE CUIVA SAN JOSE DE LA MONTANA',
                            'resolucion': '3049'}],
             'texto_hoja': '000 - LLANOS DE CUIVA SAN JOSE DE LA     MONTAÑA'},
 'BETANIA': {'cod_mpio': '091',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 2002,
                            'codigo': '017',
                            'fecha_inicio': '2002-03-31',
                            'fecha_resolucion': '2002-01-31',
                            'fecha_vencimiento': '2010-02-01',
                            'obra': 'ANDES - TAPARTO',
                            'oficinas_registro': ['004 - ANDES', '005 - CIUDAD BOLIVAR'],
                            'resolucion': '1008'}],
             'texto_hoja': '017 - ANDES - TAPARTO'},
 'BETULIA': {'cod_mpio': '093',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1998,
                            'codigo': '008',
                            'fecha_inicio': '1998-10-31',
                            'fecha_resolucion': '1998-09-16',
                            'fecha_vencimiento': '2007-05-01',
                            'obra': 'URRAO - BETULIA',
                            'oficinas_registro': ['035 - URRAO', '036 - URRAO'],
                            'resolucion': '1187'},
                           {'anio': 1996,
                            'codigo': '014',
                            'fecha_inicio': '1996-07-31',
                            'fecha_resolucion': '1996-05-28',
                            'fecha_vencimiento': '2001-06-01',
                            'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                            'oficinas_registro': ['004 - ANDES',
                                                  '024 - ANTIOQUIA',
                                                  '035 - URRAO',
                                                  '036 - URRAO'],
                            'resolucion': '0422'}],
             'texto_hoja': '008 - URRAO - BETULIA\n  014 - BOLOMBOLO - SANTA FE DE ANT'},
 'BOLIVAR': {'cod_mpio': '101',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'BRICENO': {'cod_mpio': '107',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'BURITICA': {'cod_mpio': '113',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'CACERES': {'cod_mpio': '120',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'CAICEDO': {'cod_mpio': '125',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'CALDAS': {'cod_mpio': '129',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'CAMPAMENTO': {'cod_mpio': '134',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1977,
                               'codigo': '021',
                               'fecha_resolucion': '1977-09-06',
                               'obra': 'GUADALUPE - RIO SAN PABLO',
                               'resolucion': '21'}],
                'texto_hoja': '021 - GUADALUPE - RIO SAN PABLO'},
 'CANASGORDAS': {'cod_mpio': '138',
                 'estado': '000',
                 'nota': '',
                 'proyectos': [],
                 'texto_hoja': '000 - NO GRAVADO'},
 'CARACOLI': {'cod_mpio': '142',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'CARAMANTA': {'cod_mpio': '145',
               'estado': '000',
               'nota': '',
               'proyectos': [],
               'texto_hoja': '000 - NO GRAVADO'},
 'CAREPA': {'cod_mpio': '147',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'CARMEN DE VIBORAL': {'cod_mpio': '148',
                       'estado': '000',
                       'nota': '',
                       'proyectos': [],
                       'texto_hoja': '000 - NO GRAVADO'},
 'CAROLINA': {'cod_mpio': '150',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1997,
                             'codigo': '005',
                             'fecha_inicio': '1997-09-30',
                             'fecha_resolucion': '1997-08-11',
                             'fecha_vencimiento': '2002-08-01',
                             'obra': 'PTE GABINO - GOMEZ PLATA - CAROLINA',
                             'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                             'resolucion': '1114'},
                            {'anio': 1973,
                             'codigo': '031',
                             'fecha_resolucion': '1973-03-08',
                             'obra': 'EL ROBLE - MINA VIEJA - MIRAFLORES',
                             'resolucion': '31'}],
              'texto_hoja': '005 - PTE GABINO - GOMEZ PLATA - CAROLINA\n'
                            '  031 - EL ROBLE - MINA VIEJA - MIRAFLORES'},
 'CAUCASIA': {'cod_mpio': '154',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1999,
                             'codigo': '009',
                             'fecha_inicio': '1999-03-31',
                             'fecha_resolucion': '1999-02-11',
                             'fecha_vencimiento': '2007-08-01',
                             'obra': 'CAUCASIA - NECHI',
                             'oficinas_registro': ['015 - CAUCASIA', '037 - YARUMAL'],
                             'resolucion': '1029'},
                            {'codigo': '024', 'obra': 'PUERTO TRIANA - ZARAGOZA'}],
              'texto_hoja': '009 - CAUCASIA - NECHI\n  024 - PUERTO TRIANA - ZARAGOZA'},
 'CHIGORODO': {'cod_mpio': '172',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1986,
                              'codigo': '050',
                              'fecha_resolucion': '1986-07-09',
                              'obra': 'ELECTRIF CHIGORODO - BARRANQUILLITA',
                              'resolucion': '49'},
                             {'anio': 1984,
                              'codigo': '051',
                              'fecha_resolucion': '1984-02-10',
                              'obra': 'ELECTRIF CHIGORODO - EL TRES - RIO LEON',
                              'resolucion': '3'}],
               'texto_hoja': '050 - ELECTRIF CHIGORODO - BARRANQUILLITA\n'
                             '  051 - ELECTRIF CHIGORODO - EL TRES - RIO LEON'},
 'CISNEROS': {'cod_mpio': '190',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'COCORNA': {'cod_mpio': '197',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'CONCEPCION': {'cod_mpio': '206',
                'estado': '000',
                'nota': '',
                'proyectos': [],
                'texto_hoja': '000 - NO GRAVADO'},
 'CONCORDIA': {'cod_mpio': '209',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1996,
                              'codigo': '006',
                              'fecha_inicio': '1996-09-30',
                              'fecha_resolucion': '1996-07-24',
                              'fecha_vencimiento': '2001-12-01',
                              'obra': 'LA METIDA - CONCORDIA',
                              'oficinas_registro': ['004 - ANDES'],
                              'resolucion': '0659'},
                             {'anio': 1996,
                              'codigo': '014',
                              'fecha_inicio': '1996-07-31',
                              'fecha_resolucion': '1996-05-28',
                              'fecha_vencimiento': '2001-06-01',
                              'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                              'oficinas_registro': ['004 - ANDES',
                                                    '024 - ANTIOQUIA',
                                                    '035 - URRAO',
                                                    '036 - URRAO'],
                              'resolucion': '0422'}],
               'texto_hoja': '006 - LA METIDA - CONCORDIA\n  014 - BOLOMBOLO - SANTA FE DE ANT'},
 'COPACABANA': {'cod_mpio': '212',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 2005,
                               'codigo': '020',
                               'fecha_inicio': '2005-03-31',
                               'fecha_resolucion': '2005-01-25',
                               'fecha_vencimiento': '2010-02-01',
                               'obra': 'DOBLE CALZADA BELLO - HATILLO',
                               'oficinas_registro': ['012 - GIRARDOTA', '01N - MEDELLIN Z.NORTE'],
                               'resolucion': '0707'}],
                'texto_hoja': '020 - DOBLE CALZADA BELLO - HATILLO'},
 'DABEIBA': {'cod_mpio': '234',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'DON MATIAS': {'cod_mpio': '237',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1995,
                               'codigo': '612',
                               'fecha_inicio': '1995-08-31',
                               'fecha_resolucion': '1995-10-19',
                               'fecha_vencimiento': '1999-05-01',
                               'obra': 'DON MATIAS - ENTRERRIOS - SANTA ROSA',
                               'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                               'resolucion': '612'},
                              {'anio': 2014,
                               'codigo': '100',
                               'fecha_inicio': '2014-09-01',
                               'fecha_resolucion': '2014-08-04',
                               'fecha_vencimiento': '2019-08-01',
                               'obra': 'HATILLO - BARBOSA - PRADERA',
                               'oficinas_registro': ['012 - GIRARDOTA',
                                                     '025 - SANTA ROSA DE O.',
                                                     '026 - SANTO DOMINGO',
                                                     '038 - YOLOMBO'],
                               'resolucion': '120405'}],
                'texto_hoja': '612 - DON MATIAS - ENTRERRIOS - SANTA ROSA\n'
                              '  100 - HATILLO - BARBOSA - PRADERA'},
 'EBEJICO': {'cod_mpio': '240',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1996,
                            'codigo': '014',
                            'fecha_inicio': '1996-07-31',
                            'fecha_resolucion': '1996-05-28',
                            'fecha_vencimiento': '2001-06-01',
                            'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                            'oficinas_registro': ['004 - ANDES',
                                                  '024 - ANTIOQUIA',
                                                  '035 - URRAO',
                                                  '036 - URRAO'],
                            'resolucion': '0422'},
                           {'anio': 2001,
                            'codigo': '016',
                            'fecha_resolucion': '2001-01-09',
                            'obra': 'TUNEL DE OCCIDENTE',
                            'resolucion': '020'}],
             'texto_hoja': '014 - BOLOMBOLO - SANTA FE DE ANT\n  016 - TUNEL DE OCCIDENTE'},
 'EL BAGRE': {'cod_mpio': '250',
              'estado': '001',
              'nota': 'Si en el vur no tiene valorizacion se puede hacer un no gravado.',
              'proyectos': [{'codigo': '024', 'obra': 'PUERTO TRIANA - ZARAGOZA'}],
              'texto_hoja': '024 - PUERTO TRIANA - ZARAGOZA'},
 'ENTRERRIOS': {'cod_mpio': '264',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1991,
                               'codigo': '083',
                               'fecha_resolucion': '1991-12-03',
                               'obra': 'SAN PEDRO - ENTRERRIOS',
                               'resolucion': '83'},
                              {'anio': 1995,
                               'codigo': '612',
                               'fecha_inicio': '1995-08-31',
                               'fecha_resolucion': '1995-10-19',
                               'fecha_vencimiento': '1999-05-01',
                               'obra': 'DON MATIAS - ENTRERRIOS - SANTA ROSA',
                               'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                               'resolucion': '612'}],
                'texto_hoja': '083 - SAN PEDRO - ENTRERRIOS\n  612 - DON MATIAS - ENTRERRIOS - SANTA ROSA'},
 'ENVIGADO': {'cod_mpio': '266',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1990,
                             'codigo': '036',
                             'fecha_resolucion': '1990-06-28',
                             'obra': 'ALTO PALMAS',
                             'resolucion': '36'}],
              'texto_hoja': '036 - ALTO PALMAS'},
 'FREDONIA': {'cod_mpio': '282',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1996,
                             'codigo': '004',
                             'fecha_inicio': '1997-02-28',
                             'fecha_resolucion': '1996-12-17',
                             'fecha_vencimiento': '2002-01-01',
                             'obra': 'FREDONIA - PUENTE IGLESIAS',
                             'oficinas_registro': ['010 - FREDONIA'],
                             'resolucion': '1495'},
                            {'anio': 1981,
                             'codigo': '064',
                             'fecha_resolucion': '1981-12-10',
                             'obra': 'CAMILOCE - EL CINCO - FREDONIA',
                             'resolucion': '17'}],
              'texto_hoja': '004 - FREDONIA - PUENTE IGLESIAS\n  064 - CAMILOCE - EL CINCO - FREDONIA'},
 'FRONTINO': {'cod_mpio': '284',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'GIRALDO': {'cod_mpio': '306',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'GIRARDOTA': {'cod_mpio': '308',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 2005,
                              'codigo': '020',
                              'fecha_inicio': '2005-03-31',
                              'fecha_resolucion': '2005-01-25',
                              'fecha_vencimiento': '2010-02-01',
                              'obra': 'DOBLE CALZADA BELLO - HATILLO',
                              'oficinas_registro': ['012 - GIRARDOTA', '01N - MEDELLIN Z.NORTE'],
                              'resolucion': '0707'},
                             {'anio': 1983,
                              'codigo': '052',
                              'fecha_resolucion': '1983-05-18',
                              'obra': 'GIRARDOTA - CABILDO - HATILLO',
                              'resolucion': '4'}],
               'texto_hoja': '020 - DOBLE CALZADA BELLO - HATILLO\n  052 - GIRARDOTA - CABILDO - HATILLO'},
 'GOMEZ PLATA': {'cod_mpio': '310',
                 'estado': '001',
                 'nota': '',
                 'proyectos': [{'anio': 1997,
                                'codigo': '005',
                                'fecha_inicio': '1997-09-30',
                                'fecha_resolucion': '1997-08-11',
                                'fecha_vencimiento': '2002-08-01',
                                'obra': 'PTE GABINO - GOMEZ PLATA - CAROLINA',
                                'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                                'resolucion': '1114'},
                               {'anio': 2014,
                                'codigo': '100',
                                'fecha_inicio': '2014-09-01',
                                'fecha_resolucion': '2014-08-04',
                                'fecha_vencimiento': '2019-08-01',
                                'obra': 'HATILLO - BARBOSA - PRADERA',
                                'oficinas_registro': ['012 - GIRARDOTA',
                                                      '025 - SANTA ROSA DE O.',
                                                      '026 - SANTO DOMINGO',
                                                      '038 - YOLOMBO'],
                                'resolucion': '120405'}],
                 'texto_hoja': '005 - PTE GABINO - GOMEZ PLATA - CAROLINA\n'
                               '  100 - HATILLO - BARBOSA - PRADERA'},
 'GRANADA': {'cod_mpio': '313',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'GUADALUPE': {'cod_mpio': '315',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1997,
                              'codigo': '005',
                              'fecha_inicio': '1997-09-30',
                              'fecha_resolucion': '1997-08-11',
                              'fecha_vencimiento': '2002-08-01',
                              'obra': 'PTE GABINO - GOMEZ PLATA - CAROLINA',
                              'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                              'resolucion': '1114'},
                             {'anio': 1977,
                              'codigo': '021',
                              'fecha_resolucion': '1977-09-06',
                              'obra': 'GUADALUPE - RIO SAN PABLO',
                              'resolucion': '21'},
                             {'anio': 2014,
                              'codigo': '100',
                              'fecha_inicio': '2014-09-01',
                              'fecha_resolucion': '2014-08-04',
                              'fecha_vencimiento': '2019-08-01',
                              'obra': 'HATILLO - BARBOSA - PRADERA',
                              'oficinas_registro': ['012 - GIRARDOTA',
                                                    '025 - SANTA ROSA DE O.',
                                                    '026 - SANTO DOMINGO',
                                                    '038 - YOLOMBO'],
                              'resolucion': '120405'}],
               'texto_hoja': '005 - PTE GABINO - GOMEZ PLATA - CAROLINA\n'
                             '  021 - GUADALUPE - RIO SAN PABLO\n'
                             '  100 - HATILLO - BARBOSA - PRADERA'},
 'GUARNE': {'cod_mpio': '318',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'GUATAPE': {'cod_mpio': '321',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1995,
                            'codigo': '772',
                            'fecha_inicio': '1996-02-28',
                            'fecha_resolucion': '1995-11-29',
                            'fecha_vencimiento': '1999-01-01',
                            'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                            'oficinas_registro': ['018 - MARINILLA'],
                            'resolucion': '772'}],
             'texto_hoja': '772 - GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'HELICONIA': {'cod_mpio': '347',
               'estado': '000',
               'nota': '',
               'proyectos': [],
               'texto_hoja': '000 - NO GRAVADO'},
 'HISPANIA': {'cod_mpio': '353',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'ITAGUI': {'cod_mpio': '360',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'ITUANGO': {'cod_mpio': '361',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'JARDIN': {'cod_mpio': '364',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'JERICO': {'cod_mpio': '368',
            'estado': '001',
            'nota': '',
            'proyectos': [{'anio': 1996,
                           'codigo': '002',
                           'fecha_inicio': '1996-09-30',
                           'fecha_resolucion': '1996-07-24',
                           'fecha_vencimiento': '2002-12-01',
                           'obra': 'JERICO - LA CASCADA - JAMAICA - TARSO',
                           'oficinas_registro': ['014 - JERICO'],
                           'resolucion': '0663'},
                          {'anio': 2018,
                           'codigo': '450',
                           'fecha_resolucion': '2018-08-14',
                           'obra': 'PUENTE IGLESIAS - LIBANO',
                           'resolucion': '60237812'}],
            'texto_hoja': '002 - JERICO - LA CASCADA - JAMAICA - TARSO\n  450 - PUENTE IGLESIAS - LIBANO'},
 'LA CEJA': {'cod_mpio': '376',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1988,
                            'codigo': '065',
                            'fecha_resolucion': '1988-10-19',
                            'obra': 'CARABANCHEL - LA ARCADIA',
                            'resolucion': '65'},
                           {'anio': 1988,
                            'codigo': '066',
                            'fecha_resolucion': '1988-10-19',
                            'obra': 'EL RETIRO - PANTANILLO - LA CEJA',
                            'resolucion': '66'},
                           {'anio': 1985,
                            'codigo': '990',
                            'fecha_resolucion': '1985-06-16',
                            'obra': 'RIONEGRO - LA CEJA',
                            'resolucion': '18'}],
             'texto_hoja': '065 - CARABANCHEL - LA ARCADIA\n'
                           '  066 - EL RETIRO - PANTANILLO - LA CEJA\n'
                           '  990 - RIONEGRO - LA CEJA'},
 'LA ESTRELLA': {'cod_mpio': '380',
                 'estado': '000',
                 'nota': '',
                 'proyectos': [],
                 'texto_hoja': '000 - NO GRAVADO'},
 'LA PINTADA': {'cod_mpio': '390',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1991,
                               'codigo': '085',
                               'fecha_resolucion': '1991-12-03',
                               'obra': 'PINTADA - VALPARAISO',
                               'resolucion': '85'},
                              {'anio': 1992,
                               'codigo': '086',
                               'fecha_resolucion': '1992-02-03',
                               'obra': 'PINTADA - BOLOMBOLO',
                               'resolucion': '1'}],
                'texto_hoja': '085 - PINTADA - VALPARAISO\n  086 - PINTADA - BOLOMBOLO'},
 'LA UNION': {'cod_mpio': '400',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'LIBORINA': {'cod_mpio': '411',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'MACEO': {'cod_mpio': '425', 'estado': '000', 'nota': '', 'proyectos': [], 'texto_hoja': '000 - NO GRAVADO'},
 'MARINILLA': {'cod_mpio': '440',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1989,
                              'codigo': '049',
                              'fecha_resolucion': '1989-06-06',
                              'obra': 'MARINILLA - NUEVO PENOL',
                              'resolucion': '49'}],
               'texto_hoja': '049 - MARINILLA - NUEVO PEÑOL'},
 'MEDELLIN': {'cod_mpio': '001',
              'estado': 'NO_APLICA',
              'nota': '',
              'proyectos': [],
              'texto_hoja': 'NO APLICA'},
 'MONTEBELLO': {'cod_mpio': '467',
                'estado': '000',
                'nota': '',
                'proyectos': [],
                'texto_hoja': '000 - NO GRAVADO'},
 'MURINDO': {'cod_mpio': '475',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'MUTATA': {'cod_mpio': '480',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'NARINO': {'cod_mpio': '483',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'NECHI': {'cod_mpio': '495',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1999,
                          'codigo': '009',
                          'fecha_inicio': '1999-03-31',
                          'fecha_resolucion': '1999-02-11',
                          'fecha_vencimiento': '2007-08-01',
                          'obra': 'CAUCASIA - NECHI',
                          'oficinas_registro': ['015 - CAUCASIA', '037 - YARUMAL'],
                          'resolucion': '1029'}],
           'texto_hoja': '009 - CAUCASIA - NECHI'},
 'NECOCLI': {'cod_mpio': '490',
             'estado': '001',
             'nota': 'Si en el vur no tiene valorizacion se puede hacer un no gravado.',
             'proyectos': [{'anio': 1978,
                            'codigo': '992',
                            'fecha_resolucion': '1978-03-16',
                            'obra': 'NECOCLI - RIO MULATOS',
                            'resolucion': '3'}],
             'texto_hoja': '992 - NECOCLI - RIO MULATOS'},
 'OLAYA': {'cod_mpio': '501',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1987,
                          'codigo': '042',
                          'fecha_resolucion': '1987-10-01',
                          'obra': 'SOPETRAN - RIO CAUCA',
                          'resolucion': '42'},
                         {'anio': 2001,
                          'codigo': '016',
                          'fecha_resolucion': '2001-01-09',
                          'obra': 'TUNEL DE OCCIDENTE',
                          'resolucion': '020'}],
           'texto_hoja': '042 - SOPETRAN - RIO CAUCA\n  016 - TUNEL DE OCCIDENTE'},
 'PENOL': {'cod_mpio': '541',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1989,
                          'codigo': '049',
                          'fecha_resolucion': '1989-06-06',
                          'obra': 'MARINILLA - NUEVO PENOL',
                          'resolucion': '49'},
                         {'anio': 1995,
                          'codigo': '772',
                          'fecha_inicio': '1996-02-28',
                          'fecha_resolucion': '1995-11-29',
                          'fecha_vencimiento': '1999-01-01',
                          'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                          'oficinas_registro': ['018 - MARINILLA'],
                          'resolucion': '772'}],
           'texto_hoja': '049 - MARINILLA - NUEVO PEÑOL\n  772 - GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'PEQUE': {'cod_mpio': '543', 'estado': '000', 'nota': '', 'proyectos': [], 'texto_hoja': '000 - NO GRAVADO'},
 'PUEBLO RICO': {'cod_mpio': '576',
                 'estado': '000',
                 'nota': '',
                 'proyectos': [],
                 'texto_hoja': '000 - NO GRAVADO'},
 'PUERTO BERRIO': {'cod_mpio': '579',
                   'estado': '001',
                   'nota': '',
                   'proyectos': [{'anio': 1980,
                                  'codigo': '080',
                                  'fecha_resolucion': '1980-08-19',
                                  'obra': 'SAN JUAN DE BEDOUT - BODEGAS - CHISPAS',
                                  'resolucion': '10'},
                                 {'anio': 1984,
                                  'codigo': '886',
                                  'fecha_resolucion': '1984-06-13',
                                  'obra': 'TRONCAL DE LA PAZ',
                                  'resolucion': '9'},
                                 {'anio': 2018,
                                  'codigo': '400',
                                  'fecha_resolucion': '2018-08-13',
                                  'obra': 'PUERTO TRIUNFO - PUERTO NARE',
                                  'resolucion': '60237706'}],
                   'texto_hoja': '080 - SAN JUAN DE BEDOUT - BODEGAS - CHISPAS\n'
                                 '  886 - TRONCAL DE LA PAZ\n'
                                 '  400 - PUERTO TRIUNFO - PUERTO NARE'},
 'PUERTO NARE': {'cod_mpio': '585',
                 'estado': '001',
                 'nota': '',
                 'proyectos': [{'anio': 1979,
                                'codigo': '022',
                                'fecha_resolucion': '1979-10-31',
                                'obra': 'NARICES - NARE - LA SIERRA',
                                'resolucion': '21'},
                               {'anio': 1984,
                                'codigo': '886',
                                'fecha_resolucion': '1984-06-13',
                                'obra': 'TRONCAL DE LA PAZ',
                                'resolucion': '9'},
                               {'anio': 2018,
                                'codigo': '400',
                                'fecha_resolucion': '2018-08-13',
                                'obra': 'PUERTO TRIUNFO - PUERTO NARE',
                                'resolucion': '60237706'}],
                 'texto_hoja': '022 - NARICES - NARE - LA SIERRA\n'
                               '  886 - TRONCAL DE LA PAZ\n'
                               '  400 - PUERTO TRIUNFO - PUERTO NARE'},
 'PUERTO TRIUNFO': {'cod_mpio': '591',
                    'estado': '001',
                    'nota': '',
                    'proyectos': [{'anio': 1979,
                                   'codigo': '022',
                                   'fecha_resolucion': '1979-10-31',
                                   'obra': 'NARICES - NARE - LA SIERRA',
                                   'resolucion': '21'},
                                  {'anio': 1984,
                                   'codigo': '886',
                                   'fecha_resolucion': '1984-06-13',
                                   'obra': 'TRONCAL DE LA PAZ',
                                   'resolucion': '9'},
                                  {'anio': 1985,
                                   'codigo': '885',
                                   'fecha_resolucion': '1985-04-20',
                                   'obra': 'ELECTRIF DORADAL',
                                   'resolucion': '8'},
                                  {'anio': 2018,
                                   'codigo': '400',
                                   'fecha_resolucion': '2018-08-13',
                                   'obra': 'PUERTO TRIUNFO - PUERTO NARE',
                                   'resolucion': '60237706'}],
                    'texto_hoja': '022 - NARICES - NARE - LA SIERRA\n'
                                  '  886 - TRONCAL DE LA PAZ\n'
                                  '  885 - ELECTRIF DORADAL\n'
                                  '  400 - PUERTO TRIUNFO - PUERTO NARE'},
 'REMEDIOS': {'cod_mpio': '604',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1980,
                             'codigo': '080',
                             'fecha_resolucion': '1980-08-19',
                             'obra': 'SAN JUAN DE BEDOUT - BODEGAS - CHISPAS',
                             'resolucion': '10'}],
              'texto_hoja': '080 - SAN JUAN DE BEDOUT - BODEGAS - CHISPAS'},
 'RETIRO': {'cod_mpio': '607',
            'estado': '001',
            'nota': '',
            'proyectos': [{'anio': 1988,
                           'codigo': '065',
                           'fecha_resolucion': '1988-10-19',
                           'obra': 'CARABANCHEL - LA ARCADIA',
                           'resolucion': '65'},
                          {'anio': 1988,
                           'codigo': '066',
                           'fecha_resolucion': '1988-10-19',
                           'obra': 'EL RETIRO - PANTANILLO - LA CEJA',
                           'resolucion': '66'},
                          {'anio': 2018,
                           'codigo': '300',
                           'fecha_resolucion': '2018-07-18',
                           'obra': 'CARABANCHEL - LA ARCADIA',
                           'resolucion': '60232705'},
                          {'anio': 2018,
                           'codigo': '350',
                           'fecha_resolucion': '2018-07-18',
                           'obra': 'PURO CUERO - PUENTE CHAPINEROS',
                           'resolucion': '60232705'}],
            'texto_hoja': '065 - CARABANCHEL - LA ARCADIA\n'
                          '  066 - EL RETIRO - PANTANILLO - LA CEJA\n'
                          '  300 - CARABANCHEL - LA ARCADIA\n'
                          '  350 - PURO CUERO - PUENTE CHAPINEROS'},
 'RIONEGRO': {'cod_mpio': '615',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1990,
                             'codigo': '036',
                             'fecha_resolucion': '1990-06-28',
                             'obra': 'ALTO PALMAS',
                             'resolucion': '36'},
                            {'anio': 1987,
                             'codigo': '087',
                             'fecha_resolucion': '1987-07-01',
                             'nota_cruce': 'año tomado por nombre de obra (código distinto en hoja '
                                           'Proyectos)',
                             'obra': 'TABLAZO - PAIMADO - GUAYABITO',
                             'resolucion': '017'},
                            {'anio': 1985,
                             'codigo': '990',
                             'fecha_resolucion': '1985-06-16',
                             'obra': 'RIONEGRO - LA CEJA',
                             'resolucion': '18'}],
              'texto_hoja': '036 - ALTO PALMAS\n'
                            '  087 - TABLAZO - PAIMADO - GUAYABITO\n'
                            '  990 - RIONEGRO - LA CEJA'},
 'SABANALARGA': {'cod_mpio': '628',
                 'estado': '000',
                 'nota': '',
                 'proyectos': [],
                 'texto_hoja': '000 - NO GRAVADO'},
 'SABANETA': {'cod_mpio': '631',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'SALGAR': {'cod_mpio': '642',
            'estado': '001',
            'nota': '',
            'proyectos': [{'anio': 1996,
                           'codigo': '006',
                           'fecha_inicio': '1996-09-30',
                           'fecha_resolucion': '1996-07-24',
                           'fecha_vencimiento': '2001-12-01',
                           'obra': 'LA METIDA - CONCORDIA',
                           'oficinas_registro': ['004 - ANDES'],
                           'resolucion': '0659'}],
            'texto_hoja': '006 - LA METIDA - CONCORDIA'},
 'SAN ANDRES DE C.': {'cod_mpio': '647',
                      'estado': '001',
                      'nota': '',
                      'proyectos': [{'anio': 2002,
                                     'codigo': '018',
                                     'fecha_inicio': '2002-05-31',
                                     'fecha_resolucion': '2002-03-22',
                                     'fecha_vencimiento': '2012-02-01',
                                     'obra': 'LLANOS DE CUIVA - SAN JOSE DE LA M.',
                                     'oficinas_registro': ['025 - SANTA ROSA DE O.',
                                                           '029 - SOPETRAN',
                                                           '037 - YARUMAL'],
                                     'resolucion': '3049'}],
                      'texto_hoja': '018 - LLANOS DE CUIVA - SAN JOSE DE LA M.'},
 'SAN CARLOS': {'cod_mpio': '649',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1995,
                               'codigo': '772',
                               'fecha_inicio': '1996-02-28',
                               'fecha_resolucion': '1995-11-29',
                               'fecha_vencimiento': '1999-01-01',
                               'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                               'oficinas_registro': ['018 - MARINILLA'],
                               'resolucion': '772'}],
                'texto_hoja': '772 - GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'SAN FRANCISCO': {'cod_mpio': '652',
                   'estado': '000',
                   'nota': '',
                   'proyectos': [],
                   'texto_hoja': '000 - NO GRAVADO'},
 'SAN JERONIMO': {'cod_mpio': '656',
                  'estado': '001',
                  'nota': '',
                  'proyectos': [{'anio': 2001,
                                 'codigo': '016',
                                 'fecha_resolucion': '2001-01-09',
                                 'obra': 'TUNEL DE OCCIDENTE',
                                 'resolucion': '020'},
                                {'anio': 2018,
                                 'codigo': '250',
                                 'fecha_resolucion': '2018-07-16',
                                 'obra': 'SAN JERONIMO - LOMA HERMOSA',
                                 'resolucion': '60232128'}],
                  'texto_hoja': '016 - TUNEL DE OCCIDENTE\n  250 - SAN JERONIMO - LOMA HERMOSA'},
 'SAN JOSE DE LA M.': {'cod_mpio': '658',
                       'estado': '001',
                       'nota': '',
                       'proyectos': [{'anio': 2002,
                                      'codigo': '018',
                                      'fecha_inicio': '2002-05-31',
                                      'fecha_resolucion': '2002-03-22',
                                      'fecha_vencimiento': '2012-02-01',
                                      'obra': 'LLANOS DE CUIVA - SAN JOSE DE LA M.',
                                      'oficinas_registro': ['025 - SANTA ROSA DE O.',
                                                            '029 - SOPETRAN',
                                                            '037 - YARUMAL'],
                                      'resolucion': '3049'}],
                       'texto_hoja': '018 - LLANOS DE CUIVA - SAN JOSE DE LA M.'},
 'SAN JUAN DE URABA': {'cod_mpio': '659',
                       'estado': '000',
                       'nota': '',
                       'proyectos': [],
                       'texto_hoja': '000 - NO GRAVADO'},
 'SAN LUIS': {'cod_mpio': '660',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 2018,
                             'codigo': '400',
                             'fecha_resolucion': '2018-08-13',
                             'obra': 'PUERTO TRIUNFO - PUERTO NARE',
                             'resolucion': '60237706'}],
              'texto_hoja': '400 - PUERTO TRIUNFO - PUERTO NARE'},
 'SAN PEDRO DE LOS M.': {'cod_mpio': '664',
                         'estado': '001',
                         'nota': '',
                         'proyectos': [{'anio': 2002,
                                        'codigo': '019',
                                        'fecha_inicio': '2002-08-31',
                                        'fecha_resolucion': '2002-07-05',
                                        'fecha_vencimiento': '2010-07-01',
                                        'obra': 'PAJARITO - SAN PEDRO',
                                        'oficinas_registro': ['01N - MEDELLIN Z.NORTE'],
                                        'resolucion': '10929'},
                                       {'anio': 1991,
                                        'codigo': '083',
                                        'fecha_resolucion': '1991-12-03',
                                        'obra': 'SAN PEDRO - ENTRERRIOS',
                                        'resolucion': '83'}],
                         'texto_hoja': '019 - PAJARITO - SAN PEDRO\n  083 - SAN PEDRO - ENTRERRIOS'},
 'SAN PEDRO DE URABA': {'cod_mpio': '665',
                        'estado': '000',
                        'nota': '',
                        'proyectos': [],
                        'texto_hoja': '000 - NO GRAVADO'},
 'SAN RAFAEL': {'cod_mpio': '667',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1995,
                               'codigo': '772',
                               'fecha_inicio': '1996-02-28',
                               'fecha_resolucion': '1995-11-29',
                               'fecha_vencimiento': '1999-01-01',
                               'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                               'oficinas_registro': ['018 - MARINILLA'],
                               'resolucion': '772'}],
                'texto_hoja': '772 -  GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'SAN ROQUE': {'cod_mpio': '670',
               'estado': '001',
               'nota': '',
               'proyectos': [{'anio': 1995,
                              'codigo': '772',
                              'fecha_inicio': '1996-02-28',
                              'fecha_resolucion': '1995-11-29',
                              'fecha_vencimiento': '1999-01-01',
                              'obra': 'GUATAPE - SAN RAFAEL - SAN CARLOS',
                              'oficinas_registro': ['018 - MARINILLA'],
                              'resolucion': '772'}],
               'texto_hoja': '772 -  GUATAPE - SAN RAFAEL - SAN CARLOS'},
 'SAN VICENTE': {'cod_mpio': '674',
                 'estado': '001',
                 'nota': 'Si en el vur no tiene valorizacion se puede hacer un no gravado.',
                 'proyectos': [{'anio': 1992,
                                'codigo': '991',
                                'fecha_resolucion': '1992-11-11',
                                'obra': 'AUTOPISTA - SAN VICENTE',
                                'resolucion': '5'}],
                 'texto_hoja': '991 - AUTOPISTA - SAN VICENTE'},
 'SANTA BARBARA': {'cod_mpio': '679',
                   'estado': '000',
                   'nota': '',
                   'proyectos': [],
                   'texto_hoja': '000 - NO GRAVADO'},
 'SANTA FE DE ANT': {'cod_mpio': '042',
                     'estado': '001',
                     'nota': '',
                     'proyectos': [{'anio': 1996,
                                    'codigo': '014',
                                    'fecha_inicio': '1996-07-31',
                                    'fecha_resolucion': '1996-05-28',
                                    'fecha_vencimiento': '2001-06-01',
                                    'obra': 'BOLOMBOLO - SANTA FE DE ANT',
                                    'oficinas_registro': ['004 - ANDES',
                                                          '024 - ANTIOQUIA',
                                                          '035 - URRAO',
                                                          '036 - URRAO'],
                                    'resolucion': '0422'},
                                   {'anio': 2001,
                                    'codigo': '016',
                                    'fecha_resolucion': '2001-01-09',
                                    'obra': 'TUNEL DE OCCIDENTE',
                                    'resolucion': '020'}],
                     'texto_hoja': '014 - BOLOMBOLO - SANTA FE DE ANT\n  016 - TUNEL DE OCCIDENTE'},
 'SANTA ROSA': {'cod_mpio': '686',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 2002,
                               'codigo': '018',
                               'fecha_inicio': '2002-05-31',
                               'fecha_resolucion': '2002-03-22',
                               'fecha_vencimiento': '2012-02-01',
                               'obra': 'LLANOS DE CUIVA - SAN JOSE DE LA M.',
                               'oficinas_registro': ['025 - SANTA ROSA DE O.',
                                                     '029 - SOPETRAN',
                                                     '037 - YARUMAL'],
                               'resolucion': '3049'},
                              {'anio': 1973,
                               'codigo': '031',
                               'fecha_resolucion': '1973-03-08',
                               'obra': 'EL ROBLE - MINA VIEJA - MIRAFLORES',
                               'resolucion': '31'},
                              {'anio': 1995,
                               'codigo': '612',
                               'fecha_inicio': '1995-08-31',
                               'fecha_resolucion': '1995-10-19',
                               'fecha_vencimiento': '1999-05-01',
                               'obra': 'DON MATIAS - ENTRERRIOS - SANTA ROSA',
                               'oficinas_registro': ['025 - SANTA ROSA DE O.'],
                               'resolucion': '612'},
                              {'anio': 2014,
                               'codigo': '100',
                               'fecha_inicio': '2014-09-01',
                               'fecha_resolucion': '2014-08-04',
                               'fecha_vencimiento': '2019-08-01',
                               'obra': 'HATILLO - BARBOSA - PRADERA',
                               'oficinas_registro': ['012 - GIRARDOTA',
                                                     '025 - SANTA ROSA DE O.',
                                                     '026 - SANTO DOMINGO',
                                                     '038 - YOLOMBO'],
                               'resolucion': '120405'}],
                'texto_hoja': '018 - LLANOS DE CUIVA - SAN JOSE DE LA M.\n'
                              '  031 - EL ROBLE - MINA VIEJA - MIRAFLORES\n'
                              '  612 - DON MATIAS - ENTRERRIOS - SANTA ROSA\n'
                              '  100 - HATILLO - BARBOSA - PRADERA'},
 'SANTO DOMINGO': {'cod_mpio': '690',
                   'estado': '001',
                   'nota': '',
                   'proyectos': [{'anio': 2014,
                                  'codigo': '100',
                                  'fecha_inicio': '2014-09-01',
                                  'fecha_resolucion': '2014-08-04',
                                  'fecha_vencimiento': '2019-08-01',
                                  'obra': 'HATILLO - BARBOSA - PRADERA',
                                  'oficinas_registro': ['012 - GIRARDOTA',
                                                        '025 - SANTA ROSA DE O.',
                                                        '026 - SANTO DOMINGO',
                                                        '038 - YOLOMBO'],
                                  'resolucion': '120405'}],
                   'texto_hoja': '100 - HATILLO - BARBOSA - PRADERA'},
 'SANTUARIO': {'cod_mpio': '697',
               'estado': '000',
               'nota': '',
               'proyectos': [],
               'texto_hoja': '000 - NO GRAVADO'},
 'SEGOVIA': {'cod_mpio': '736',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'SONSON': {'cod_mpio': '756',
            'estado': '001',
            'nota': '',
            'proyectos': [{'anio': 1985,
                           'codigo': '885',
                           'fecha_resolucion': '1985-04-20',
                           'obra': 'ELECTRIF DORADAL',
                           'resolucion': '8'}],
            'texto_hoja': '885 - ELECTRIF DORADAL'},
 'SOPETRAN': {'cod_mpio': '761',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1987,
                             'codigo': '042',
                             'fecha_resolucion': '1987-10-01',
                             'obra': 'SOPETRAN - RIO CAUCA',
                             'resolucion': '42'},
                            {'anio': 2001,
                             'codigo': '016',
                             'fecha_resolucion': '2001-01-09',
                             'obra': 'TUNEL DE OCCIDENTE',
                             'resolucion': '020'}],
              'texto_hoja': '042 - SOPETRAN - RIO CAUCA\n  016 - TUNEL DE OCCIDENTE'},
 'TAMESIS': {'cod_mpio': '789',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1996,
                            'codigo': '003',
                            'fecha_inicio': '1996-08-31',
                            'fecha_resolucion': '1996-06-26',
                            'fecha_vencimiento': '2001-07-01',
                            'obra': 'LA FABIANA - LA MESA',
                            'oficinas_registro': ['032- TAMESIS'],
                            'resolucion': '0525'},
                           {'anio': 2018,
                            'codigo': '450',
                            'fecha_resolucion': '2018-08-14',
                            'obra': 'PUENTE IGLESIAS - LIBANO',
                            'resolucion': '60237812'}],
             'texto_hoja': '003 - LA FABIANA - LA MESA\n  450 - PUENTE IGLESIAS - LIBANO'},
 'TARAZA': {'cod_mpio': '790',
            'estado': '001',
            'nota': 'Si en el vur no tiene valorizacion se puede hacer un no gravado.',
            'proyectos': [{'anio': 1979,
                           'codigo': '015',
                           'fecha_resolucion': '1979-08-16',
                           'obra': 'TARAZA - CANON DE IGLESIAS',
                           'resolucion': '15'}],
            'texto_hoja': '015 - TARAZA - CAÑON DE IGLESIAS'},
 'TARSO': {'cod_mpio': '792',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1996,
                          'codigo': '002',
                          'fecha_inicio': '1996-09-30',
                          'fecha_resolucion': '1996-07-24',
                          'fecha_vencimiento': '2002-12-01',
                          'obra': 'JERICO - LA CASCADA - JAMAICA - TARSO',
                          'oficinas_registro': ['014 - JERICO'],
                          'resolucion': '0663'}],
           'texto_hoja': '002 - JERICO - LA CASCADA - JAMAICA - TARSO'},
 'TITIRIBI': {'cod_mpio': '809',
              'estado': '001',
              'nota': '',
              'proyectos': [{'anio': 1991,
                             'codigo': '033',
                             'fecha_resolucion': '1991-04-25',
                             'obra': 'ALBANIA - TITIRIBI',
                             'resolucion': '20'}],
              'texto_hoja': '033 - ALBANIA - TITIRIBI'},
 'TOLEDO': {'cod_mpio': '819',
            'estado': '000',
            'nota': '',
            'proyectos': [],
            'texto_hoja': '000 - NO GRAVADO'},
 'TURBO': {'cod_mpio': '837', 'estado': '000', 'nota': '', 'proyectos': [], 'texto_hoja': '000 - NO GRAVADO'},
 'URAMITA': {'cod_mpio': '842',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'URRAO': {'cod_mpio': '847',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1998,
                          'codigo': '008',
                          'fecha_inicio': '1998-10-31',
                          'fecha_resolucion': '1998-09-16',
                          'fecha_vencimiento': '2007-05-01',
                          'obra': 'URRAO - BETULIA',
                          'oficinas_registro': ['035 - URRAO', '036 - URRAO'],
                          'resolucion': '1187'}],
           'texto_hoja': '008 - URRAO - BETULIA'},
 'VALDIVIA': {'cod_mpio': '854',
              'estado': '000',
              'nota': '',
              'proyectos': [],
              'texto_hoja': '000 - NO GRAVADO'},
 'VALPARAISO': {'cod_mpio': '856',
                'estado': '001',
                'nota': '',
                'proyectos': [{'anio': 1991,
                               'codigo': '085',
                               'fecha_resolucion': '1991-12-03',
                               'obra': 'PINTADA - VALPARAISO',
                               'resolucion': '85'}],
                'texto_hoja': '085 - PINTADA - VALPARAISO'},
 'VEGACHI': {'cod_mpio': '858',
             'estado': '000',
             'nota': '',
             'proyectos': [],
             'texto_hoja': '000 - NO GRAVADO'},
 'VENECIA': {'cod_mpio': '861',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 1996,
                            'codigo': '001',
                            'fecha_inicio': '1997-02-01',
                            'fecha_resolucion': '1996-12-18',
                            'fecha_vencimiento': '2002-01-01',
                            'obra': 'EL CINCO - VENECIA - BOLOMBOLO',
                            'oficinas_registro': ['010 - FREDONIA'],
                            'resolucion': '1500'},
                           {'anio': 1981,
                            'codigo': '064',
                            'fecha_resolucion': '1981-12-10',
                            'obra': 'CAMILOCE - EL CINCO - FREDONIA',
                            'resolucion': '17'}],
             'texto_hoja': '001 - EL CINCO - VENECIA - BOLOMBOLO\n  064 - CAMILOCE - EL CINCO - FREDONIA'},
 'VIGIA DEL FUERTE': {'cod_mpio': '873',
                      'estado': '000',
                      'nota': '',
                      'proyectos': [],
                      'texto_hoja': '000 - NO GRAVADO'},
 'YALI': {'cod_mpio': '885', 'estado': '000', 'nota': '', 'proyectos': [], 'texto_hoja': '000 - NO GRAVADO'},
 'YARUMAL': {'cod_mpio': '887',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 2002,
                            'codigo': '018',
                            'fecha_inicio': '2002-05-31',
                            'fecha_resolucion': '2002-03-22',
                            'fecha_vencimiento': '2012-02-01',
                            'obra': 'LLANOS DE CUIVA - SAN JOSE DE LA M.',
                            'oficinas_registro': ['025 - SANTA ROSA DE O.',
                                                  '029 - SOPETRAN',
                                                  '037 - YARUMAL'],
                            'resolucion': '3049'}],
             'texto_hoja': '018 - LLANOS DE CUIVA - SAN JOSE DE LA M.'},
 'YOLOMBO': {'cod_mpio': '890',
             'estado': '001',
             'nota': '',
             'proyectos': [{'anio': 2014,
                            'codigo': '100',
                            'fecha_inicio': '2014-09-01',
                            'fecha_resolucion': '2014-08-04',
                            'fecha_vencimiento': '2019-08-01',
                            'obra': 'HATILLO - BARBOSA - PRADERA',
                            'oficinas_registro': ['012 - GIRARDOTA',
                                                  '025 - SANTA ROSA DE O.',
                                                  '026 - SANTO DOMINGO',
                                                  '038 - YOLOMBO'],
                            'resolucion': '120405'}],
             'texto_hoja': '100 - HATILLO - BARBOSA - PRADERA'},
 'YONDO': {'cod_mpio': '893',
           'estado': '001',
           'nota': '',
           'proyectos': [{'anio': 1980,
                          'codigo': '080',
                          'fecha_resolucion': '1980-08-19',
                          'obra': 'SAN JUAN DE BEDOUT - BODEGAS - CHISPAS',
                          'resolucion': '10'}],
           'texto_hoja': '080 - SAN JUAN DE BEDOUT - BODEGAS - CHISPAS'},
 'ZARAGOZA': {'cod_mpio': '895',
              'estado': '001',
              'nota': 'Si en el vur no tiene valorizacion se puede hacer un no gravado.',
              'proyectos': [{'codigo': '024', 'obra': 'PUERTO TRIANA - ZARAGOZA'}],
              'texto_hoja': '024 - PUERTO TRIANA - ZARAGOZA'}}
# ==== FIN DATOS DE PROYECTOS ====


def registrar_error_inicio():
    """Guarda el error de arranque en un archivo y lo muestra (con pythonw no hay consola donde verlo)."""
    import traceback
    detalle = traceback.format_exc()
    ruta = CARPETA_USUARIO / "error_inicio.txt"
    try:
        CARPETA_USUARIO.mkdir(parents=True, exist_ok=True)
        ruta.write_text(f"{datetime.now():%Y-%m-%d %H:%M:%S} · versión {VERSION_APLICATIVO}\n\n{detalle}", encoding="utf-8")
    except OSError:
        pass
    try:
        aviso = tk.Tk()
        aviso.withdraw()
        messagebox.showerror("Consulta VUR", "El programa no pudo iniciar.\n\n"
                             f"{detalle.strip().splitlines()[-1]}\n\nEl detalle quedó en:\n{ruta}\n"
                             "Envíe ese archivo al desarrollador.", parent=aviso)
        aviso.destroy()
    except Exception:  # noqa: BLE001 - si ni siquiera Tk abre, queda el archivo
        pass


def principal():
    import sys
    try:
        _principal(sys)
    except SystemExit:
        raise
    except BaseException:
        registrar_error_inicio()
        raise


def _principal(sys):
    if "--probar" in sys.argv:
        sys.exit(0 if ejecutar_autopruebas() else 1)
    if "--recursos" in sys.argv:
        print(preparar_recursos() or "No se pudieron crear los recursos")
        return
    preparar_recursos()
    cargar_fuentes_privadas()
    raiz = tk.Tk()
    raiz.withdraw()           # la ventana principal aparece después de elegir la carpeta
    primer_inicio(raiz)
    raiz.deiconify()
    App(raiz)
    raiz.mainloop()


if __name__ == "__main__":
    principal()
