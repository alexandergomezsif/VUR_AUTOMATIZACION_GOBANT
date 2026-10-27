"""
Pruebas automáticas de las reglas de análisis de vur_valorizacion.py.

Ejecútelas después de CUALQUIER cambio al programa (doble clic en
"Probar reglas.bat"). Si alguna falla, no use el programa hasta corregirlo.

Cada caso reproduce el formato real de las anotaciones de VUR, con nombres y
documentos ficticios. Cuando aparezca un error nuevo en producción, agregue
aquí un caso que lo reproduzca antes de corregirlo.
"""
import queue
import sys
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import vur_valorizacion as v  # noqa: E402

P = "PERSONAS QUE INTERVIENEN EN EL ACTO (X-Titular de derecho real de dominio,I-Titular de dominio incompleto)"


def an(nro, fecha, espec, lineas, cancela=None):
    """Arma una anotación con el mismo formato HTML que entrega VUR."""
    c = f"<br>    Se cancela anotación No: {cancela}" if cancela else ""
    return (f"ANOTACION: Nro {nro} Fecha: {fecha} Radicación: 1<br>Doc: OFICIO 1 VALOR ACTO: ${c}<br/>"
            f"ESPECIFICACION: {espec}<br/>{P}<br/>" + "<br/>".join(lineas))


COMPRA = an(1, "01-01-1980", "101 COMPRAVENTA (MODO DE ADQUISICION)", ["DE: VENDEDOR FICTICIO", "A: TITULAR FICTICIO   CC 1  X"])
HIPOTECA = an(2, "12-04-1996", "210 HIPOTECA ABIERTA CUANTIA INDETERMINADA NATURALEZA Y N: ESC. 446\r\nNATURALEZA JURIDICA DEL ACTO: GRAVAMEN (GRAVAMEN)",
              ["DE: TITULAR FICTICIO", "A: BANCO FICTICIO"])
VAL_430 = an(5, "13-05-1999", "430 GRAVAMEN POR CONTRIBUCION DE VALORIZACION (GRAVAMEN)", ["DE: DEPARTAMENTO DE ANTIOQUIA", "A: TITULAR FICTICIO"])
CAN_5 = an(6, "16-06-2009", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA (CANCELACION)", ["A: TITULAR FICTICIO"], cancela=5)
VAL_BELLO = an(5, "28-12-1999", "430 VALORIZACION OBRA 808 RESOLUCION 314 DE 01-12-99 (GRAVAMEN)", ["DE: IMPES DE BELLO"])
CAN_BELLO = an(6, "10-11-2006", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA DE VALORIZACION OBRA 808. (CANCELACION)", ["DE: MUNICIPIO DE BELLO"])
AFECT = an(8, "08-08-2002", "380 AFECTACION DE INENAJENABILIDAD POR VALORIZACION (MEDIDA CAUTELAR)", ["DE: PLANEACION Y VALORIZACION", "A: TITULAR FICTICIO"])
CAN_AFECT = an(11, "24-08-2026", "0842 CANCELACION PROVIDENCIA ADMINISTRATIVA DE AFECTACION DE INENAJENABILIDAD POR VALORIZACION (CANCELACION)",
               ["DE: ALCALDIA"], cancela=8)


class API:
    """Sustituye a VUR: devuelve respuestas fijas por matrícula."""
    def __init__(self, db, ej):
        self.db, self.ej = db, ej

    def datos_basicos(self, c, n):
        return self.db.get(f"{c}-{n}")

    def estado_juridico(self, c, n):
        return self.ej.get(f"{c}-{n}")


def folio(anotaciones):
    return v.analizar_folio([v.parsear_anotacion(x) for x in anotaciones])


def consultar(db, ej, circulo, numero):
    app = v.App.__new__(v.App)
    app.cola = queue.Queue()
    reg = {c: "" for c in v.COLUMNAS_SALIDA}
    return app._consultar(API(db, ej), circulo, numero, reg)


def db_basico(mpio, apertura="20/04/1981", matrices=(), propietarios=True):
    return {"circulo": "008", "numeroMatricula": "1", "municipio": {"nombre": mpio}, "fechaAperturaFolio": apertura,
            "matriculasMatriz": list(matrices),
            "propietarios": [{"nombre": "TITULAR FICTICIO", "numeroDocumento": "1", "titular": "X"}] if propietarios else []}


class Clasificacion(unittest.TestCase):
    def test_hipoteca_no_es_valorizacion(self):
        self.assertEqual(folio([COMPRA, HIPOTECA])["estado"], "NO GRAVADA")

    def test_gravamen_vigente(self):
        self.assertEqual(folio([COMPRA, VAL_430])["estado"], "GRAVADA")

    def test_cancelacion_con_numero(self):
        r = folio([COMPRA, VAL_430, CAN_5])
        self.assertEqual(r["estado"], "GRAVAMEN CANCELADO")
        self.assertEqual(r["valorizacion"][0]["vinculo"], "numero")

    def test_cancelacion_sin_numero_bello(self):
        """Error del 23/10/2026: la cancelación sin número se tomaba como otro gravamen."""
        r = folio([COMPRA, VAL_BELLO, CAN_BELLO])
        self.assertEqual(r["estado"], "GRAVAMEN CANCELADO")
        self.assertEqual(r["valorizacion"][0]["vinculo"], "texto")

    def test_afectacion_no_define_estado(self):
        r = folio([COMPRA, AFECT, CAN_AFECT])
        self.assertEqual(r["estado"], "NO GRAVADA")
        self.assertEqual(len(r["afectaciones"]), 1)
        self.assertEqual(r["afectaciones"][0]["cancelada_por"], 11)

    def test_cancelacion_codigo_0723(self):
        """Caso real 024-12609 (23/10/2026): cancelación con código 0723, no 0842."""
        val = an(2, "23-07-2001", "0212 VALORIZACION (GRAVAMEN)", ["DE: MUNICIPIO"])
        can = an(3, "29-05-2003", "0723 CANCELACION VALORIZACION (CANCELACION)", ["DE: MUNICIPIO"], cancela=2)
        self.assertEqual(folio([COMPRA, val, can])["estado"], "GRAVAMEN CANCELADO")
        can_sin_marca = an(3, "29-05-2003", "0723 CANCELACION VALORIZACION", ["DE: MUNICIPIO"])
        self.assertEqual(folio([COMPRA, val, can_sin_marca])["estado"], "GRAVAMEN CANCELADO")

    def test_codigo_desconocido_es_dudosa(self):
        a = v.parsear_anotacion(an(7, "01-01-2010", "0999 OTRA ANOTACION DE VALORIZACION (OTRO)", ["DE: MUNICIPIO"]))
        self.assertEqual(v.tipo_valorizacion(a), "dudosa")


class Resultados(unittest.TestCase):
    def test_confiabilidad_alta(self):
        r = consultar({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430, CAN_5]}}, "008", "1")
        self.assertEqual((r["Estado"], r["Confiabilidad"]), ("GRAVAMEN CANCELADO", "ALTA"))

    def test_afectacion_vigente_queda_para_revision(self):
        r = consultar({"008-1": db_basico("APARTADO")}, {"008-1": {"textoAnotaciones": [COMPRA, AFECT]}}, "008", "1")
        self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "REVISAR"))
        self.assertIn("afectación por valorización VIGENTE", r["Motivos_Revision"])
        self.assertIn("380 AFECTACION", r["Afectaciones_Valorizacion"])

    def test_bello_en_matriz_con_propietario_deducido(self):
        hijo = db_basico("BELLO", "19/12/2024", ["008-2"], propietarios=False)
        madre = db_basico("BELLO", "01/01/1990")
        r = consultar({"008-1": hijo, "008-2": madre},
                      {"008-1": {"textoAnotaciones": [an(2, "20-12-2024", "0125 COMPRAVENTA (MODO DE ADQUISICION)", ["A: COMPRADOR FICTICIO  CC 2  X"])]},
                       "008-2": {"textoAnotaciones": [COMPRA, VAL_BELLO, CAN_BELLO]}}, "008", "1")
        self.assertEqual(r["Estado"], "GRAVAMEN CANCELADO")
        self.assertEqual(r["Matricula_Con_Anotacion"], "008-2")
        self.assertEqual(r["Confiabilidad"], "REVISAR")
        self.assertIn("anotación 1", r["Propietario_Gravado"])
        self.assertIn("anotación 2", r["Propietarios_Actuales"])

    def test_municipio_000_no_consulta(self):
        r = consultar({"008-1": db_basico("GUARNE")}, {}, "008", "1")
        self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "ALTA"))

    def test_reintento_si_vur_responde_vacio(self):
        """Caso real 024-0024613: el 22/10 VUR respondió vacío y el 23/10 con datos."""
        import time
        class APIFalla(API):
            def __init__(self, *a):
                super().__init__(*a)
                self.llamadas = 0
            def datos_basicos(self, c, n):
                self.llamadas += 1
                return None if self.llamadas == 1 else super().datos_basicos(c, n)
        dormir, time.sleep = time.sleep, lambda s: None
        try:
            app = v.App.__new__(v.App)
            app.cola = queue.Queue()
            reg = {c: "" for c in v.COLUMNAS_SALIDA}
            api = APIFalla({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430]}})
            r = app._consultar(api, "008", "1", reg)
        finally:
            time.sleep = dormir
        self.assertEqual(r["Estado"], "GRAVADA")

    def test_sin_datos(self):
        import time
        dormir, time.sleep = time.sleep, lambda s: None
        try:
            r = consultar({}, {}, "008", "1")
        finally:
            time.sleep = dormir
        self.assertEqual((r["Estado"], r["Confiabilidad"]), ("SIN DATOS EN VUR", "REVISAR"))


class SegundaPasada(unittest.TestCase):
    def test_seleccion(self):
        f = v.necesita_segunda_pasada
        self.assertTrue(f({"Estado": "SIN DATOS EN VUR"}))
        self.assertTrue(f({"Estado": "ERROR", "Notas_Proceso": "Falló tras 3 intentos"}))
        self.assertFalse(f({"Estado": "ERROR", "Notas_Proceso": "Formato de matrícula inválido"}))
        self.assertTrue(f({"Estado": v.ESTADO_INCOMPLETO,
                           "Notas_Proceso": "Matriz 017-7991: VUR no devolvió datos (dos intentos); no se pudo revisar"}))
        self.assertFalse(f({"Estado": v.ESTADO_INCOMPLETO, "Notas_Proceso": "Tope de seguridad de 300 folios alcanzado"}))
        self.assertFalse(f({"Estado": "NO GRAVADA", "Confiabilidad": "REVISAR",
                            "Motivos_Revision": "Tiene afectación por valorización"}))
        self.assertFalse(f({"Estado": "GRAVADA"}))

    def test_reconsulta_y_reemplaza(self):
        """Caso 024-0024613: primera pasada sin datos, segunda con valorización."""
        import threading
        import time
        import pandas as pd

        class APIRecupera(API):
            def __init__(self, *a):
                super().__init__(*a)
                self.activa = False
            def datos_basicos(self, c, n):
                return super().datos_basicos(c, n) if self.activa else None

        api = APIRecupera({"008-1": db_basico("NECHI")}, {"008-1": {"textoAnotaciones": [COMPRA, VAL_430]}})
        app = v.App.__new__(v.App)
        app.cola, app.lock, app.detener, app.ruta_excel = queue.Queue(), threading.Lock(), threading.Event(), None
        app.df = pd.DataFrame([{"Matricula_Original": "MI 008-0000001", "Matricula_Limpia": "008-0000001"}])
        import tempfile
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        app.historico, app._previo, app.usuario, app.equipo = v.Historico(temporal.name), {}, "prueba", "PC"
        dormir, espera, time.sleep = time.sleep, v.ESPERA_SEGUNDA_PASADA, lambda s: None
        v.ESPERA_SEGUNDA_PASADA = 0
        try:
            app.resultados = [app._procesar_matricula(api, app.df.iloc[0])]
            self.assertEqual(app.resultados[0]["Estado"], "SIN DATOS EN VUR")
            api.activa = True                      # VUR se recupera
            app._segunda_pasada(api)
        finally:
            time.sleep, v.ESPERA_SEGUNDA_PASADA = dormir, espera
        self.assertEqual(app.resultados[0]["Estado"], "GRAVADA")
        self.assertIn("cambió de SIN DATOS EN VUR a GRAVADA", app.resultados[0]["Segunda_Pasada"])


class Matrices(unittest.TestCase):
    def test_ocho_niveles(self):
        db, ej = {}, {}
        for i in range(9):
            db[f"020-{100 + i}"] = {"fechaAperturaFolio": f"01/01/{2020 - i * 5}",
                                    "matriculasMatriz": [f"020-{101 + i}"] if i < 8 else []}
        ej["020-107"] = {"textoAnotaciones": [COMPRA, VAL_430]}
        r = v.revisar_cadena(API(db, ej), "020", "100", db["020-100"], date(1985, 6, 16))
        self.assertEqual((r["estado"], r["matricula"], r["nivel"]), ("GRAVADA", "020-107", 7))

    def test_corte_por_fecha(self):
        db = {"025-1": {"fechaAperturaFolio": "01/01/1970", "matriculasMatriz": ["025-9"]}}
        r = v.revisar_cadena(API(db, {}), "025", "1", db["025-1"], date(1973, 3, 8))
        self.assertEqual(len(r["cadena"]), 1)

    def test_matriz_sin_respuesta_es_incompleta(self):
        db = {"017-1": {"fechaAperturaFolio": "19/12/2023", "matriculasMatriz": ["017-2"]}}
        r = v.revisar_cadena(API(db, {}), "017", "1", db["017-1"], date(1988, 10, 19))
        self.assertEqual(r["estado"], v.ESTADO_INCOMPLETO)


class Mejoras261026(unittest.TestCase):
    """Cambios del 26/10/2026 tras comparar el lote 'matriculas 2 monica' (23/10 vs 26/10)."""

    def test_anotacion_sin_especificacion_usa_naturaleza(self):
        antigua = ("ANOTACION: Nro 4 Fecha: 01-01-1985 Radicación: 1<br>Doc: RESOLUCION 1<br/>"
                   "NATURALEZA JURIDICA DEL ACTO: 430 GRAVAMEN POR CONTRIBUCION DE VALORIZACION<br/>DE: MUNICIPIO")
        self.assertEqual(folio([COMPRA, antigua])["estado"], "GRAVADA")

    def test_no_reconocida_dice_cual_y_si_menciona_valorizacion(self):
        rota = "Anotación ilegible de VALORIZACION sin número<br/>DE: MUNICIPIO"
        r = consultar({"017-1": db_basico("LA CEJA")}, {"017-1": {"textoAnotaciones": [COMPRA, rota, "  "]}},
                      "017", "1")
        self.assertEqual(r["Confiabilidad"], "REVISAR")
        self.assertIn("entrada 2 sin número de anotación y MENCIONA VALORIZACIÓN", r["Motivos_Revision"])
        self.assertNotIn("entrada 3", r["Motivos_Revision"])  # la entrada vacía se ignora
        self.assertIn("Anotación ilegible", r["Notas_Proceso"])

    def test_afectacion_cancelada_trae_texto_de_la_cancelacion(self):
        r = consultar({"008-1": db_basico("APARTADO")}, {"008-1": {"textoAnotaciones": [COMPRA, AFECT, CAN_AFECT]}},
                      "008", "1")
        self.assertIn("--- Cancelación: 008-1 anotación 11", r["Afectaciones_Valorizacion"])
        # Decisión del usuario (26/10/2026): afectación cancelada con número -> ALTA, pero visible.
        self.assertEqual((r["Estado"], r["Confiabilidad"]), ("NO GRAVADA", "ALTA"))
        self.assertIn("afectación por valorización ya cancelada", r["Notas_Proceso"])

    def test_propietario_gravado_desde_la_matriz(self):
        """Caso Bello 01N-5052127: el gravamen no nombra dueño y el folio no tiene compra anterior."""
        hijo = db_basico("BELLO", "19/12/2024", ["008-2"])
        con_gravamen = db_basico("BELLO", "24/06/1992", ["008-3"])
        abuela = db_basico("BELLO", "01/01/1970")
        venta_posterior = an(9, "01-01-1995", "0125 COMPRAVENTA (MODO DE ADQUISICION)", ["A: OTRO DUEÑO FICTICIO"])
        r = consultar({"008-1": hijo, "008-2": con_gravamen, "008-3": abuela},
                      {"008-1": {"textoAnotaciones": []},
                       "008-2": {"textoAnotaciones": [VAL_BELLO, CAN_BELLO]},
                       "008-3": {"textoAnotaciones": [COMPRA, venta_posterior]}}, "008", "1")
        self.assertEqual(r["Matricula_Con_Anotacion"], "008-2")
        self.assertIn("TITULAR FICTICIO", r["Propietario_Gravado"])
        self.assertIn("(dueño de la matriz 008-3 según su anotación 1, deducido)", r["Propietario_Gravado"])
        self.assertNotIn("OTRO DUEÑO", r["Propietario_Gravado"])  # venta posterior a la apertura del folio
        self.assertEqual(r["Confiabilidad"], "REVISAR")
        self.assertIn("deducido de la matriz", r["Motivos_Revision"])

    def test_excel_con_hoja_resumen(self):
        import tempfile
        import openpyxl
        regs = [dict(Matricula_Limpia="008-1", Estado="NO GRAVADA", Confiabilidad="ALTA"),
                dict(Matricula_Limpia="008-2", Estado="GRAVADA", Confiabilidad="REVISAR", Motivos_Revision="=prueba")]
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "r.xlsx"
            v.exportar_excel(regs, ruta)
            wb = openpyxl.load_workbook(ruta)
            self.assertEqual(wb.sheetnames, ["Resultados", "Resumen"])
            self.assertEqual(wb.active.title, "Resumen")
            valores = [c for fila in wb["Resumen"].iter_rows(values_only=True) for c in fila if c is not None]
            self.assertIn("Requieren revisión del profesional", valores)
            self.assertIn("=prueba", valores)  # texto, no fórmula


class Comparacion(unittest.TestCase):
    def _excel(self, d, nombre, regs):
        ruta = Path(d) / nombre
        v.exportar_excel(regs, ruta)
        return ruta

    def test_detecta_cambios_y_causa(self):
        import tempfile
        import openpyxl
        base = dict(Matricula_Limpia="01N-1", Municipio="BELLO", Confiabilidad="ALTA",
                    Cadena_Revisada="01N-1 (nivel 0, apertura ?, 5 anotaciones: GRAVADA)", Version_Aplicativo="A")
        with tempfile.TemporaryDirectory() as d:
            antes = self._excel(d, "a.xlsx", [dict(base, Estado="GRAVADA"), dict(base, Matricula_Limpia="01N-2", Estado="NO GRAVADA")])
            ahora = self._excel(d, "b.xlsx", [dict(base, Estado="GRAVAMEN CANCELADO",
                                                   Cadena_Revisada="01N-1 (nivel 0, apertura ?, 6 anotaciones: X)"),
                                              dict(base, Matricula_Limpia="01N-2", Estado="NO GRAVADA"),
                                              dict(base, Matricula_Limpia="01N-3", Estado="NO GRAVADA")])
            r = v.comparar_resultados(antes, ahora, Path(d) / "c.xlsx")
            self.assertEqual((r["comunes"], r["iguales"], r["cambio_estado"], r["solo_ahora"]), (2, 1, 1, 1))
            fila = list(openpyxl.load_workbook(Path(d) / "c.xlsx")["Cambios"].iter_rows(min_row=2, values_only=True))[0]
            self.assertEqual(fila[0], "01N-1")
            self.assertIn("inscripciones nuevas", fila[7])

    def test_lotes_distintos(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            a = self._excel(d, "a.xlsx", [dict(Matricula_Limpia="001-1", Estado="NO GRAVADA")])
            b = self._excel(d, "b.xlsx", [dict(Matricula_Limpia="002-1", Estado="NO GRAVADA")])
            with self.assertRaises(ValueError):
                v.comparar_resultados(a, b, Path(d) / "c.xlsx")


class HistoricoConsultas(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def reg(self, m="008-1", estado="NO GRAVADA", conf="ALTA", dias=1, **extra):
        from datetime import datetime, timedelta
        r = {"Matricula_Limpia": m, "Estado": estado, "Confiabilidad": conf,
             "Fecha_Consulta": f"{datetime.now() - timedelta(days=dias):%Y-%m-%d %H:%M:%S}", "Funcionario": "jperez",
             "Version_Reglas": v.VERSION_REGLAS}
        r.update(extra)
        return r

    def test_reglas_de_reutilizacion(self):
        e = v.evaluar_historico
        self.assertTrue(e(self.reg())[0])
        self.assertTrue(e(self.reg(estado="GRAVADA"))[0])
        self.assertFalse(e(self.reg(dias=v.VIGENCIA_HISTORICO_DIAS + 1))[0])       # vencida
        self.assertFalse(e(self.reg(conf="REVISAR"))[0])                           # requería revisión
        self.assertFalse(e(self.reg(estado="SIN DATOS EN VUR"))[0])
        self.assertFalse(e(self.reg(estado=v.ESTADO_INCOMPLETO))[0])
        self.assertFalse(e(self.reg(Version_Reglas="2026.01.01-1"))[0])            # reglas anteriores
        self.assertIn("jperez", e(self.reg())[1])

    def test_guardar_y_recargar_con_varios_usuarios(self):
        a = v.Historico(self.tmp.name)
        a.agregar(self.reg("008-1", dias=3))
        b = v.Historico(self.tmp.name)        # otra ventana / otro funcionario
        b._archivo_sesion = Path(self.tmp.name) / "registros" / "otro.jsonl"
        b.agregar(self.reg("008-1", estado="GRAVADA", dias=0))
        b.agregar(self.reg("008-2", estado="ERROR", conf="REVISAR", dias=0))
        (Path(self.tmp.name) / "registros" / "danado.jsonl").write_text('{"roto": \n', encoding="utf-8")
        h = v.Historico(self.tmp.name)
        self.assertEqual(h.cargar(), 2)
        self.assertEqual(h.buscar("008-1")["Estado"], "GRAVADA")            # el más reciente
        self.assertTrue(h.exportar_excel())
        self.assertTrue(h.archivo_excel.exists())

    def test_un_error_no_tapa_un_resultado_valido(self):
        h = v.Historico(self.tmp.name)
        h.agregar(self.reg("008-1", estado="NO GRAVADA", dias=2))
        h.agregar(self.reg("008-1", estado="ERROR", conf="REVISAR", dias=0))
        self.assertEqual(h.buscar("008-1")["Estado"], "NO GRAVADA")


if __name__ == "__main__":
    resultado = unittest.main(exit=False, verbosity=2).result
    print("\n" + ("TODAS LAS PRUEBAS PASARON" if resultado.wasSuccessful()
                  else "HAY PRUEBAS FALLIDAS: no use el programa hasta corregirlo"))
