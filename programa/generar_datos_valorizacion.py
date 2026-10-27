"""
Genera datos_valorizacion.py a partir de INFORMACION PROYECTOS DE VALORIZACION.xlsx.

Uso (solo cuando cambie el Excel):
    python generar_datos_valorizacion.py "INFORMACION PROYECTOS DE VALORIZACION (1).xlsx"

Hojas leídas: "Proy x Mpio", "Proyectos Valorización" y "Vencimiento".
El resultado queda junto a este script y vur_valorizacion.py lo importa, así
no se abre el Excel en cada consulta.
"""
import datetime
import pprint
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

NOMBRE_EXCEL = "INFORMACION PROYECTOS DE VALORIZACION (1).xlsx"
_aqui = Path(__file__).resolve().parent
_candidatos = [_aqui / NOMBRE_EXCEL, _aqui.parent / "datos" / NOMBRE_EXCEL]  # carpeta del script o ..\datos
RUTA = sys.argv[1] if len(sys.argv) > 1 else str(next((p for p in _candidatos if p.exists()), _candidatos[0]))
wb = openpyxl.load_workbook(RUTA, data_only=True)

def norm(t):
    t = unicodedata.normalize("NFD", str(t or ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()

def cod(v):
    s = str(v).strip()
    return s.zfill(3) if s.isdigit() else s

def fecha(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime.datetime) else None

# Proyectos (hoja Proyectos Valorización): código -> lista (pueden repetirse códigos)
proyectos = {}
for r in wb['Proyectos Valorización'].iter_rows(min_row=3, values_only=True):
    if r[0] is None: continue
    proyectos.setdefault(cod(r[0]), []).append({"obra": norm(r[3]), "resolucion": str(r[1]).strip(), "fecha_resolucion": fecha(r[2]), "anio": r[4]})

venc = {}
for r in wb['Vencimiento'].iter_rows(min_row=4, values_only=True):
    if r[0] is None: continue
    venc[cod(r[0])] = {"obra_venc": norm(r[1]), "resolucion_venc": str(r[2]).strip(), "fecha_inicio": fecha(r[3]),
                       "fecha_vencimiento": fecha(r[4]), "oficinas_registro": [norm(x) for x in str(r[5] or "").splitlines() if x.strip()]}

municipios = {}
for r in wb['Proy x Mpio'].iter_rows(min_row=3, values_only=True):
    if not r[1]: continue
    mun = norm(r[1]); texto = str(r[2] or "").strip(); nota = (str(r[3]).strip() if r[3] else "")
    lineas = [norm(l) for l in texto.splitlines() if l.strip()]
    lista = []
    for l in lineas:
        m = re.match(r"^(\d{3})\s*-?\s*(.*)$", l)
        if m: lista.append((m.group(1), m.group(2).strip(" -")))
        else: lista.append((None, l))
    if lineas == ["000 - NO GRAVADO"]:
        estado = "000"
    elif lineas == ["NO APLICA"]:
        estado = "NO_APLICA"
    elif "CONDONADA" in norm(nota):
        estado = "CONDONADA"
    else:
        estado = "001"
    projs = []
    for c, nombre in lista:
        if estado in ("000", "NO_APLICA"): break
        info = {"codigo": c, "obra": nombre}
        cands = proyectos.get(c, [])
        palabras = set(re.findall(r"[A-Z]{4,}", nombre))
        mejor = max(cands, key=lambda p: len(palabras & set(re.findall(r"[A-Z]{4,}", p["obra"]))), default=None) if cands else None
        comunes = lambda p: len(palabras & set(re.findall(r"[A-Z]{4,}", p["obra"])))
        if not mejor or (comunes(mejor) == 0 and mejor["obra"]):
            # Código distinto entre hojas (p. ej. 087 vs 017 TABLAZO): se busca por nombre.
            todos = [p for lst in proyectos.values() for p in lst]
            alt = max(todos, key=comunes)
            mejor = alt if comunes(alt) >= 2 else None
            if mejor: info["nota_cruce"] = "año tomado por nombre de obra (código distinto en hoja Proyectos)"
        if mejor: info.update({k: mejor[k] for k in ("resolucion", "fecha_resolucion", "anio")})
        if c in venc: info.update({k: venc[c][k] for k in ("fecha_inicio", "fecha_vencimiento", "oficinas_registro")})
        projs.append(info)
    municipios[mun] = {"cod_mpio": cod(r[0]), "estado": estado, "proyectos": projs, "nota": nota, "texto_hoja": texto}

salida = Path(__file__).with_name("datos_valorizacion.py")
encabezado = (
    '"""Datos de valorización generados automáticamente. No editar a mano.\n\n'
    f"Fuente: {Path(RUTA).name}\n"
    f"Generado: {datetime.datetime.now():%Y-%m-%d %H:%M}\n"
    "Regenerar con: python generar_datos_valorizacion.py <ruta del Excel>\n\n"
    "estado: '000' no gravado | '001' tiene proyectos | 'NO_APLICA' | 'CONDONADA'\n"
    '"""\n\n'
    f"FUENTE_DATOS = {Path(RUTA).name!r}\n\n"
)
salida.write_text(encabezado + "MUNICIPIOS_VALORIZACION = " + pprint.pformat(municipios, width=110, sort_dicts=True) + "\n", encoding="utf-8")
sin_anio = [(k, p["codigo"], p["obra"]) for k, m in municipios.items() for p in m["proyectos"] if "anio" not in p]
print(f"{len(municipios)} municipios escritos en {salida}")
for k, c, o in sin_anio:
    print(f"  Aviso: {k}: proyecto {c} {o} sin año en la hoja 'Proyectos Valorización'")
